from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd


DEFAULT_OUTPUT = Path("output") / "empirical"


def parse_cols(value: str | None) -> list[str]:
    if value is None:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def parse_semicolon(value: str | None) -> list[str]:
    if value is None:
        return []
    return [item.strip() for item in value.split(";") if item.strip()]


def ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def write_json(obj: Any, path: Path) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=True), encoding="utf-8")


def read_data(path: str | Path, sheet: str | None = None) -> pd.DataFrame:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".csv", ".txt"}:
        return pd.read_csv(path)
    if suffix == ".tsv":
        return pd.read_csv(path, sep="\t")
    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(path, sheet_name=sheet or 0)
    if suffix == ".parquet":
        return pd.read_parquet(path)
    if suffix == ".json":
        return pd.read_json(path)
    if suffix == ".jsonl":
        return pd.read_json(path, lines=True)
    if suffix == ".dta":
        return pd.read_stata(path)
    raise ValueError(f"Unsupported input extension: {suffix}")


def write_data(df: pd.DataFrame, path: str | Path) -> None:
    path = Path(path)
    ensure_dir(path.parent)
    suffix = path.suffix.lower()
    if suffix in {"", ".csv", ".txt"}:
        df.to_csv(path, index=False)
    elif suffix == ".tsv":
        df.to_csv(path, index=False, sep="\t")
    elif suffix in {".xlsx", ".xls"}:
        df.to_excel(path, index=False)
    elif suffix == ".parquet":
        df.to_parquet(path, index=False)
    elif suffix == ".json":
        df.to_json(path, orient="records", indent=2)
    elif suffix == ".jsonl":
        df.to_json(path, orient="records", lines=True)
    elif suffix == ".dta":
        df.to_stata(path, write_index=False)
    else:
        raise ValueError(f"Unsupported output extension: {suffix}")


def require_columns(df: pd.DataFrame, columns: Iterable[str], label: str = "columns") -> None:
    missing = [col for col in columns if col not in df.columns]
    if missing:
        raise ValueError(f"Missing {label}: {missing}")


def numeric_columns(df: pd.DataFrame, requested: list[str] | None = None) -> list[str]:
    if requested:
        require_columns(df, requested)
        return [col for col in requested if pd.api.types.is_numeric_dtype(df[col])]
    return list(df.select_dtypes(include=[np.number]).columns)


def write_docx_table(df: pd.DataFrame, path: Path) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Install python-docx to export docx tables.") from exc
    document = Document()
    table = document.add_table(rows=1, cols=len(df.columns))
    for idx, col in enumerate(df.columns):
        table.rows[0].cells[idx].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = "" if pd.isna(value) else str(value)
    ensure_dir(path.parent)
    document.save(path)


def export_table(df: pd.DataFrame, stem: Path, formats: list[str] | None = None) -> list[str]:
    formats = formats or ["csv", "xlsx", "tex"]
    ensure_dir(stem.parent)
    written: list[str] = []
    for fmt in formats:
        fmt = fmt.lower().lstrip(".")
        path = stem.with_suffix(f".{fmt}")
        if fmt == "csv":
            df.to_csv(path, index=False)
        elif fmt == "xlsx":
            df.to_excel(path, index=False)
        elif fmt == "tex":
            path.write_text(df.to_latex(index=False, escape=False), encoding="utf-8")
        elif fmt == "html":
            path.write_text(df.to_html(index=False, escape=False), encoding="utf-8")
        elif fmt == "json":
            path.write_text(df.to_json(orient="records", indent=2), encoding="utf-8")
        elif fmt == "docx":
            write_docx_table(df, path)
        else:
            raise ValueError(f"Unsupported table format: {fmt}")
        written.append(str(path))
    return written


def save_figure(fig: Any, stem: Path, png: bool = True, pdf: bool = True) -> list[str]:
    ensure_dir(stem.parent)
    written: list[str] = []
    if png:
        png_path = stem.with_suffix(".png")
        fig.savefig(png_path, dpi=300, bbox_inches="tight")
        written.append(str(png_path))
    if pdf:
        pdf_path = stem.with_suffix(".pdf")
        fig.savefig(pdf_path, bbox_inches="tight")
        written.append(str(pdf_path))
    return written


def import_matplotlib() -> Any:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    return plt


def add_common_io(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--input", required=True, help="Input data file.")
    parser.add_argument("--sheet", default=None, help="Excel sheet name or index if reading xlsx/xls.")


def add_table_format_arg(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--formats", default="csv,xlsx,tex", help="Comma-separated table formats.")


def run_with_errors(func: Any, argv: list[str] | None = None) -> int:
    try:
        func(argv)
    except Exception as exc:
        print(f"ERROR: {exc}", flush=True)
        return 1
    return 0
