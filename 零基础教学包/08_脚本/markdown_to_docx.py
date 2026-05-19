#!/usr/bin/env python
import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_code_block(doc, lines):
    if not lines:
        return
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_table(doc, rows):
    split_rows = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        split_rows.append(cells)
    if len(split_rows) < 2:
        return False
    if all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in split_rows[1]):
        body = [split_rows[0]] + split_rows[2:]
    else:
        body = split_rows
    if not body:
        return False
    width = max(len(r) for r in body)
    table = doc.add_table(rows=len(body), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(body):
        for j in range(width):
            table.cell(i, j).text = row[j] if j < len(row) else ""
    return True


def convert(md_path, docx_path):
    text = Path(md_path).read_text(encoding="utf-8")
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)

    in_code = False
    code_lines = []
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []

    for raw in text.splitlines():
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if line.startswith("|") and line.endswith("|"):
            table_lines.append(line)
            continue
        flush_table()

        if not line.strip():
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line).strip(), style="List Number")
        elif line == "---":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)

    flush_table()
    if code_lines:
        add_code_block(doc, code_lines)

    out = Path(docx_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_md")
    parser.add_argument("output_docx")
    args = parser.parse_args()
    convert(args.input_md, args.output_docx)


if __name__ == "__main__":
    main()

