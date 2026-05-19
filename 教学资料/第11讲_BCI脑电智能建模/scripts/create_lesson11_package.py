from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"


def write_texts():
    files = {
        "第11讲_教学大纲.md": """# 第11讲 BCI 脑电智能建模：从 EEG 特征到可复核解码系统

## 这节课解决什么科研问题和需求

本讲解决的是：已经有 EEG/ERP 数据或公开 BCI 数据集，但不知道如何把它变成可发表、可复现、不会训练测试泄漏的智能建模研究。

第10讲主要讲 EEG/ERP 实验设计、预处理、ERP、频域和连接分析。本讲进一步进入 BCI 和 EEG machine learning：数据结构、公开 benchmark、特征工程、经典机器学习、深度学习、跨被试迁移、模型解释、在线或伪在线解码，以及论文结果报告。

学生学完后应该能独立完成：

1. 建立 BCI 项目数据结构。
2. 选择合适的公开 benchmark 或自有 EEG 数据。
3. 明确预测任务、标签来源和分析单位。
4. 设计 EEG 特征工程方案。
5. 建立 classical ML baseline。
6. 判断什么时候可以使用 deep learning。
7. 设计 within-subject、cross-session、cross-subject 或 online 评估。
8. 审计训练测试泄漏。
9. 报告 subject-wise 指标、置信区间、混淆矩阵和失败被试。
10. 把 BCI 建模结果写成论文 Methods、Results 和 Limitations。

## 需要哪些 Skills

本讲建议复制以下 Skills 到第11讲工作区：

```text
bci-data-structure
bci-benchmark-datasets
eeg-feature-engineering-bci
eeg-ml-classical-bci
eeg-deep-learning-bci
eeg-cross-subject-transfer-bci
eeg-model-evaluation-leakage
eeg-model-interpretability-bci
bci-online-decoding
bci-results-reporting
```

这些 Skills 不要求安装到 Codex 根目录。课堂要求复制到工作区的 `skills/` 文件夹，便于学生迁移到自己的项目。

## 输入材料是什么

最低输入材料：

1. 研究问题：例如 AI 信任、认知负荷、情绪识别、运动想象、P300、SSVEP、错误相关电位。
2. EEG 数据说明：被试、session、run、trial、channel、sampling rate、事件码、标签。
3. 标签来源：实验条件、行为反应、主观量表、任务状态、外部事件。
4. 目标声明：被试内预测、跨 session 泛化、跨被试泛化、公开 benchmark 对比、伪在线或在线 BCI。
5. 已有预处理结果或计划：第10讲产出的 clean epochs、ERP/频域特征或 raw EEG。
6. 若使用公开数据：数据集链接、license、引用要求、数据格式和下载记录。

## Agent 应该怎么执行

Agent 的执行顺序必须固定：

1. 先建立数据结构和 metadata。
2. 再确认预测任务、标签、分析单位和 split unit。
3. 再做公开数据集或自有数据的适用性审计。
4. 再写特征工程方案。
5. 先建立 classical baseline。
6. 只有在样本量、验证设计和对照充分时才设计 deep learning。
7. 评估方案必须先于模型结果确定。
8. 所有 scaler、CSP、PCA、feature selection 都必须在训练 fold 内 fit。
9. 结果报告必须包含 subject-wise 指标和泄漏审计。
10. 在线或伪在线结果必须明确 latency、window、decision rule 和 offline-online gap。

## 输出文件应该长什么样

课堂最终输出包：

```text
第11讲_BCI脑电智能建模工作区/
├─ AGENTS.md
├─ research_log.md
├─ data/
│  ├─ raw/
│  ├─ events/
│  ├─ labels/
│  ├─ features/
│  └─ metadata/
├─ bci/
│  ├─ bci_project_structure.md
│  ├─ bci_dataset_selection.md
│  ├─ bci_prediction_task.md
│  ├─ bci_feature_engineering_plan.md
│  ├─ bci_classical_ml_plan.md
│  ├─ bci_deep_learning_plan.md
│  ├─ bci_cross_subject_transfer_plan.md
│  ├─ bci_online_decoding_plan.md
│  └─ bci_results_report.md
├─ qc/
│  ├─ bci_leakage_audit.md
│  ├─ bci_model_interpretability_audit.md
│  └─ reviewer_risk_register.md
├─ output/
│  ├─ model_comparison.xlsx
│  ├─ subject_metrics.xlsx
│  └─ figures/
└─ skills/
```

## 哪些地方必须人工核验

必须人工核验：

1. 标签是否在预测时刻真实可用。
2. 滑窗是否把同一 trial 切片分到训练集和测试集。
3. 是否用 target test labels 做了 domain adaptation。
4. 标准化、CSP、PCA、特征筛选是否只在训练 fold 内 fit。
5. 公开数据集 license 和 citation 是否正确。
6. 指标是否报告 subject-wise，而不是只报 pooled accuracy。
7. 深度学习是否有 classical baseline 对照。
8. online 结果是否真实在线，而不是离线准确率包装。
9. 模型解释是否只解释模型证据，而不是直接解释神经机制。
10. 论文里是否隐藏失败被试、失败 seed 或非显著结果。

## 学生课后作业

选择一个 EEG/BCI 研究任务，完成一个可审计的 BCI 建模方案。可以用自己的 EEG 数据，也可以选择公开数据集。

提交内容：

1. `bci_project_structure.md`
2. `bci_dataset_selection.md`
3. `bci_prediction_task.md`
4. `bci_feature_engineering_plan.md`
5. `bci_classical_ml_plan.md`
6. `bci_deep_learning_plan.md`，若不适合使用深度学习，写明不使用理由。
7. `bci_leakage_audit.md`
8. `bci_results_report.md`
9. 至少 1 个 Excel 模板填写结果。
10. 300 字反思：你的模型最可能被审稿人质疑什么，准备用什么补充分析回应。
""",
        "第11讲_零基础实操教程.md": """# 第11讲 零基础实操教程

## 0. 你今天要学会什么

你不需要一开始就会脑电建模。本教程按最小闭环来做：先把项目结构搭好，再让 Agent 生成每一步的计划、表格和审计文件。你要掌握的不是“马上跑出高准确率”，而是“让模型结果可信”。

最重要的判断是：你的验证方式决定你能说什么。

- 被试内划分只能说对同一被试的离线预测有效。
- 跨 session 划分才接近“同一人不同时间是否能泛化”。
- 跨被试划分才可以讨论新被试泛化。
- 在线 BCI 必须报告延迟、反馈、决策规则和用户负担。

## 1. 创建工作区

打开 PowerShell，执行：

```powershell
New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第11讲_BCI脑电智能建模工作区"
cd "D:\\AI科研训练营\\第11讲_BCI脑电智能建模工作区"

New-Item -ItemType Directory -Force -Path `
  data,data\\raw,data\\events,data\\labels,data\\features,data\\metadata,`
  bci,bci\\scripts,bci\\models,bci\\figures,qc,output,skills,input

New-Item -ItemType File -Force -Path AGENTS.md,research_log.md
```

检查你是否看到：

```text
data/
bci/
qc/
output/
skills/
AGENTS.md
research_log.md
```

## 2. 复制本讲 Skills

```powershell
cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"

powershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `
  -SkillName bci-data-structure,bci-benchmark-datasets,eeg-feature-engineering-bci,eeg-ml-classical-bci,eeg-deep-learning-bci,eeg-cross-subject-transfer-bci,eeg-model-evaluation-leakage,eeg-model-interpretability-bci,bci-online-decoding,bci-results-reporting `
  -Workspace "D:\\AI科研训练营\\第11讲_BCI脑电智能建模工作区"
```

复制完成后检查 `skills/`。如果缺少某个 Skill，不要强行继续，先在 `research_log.md` 记录缺失项，并用现有 Skill 的保守流程替代。

## 3. 写 AGENTS.md

把 `lesson11_workspace_AGENTS_template.md` 的内容复制到你的工作区 `AGENTS.md`。

核心约束必须保留：

```text
- 禁止覆盖 raw EEG 数据。
- 禁止把同一 trial 的滑窗切片分到 train/test 两边。
- 禁止在 test set 上调参。
- 禁止用 target test labels 做迁移。
- 禁止只报告 pooled accuracy。
- 禁止把模型解释直接写成神经机制。
```

## 4. 准备输入材料

把材料放到：

```text
input/研究问题.md
data/raw/原始脑电文件或说明
data/events/事件码表或 events.tsv
data/labels/标签说明
```

若没有真实 EEG 数据，也可以先用公开 benchmark 练习。你要让 Agent 生成 `bci_dataset_selection.md`，列出可用数据集、subjects、sessions、channels、labels、license、citation 和限制。

## 5. 第一次给 Agent 的任务：建立数据结构

复制给 Agent：

```text
请使用本工作区 skills/bci-data-structure。

任务：基于 input/研究问题.md、data/raw、data/events、data/labels，生成 bci/bci_project_structure.md。

要求：
1. 定义 subject -> session -> run -> trial -> window 层级。
2. 设计 metadata 表字段。
3. 写清 label schema 和预测时刻。
4. 先决定 split unit。
5. 审计 class balance 和泄漏风险。
6. 原始数据只读保存，不覆盖 raw 文件。

输出：
- 项目文件夹结构
- Metadata schema
- Label schema
- Split unit
- Class-balance audit
- Leakage risks
- Next modeling step
```

人工核验：

- subject、session、trial 是否能唯一定位。
- 标签是否来自预测时刻之后的信息。
- 同一 trial 的多个 window 是否可能被拆开。

## 6. 选择公开数据集或自有数据

复制给 Agent：

```text
请使用 skills/bci-benchmark-datasets。

我的任务是：[写你的任务，例如 motor imagery / P300 / SSVEP / cognitive workload / emotion / AI trust EEG]。

请生成 bci/bci_dataset_selection.md。比较 3-5 个候选公开数据集或说明自有数据是否足够。

比较字段：
dataset、task、subjects、sessions、channels、sampling_rate、labels、license、citation、MOABB/OpenNeuro/PhysioNet/BCI Competition 是否支持、适合的 split、主要限制。

注意：
- 不要编造数据集细节。
- 不确定 URL 或 license 时标记“待人工核验”。
- 不要把不匹配标签的数据集说成可用。
```

人工核验：

- 数据集文档是否真实存在。
- license 是否允许教学、研究和再发布衍生特征。
- 公开数据集是否真的支持你的任务标签。

## 7. 定义预测任务

复制给 Agent：

```text
请生成 bci/bci_prediction_task.md。

字段包括：
- prediction target
- labels and classes
- prediction time
- input unit: trial / epoch / sliding window / session / subject
- evaluation claim: within-subject / cross-session / cross-subject / online
- class balance
- chance level
- exclusion criteria
- cannot-claim boundary
```

人工核验：

- 二分类、多分类或回归是否明确。
- chance level 是否按类别分布而不是默认 50%。
- 如果是跨被试泛化，split 必须按 subject。

## 8. 设计特征工程

复制给 Agent：

```text
请使用 skills/eeg-feature-engineering-bci。

基于 bci_prediction_task.md，生成 bci/bci_feature_engineering_plan.md。

请比较：
1. band power / PSD
2. ERP amplitude 或 latency
3. CSP / FBCSP
4. Riemannian covariance features
5. connectivity features
6. time-window features

每类特征写清输入、窗口、通道、频段、fit 位置、输出表字段、适用模型和泄漏风险。
```

人工核验：

- scaling、CSP、feature selection、PCA 是否只在训练 fold 中 fit。
- 频段、窗口、通道是否有理论或 benchmark 依据。
- 特征数量是否远大于样本数。

## 9. 先做 classical baseline

复制给 Agent：

```text
请使用 skills/eeg-ml-classical-bci。

生成 bci/bci_classical_ml_plan.md。

模型至少包括：
- chance baseline
- majority baseline
- logistic regression 或 LDA
- SVM
- CSP+LDA 或 FBCSP
- Riemannian classifier

写清：
pipeline、nested CV、hyperparameters、metrics、subject-wise reporting、permutation test 和 model comparison。
```

人工核验：

- 所有模型必须用同一 split。
- 不允许只报告 best fold。
- 不允许不同模型用不同预处理结果比较。

## 10. 判断是否需要 deep learning

复制给 Agent：

```text
请使用 skills/eeg-deep-learning-bci。

请基于数据量、任务、split 和 baseline，生成 bci/bci_deep_learning_plan.md。

候选模型包括 EEGNet、DeepConvNet、ShallowConvNet、CNN-LSTM、TCN、Transformer 或 self-supervised pretraining。

必须输出：
- 是否建议使用 deep learning
- tensor schema: shape、channel、time、batch
- augmentation 规则
- train/validation/test 划分
- baseline 对照
- early stopping
- metrics
- failure risks
```

人工核验：

- augmentation 必须在 split 之后。
- 同一 trial 的 window 不能跨集合。
- 小样本情况下不要强行上 Transformer。

## 11. 跨被试迁移

复制给 Agent：

```text
请使用 skills/eeg-cross-subject-transfer-bci。

生成 bci/bci_cross_subject_transfer_plan.md。

写清：
- target deployment: new subject / new session / new device / new task
- validation: LOSO / leave-one-session-out / cross-dataset / calibration split
- source and target split
- adaptation method: normalization / covariance alignment / Riemannian alignment / fine-tuning / domain adaptation / few-shot calibration
- calibration policy
- subject-wise metrics
- failure-case analysis
```

人工核验：

- zero-calibration 不能用 target subject trials 训练。
- labeled calibration 必须写清用了多少 target labels。
- 不隐藏失败被试。

## 12. 在线或伪在线解码

复制给 Agent：

```text
请使用 skills/bci-online-decoding。

生成 bci/bci_online_decoding_plan.md。

写清：
- offline / pseudo-online / real online
- sampling rate
- window length and step
- causal preprocessing
- model loading
- smoothing / threshold / majority voting
- latency budget
- feedback timing
- online metrics: accuracy、latency、false positives、time-to-decision、ITR、fatigue
- offline-online gap
```

人工核验：

- 在线滤波不能使用未来样本。
- 离线准确率不能写成 online performance。
- threshold 不能在 online test data 上调。

## 13. 泄漏审计

复制给 Agent：

```text
请使用 skills/eeg-model-evaluation-leakage。

读取 bci_project_structure、prediction_task、feature_engineering_plan、classical_ml_plan、deep_learning_plan、cross_subject_transfer_plan，生成 qc/bci_leakage_audit.md。

逐项检查：
- window leakage
- subject leakage
- session leakage
- normalization leakage
- CSP/PCA/feature selection leakage
- hyperparameter tuning leakage
- target test label leakage
- augmentation leakage
- benchmark comparability

每项给 pass/revise/fail 和 required fix。
```

人工核验：

- 如果审计 fail，不能进入结果报告。
- 如果只能 revise，先改 split 和 pipeline，再继续。

## 14. 结果报告

复制给 Agent：

```text
请使用 skills/bci-results-reporting。

生成 bci/bci_results_report.md。

必须包含：
- split disclosure
- preprocessing and feature extraction disclosure
- baseline and chance level
- subject-wise metrics
- aggregate metrics with uncertainty
- confusion matrix
- class-level metrics
- permutation or paired tests
- model comparison table
- leakage safeguards
- failed subjects and limitations
- claim boundary
```

人工核验：

- 报告 balanced accuracy、AUC、F1、confusion matrix，不只 accuracy。
- 类别不平衡时必须报告 class-level metrics。
- 不要把 offline BCI 写成 real-time BCI。

## 15. 最终提交检查

你最终的文件夹至少应该有：

```text
bci/bci_project_structure.md
bci/bci_dataset_selection.md
bci/bci_prediction_task.md
bci/bci_feature_engineering_plan.md
bci/bci_classical_ml_plan.md
bci/bci_deep_learning_plan.md
bci/bci_cross_subject_transfer_plan.md
bci/bci_online_decoding_plan.md
bci/bci_results_report.md
qc/bci_leakage_audit.md
qc/bci_model_interpretability_audit.md
output/model_comparison.xlsx
output/subject_metrics.xlsx
```

如果没有真实模型结果，允许提交“建模方案版”。但不能编造准确率、AUC、F1 或 p 值。
""",
        "第11讲_课堂任务单.md": """# 第11讲 课堂任务单

## 课堂目标

本节课结束前，每位学生完成一个 BCI 脑电智能建模工作区，并产出一套可以继续跑模型的计划文件和审计表。

## 任务 1：创建工作区

完成后勾选：

- [ ] 已创建第11讲工作区。
- [ ] 已创建 `data/raw`、`data/events`、`data/labels`、`data/features`、`bci`、`qc`、`output`、`skills`。
- [ ] 已创建 `AGENTS.md` 和 `research_log.md`。
- [ ] 已复制本讲 Skills。

## 任务 2：填写 AGENTS.md

必须包含：

- [ ] raw EEG 不覆盖。
- [ ] 同一 trial 的 windows 不跨 train/test。
- [ ] test set 不调参。
- [ ] target test labels 不用于迁移。
- [ ] reporting 必须 subject-wise。
- [ ] online/offline 区分清楚。

## 任务 3：建立数据结构

输出：`bci/bci_project_structure.md`

人工检查：

- [ ] subject、session、run、trial、window 层级清楚。
- [ ] metadata 字段足够复现。
- [ ] label schema 说明标签来源和预测时刻。
- [ ] split unit 已在建模前确定。
- [ ] class balance 有检查。
- [ ] 已列泄漏风险。

## 任务 4：选择数据集或说明自有数据

输出：`bci/bci_dataset_selection.md`

人工检查：

- [ ] 至少比较 3 个候选公开数据集，或说明自有数据为什么可用。
- [ ] 数据集 task、subjects、sessions、channels、labels、license、citation 不缺失。
- [ ] 不确定信息已标“待人工核验”。
- [ ] 没有把不匹配标签的数据集硬说成可用。

## 任务 5：定义预测任务

输出：`bci/bci_prediction_task.md`

人工检查：

- [ ] target 和 classes 清楚。
- [ ] chance level 清楚。
- [ ] 输入单位清楚。
- [ ] 评估声明清楚：within-subject / cross-session / cross-subject / online。
- [ ] 不能声称的内容写清楚。

## 任务 6：特征工程方案

输出：`bci/bci_feature_engineering_plan.md`

人工检查：

- [ ] band power、ERP、CSP/FBCSP、Riemannian、connectivity 至少比较 3 类。
- [ ] 窗口、通道、频段清楚。
- [ ] scaler、CSP、PCA、feature selection 在 fold 内 fit。
- [ ] 输出特征表字段清楚。

## 任务 7：classical baseline

输出：`bci/bci_classical_ml_plan.md`

人工检查：

- [ ] chance baseline 和 majority baseline 存在。
- [ ] LDA/logistic/SVM/CSP/Riemannian 至少 3 类模型。
- [ ] 同一 split 比较模型。
- [ ] nested CV 或 validation set 明确。
- [ ] subject-wise metrics 明确。

## 任务 8：deep learning 判断

输出：`bci/bci_deep_learning_plan.md`

人工检查：

- [ ] 写明是否适合 deep learning。
- [ ] tensor schema 清楚。
- [ ] augmentation 在 split 之后。
- [ ] 有 classical baseline 对照。
- [ ] 没有夸大深度模型价值。

## 任务 9：跨被试与在线计划

输出：

- `bci/bci_cross_subject_transfer_plan.md`
- `bci/bci_online_decoding_plan.md`

人工检查：

- [ ] LOSO 或 cross-session 逻辑清楚。
- [ ] calibration policy 清楚。
- [ ] online latency 和 window 设置清楚。
- [ ] 离线、伪在线、真实在线区分清楚。

## 任务 10：泄漏审计与结果报告

输出：

- `qc/bci_leakage_audit.md`
- `bci/bci_results_report.md`

人工检查：

- [ ] 泄漏审计每项有 pass/revise/fail。
- [ ] revise/fail 有 required fix。
- [ ] 结果报告包含 split、baseline、subject-wise、confusion matrix、uncertainty 和 limitations。
- [ ] 没有编造模型指标。
""",
        "第11讲_教师带做讲稿.md": """# 第11讲 教师带做讲稿

## 开场 0-10 分钟

今天这一讲不要让学生追求“高准确率”。先告诉学生：BCI 论文最常见的问题不是模型不够新，而是验证设计不支持论文声称。

可以用一句话开场：

> EEG 建模里，split 决定了你能说什么；泄漏会让所有漂亮准确率失去意义。

接着解释第10讲和第11讲的区别：

- 第10讲：把 EEG/ERP 实验做得可复核。
- 第11讲：把 EEG 数据变成可审计的 BCI 智能建模系统。

## 10-25 分钟：工作区与 Skills

教师屏幕操作：

1. 创建第11讲工作区。
2. 复制 10 个 BCI Skills。
3. 打开 `lesson11_workspace_AGENTS_template.md`。
4. 让学生把约束写入 `AGENTS.md`。

强调：

- Skills 放工作区，不放根目录。
- 本讲所有模型计划都必须被 AGENTS.md 的边界约束。
- API key 不写进任何文件。

## 25-45 分钟：数据结构和预测任务

教师讲解层级：

```text
subject -> session -> run -> trial -> window
```

现场让学生判断：

1. 如果同一 trial 切成 20 个 window，能不能随机分 train/test？
2. 如果目标是新被试泛化，能不能同一被试部分 trial 用于训练？
3. 如果标签来自 trial 结束后的量表，能不能用于实时预测？

标准答案：

1. 不能，window leakage。
2. 不能，subject leakage。
3. 取决于预测时刻；若预测时不可用，就是 label leakage。

## 45-60 分钟：数据集选择

教师展示 `bci-benchmark-datasets` 的用途：不要凭记忆说某数据集能用，必须看 task、labels、license、citation。

现场任务：

- 学生选一个任务：motor imagery、P300、SSVEP、workload、emotion、AI trust。
- Agent 生成 `bci_dataset_selection.md`。
- 学生标出“推荐数据集”和“不可用数据集”。

教师提醒：

- 不要把 emotion 数据集拿来做 cognitive workload，除非标签确实支持。
- 不同 benchmark split 不可直接比较。

## 60-80 分钟：特征工程

教师讲解五类常用特征：

1. band power / PSD
2. ERP amplitude / latency
3. CSP / FBCSP
4. Riemannian covariance
5. connectivity

强调：

- 特征是测量选择，不是模型装饰。
- CSP、PCA、feature selection 必须在训练 fold 内 fit。
- 频段、通道、窗口要和任务有关。

课堂纠错：

如果学生写“全通道、全频段、全窗口全部扔给模型”，要指出这会带来维度灾难和后验选择风险。

## 80-100 分钟：classical baseline

教师让学生先做 baseline，不要直接 deep learning。

解释 baseline 层次：

```text
chance -> majority -> LDA/logistic -> SVM -> CSP+LDA -> Riemannian
```

教师强调：

- 没有 baseline，深度学习结果无法解释。
- 模型比较必须同一 split。
- 指标要有 balanced accuracy、AUC、F1、confusion matrix。

## 100-120 分钟：deep learning 和跨被试迁移

教师讲深度学习适用条件：

- 数据量足够。
- 有独立 validation/test。
- 有 classical baseline。
- 有不泄漏的数据增强。
- 报告 seed 或重复结果。

跨被试迁移重点：

- zero-calibration：不能用 target subject labeled trials。
- few-shot calibration：必须说明用了多少 labeled target data。
- LOSO 必须报告每个 left-out subject。

## 120-135 分钟：online / pseudo-online

教师讲清三种结果：

- Offline：所有数据事后分析。
- Pseudo-online：按时间顺序模拟在线，但仍在离线数据上。
- Real online：实时采集、实时推理、实时反馈。

强调：

- 在线滤波不能用未来数据。
- latency、window、step、decision rule 必须报告。
- 离线准确率不能写成实时系统性能。

## 135-150 分钟：泄漏审计与结果报告

让学生运行 prompt bank 中的“泄漏审计”任务。

教师逐项看：

1. window leakage
2. subject leakage
3. normalization leakage
4. hyperparameter tuning leakage
5. augmentation leakage
6. target test label leakage

最后展示结果报告模板，强调：

- 先报告 split，因为 split 决定 claim。
- 再报告 baseline 和 chance。
- 再报告 subject-wise。
- 最后才报告 aggregate。

## 课堂收尾

总结语：

> 第11讲的核心不是把模型变复杂，而是让 BCI 结论和验证设计一一对应。只要 split 错了，模型再高级也不能救论文。
""",
        "第11讲_课后作业.md": """# 第11讲 课后作业

## 作业目标

基于一个 EEG/BCI 任务，完成一个可复现、可审计、可写入论文的建模方案。

## 任务要求

任选一种：

1. 使用自己的 EEG/ERP 数据。
2. 使用公开 BCI benchmark。
3. 如果暂时没有数据，使用假想项目，但不能编造模型结果，只能提交建模方案。

## 必交文件

```text
AGENTS.md
bci/bci_project_structure.md
bci/bci_dataset_selection.md
bci/bci_prediction_task.md
bci/bci_feature_engineering_plan.md
bci/bci_classical_ml_plan.md
bci/bci_deep_learning_plan.md
bci/bci_cross_subject_transfer_plan.md
bci/bci_online_decoding_plan.md
qc/bci_leakage_audit.md
qc/bci_model_interpretability_audit.md
bci/bci_results_report.md
```

Excel 至少提交 3 个：

```text
bci_dataset_structure_template.xlsx
bci_feature_plan_template.xlsx
bci_split_evaluation_template.xlsx
bci_model_metrics_template.xlsx
bci_leakage_audit_template.xlsx
bci_training_log_template.xlsx
```

## 评分标准

| 项目 | 占比 |
|---|---:|
| 数据结构与标签定义 | 20% |
| 评估声明与 split 设计 | 25% |
| 特征工程与 baseline | 20% |
| deep learning / transfer / online 的适用性判断 | 15% |
| 泄漏审计与结果报告 | 20% |

## 一票否决

出现以下任一问题，作业直接退回：

1. 编造准确率、AUC、F1、p 值或数据集信息。
2. 同一 trial 的滑窗跨 train/test。
3. 声称跨被试泛化但 split 没有按 subject。
4. 在 test set 上调参。
5. 只报告 pooled accuracy，不报告 subject-wise。
6. 把 offline 结果写成 online BCI。

## 反思题

用 300-500 字回答：

1. 你的模型最容易发生哪一种泄漏？
2. 你的验证设计最多能支持什么 claim？
3. 如果审稿人质疑泛化能力，你准备补什么分析？
4. 如果模型解释与神经机制不一致，你怎么写限制？
""",
        "lesson11_workspace_AGENTS_template.md": """# 第11讲 BCI 脑电智能建模工作区 Agent 指令

你是 BCI 脑电智能建模助手，服务于管理学、心理学、神经科学和计算机科学交叉科研。

## 基本原则

- 禁止覆盖、移动或重写 raw EEG 原始数据。
- 所有衍生数据必须写入 `data/features`、`bci/models`、`output` 或 `qc`。
- 所有模型计划必须先说明 prediction target、label source、prediction time 和 split unit。
- 如果信息不足，标注“待人工核验”，不要编造。

## 建模边界

- 禁止把同一 trial 的滑窗切片分到 train/test 两边。
- 禁止在 test set 上调参或选择模型。
- 禁止在全数据上 fit scaler、PCA、CSP、feature selection 后再划分。
- 禁止用 target test labels 做 transfer learning 或 domain adaptation。
- 禁止只报告 pooled accuracy。
- 禁止隐藏失败被试、失败 seed 或不显著结果。

## 论文表达边界

- 模型预测结果不能直接证明神经机制。
- offline、pseudo-online、real online 必须区分。
- 跨被试、跨 session、跨数据集的 claim 必须由对应 split 支持。
- 模型解释只能解释模型依赖的信号模式，不得直接写成心理构念或脑机制。

## 输出要求

所有输出优先写入：

```text
bci/
qc/
output/
```

每个结果文件必须包含：

1. 输入文件。
2. 处理步骤。
3. 参数或决策规则。
4. 输出文件。
5. 人工核验项。
6. 不能支持的结论。
""",
        "lesson11_agent_prompt_bank.md": """# 第11讲 Agent Prompt Bank

## Prompt 1：建立 BCI 项目结构

```text
请使用 skills/bci-data-structure。基于 input/研究问题.md、data/raw、data/events、data/labels，生成 bci/bci_project_structure.md。定义 subject-session-run-trial-window 层级、metadata schema、label schema、split unit、class-balance audit 和 leakage risks。禁止覆盖 raw EEG。
```

## Prompt 2：选择 benchmark 数据集

```text
请使用 skills/bci-benchmark-datasets。我的任务是：[填写任务]。请生成 bci/bci_dataset_selection.md，比较 3-5 个候选公开数据集或自有数据是否可用，字段包括 subjects、sessions、channels、sampling_rate、labels、license、citation、split、限制。不要编造 URL、license 或 citation。
```

## Prompt 3：定义预测任务

```text
请生成 bci/bci_prediction_task.md。写清 prediction target、class labels、label source、prediction time、input unit、evaluation claim、class balance、chance level、exclusion criteria 和 cannot-claim boundary。
```

## Prompt 4：生成特征工程计划

```text
请使用 skills/eeg-feature-engineering-bci。生成 bci/bci_feature_engineering_plan.md。比较 band power、ERP、CSP/FBCSP、Riemannian、connectivity、time-window features。说明窗口、通道、频段、fit 位置、输出字段和泄漏风险。
```

## Prompt 5：生成 classical baseline 计划

```text
请使用 skills/eeg-ml-classical-bci。生成 bci/bci_classical_ml_plan.md。至少包括 chance、majority、LDA/logistic、SVM、CSP+LDA、Riemannian。写清 pipeline、nested CV、hyperparameter plan、metrics、subject-wise report、permutation test 和 model comparison。
```

## Prompt 6：判断 deep learning 是否适合

```text
请使用 skills/eeg-deep-learning-bci。基于数据量、split、baseline 和任务，生成 bci/bci_deep_learning_plan.md。候选 EEGNet、DeepConvNet、ShallowConvNet、CNN-LSTM、TCN、Transformer。必须给出是否建议使用 deep learning、tensor schema、augmentation、early stopping、metrics、failure risks。
```

## Prompt 7：跨被试迁移计划

```text
请使用 skills/eeg-cross-subject-transfer-bci。生成 bci/bci_cross_subject_transfer_plan.md。说明 target deployment、validation、source/target split、adaptation method、calibration policy、subject-wise metrics、failure-case analysis。
```

## Prompt 8：在线或伪在线计划

```text
请使用 skills/bci-online-decoding。生成 bci/bci_online_decoding_plan.md。说明 offline/pseudo-online/real online、window length、step size、causal preprocessing、decision rule、latency budget、online metrics、offline-online gap 和 deployment risks。
```

## Prompt 9：泄漏审计

```text
请使用 skills/eeg-model-evaluation-leakage。读取 bci/ 下全部计划，生成 qc/bci_leakage_audit.md。逐项检查 window、subject、session、normalization、CSP/PCA/feature selection、hyperparameter tuning、target test label、augmentation、benchmark comparability。每项给 pass/revise/fail 和 required fix。
```

## Prompt 10：模型解释审计

```text
请使用 skills/eeg-model-interpretability-bci。生成 qc/bci_model_interpretability_audit.md。说明可用解释方法、适合解释的层级、不能支持的神经机制 claim、需要的人类核验项和论文写作边界。
```

## Prompt 11：结果报告

```text
请使用 skills/bci-results-reporting。生成 bci/bci_results_report.md。必须先报告 split，再报告 baseline、chance level、subject-wise metrics、aggregate metrics with uncertainty、confusion matrix、class-level metrics、leakage safeguards、failed subjects、limitations 和 claim boundary。没有真实结果时，输出表格模板，不要编造指标。
```
""",
        "bci_model_plan_template.md": """# BCI Model Plan Template

## 1. Prediction Task

- Target:
- Label source:
- Prediction time:
- Classes or regression target:
- Input unit:
- Evaluation claim:

## 2. Data Structure

- Subjects:
- Sessions:
- Runs:
- Trials:
- Windows:
- Metadata file:

## 3. Split Design

- Split unit:
- Train:
- Validation:
- Test:
- Nested CV:
- Claim supported:
- Claim not supported:

## 4. Feature Engineering

| Feature | Window | Channels | Frequency | Fit inside fold | Output |
|---|---|---|---|---|---|
| band power |  |  |  | yes/no |  |
| ERP |  |  |  | yes/no |  |
| CSP/FBCSP |  |  |  | yes/no |  |
| Riemannian |  |  |  | yes/no |  |

## 5. Models

- Chance:
- Majority:
- LDA/logistic:
- SVM:
- CSP+LDA:
- Riemannian:
- Deep learning, if justified:

## 6. Metrics

- Accuracy:
- Balanced accuracy:
- AUC:
- F1:
- Confusion matrix:
- Subject-wise:
- Uncertainty:

## 7. Risks

- Leakage:
- Class imbalance:
- Failed subjects:
- Overfitting:
- Online gap:
""",
        "bci_leakage_audit_template.md": """# BCI Leakage Audit Template

| Risk | Question | Evidence | Status | Required Fix |
|---|---|---|---|---|
| Window leakage | Same trial windows split across train/test? |  | pass/revise/fail |  |
| Subject leakage | Same subject in train/test for cross-subject claim? |  | pass/revise/fail |  |
| Session leakage | Same session mixed when claiming cross-session? |  | pass/revise/fail |  |
| Normalization leakage | Scaler fit on all data? |  | pass/revise/fail |  |
| CSP/PCA leakage | CSP/PCA fit outside training fold? |  | pass/revise/fail |  |
| Feature selection leakage | Selected features using all labels? |  | pass/revise/fail |  |
| Tuning leakage | Test set used for hyperparameter selection? |  | pass/revise/fail |  |
| Augmentation leakage | Augmented samples cross split boundary? |  | pass/revise/fail |  |
| Target label leakage | Target test labels used in adaptation? |  | pass/revise/fail |  |
| Benchmark mismatch | Compared results use incompatible split? |  | pass/revise/fail |  |

## Final Decision

- Overall: pass / revise / fail
- Must fix before modeling:
- Must fix before manuscript:
""",
        "bci_results_report_template.md": """# BCI Results Report Template

## Split Disclosure

Describe the split first. State whether the result is within-subject, cross-session, cross-subject, cross-dataset, pseudo-online, or online.

## Baseline And Chance

| Model | Split | Chance | Balanced Accuracy | AUC | F1 | Notes |
|---|---|---:|---:|---:|---:|---|

## Subject-Wise Metrics

| Subject | Accuracy | Balanced Accuracy | AUC | F1 | Failed? | Notes |
|---|---:|---:|---:|---:|---|---|

## Confusion Matrix

Fill one matrix per main model. Include class labels and sample size.

## Statistical Comparison

- Paired test or permutation test:
- Correction:
- Effect size:
- CI:

## Leakage Safeguards

- Split rule:
- Fold-internal preprocessing:
- Tuning:
- Augmentation:
- Transfer policy:

## Limitations

- Offline/online boundary:
- Subject variability:
- Dataset limitation:
- Model interpretability limitation:
- Claim not supported:
""",
        "bci_online_decoding_plan_template.md": """# BCI Online Decoding Plan Template

## System Type

- Offline:
- Pseudo-online:
- Real online:

## Streaming Constraints

- Sampling rate:
- Window length:
- Step size:
- Buffer:
- Hardware:

## Causal Preprocessing

- Filter:
- Referencing:
- Artifact handling:
- Feature extraction:
- Future samples used? Must be no.

## Decision Rule

- Model:
- Probability threshold:
- Smoothing:
- Majority voting:
- Rejection class:
- Feedback timing:

## Latency Budget

| Component | Expected ms |
|---|---:|
| Acquisition buffer |  |
| Preprocessing |  |
| Feature extraction |  |
| Inference |  |
| Decision smoothing |  |
| Feedback display |  |

## Online Metrics

- Accuracy:
- Balanced accuracy:
- False positive rate:
- Time-to-decision:
- ITR:
- User fatigue:

## Offline-Online Gap

Explain why offline performance may not transfer to online deployment.
""",
    }
    for name, text in files.items():
        (ROOT / name).write_text(text, encoding="utf-8")


def style_sheet(ws, widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    side = Side(style="thin", color="D2D8DE")
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, color="202A36")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="0D746E")
                cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7F8")
    ws.row_dimensions[1].height = 36
    for idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1] if widths and idx <= len(widths) else 20


def make_wb(filename, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows, widths in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_sheet(ws, widths)
    path = ROOT / filename
    wb.save(path)
    load_workbook(path).close()


def create_excels():
    status = '"pass,revise,fail,to_check,not_applicable"'
    yesno = '"yes,no,to_check,not_applicable"'
    outputs = {
        "bci_dataset_structure_template.xlsx": [
            ("subjects_sessions", ["subject_id", "session_id", "run_id", "task", "date", "device", "n_channels", "sampling_rate", "notes"], [["sub-001", "ses-01", "run-01", "to_check", "YYYY-MM-DD", "to_check", "", "", ""]], [16, 16, 16, 24, 16, 24, 16, 18, 42]),
            ("trials_epochs", ["subject_id", "session_id", "run_id", "trial_id", "event_code", "condition", "label", "window_id", "start_ms", "end_ms", "file_path"], [["sub-001", "ses-01", "run-01", "trial-001", "", "", "", "win-001", "", "", ""]], [16, 16, 16, 16, 16, 22, 18, 16, 16, 16, 44]),
            ("labels", ["label_name", "class_value", "source", "prediction_time_available", "leakage_risk", "notes"], [["workload", "high/low", "task condition", "yes", "to_check", ""]], [22, 18, 28, 28, 28, 44]),
            ("file_lineage", ["file", "source", "created_by", "created_at", "do_not_overwrite", "notes"], [["data/raw/sub-001.vhdr", "acquisition", "human", "YYYY-MM-DD", "yes", "raw EEG"]], [42, 30, 18, 18, 18, 44]),
        ],
        "bci_benchmark_dataset_template.xlsx": [
            ("dataset_candidates", ["dataset", "task", "subjects", "sessions", "channels", "sampling_rate", "labels", "repository", "license", "citation", "fit_to_task", "limitations"], [["to_check", "motor imagery/P300/SSVEP/workload", "", "", "", "", "", "MOABB/OpenNeuro/PhysioNet", "to_check", "to_check", "to_check", ""]], [26, 34, 14, 14, 14, 16, 30, 30, 22, 42, 18, 46]),
            ("comparability_audit", ["comparison", "same_split", "same_preprocessing", "same_metric", "status", "notes"], [["baseline vs proposed", "to_check", "to_check", "to_check", "to_check", ""]], [32, 18, 22, 18, 18, 46]),
        ],
        "bci_feature_plan_template.xlsx": [
            ("feature_dictionary", ["feature_name", "feature_family", "window", "channels", "frequency", "fit_inside_fold", "parameters", "output_column", "interpretation", "risk"], [["alpha_power_parietal", "band_power", "to_check", "Pz/Oz/to_check", "8-12 Hz", "yes", "", "", "neural signal feature, not direct construct", "to_check"]], [26, 22, 20, 24, 18, 20, 36, 26, 44, 34]),
            ("transform_pipeline", ["step", "operation", "fit_on", "transform_on", "inside_cv", "notes"], [["scaler", "StandardScaler", "train fold", "valid/test fold", "yes", ""]], [22, 28, 24, 24, 18, 46]),
        ],
        "bci_split_evaluation_template.xlsx": [
            ("split_design", ["claim", "split_unit", "train", "validation", "test", "allowed", "notes"], [["cross-subject", "subject", "source subjects", "source validation subjects", "held-out subject", "yes", "LOSO or group split"]], [26, 22, 30, 30, 30, 16, 46]),
            ("generalization_claims", ["result_type", "can_claim", "cannot_claim", "required_evidence"], [["within-subject", "same-subject offline prediction", "new subject generalization", "subject-specific split and report"], ["cross-subject", "new subject offline generalization", "real-time online BCI", "held-out subjects"], ["online", "real-time system performance", "if not actually online", "latency and feedback logs"]], [24, 38, 40, 46]),
            ("nested_cv_plan", ["outer_split", "inner_split", "tuned_parameters", "test_usage", "status"], [["LOSO", "source-subject CV", "C/gamma/features", "final evaluation only", "to_check"]], [24, 24, 34, 34, 18]),
        ],
        "bci_model_metrics_template.xlsx": [
            ("model_comparison", ["model", "split", "accuracy", "balanced_accuracy", "auc", "f1", "kappa", "ci_or_sd", "chance", "notes"], [["SVM", "LOSO/to_check", "", "", "", "", "", "", "", ""]], [24, 22, 14, 20, 14, 14, 14, 18, 16, 44]),
            ("subject_metrics", ["subject_id", "model", "accuracy", "balanced_accuracy", "auc", "f1", "n_trials", "failed", "failure_reason"], [["sub-001", "SVM", "", "", "", "", "", "to_check", ""]], [16, 22, 14, 20, 14, 14, 16, 16, 44]),
            ("confusion_matrix", ["true_label", "predicted_label", "count", "model", "split", "notes"], [["class_0", "class_0", "", "to_check", "to_check", ""]], [24, 24, 14, 22, 22, 44]),
        ],
        "bci_leakage_audit_template.xlsx": [
            ("leakage_audit", ["risk", "question", "evidence_file", "status", "required_fix", "notes"], [["window_leakage", "same trial windows across train/test?", "split_design", "to_check", "", ""]], [26, 56, 34, 18, 42, 42]),
            ("required_fixes", ["priority", "issue", "fix", "owner", "deadline", "status"], [["high", "to_check", "to_check", "student", "YYYY-MM-DD", "to_check"]], [16, 44, 44, 18, 18, 18]),
        ],
        "bci_training_log_template.xlsx": [
            ("training_run_log", ["run_id", "date", "data_version", "feature_version", "split_version", "model", "parameters", "seed", "result_file", "status", "notes"], [["run_001", "YYYY-MM-DD", "v1", "v1", "v1", "SVM", "", "42", "", "to_check", ""]], [16, 16, 18, 18, 18, 24, 40, 12, 36, 18, 44]),
            ("environment", ["item", "value", "source", "notes"], [["python_version", "", "python -V", ""], ["mne_version", "", "import mne", ""], ["sklearn_version", "", "import sklearn", ""]], [24, 34, 34, 44]),
        ],
        "bci_online_decoding_template.xlsx": [
            ("latency_budget", ["component", "expected_ms", "measured_ms", "status", "notes"], [["buffer", "", "", "to_check", ""], ["preprocessing", "", "", "to_check", ""], ["feature_extraction", "", "", "to_check", ""], ["inference", "", "", "to_check", ""], ["feedback", "", "", "to_check", ""]], [28, 18, 18, 18, 46]),
            ("online_metrics", ["session", "accuracy", "balanced_accuracy", "false_positive_rate", "time_to_decision_ms", "itr", "fatigue_note", "notes"], [["ses-01", "", "", "", "", "", "", ""]], [18, 14, 20, 22, 22, 16, 34, 44]),
        ],
    }
    for filename, sheets in outputs.items():
        make_wb(filename, sheets)
        wb = load_workbook(ROOT / filename)
        dv_status = DataValidation(type="list", formula1=status, allow_blank=True)
        dv_yesno = DataValidation(type="list", formula1=yesno, allow_blank=True)
        for ws in wb.worksheets:
            ws.add_data_validation(dv_status)
            ws.add_data_validation(dv_yesno)
            dv_status.add(f"A2:Z500")
            dv_yesno.add(f"A2:Z500")
        wb.save(ROOT / filename)
        load_workbook(ROOT / filename).close()


COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, level = item if isinstance(item, tuple) else (item, 0)
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=8.6):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 11 讲 · BCI 脑电智能建模", 0.45, 0.055, 3.5, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.0, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=9.5):
    gap = 0.12
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.01), Inches(y + h / 2), Inches(bx - 0.02), Inches(y + h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"])
    no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.16, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.18, accent=COLORS["amber"], body_size=11.2, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.72, size=8.2)
    textbox(slide, "人工核验", 4.65, 4.82, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.3)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 11 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "BCI 脑电\n智能建模", 0.52, 1.70, 2.25, 1.20, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "从 EEG 特征到可审计解码系统", 0.55, 6.10, 2.25, 0.45, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "让脑电机器学习结果能泛化、可复现、经得起审稿", 3.55, 1.10, 8.75, 0.72, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "Data Structure · Features · Baselines · Deep Learning · Transfer · Online · Leakage Audit", 3.58, 1.95, 8.75, 0.35, size=14, color=COLORS["muted"])
    flow(slide, ["数据", "标签", "特征", "Baseline", "深度模型", "迁移", "在线", "报告"], 3.60, 3.10, 8.85, h=0.62, size=9.0)
    card(slide, "本讲交付", "bci_project_structure、dataset_selection、prediction_task、feature_engineering_plan、classical_ml_plan、deep_learning_plan、transfer_plan、online_plan、leakage_audit、results_report。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=11.4)
    card(slide, "底线", "split 决定 claim；没有泄漏审计的高准确率不能写进论文主结论。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def create_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, ["有 EEG 数据，但不知道怎么变成 BCI 建模研究", "模型准确率很高，但可能只是训练测试泄漏", "不清楚 within-subject、cross-session、cross-subject 各自能声称什么", "想用 EEGNet/Transformer，但没有 baseline 和验证边界", "想做在线 BCI，却把离线结果包装成实时系统", "论文里缺 subject-wise 指标、失败被试和模型限制"], 0.85, 1.55, 5.95, 3.20, size=15.0)
    card(slide, "本讲不做", "不追求漂亮准确率；不把模型解释直接写成神经机制；不隐藏失败被试。", 7.05, 1.72, 5.10, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "建立一套可复现、可审计、可写进论文的 EEG/BCI 智能建模流程。", 7.05, 3.35, 5.10, 1.20, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["结构", "标签", "特征", "模型", "验证", "泄漏", "报告"], 0.95, 5.65, 11.35, h=0.62, size=10.5)

    slide = new("本讲最终产出文件")
    rows = [["阶段", "文件", "作用"], ["结构", "bci_project_structure.md", "subject/session/run/trial/window 与 metadata"], ["数据集", "bci_dataset_selection.md", "公开 benchmark 或自有数据适用性"], ["任务", "bci_prediction_task.md", "标签、预测时刻、split 和 claim"], ["特征", "bci_feature_engineering_plan.md", "band power、ERP、CSP、Riemannian、连接"], ["模型", "classical_ml_plan / deep_learning_plan", "baseline、深度模型、调参和对照"], ["迁移/在线", "transfer_plan / online_decoding_plan", "跨被试、校准、延迟、反馈"], ["审计/报告", "leakage_audit / results_report", "防泄漏、subject-wise 指标和论文写作"]]
    table(slide, rows, 0.70, 1.46, 11.95, 4.98, col_widths=[1.55, 4.75, 5.65], font_size=8.5)
    card(slide, "提交底线", "没有泄漏审计、没有 split disclosure、没有 subject-wise metrics 的结果不能作为主结论。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [["功能", "Skills", "用途"], ["结构", "bci-data-structure", "层级、metadata、label schema、split unit"], ["数据集", "bci-benchmark-datasets", "MOABB、PhysioNet、OpenNeuro、BCI Competition"], ["特征", "eeg-feature-engineering-bci", "PSD、ERP、CSP、FBCSP、Riemannian、连接"], ["Baseline", "eeg-ml-classical-bci", "LDA、SVM、CSP+LDA、Riemannian"], ["深度模型", "eeg-deep-learning-bci", "EEGNet、DeepConvNet、TCN、Transformer"], ["迁移", "eeg-cross-subject-transfer-bci", "LOSO、domain adaptation、few-shot calibration"], ["在线/报告", "bci-online-decoding / bci-results-reporting", "latency、ITR、subject-wise、limitations"], ["审计", "eeg-model-evaluation-leakage", "window、subject、normalization、tuning 泄漏"]]
    table(slide, rows, 0.68, 1.42, 12.0, 5.10, col_widths=[1.55, 4.55, 5.9], font_size=8.4)

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第11讲_BCI脑电智能建模工作区"\ncd "D:\\AI科研训练营\\第11讲_BCI脑电智能建模工作区"\n\nNew-Item -ItemType Directory -Force -Path data,data\\raw,data\\events,data\\labels,data\\features,data\\metadata,bci,bci\\scripts,bci\\models,bci\\figures,qc,output,skills,input\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md', 0.85, 1.55, 11.65, 2.50, size=8.9)
    card(slide, "检查标准", "raw 只放原始 EEG；features 放衍生特征；models 放训练产物；qc 放审计。", 0.95, 4.45, 5.60, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "输入示例", "clean epochs、events.tsv、labels.csv、公开 benchmark 文档、预处理 QC。", 6.85, 4.45, 5.60, 1.05, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName bci-data-structure,bci-benchmark-datasets,eeg-feature-engineering-bci,eeg-ml-classical-bci,eeg-deep-learning-bci,eeg-cross-subject-transfer-bci,eeg-model-evaluation-leakage,eeg-model-interpretability-bci,bci-online-decoding,bci-results-reporting `\n  -Workspace "D:\\AI科研训练营\\第11讲_BCI脑电智能建模工作区"', 0.80, 1.48, 11.75, 2.92, size=8.0)
    bullets(slide, ["复制后检查 skills/ 文件夹。", "本讲 Skills 只放工作区，后续按需复制到自己的项目。", "缺失 Skill 要写入 research_log.md，不要假装已使用。"], 0.95, 4.75, 11.1, 1.0, size=14.2)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第11讲工作区 Agent 指令\n\n你是 BCI 脑电智能建模助手。\n\n原则：\n- 禁止覆盖 raw EEG。\n- 禁止同一 trial windows 跨 train/test。\n- 禁止 test set 调参。\n- 禁止全数据 fit scaler/CSP/PCA。\n- 禁止 target test labels 做迁移。\n- 必须报告 subject-wise metrics。\n- offline、pseudo-online、real online 必须区分。", 0.85, 1.55, 6.45, 3.50, size=9.2)
    card(slide, "为什么要写", "BCI 建模自由度很高，先把不可碰的边界写进工作区。", 7.62, 1.70, 4.45, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson11_workspace_AGENTS_template.md。", 7.62, 3.20, 4.45, 0.80, accent=COLORS["blue"])
    card(slide, "不能省略", "没有边界时，Agent 容易把泄漏结果包装成模型进步。", 7.62, 4.45, 4.45, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    slide = new("BCI 建模全流程地图")
    flow(slide, ["研究任务", "数据结构", "标签", "split", "特征", "baseline", "深度模型", "迁移", "在线", "报告"], 0.74, 1.48, 11.9, h=0.62, size=8.4)
    rows = [["如果缺失", "后果"], ["metadata 不清", "无法追溯 subject/session/trial"], ["标签时刻不清", "可能 label leakage"], ["split unit 不清", "claim 不成立"], ["baseline 缺失", "无法判断新模型价值"], ["泄漏审计缺失", "高准确率不可信"]]
    table(slide, rows, 1.05, 2.75, 11.10, 2.85, col_widths=[3.2, 7.9], font_size=10.2)
    card(slide, "原则", "先定义 split 和 claim，再训练模型；不要看结果后倒推验证方式。", 1.05, 6.08, 11.10, 0.58, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)

    step_slide(prs, 4, "实操 4：建立数据结构", "定义 subject、session、run、trial、window 和 label schema。", "请使用 skills/bci-data-structure，生成 bci/bci_project_structure.md。写清 metadata、label schema、split unit、class balance 和 leakage risks。", "bci/bci_project_structure.md", ["层级唯一", "标签来源清楚", "split unit 先定", "class balance 可查", "raw 不覆盖"], "大多数 BCI 泄漏从数据结构混乱开始。")

    slide = new("数据结构应长什么样")
    rows = [["层级", "例子", "建模意义"], ["subject", "sub-001", "跨被试泛化必须隔离"], ["session", "ses-01", "跨 session 泛化必须隔离"], ["run", "run-01", "采集批次和漂移"], ["trial", "trial-001", "滑窗不能跨 split"], ["window", "win-001", "模型输入单位"], ["label", "high workload", "预测目标和 chance level"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.45, col_widths=[2.0, 3.0, 6.75], font_size=10.0)
    card(slide, "课堂判断", "如果同一 trial 被切成多个 window，它们必须一起进入 train 或 test。", 0.95, 6.22, 11.35, 0.62, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(prs, 5, "实操 5：选择 benchmark 或自有数据", "确认数据集是否真正支持研究任务和论文 claim。", "请使用 skills/bci-benchmark-datasets，生成 bci/bci_dataset_selection.md。比较 task、subjects、sessions、channels、labels、license、citation、split 和限制。", "bci/bci_dataset_selection.md", ["数据集真实", "license 可用", "标签匹配", "split 可比", "citation 待核验"], "不确定数据集细节时必须标待人工核验。")

    step_slide(prs, 6, "实操 6：定义预测任务", "把研究问题变成明确的 prediction target 和 evaluation claim。", "请生成 bci/bci_prediction_task.md。包括 target、class labels、label source、prediction time、input unit、evaluation claim、class balance、chance level、cannot-claim boundary。", "bci/bci_prediction_task.md", ["target 清楚", "标签时刻可用", "chance level 正确", "claim 与 split 匹配", "不能声称写清楚"], "预测任务定义不好，后面所有模型都会失焦。")

    slide = new("Split 决定 Claim")
    rows = [["Split", "能说", "不能说"], ["random epoch", "同一数据内部分类", "跨被试/真实泛化"], ["within-subject", "同一被试离线预测", "新被试泛化"], ["cross-session", "同一被试跨时间泛化", "新被试泛化"], ["LOSO", "新被试离线泛化", "真实在线表现"], ["online", "实时系统表现", "如果没有延迟和反馈记录就不能说"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.05, col_widths=[2.4, 4.7, 4.65], font_size=9.8)
    card(slide, "课堂底线", "声称跨被试泛化，却随机划分 epoch，是 BCI 论文中非常严重的问题。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(prs, 7, "实操 7：生成特征工程计划", "选择可解释、可复现、适合任务的 EEG 特征。", "请使用 skills/eeg-feature-engineering-bci，生成 bci/bci_feature_engineering_plan.md。比较 band power、ERP、CSP/FBCSP、Riemannian、connectivity、time-window features。", "bci/bci_feature_engineering_plan.md", ["窗口/通道/频段清楚", "fit 在训练 fold", "输出字段清楚", "特征与任务匹配", "泄漏风险写明"], "特征工程是测量选择，不是模型装饰。")

    slide = new("常用 BCI 特征")
    rows = [["特征", "适合", "注意事项"], ["Band power / PSD", "负荷、注意、情绪、MI", "频段和归一化要先验"], ["ERP amplitude", "P300、ERN、FRN 等", "窗口和 ROI 不能后验"], ["CSP / FBCSP", "运动想象等空间模式", "CSP 必须 fold 内 fit"], ["Riemannian", "协方差模式", "split 和正则化要清楚"], ["Connectivity", "耦合模式探索", "不能写成因果方向"]]
    table(slide, rows, 0.78, 1.48, 11.75, 3.85, col_widths=[2.6, 4.15, 4.95], font_size=9.8)
    card(slide, "人工核验", "特征数量远大于样本数时，要先降维、正则化或改为更稳定的特征。", 0.95, 5.88, 11.35, 0.75, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.3)

    step_slide(prs, 8, "实操 8：建立 Classical Baseline", "先用简单模型建立可信对照，再讨论复杂模型。", "请使用 skills/eeg-ml-classical-bci，生成 bci/bci_classical_ml_plan.md。至少包括 chance、majority、LDA/logistic、SVM、CSP+LDA、Riemannian。", "bci/bci_classical_ml_plan.md", ["同一 split", "nested CV", "不只 best run", "subject-wise", "permutation test"], "没有 baseline 的 deep learning 结果说服力很弱。")

    slide = new("Baseline 层次")
    flow(slide, ["chance", "majority", "LDA/logistic", "SVM", "CSP+LDA", "Riemannian", "proposed"], 0.82, 1.45, 11.7, h=0.62, size=9.4)
    rows = [["模型", "价值"], ["chance / majority", "确认任务不是类别比例造成"], ["LDA / logistic", "线性可解释 baseline"], ["SVM", "传统强 baseline"], ["CSP+LDA", "运动想象常用对照"], ["Riemannian", "协方差特征强 baseline"], ["proposed", "必须超过合理对照才有意义"]]
    table(slide, rows, 1.05, 2.65, 11.10, 3.25, col_widths=[3.0, 8.1], font_size=10.0)

    step_slide(prs, 9, "实操 9：判断深度学习是否必要", "在数据量、验证设计和 baseline 足够时才引入深度模型。", "请使用 skills/eeg-deep-learning-bci，生成 bci/bci_deep_learning_plan.md。包括 EEGNet、DeepConvNet、ShallowConvNet、CNN-LSTM、TCN、Transformer 适用性、tensor schema、augmentation 和 early stopping。", "bci/bci_deep_learning_plan.md", ["数据量足够", "tensor schema 清楚", "augmentation split 后", "有 baseline", "不夸大机制"], "小数据直接上 Transformer 通常不是好主意。")

    slide = new("Deep Learning 适用边界")
    rows = [["情况", "建议"], ["小样本、少被试", "先 classical baseline；深度模型只做探索"], ["多被试、多 session", "可考虑 EEGNet/DeepConvNet"], ["时间动态强", "可考虑 TCN/CNN-LSTM"], ["跨被试迁移", "需要 LOSO、校准和 subject-wise report"], ["Transformer", "需要足够数据、正则和对照，不做默认选择"]]
    table(slide, rows, 0.78, 1.48, 11.75, 3.85, col_widths=[3.2, 8.55], font_size=10.0)
    card(slide, "底线", "深度模型要和 classical baseline 在同一 split 下比较，不能只报最好 seed。", 0.95, 5.88, 11.35, 0.75, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    step_slide(prs, 10, "实操 10：跨被试迁移计划", "设计新被试、新 session 或新设备的泛化验证。", "请使用 skills/eeg-cross-subject-transfer-bci，生成 bci/bci_cross_subject_transfer_plan.md。说明 deployment、validation、source/target split、adaptation、calibration、subject-wise metrics。", "bci/bci_cross_subject_transfer_plan.md", ["LOSO 清楚", "target labels 不泄漏", "calibration 数量写明", "失败被试报告", "不只 pooled"], "zero-calibration 和 few-shot calibration 必须严格区分。")

    slide = new("跨被试迁移常见路线")
    rows = [["路线", "用法", "风险"], ["LOSO", "每次留一被试做 test", "subject variance 大"], ["covariance alignment", "对齐协方差分布", "不能用 target test labels"], ["Riemannian alignment", "协方差空间对齐", "参数估计要透明"], ["fine-tuning", "少量 target calibration", "必须说明 target labels"], ["domain adversarial", "跨域学习", "需要严格验证和 baseline"]]
    table(slide, rows, 0.78, 1.48, 11.75, 3.85, col_widths=[2.8, 4.4, 4.55], font_size=9.7)
    card(slide, "报告要求", "每个 left-out subject 都要有指标；失败被试不能删除后只报平均。", 0.95, 5.88, 11.35, 0.75, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    step_slide(prs, 11, "实操 11：在线或伪在线解码", "把离线模型转成可讨论延迟、窗口和反馈的系统计划。", "请使用 skills/bci-online-decoding，生成 bci/bci_online_decoding_plan.md。包括 offline/pseudo-online/real online、window、step、causal preprocessing、decision rule、latency、metrics。", "bci/bci_online_decoding_plan.md", ["不用未来样本", "latency 清楚", "threshold 不用 test 调", "反馈时序清楚", "offline-online gap 写明"], "离线准确率不能写成实时 BCI 性能。")

    slide = new("Online BCI 结果怎么写")
    rows = [["类型", "需要报告"], ["Offline", "split、features、metrics、limitations"], ["Pseudo-online", "时间顺序、滑窗、因果处理、模拟决策"], ["Real online", "硬件、延迟、反馈、ITR、false positives、疲劳"], ["All", "decision rule、threshold、rejection class、失败 session"]]
    table(slide, rows, 0.78, 1.48, 11.75, 3.20, col_widths=[2.4, 9.35], font_size=10.8)
    card(slide, "人工核验", "如果滤波或特征用了未来样本，它不是严格在线流程。", 0.95, 5.25, 11.35, 0.82, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.8)

    step_slide(prs, 12, "实操 12：泄漏审计", "在结果报告之前逐项排除训练测试泄漏。", "请使用 skills/eeg-model-evaluation-leakage，生成 qc/bci_leakage_audit.md。检查 window、subject、session、normalization、CSP/PCA、feature selection、tuning、augmentation、target labels。", "qc/bci_leakage_audit.md", ["每项 pass/revise/fail", "fail 不进结果", "修正方案清楚", "证据文件可查", "claim 重新匹配"], "泄漏审计不过，不建议继续写论文结果。")

    slide = new("泄漏类型速查")
    rows = [["泄漏", "例子", "修正"], ["window", "同一 trial windows 跨 train/test", "按 trial 分组"], ["subject", "同一被试同时训练和测试", "按 subject 分组"], ["normalization", "全数据 fit scaler", "fold 内 fit"], ["feature selection", "全标签选特征", "训练 fold 内选"], ["tuning", "test 上选参数", "nested CV"], ["augmentation", "增强样本跨集合", "split 后只增强 train"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.40, col_widths=[2.3, 4.85, 4.60], font_size=9.7)

    slide = new("模型解释：能说什么，不能说什么")
    rows = [["解释方法", "能说", "不能说"], ["feature importance", "模型依赖哪些特征", "该脑区产生心理机制"], ["saliency", "输入时间/通道敏感区域", "真实神经因果"], ["ablation", "去掉特征后性能变化", "构念被证明"], ["topomap", "空间模式描述", "直接定位脑源"], ["SHAP", "预测贡献", "生理机制结论"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.05, col_widths=[2.6, 4.7, 4.45], font_size=9.7)
    card(slide, "写作边界", "模型解释服务于模型透明度，不等同神经机制证据。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(prs, 13, "实操 13：结果报告", "把模型结果写成可投稿的 BCI Results 结构。", "请使用 skills/bci-results-reporting，生成 bci/bci_results_report.md。先报告 split，再报告 baseline、chance、subject-wise、aggregate、confusion matrix、统计比较、leakage safeguards、limitations。", "bci/bci_results_report.md", ["split 先写", "subject-wise", "失败被试不隐藏", "不只 accuracy", "offline/online 边界"], "没有真实结果时只输出模板，不能编造指标。")

    slide = new("BCI 结果表应该包含")
    rows = [["表", "必须字段"], ["Model comparison", "model、split、balanced accuracy、AUC、F1、CI/SD、chance"], ["Subject-wise", "subject、metrics、failed、failure reason"], ["Confusion matrix", "true labels、predicted labels、class counts"], ["Statistical tests", "paired/permutation、effect size、CI"], ["Leakage disclosure", "split、fold-internal transform、tuning、augmentation"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.10, col_widths=[3.0, 8.75], font_size=9.8)
    card(slide, "审稿人会看", "你的结论是否由 split 支持，失败被试是否透明，指标是否处理类别不平衡。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    slide = new("课堂示范：AI 信任 EEG BCI")
    rows = [["模块", "示范内容"], ["任务", "根据 EEG 预测用户是否处于高认知负荷或低信任状态"], ["数据", "AI 推荐阅读任务、事件码、行为采纳、量表"], ["特征", "额中 theta、顶区 alpha、P3/LPP、Riemannian covariance"], ["模型", "SVM / Riemannian baseline；EEGNet 作为探索"], ["split", "被试内用于个体适配；LOSO 才讨论新用户泛化"], ["边界", "模型预测低信任状态，不等于证明信任神经机制"]]
    table(slide, rows, 0.78, 1.45, 11.75, 4.95, col_widths=[2.0, 9.75], font_size=9.6)
    card(slide, "最大风险", "标签来自事后量表时，不能直接说实时预测用户信任。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("常见错误与修正")
    rows = [["错误", "后果", "修正"], ["随机 epoch split", "跨被试 claim 不成立", "Group split / LOSO"], ["全数据标准化", "normalization leakage", "Pipeline fold 内 fit"], ["只报最好 seed", "结果不稳定", "重复种子和均值/CI"], ["无 baseline", "无法证明改进", "classical baseline"], ["隐藏失败被试", "泛化性被夸大", "subject-wise 报告"], ["offline 写成 online", "部署 claim 错误", "区分结果类型"]]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.55, 4.10, 5.10], font_size=9.5)

    slide = new("课堂 150 分钟带做安排")
    rows = [["时间", "教师带做", "学生产出"], ["0-15 分钟", "BCI claim 与 split", "理解底线"], ["15-30 分钟", "工作区、Skills、AGENTS", "可运行工作区"], ["30-50 分钟", "数据结构和预测任务", "project_structure / task"], ["50-65 分钟", "benchmark 选择", "dataset_selection"], ["65-85 分钟", "特征工程", "feature_plan"], ["85-105 分钟", "classical baseline", "classical_ml_plan"], ["105-120 分钟", "deep learning 判断", "deep_learning_plan"], ["120-135 分钟", "迁移和在线", "transfer / online"], ["135-150 分钟", "泄漏审计和报告", "audit / report"]]
    table(slide, rows, 0.78, 1.45, 11.75, 5.05, col_widths=[1.8, 5.0, 4.95], font_size=8.6)

    slide = new("课后提交要求")
    bullets(slide, ["提交文件夹：姓名_第11讲_BCI脑电智能建模", "必须包含 project_structure、dataset_selection、prediction_task、feature_plan、classical_ml_plan、deep_learning_plan、transfer_plan、online_plan、leakage_audit、results_report", "至少填写 3 个 Excel 模板", "附 300-500 字反思：你的验证设计最多支持什么 claim，最大泄漏风险是什么"], 0.90, 1.55, 11.3, 1.80, size=14.3)
    rows = [["评分项", "占比"], ["数据结构与标签定义", "20%"], ["split 与 claim 匹配", "25%"], ["特征和 baseline", "20%"], ["deep/transfer/online 判断", "15%"], ["泄漏审计和报告", "20%"]]
    table(slide, rows, 2.15, 3.85, 8.85, 2.10, col_widths=[6.6, 2.25], font_size=11.3)
    card(slide, "一票否决", "编造指标、window 泄漏、subject 泄漏、test 调参、只报 pooled accuracy。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.1)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是得到一个高准确率，而是建立一套让 BCI 结果和科研声称严格匹配的建模系统。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["结构", "标签", "split", "特征", "baseline", "deep", "transfer", "online", "audit", "report"], 1.00, 3.50, 11.20, h=0.70, size=8.5)
    card(slide, "下一步", "学生把第11讲结果接入第12讲论文写作、审稿自查和投稿排版流程。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(ROOT / "第11讲_课件.pptx")


def main():
    write_texts()
    create_excels()
    create_ppt()
    print(ROOT)


if __name__ == "__main__":
    main()
