from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第03讲_课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(247, 248, 245),
    "ink": RGBColor(31, 41, 51),
    "muted": RGBColor(91, 103, 112),
    "teal": RGBColor(15, 118, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(180, 83, 9),
    "red": RGBColor(185, 28, 28),
    "line": RGBColor(210, 215, 220),
    "white": RGBColor(255, 255, 255),
    "pale_teal": RGBColor(219, 245, 241),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(254, 243, 199),
    "pale_red": RGBColor(254, 226, 226),
    "soft": RGBColor(236, 239, 241),
    "dark": RGBColor(33, 43, 54),
}

FONT = "Microsoft YaHei"


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color=COLORS["line"], width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def clear_line(shape):
    shape.line.fill.background()


def add_text(slide, text, x, y, w, h, size=20, color=COLORS["ink"], bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=COLORS["ink"], gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def add_code(slide, text, x, y, w, h, size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, COLORS["dark"])
    clear_line(shape)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(235, 239, 245)
    return shape


def add_header(slide, title, section="第 03 讲", num=None):
    set_fill(slide.background, COLORS["bg"])
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34)
    )
    set_fill(bar, COLORS["teal"])
    clear_line(bar)
    add_text(slide, section, 0.45, 0.05, 1.6, 0.22, size=8, color=COLORS["white"], bold=True)
    if num is not None:
        add_text(slide, f"{num:02d}", 12.05, 0.04, 0.7, 0.25, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.62, 11.8, 0.55, size=25, color=COLORS["ink"], bold=True)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.21), Inches(1.1), Inches(0.06)
    )
    set_fill(line, COLORS["amber"])
    clear_line(line)


def add_card(slide, title, body, x, y, w, h, fill=COLORS["white"], accent=COLORS["teal"],
             title_size=15, body_size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, COLORS["line"], 0.8)
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    set_fill(stripe, accent)
    clear_line(stripe)
    add_text(slide, title, x + 0.18, y + 0.15, w - 0.28, 0.35, size=title_size, color=accent, bold=True)
    add_text(slide, body, x + 0.18, y + 0.62, w - 0.32, h - 0.72, size=body_size, color=COLORS["ink"])
    return shape


def add_table(slide, rows, x, y, w, h, col_widths=None, font_size=10, header_fill=COLORS["teal"]):
    nrows, ncols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size if r else font_size + 0.5)
                p.font.bold = bool(r == 0)
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(242, 245, 247))
    return table_shape


def add_flow(slide, labels, x, y, w, box_h=0.58, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11):
    gap = 0.18
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(box_h)
        )
        set_fill(shape, color)
        set_line(shape, accent, 1)
        add_text(slide, label, bx + 0.08, y + 0.14, box_w - 0.16, box_h - 0.2,
                 size=font_size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            x1 = Inches(bx - gap + 0.02)
            x2 = Inches(bx - 0.02)
            ymid = Inches(y + box_h / 2)
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, ymid, x2, ymid)
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.2)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_fill(slide.background, COLORS["bg"])
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.0), Inches(7.5))
    set_fill(left, COLORS["teal"])
    clear_line(left)
    add_text(slide, "真实文献检索\n引用核验\n平台联动", 0.5, 0.95, 2.2, 2.05, size=24, color=COLORS["white"], bold=True)
    add_text(slide, "第 03 讲", 0.6, 5.95, 1.4, 0.35, size=14, color=COLORS["white"], bold=True)
    add_text(slide, "AMiner · AI4Scholar · BibTeX · PDF · 飞书文献日报", 3.45, 1.22, 9.0, 0.7, size=28, color=COLORS["ink"], bold=True)
    add_text(slide, "真实文献、自动引用、学者画像、科学画图、目标期刊 PDF 下载与团队协作", 3.5, 2.1, 8.7, 0.45, size=15, color=COLORS["muted"])
    add_flow(slide, ["检索", "核验", "BibTeX", "PDF", "日报", "飞书"], 3.55, 3.25, 8.8, box_h=0.66, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_card(slide, "本讲交付", "candidate_papers.xlsx、verified_references.bib、pdf_download_log.xlsx、daily report、AMiner scholar profile、sci_draw prompt、飞书表字段", 3.55, 4.65, 8.75, 1.35, fill=COLORS["white"], accent=COLORS["amber"], title_size=14, body_size=13)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    title_slide(prs)

    def new_slide(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_header(slide, title, num=len(prs.slides))
        return slide

    slide = new_slide("这节课解决什么科研问题")
    add_bullets(slide, [
        "让 AI 找真实文献，而不是编造引用",
        "生成可核验 BibTeX，并做句子级支持复核",
        "按目标期刊和主题追踪新论文",
        "合法下载开放 PDF 或已有权限全文到本地",
        "用 AMiner 做学者画像、同名消歧和团队分析",
        "用 AI4Scholar sci_draw 生成科研机制图或流程图",
        "把每日文献日报同步到飞书表格，服务课题组协作",
    ], 0.9, 1.55, 6.1, 3.8, size=17)
    add_card(slide, "本讲核心", "文献工作流必须有证据链：检索来源、稳定 ID、BibTeX、PDF 状态、人工核验、日志。", 7.25, 1.8, 4.95, 1.65, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)
    add_card(slide, "课堂底线", "没有 DOI/ID/URL/平台证据的文献，不能写成已核验。", 7.25, 4.05, 4.95, 1.25, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=14)

    slide = new_slide("本讲固定模板")
    rows = [
        ["模块", "第三讲怎么落实"],
        ["科研问题和需求", "真实文献、BibTeX、PDF、日报、学者画像、飞书联动"],
        ["需要哪些 Skills", "AI4Scholar、AMiner、OpenAlex、citation、lark-sheets、pdf/xlsx"],
        ["输入材料", "主题、关键词、种子论文、目标期刊、学者、平台凭据"],
        ["Agent 怎么执行", "检索、核验、下载、制表、日报、同步"],
        ["输出文件", "candidate_papers、bib、pdf log、daily report、scholar profile"],
        ["人工核验", "引用支持、元数据、同名学者、PDF 合法性"],
        ["课后作业", "20 篇候选、10 条 BibTeX、5 个 PDF、1 份日报"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.9, col_widths=[2.65, 9.2], font_size=12)
    add_text(slide, "第三讲开始，文献输出必须能被导师或审稿人追溯。", 0.9, 6.55, 11.2, 0.35, size=15, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("真实文献生产线")
    add_flow(slide, ["研究问题", "检索式", "真实平台", "元数据核验", "BibTeX", "PDF", "日报/飞书"], 0.7, 1.55, 12.0, box_h=0.7, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11)
    add_card(slide, "输入", "研究主题、关键词、种子论文、目标期刊、目标学者、时间范围。", 0.85, 2.95, 3.7, 1.3, fill=COLORS["white"], accent=COLORS["teal"])
    add_card(slide, "平台", "AI4Scholar、AMiner、OpenAlex、Semantic Scholar、PubMed、arXiv、CNKI。", 4.85, 2.95, 3.7, 1.3, fill=COLORS["white"], accent=COLORS["blue"])
    add_card(slide, "输出", "候选文献表、BibTeX、PDF、文献日报、飞书表格、学者画像。", 8.85, 2.95, 3.7, 1.3, fill=COLORS["white"], accent=COLORS["amber"])
    add_text(slide, "模型负责组织流程，平台负责提供证据，人工负责最终核验。", 0.9, 5.65, 11.5, 0.45, size=18, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("平台功能侧重")
    rows = [
        ["平台", "最擅长", "本讲用途"],
        ["AI4Scholar", "论文级检索、详情、引用网络、PDF、auto_cite、sci_draw", "找真实论文、BibTeX、PDF、画图"],
        ["AMiner", "学者画像、机构团队、同名消歧、专利、知识图谱", "专家发现和学者核验"],
        ["OpenAlex", "开放学术元数据、趋势、引用信息", "交叉核验和批量元数据"],
        ["Semantic Scholar", "AI/CS/认知科学、开放 PDF 链接", "补充检索和 PDF 来源"],
        ["PubMed", "医学、神经科学、心理生理", "神经科学/BCI 文献"],
        ["飞书表格", "团队协作、筛选、任务跟踪", "每日文献日报和核验表"],
    ]
    add_table(slide, rows, 0.6, 1.45, 12.15, 4.95, col_widths=[2.0, 5.2, 4.95], font_size=10.5)

    slide = new_slide("需要哪些 Skills")
    rows = [
        ["功能", "Skills"],
        ["AI4Scholar 总入口", "ai4scholar-research"],
        ["论文搜索/详情", "ai4scholar-paper-search, ai4scholar-paper-detail-batch, academic-research-openalex"],
        ["引用网络", "ai4scholar-citation-network"],
        ["PDF/全文", "ai4scholar-pdf-fulltext-reading, pdf, markitdown"],
        ["BibTeX/自动引用", "ai4scholar-auto-citation-bibtex, citation-management"],
        ["学者画像", "aminer-mcp-research, ai4scholar-author-intelligence"],
        ["文献推荐", "ai4scholar-paper-recommendation"],
        ["科学画图", "ai4scholar-sci-draw"],
        ["飞书联动", "lark-sheets, lark-base, lark-doc"],
    ]
    add_table(slide, rows, 0.65, 1.35, 12.05, 5.45, col_widths=[3.0, 9.05], font_size=10.8)

    slide = new_slide("工作区结构")
    add_code(slide, "第03讲_真实文献工作区/\n├─ AGENTS.md\n├─ configs/mcp.example.json\n├─ skills/\n├─ literature/\n│  ├─ search_strategy.md\n│  ├─ candidate_papers.xlsx\n│  ├─ verified_references.bib\n│  ├─ bibtex_audit.md\n│  ├─ pdfs/\n│  ├─ pdf_download_log.xlsx\n│  ├─ daily_reports/\n│  └─ aminer_scholar_profiles.md\n├─ figures/\n├─ feishu/\n├─ qc/\n└─ research_log.md", 0.85, 1.45, 5.8, 5.45, size=11.5)
    add_bullets(slide, [
        "literature：文献、PDF、BibTeX、日报",
        "figures：sci_draw 提示词和图形审查",
        "feishu：飞书字段和待写入数据",
        "qc：人工核验清单",
        "configs：只放示例，不放真实密钥",
    ], 7.05, 1.75, 5.15, 2.75, size=18)
    add_card(slide, "关键原则", "每个文献条目都必须有来源、ID、状态和下一步动作。", 7.05, 5.2, 5.15, 0.9, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("MCP 配置：AI4Scholar 与 AMiner")
    add_code(slide, "{\n  \"mcpServers\": {\n    \"ai4scholar\": {\n      \"url\": \"https://mcp.ai4scholar.net/sse\",\n      \"headers\": {\n        \"Authorization\": \"Bearer ${AI4SCHOLAR_API_KEY}\"\n      }\n    },\n    \"aminer\": {\n      \"url\": \"https://mcp.aminer.cn/sse\",\n      \"headers\": {\n        \"Authorization\": \"Bearer ${AMINER_MCP_TOKEN}\"\n      }\n    }\n  }\n}", 0.85, 1.45, 6.2, 4.9, size=11)
    add_bullets(slide, [
        "真实 Token 只放环境变量或私有配置",
        "配置后重启 Agent 或 gateway",
        "先 list tools，再按实际 schema 调用",
        "MCP 不可用时输出手动查询计划",
        "不要编造工具返回结果",
    ], 7.45, 1.8, 4.65, 2.8, size=17)

    slide = new_slide("检索策略：先写检索式")
    add_code(slide, "(\"AI agent\" OR \"LLM agent\" OR \"autonomous agent\")\nAND (\"literature review\" OR \"scholarly search\" OR \"research workflow\")\nAND (\"human-AI collaboration\" OR \"scientific discovery\")", 0.85, 1.55, 6.1, 1.45, size=13)
    add_bullets(slide, [
        "至少 4 条英文检索式和 2 条中文检索式",
        "覆盖理论、方法、数据和应用场景",
        "限定时间范围、目标期刊或数据库",
        "不要检索前就编造论文清单",
    ], 7.35, 1.65, 4.85, 2.15, size=18)
    add_table(slide, [
        ["输出", "literature/search_strategy.md"],
        ["字段", "topic, query, source, time_range, screening_criteria"],
        ["核验", "检索式是否过窄、过宽或只含流行词"],
    ], 0.85, 4.25, 11.8, 1.45, col_widths=[2.3, 9.5], font_size=12)

    slide = new_slide("候选文献表")
    rows = [
        ["字段", "说明"],
        ["title / authors / year / venue", "基础元数据"],
        ["DOI_or_ID / url", "核验路径"],
        ["abstract", "初筛依据"],
        ["source", "AI4Scholar / OpenAlex / PubMed / arXiv 等"],
        ["relevance_score", "1-5 分"],
        ["method_tag", "experiment / EEG / ML / text mining / review 等"],
        ["verification_status", "unverified / partial / verified / rejected"],
        ["next_action", "read / cite / download_pdf / check_author / ignore"],
    ]
    add_table(slide, rows, 0.7, 1.45, 11.95, 5.15, col_widths=[3.1, 8.85], font_size=11.5)
    add_text(slide, "没有稳定 ID 的文献不是不能用，但必须明确标注“需人工核验”。", 0.9, 6.65, 11.2, 0.35, size=15, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("怎么让 AI 生成真实 BibTeX")
    add_flow(slide, ["句子拆分", "判断需引用主张", "auto_cite", "生成 BibTeX", "元数据核验", "支持强度核验"], 0.8, 1.55, 11.8, box_h=0.7, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=11)
    add_code(slide, "Use AI4Scholar auto_cite to add APA 7th citations.\nReturn cited text, references, BibTeX, and a sentence-level citation-support audit table.\nDo not add citations to claims that cannot be supported.", 0.9, 2.85, 11.55, 1.25, size=13)
    add_card(slide, "核心规则", "auto_cite 可以降低假引用风险，但不能替代人工复核。每条引用要检查 DOI 和句子支持关系。", 0.9, 4.65, 11.55, 1.05, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("BibTeX 复核表")
    rows = [
        ["核验项", "问题"],
        ["title_match", "题名是否与 DOI 页面一致"],
        ["authors_match", "作者和顺序是否一致"],
        ["year_match", "在线优先和正式出版年份是否区分"],
        ["venue_match", "期刊/会议是否真实"],
        ["DOI_match", "DOI 是否属于同一篇论文"],
        ["required_fields_missing", "journal, volume, number, pages, doi 是否缺失"],
        ["support_strength", "引用是否真正支撑对应句子"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 4.8, col_widths=[3.0, 8.85], font_size=12)
    add_text(slide, "通过元数据核验，不等于通过句子支持核验。", 0.9, 6.45, 11.2, 0.35, size=16, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("句子级引用支持核验")
    rows = [
        ["支持强度", "含义", "动作"],
        ["strong", "论文直接支持该句具体主张", "保留"],
        ["partial", "支持相近主张，但不完全对应", "修改句子或替换引用"],
        ["weak", "只相关，不足以支撑", "替换"],
        ["reject", "不支持或错引", "删除"],
    ]
    add_table(slide, rows, 0.85, 1.55, 11.6, 3.15, col_widths=[2.0, 6.2, 3.4], font_size=12)
    add_bullets(slide, [
        "方法论文不能随便支持理论机制",
        "综述论文不能替代关键一手证据",
        "预印本不能写成已发表论文",
        "引用支持的是“句子”，不是段落气氛",
    ], 1.0, 5.15, 11.1, 1.35, size=17)

    slide = new_slide("目标期刊 PDF 自动下载")
    add_flow(slide, ["目标期刊", "检索", "核验 DOI", "开放 PDF", "下载本地", "日志"], 0.8, 1.55, 11.8, box_h=0.7, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11)
    add_bullets(slide, [
        "优先：arXiv / bioRxiv / medRxiv",
        "然后：Semantic Scholar open-access PDF",
        "然后：publisher open-access",
        "最后：DOI access when institutional access permits",
        "不可用：记录 failure，不绕过付费墙",
    ], 0.95, 2.85, 5.45, 2.7, size=17)
    add_card(slide, "本地命名", "literature/pdfs/{year}_{first_author}_{short_title}.pdf", 6.95, 2.95, 5.35, 0.95, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)
    add_card(slide, "失败日志", "记录 DOI、URL、尝试路径、失败原因和下一步动作。无法下载不是失败，没记录才是失败。", 6.95, 4.35, 5.35, 1.15, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("PDF 下载日志字段")
    rows = [
        ["字段", "说明"],
        ["paper_id / title", "论文编号和标题"],
        ["authors / year / venue", "基础元数据"],
        ["DOI", "核验标识"],
        ["access_route", "arXiv / Semantic Scholar OA / DOI / publisher OA / unavailable"],
        ["pdf_status", "downloaded / failed / no_access / unreadable"],
        ["local_path", "本地保存路径"],
        ["extraction_quality", "good / partial / poor"],
        ["manual_check", "需人工核验项"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.2, col_widths=[3.15, 8.7], font_size=11.5)

    slide = new_slide("每日文献日报")
    add_bullets(slide, [
        "追踪关键词、目标期刊、目标学者和研究团队",
        "标注今日新增论文、PDF 状态、BibTeX 状态",
        "选出最值得读的 3-5 篇",
        "列出人工核验项和明日追踪任务",
        "同步到飞书表格后用于导师组或课题组协作",
    ], 0.9, 1.55, 5.8, 3.2, size=18)
    rows = [
        ["日报模块", "内容"],
        ["新增论文", "title, authors, venue, DOI, relevance"],
        ["重点阅读", "why_read, method, possible_use"],
        ["学者动态", "AMiner evidence, coauthor network"],
        ["PDF/BibTeX", "download, bibtex, verification status"],
        ["明日任务", "read / cite / download / check author"],
    ]
    add_table(slide, rows, 7.0, 1.55, 5.3, 3.65, col_widths=[1.8, 3.5], font_size=11)

    slide = new_slide("飞书表格联动")
    add_flow(slide, ["本地日报", "字段清洗", "lark-sheets", "飞书表格", "筛选协作"], 1.0, 1.55, 11.2, box_h=0.7, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_code(slide, "lark-cli sheets +create --title \"AI Agent 科研每日文献日报\" --headers \"report_date,topic,query,source,title,authors,year,venue,DOI,url,relevance_score,pdf_status,bibtex_status,verification_status,next_action\"\n\nlark-cli sheets +append --spreadsheet-token \"<spreadsheet_token>\" --sheet-id \"<sheet_id>\" --values \"<json rows>\"", 0.85, 2.65, 11.8, 1.65, size=11)
    add_card(slide, "安全", "飞书表格不写 API Key，不公开机构权限 PDF，不把未核验状态改成已确认。", 1.0, 4.85, 11.15, 0.95, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=14)

    slide = new_slide("飞书文献表字段设计")
    rows = [
        ["字段", "用途"],
        ["report_date / topic / query", "日报日期、主题、检索式"],
        ["source / title / authors / year / venue", "来源和基础信息"],
        ["DOI / url / abstract", "核验和初筛"],
        ["relevance_score / method_tag / theory_tag", "筛选和分类"],
        ["pdf_status / pdf_local_path", "PDF 获取状态"],
        ["bibtex_status / verification_status", "引用和元数据状态"],
        ["manual_check_notes / next_action", "人工核验和下一步"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 4.95, col_widths=[4.0, 7.85], font_size=12)

    slide = new_slide("AMiner 学者画像")
    add_table(slide, [
        ["输出模块", "内容"],
        ["Query target", "姓名、机构、研究领域"],
        ["AMiner evidence", "AMiner 返回证据"],
        ["Cross-check evidence", "AI4Scholar / OpenAlex / 官网 / ORCID"],
        ["Likely identity", "是否为目标学者"],
        ["Representative papers", "代表作"],
        ["Patents or outputs", "专利或应用成果"],
        ["Coauthor network", "合作者和机构网络"],
        ["Unverified items", "待核验项"],
    ], 0.75, 1.45, 6.0, 5.15, col_widths=[2.4, 3.6], font_size=11)
    add_bullets(slide, [
        "中文同名非常常见",
        "不能把第一条结果直接当目标学者",
        "必须用机构、领域、代表作和合作者消歧",
        "不要用 h-index 或引用量单独判断质量",
    ], 7.15, 1.85, 4.95, 2.25, size=18)
    add_card(slide, "课堂用途", "专家发现、审稿人分析、团队画像、研究趋势判断。", 7.15, 4.75, 4.95, 0.9, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=14)

    slide = new_slide("AI4Scholar sci_draw 画图")
    add_bullets(slide, [
        "适合：机制图、实验流程图、BCI pipeline、EEG/ERP 流程、图形摘要",
        "不适合：编造数据结果、脑区激活、统计显著性",
        "真实数据图先用 Python/R 画，再用 sci_draw 做机制图或风格润色",
        "每次生成都要保留提示词、版本和审查记录",
    ], 0.9, 1.55, 5.75, 3.1, size=17)
    add_code(slide, "Use AI4Scholar sci_draw to generate a publication-style scientific figure.\nFigure type: mechanism diagram / workflow / BCI pipeline / EEG experiment procedure\nResearch topic:\nRequired elements:\nForbidden elements:\nLanguage:\nStyle:\nOutput requirement: PNG / SVG / editable vector preferred", 7.0, 1.55, 5.4, 3.25, size=10.8)
    add_card(slide, "审查重点", "科学准确、标签正确、逻辑清楚、无虚构数据、适合目标论文或汇报。", 1.0, 5.25, 11.15, 0.85, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("课堂实操 1：创建工作区与复制 Skills")
    add_code(slide, "New-Item -ItemType Directory -Force -Path \"D:\\AI科研训练营\\第03讲_真实文献工作区\"\ncd \"D:\\AI科研训练营\\第03讲_真实文献工作区\"\nNew-Item -ItemType Directory -Force -Path skills,literature,figures,feishu,configs,output,qc\nNew-Item -ItemType Directory -Force -Path literature\\pdfs,literature\\bib,literature\\daily_reports,literature\\search_logs\n\ncd \"D:\\desk\\AI agent科研资料\\本地Skills功能分类库\"\ncopy_skill_to_workspace.ps1 -SkillName ai4scholar-research,ai4scholar-paper-search,ai4scholar-paper-detail-batch,ai4scholar-pdf-fulltext-reading,ai4scholar-auto-citation-bibtex,aminer-mcp-research,citation-management,pdf,markitdown,xlsx -Workspace \"D:\\AI科研训练营\\第03讲_真实文献工作区\"", 0.75, 1.45, 12.1, 4.15, size=9.5)
    add_text(slide, "复制命令可按本地脚本实际参数调整，原则是按项目复制，不污染全局根目录。", 0.9, 6.15, 11.3, 0.35, size=15, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("课堂实操 2：搜索与候选表")
    add_code(slide, "请用 AI4Scholar / OpenAlex / Semantic Scholar 等真实来源检索文献。\n主题：[主题]\n时间范围：[年份]\n目标期刊：[期刊列表]\n输出 literature/candidate_papers.xlsx。\n每篇至少包含 title, authors, year, venue, DOI_or_ID, url, abstract, source, verification_status。\n没有 DOI 或稳定 ID 的文献必须标注“需人工核验”。", 0.85, 1.55, 11.8, 2.15, size=12.5)
    add_table(slide, [
        ["验收", "至少 20 篇候选文献；去重；有稳定 ID 或未核验标注"],
        ["风险", "关键词相似但主题不相关；预印本误作正式论文"],
        ["输出", "literature/candidate_papers.xlsx"],
    ], 0.85, 4.25, 11.8, 1.45, col_widths=[2.0, 9.8], font_size=12)

    slide = new_slide("课堂实操 3：BibTeX 与引用复核")
    add_code(slide, "请对 candidate_papers.xlsx 中筛选出的 10 篇高相关文献生成 BibTeX。\n优先使用 DOI 或稳定 ID 查询元数据。\n输出 literature/verified_references.bib 和 literature/bibtex_audit.md。\n不要凭题名编造 BibTeX。\n\n随后对 5 个句子做 sentence-level citation-support audit。", 0.85, 1.55, 11.8, 2.15, size=12.5)
    add_bullets(slide, [
        "元数据核验：title, authors, year, venue, DOI",
        "支持核验：strong / partial / weak / reject",
        "弱引用要替换或删除",
    ], 1.0, 4.25, 11.1, 1.3, size=18)

    slide = new_slide("课堂实操 4：PDF 下载与日报")
    add_code(slide, "Download only open-access PDFs or PDFs available through permitted institutional access.\nSave PDFs to literature/pdfs/ using {year}_{first_author}_{short_title}.pdf.\nCreate literature/pdf_download_log.xlsx and literature/pdf_download_failures.md.\nDo not use piracy sites or bypass paywalls.\n\nThen generate literature/daily_reports/YYYY-MM-DD.md.", 0.85, 1.55, 11.8, 2.25, size=12.5)
    add_bullets(slide, [
        "下载成功要记录 local_path",
        "下载失败要记录原因",
        "日报要写 PDF 和 BibTeX 状态",
        "未核验项不能写成已确认",
    ], 1.0, 4.4, 11.1, 1.4, size=18)

    slide = new_slide("课堂实操 5：AMiner 与 sci_draw")
    add_table(slide, [
        ["任务", "输出"],
        ["AMiner 学者画像", "literature/aminer_scholar_profiles.md"],
        ["同名消歧", "机构、领域、代表作、合作者、ORCID/主页"],
        ["sci_draw 提示词", "figures/sci_draw_prompts.md"],
        ["图形审查", "figures/figure_review_checklist.md"],
    ], 0.85, 1.55, 11.8, 2.65, col_widths=[3.0, 8.8], font_size=12)
    add_card(slide, "共同底线", "AMiner 不编造学者结果，sci_draw 不编造数据结果。", 1.0, 4.8, 11.15, 0.95, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=15)

    slide = new_slide("课堂实操 6：飞书表格同步")
    add_code(slide, "请根据 feishu_literature_sheet_schema.md 生成 feishu/daily_literature_rows.xlsx。\n如果飞书已授权，先写入 1-2 行测试。\n字段必须包含 report_date, topic, source, title, DOI, relevance_score, pdf_status, bibtex_status, verification_status, next_action。\n不要写入 API Key、Token 或未授权 PDF 链接。", 0.85, 1.55, 11.8, 2.0, size=12.5)
    add_table(slide, [
        ["同步前检查", "说明"],
        ["权限", "bot/user 身份是否能访问目标表格"],
        ["字段", "本地 Excel 和飞书字段一致"],
        ["状态", "unverified/partial/verified 不丢失"],
        ["安全", "不含密钥和敏感数据"],
    ], 0.85, 4.05, 11.8, 1.95, col_widths=[3.0, 8.8], font_size=12)

    slide = new_slide("必须人工核验的地方")
    rows = [
        ["风险", "必须核验"],
        ["AI 生成引用", "是否真实、是否支持对应句子"],
        ["BibTeX", "DOI、题名、作者、年份、期刊、卷期页码"],
        ["检索结果", "是否主题相关，是否重复，是否只是关键词相似"],
        ["AMiner 学者", "同名、机构、代表作、合作网络"],
        ["PDF 下载", "是否合法、是否目标论文、是否完整"],
        ["文献日报", "未核验论文不能写成已确认"],
        ["sci_draw", "不能出现虚构实验结果或脑区激活"],
        ["飞书表格", "不能包含密钥、隐私数据、未授权 PDF"],
    ]
    add_table(slide, rows, 0.65, 1.4, 12.05, 5.3, col_widths=[2.65, 9.4], font_size=11.2)

    slide = new_slide("本讲提交物")
    rows = [
        ["提交物", "最低要求"],
        ["search_strategy.md", "4 条英文检索式 + 2 条中文检索式"],
        ["candidate_papers.xlsx", "至少 20 篇候选文献"],
        ["verified_references.bib", "至少 10 条 BibTeX"],
        ["bibtex_audit.md", "元数据核验 + 句子级支持核验"],
        ["pdf_download_log.xlsx", "至少尝试 10 篇，记录成功/失败"],
        ["daily_reports/YYYY-MM-DD.md", "1 份每日文献日报"],
        ["aminer_scholar_profiles.md", "1 份学者画像"],
        ["sci_draw_prompts.md", "1 个完整画图提示词"],
        ["feishu schema/rows", "飞书字段设计或写入数据"],
    ]
    add_table(slide, rows, 0.65, 1.35, 12.05, 5.5, col_widths=[3.65, 8.4], font_size=10.8)

    slide = new_slide("评分标准")
    rows = [
        ["项目", "分值", "标准"],
        ["检索策略", "10", "检索式清晰，平台和筛选标准明确"],
        ["候选文献表", "15", "至少 20 篇，元数据可追踪"],
        ["BibTeX 与复核", "20", "至少 10 条 BibTeX，含句子级支持核验"],
        ["PDF 下载", "15", "有合法下载或完整失败日志"],
        ["文献日报", "15", "内容完整，可用于每日追踪"],
        ["AMiner 学者画像", "10", "有同名消歧和交叉核验"],
        ["sci_draw 与飞书", "10", "有图形提示词和飞书字段/数据"],
        ["日志与反思", "5", "过程可追踪，反思具体"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 5.05, col_widths=[2.7, 1.05, 8.3], font_size=11)

    slide = new_slide("直接不合格的情况")
    add_bullets(slide, [
        "编造文献、DOI、BibTeX 或学者画像",
        "使用盗版网站或绕过付费墙下载 PDF",
        "在作业、飞书或日志中暴露 API Key / Token",
        "把 AMiner 同名学者第一条结果直接当目标学者",
        "没有人工核验清单",
        "只有聊天截图，没有本地输出文件",
    ], 1.0, 1.7, 11.1, 3.0, size=21, color=COLORS["red"])
    add_card(slide, "科研引用的底线", "引用不是装饰；每一条引用都要能经得起审稿人、导师和数据库复核。", 1.0, 5.35, 11.15, 0.95, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("本讲一句话总结")
    add_text(slide, "真实文献工作流的目标，\n不是让 AI 写出“像论文”的引用，\n而是让每条文献都有来源、标识、PDF 状态、BibTeX 和核验记录。", 1.15, 1.9, 11.0, 2.0, size=27, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
    add_flow(slide, ["真实来源", "稳定 ID", "合法 PDF", "可核验 BibTeX", "团队日报"], 1.35, 4.75, 10.55, box_h=0.75, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=13)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
