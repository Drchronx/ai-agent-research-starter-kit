from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第04讲_课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(247, 248, 245),
    "ink": RGBColor(31, 41, 51),
    "muted": RGBColor(91, 103, 112),
    "teal": RGBColor(15, 118, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(180, 83, 9),
    "red": RGBColor(185, 28, 28),
    "line": RGBColor(210, 215, 220),
    "white": RGBColor(255, 255, 255),
    "pale_teal": RGBColor(219, 245, 241),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(254, 243, 199),
    "pale_red": RGBColor(254, 226, 226),
    "soft": RGBColor(236, 239, 241),
    "dark": RGBColor(33, 43, 54),
}

FONT = "Microsoft YaHei"


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color=COLORS["line"], width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def clear_line(shape):
    shape.line.fill.background()


def add_text(slide, text, x, y, w, h, size=20, color=COLORS["ink"], bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=COLORS["ink"], gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def add_code(slide, text, x, y, w, h, size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, COLORS["dark"])
    clear_line(shape)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(235, 239, 245)
    return shape


def add_header(slide, title, section="第 04 讲", num=None):
    set_fill(slide.background, COLORS["bg"])
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34)
    )
    set_fill(bar, COLORS["teal"])
    clear_line(bar)
    add_text(slide, section, 0.45, 0.05, 1.6, 0.22, size=8, color=COLORS["white"], bold=True)
    if num is not None:
        add_text(slide, f"{num:02d}", 12.05, 0.04, 0.7, 0.25, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.62, 11.8, 0.55, size=25, color=COLORS["ink"], bold=True)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.21), Inches(1.1), Inches(0.06)
    )
    set_fill(line, COLORS["amber"])
    clear_line(line)


def add_card(slide, title, body, x, y, w, h, fill=COLORS["white"], accent=COLORS["teal"],
             title_size=15, body_size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, COLORS["line"], 0.8)
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    set_fill(stripe, accent)
    clear_line(stripe)
    add_text(slide, title, x + 0.18, y + 0.15, w - 0.28, 0.35, size=title_size, color=accent, bold=True)
    add_text(slide, body, x + 0.18, y + 0.62, w - 0.32, h - 0.72, size=body_size, color=COLORS["ink"])
    return shape


def add_table(slide, rows, x, y, w, h, col_widths=None, font_size=10, header_fill=COLORS["teal"]):
    nrows, ncols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size if r else font_size + 0.5)
                p.font.bold = bool(r == 0)
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(242, 245, 247))
    return table_shape


def add_flow(slide, labels, x, y, w, box_h=0.58, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11):
    gap = 0.18
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(box_h)
        )
        set_fill(shape, color)
        set_line(shape, accent, 1)
        add_text(slide, label, bx + 0.08, y + 0.14, box_w - 0.16, box_h - 0.2,
                 size=font_size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            x1 = Inches(bx - gap + 0.02)
            x2 = Inches(bx - 0.02)
            ymid = Inches(y + box_h / 2)
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, ymid, x2, ymid)
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.2)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_fill(slide.background, COLORS["bg"])
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.0), Inches(7.5))
    set_fill(left, COLORS["teal"])
    clear_line(left)
    add_text(slide, "Zotero / Obsidian\n长期知识库", 0.45, 1.05, 2.35, 1.7, size=25, color=COLORS["white"], bold=True)
    add_text(slide, "第 04 讲", 0.6, 5.95, 1.4, 0.35, size=14, color=COLORS["white"], bold=True)
    add_text(slide, "从文献数据库到可复用科研知识资产", 3.45, 1.22, 9.0, 0.7, size=28, color=COLORS["ink"], bold=True)
    add_text(slide, "Zotero · Better BibTeX · Obsidian Vault · Paper Cards · Theory-Variable-Method Matrix · Knowledge Graph", 3.5, 2.1, 8.7, 0.45, size=14, color=COLORS["muted"])
    add_flow(slide, ["收文献", "补元数据", "制卡", "建矩阵", "画图谱", "日维护"], 3.55, 3.25, 8.8, box_h=0.66, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_card(slide, "本讲交付", "Zotero 10 篇文献、3 张 Obsidian 论文卡、理论-变量-方法矩阵、研究问题知识图谱、可复用启发库", 3.55, 4.65, 8.75, 1.35, fill=COLORS["white"], accent=COLORS["amber"], title_size=14, body_size=13)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new_slide(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_header(slide, title, num=len(prs.slides))
        return slide

    slide = new_slide("这节课解决什么问题")
    add_bullets(slide, [
        "文献越来越多，但不知道哪些读过、哪些可引用",
        "PDF、Zotero 条目、BibTeX 和笔记没有对应",
        "Obsidian 只是摘抄，没有沉淀理论、变量、方法",
        "写综述时找不到某个理论、量表、方法来自哪篇论文",
        "每次写论文都重新找文献，不能复用阅读成果",
        "研究问题没有图谱，读文献无法沉淀成选题",
    ], 0.9, 1.55, 6.0, 3.4, size=18)
    add_card(slide, "本讲核心", "Zotero 管文献事实，Obsidian 管研究理解，矩阵管可比较结构，知识图谱管问题关系。", 7.25, 1.85, 4.95, 1.65, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)
    add_card(slide, "课堂目标", "学生照着 PPT 完成一个能长期维护的 Zotero / Obsidian 工作区。", 7.25, 4.1, 4.95, 1.2, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("本讲固定模板")
    rows = [
        ["模块", "第四讲怎么落实"],
        ["科研问题和需求", "防止文献、PDF、笔记、引用再次变乱"],
        ["需要哪些 Skills", "zotero sync, obsidian card, matrix, graph, insight bank"],
        ["输入材料", "Zotero、Obsidian、10 篇文献、3 篇核心论文"],
        ["Agent 怎么执行", "审计 Zotero、生成卡片、矩阵、图谱、启发库"],
        ["输出文件", "zotero exports, paper cards, matrix, graph, audit"],
        ["人工核验", "元数据、PDF、理论机制、变量测量、图谱边"],
        ["课后作业", "10 篇文献、3 张卡、1 矩阵、1 图谱、1 启发库"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.9, col_widths=[2.65, 9.2], font_size=12)
    add_text(slide, "这讲不是教软件按钮，而是教一套长期可维护的科研知识系统。", 0.9, 6.55, 11.2, 0.35, size=15, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("Zotero 与 Obsidian 的分工")
    rows = [
        ["工具", "负责什么", "不负责什么"],
        ["Zotero", "文献事实：题名、作者、年份、期刊、DOI、PDF、BibTeX", "不负责写你的理论理解"],
        ["Obsidian", "研究理解：论文卡片、理论、变量、方法、项目、双链", "不替代文献数据库"],
        ["矩阵", "把多篇论文变成可比较结构", "不是复制摘要"],
        ["知识图谱", "显示问题、理论、变量、方法和证据关系", "不是装饰性关系图"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 3.9, col_widths=[2.0, 5.65, 4.4], font_size=12)
    add_card(slide, "原则", "Zotero 中保留事实，Obsidian 中保留解释；解释必须能回到事实。", 1.0, 5.85, 11.15, 0.75, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)

    slide = new_slide("本讲需要哪些 Skills")
    rows = [
        ["Skill", "用途"],
        ["zotero-library-sync", "读取/审计 Zotero 条目、导出元数据、PDF 附件状态"],
        ["pyzotero", "通过 Zotero Web API 读取、搜索、导出、管理文献"],
        ["obsidian-paper-card", "生成结构化 Obsidian 论文卡片"],
        ["theory-variable-method-matrix", "生成理论-变量-方法矩阵"],
        ["research-question-knowledge-graph", "生成研究问题知识图谱和 Mermaid 图"],
        ["paper-reusable-insight-bank", "沉淀可复用理论、变量、方法、写作启发"],
        ["citation-management / pdf / xlsx", "引用核验、PDF 读取、矩阵表格输出"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 5.15, col_widths=[3.4, 8.65], font_size=11.5)

    slide = new_slide("操作总路线")
    add_flow(slide, ["安装软件", "建工作区", "建 Zotero 分类", "导入文献", "导出审计", "建 Vault", "生成卡片", "矩阵图谱"], 0.55, 1.55, 12.25, box_h=0.7, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=10)
    add_table(slide, [
        ["阶段", "学生要做什么"],
        ["安装", "Zotero、Connector、Obsidian、Better BibTeX"],
        ["文献库", "10 篇文献、5 个 PDF、Collection、Tags"],
        ["导出", "references.bib、zotero_items.csv、metadata_audit.xlsx"],
        ["笔记库", "Vault 文件夹、模板、3 张论文卡片"],
        ["结构化", "理论-变量-方法矩阵、研究问题图谱"],
        ["维护", "每日 20-40 分钟更新流程"],
    ], 0.8, 2.75, 11.85, 3.5, col_widths=[2.2, 9.65], font_size=12)

    slide = new_slide("Step 1：安装 Zotero")
    add_code(slide, "官方入口：\nhttps://www.zotero.org/download/", 0.85, 1.55, 5.5, 0.85, size=15)
    add_bullets(slide, [
        "下载 Zotero 桌面端",
        "安装 Zotero Connector 浏览器插件",
        "创建或登录 Zotero 账号",
        "打开同步设置",
        "用一个论文网页测试 Connector 是否能保存条目",
    ], 0.95, 2.85, 5.6, 2.4, size=18)
    add_card(slide, "检查点", "Zotero 中能出现新条目；Connector 保存时能自动识别题名、作者、DOI 或网页信息。", 7.0, 1.75, 5.3, 1.35, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)
    add_card(slide, "常见问题", "只保存了网页快照，没有保存论文元数据。解决：换 DOI 页面、出版社页面或用 DOI 魔法棒添加。", 7.0, 3.75, 5.3, 1.35, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("Step 2：安装 Obsidian")
    add_code(slide, "官方入口：\nhttps://obsidian.md/download", 0.85, 1.55, 5.5, 0.85, size=15)
    add_bullets(slide, [
        "下载安装 Obsidian",
        "先不要急着写笔记",
        "理解 Vault：它就是一个本地文件夹",
        "理解 note：它就是 Markdown 文件",
        "本课程后面会用已有文件夹打开 Vault",
    ], 0.95, 2.85, 5.6, 2.4, size=18)
    add_card(slide, "检查点", "你能说清楚：Obsidian 不是云数据库，而是本地 Markdown 知识库。", 7.0, 1.85, 5.3, 1.2, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=14)
    add_card(slide, "课堂原则", "不要让 Obsidian 替代 Zotero。文献事实仍以 Zotero 为准。", 7.0, 3.65, 5.3, 1.1, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("Step 3：安装推荐插件")
    add_table(slide, [
        ["插件", "装在哪里", "用途"],
        ["Better BibTeX", "Zotero", "稳定 citekey、BibTeX 自动导出"],
        ["Templates", "Obsidian", "插入论文卡片模板"],
        ["Dataview", "Obsidian", "后续按字段检索笔记"],
        ["Zotero Integration", "Obsidian", "从 Zotero 条目生成 note"],
    ], 0.75, 1.55, 11.85, 3.1, col_widths=[3.0, 2.0, 6.85], font_size=12)
    add_code(slide, "建议 citekey：\nauth.lower + year + shorttitle(3,3)\n\n示例：\nsmith2024aitrust", 0.95, 5.0, 5.6, 1.25, size=13)
    add_card(slide, "课堂提醒", "插件界面可能随版本变化。掌握逻辑比背按钮更重要：稳定 key、模板制卡、字段检索。", 7.0, 5.05, 5.2, 1.1, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("Step 4：创建本地工作区")
    add_code(slide, "New-Item -ItemType Directory -Force -Path \"D:\\AI科研训练营\\第04讲_知识库工作区\"\ncd \"D:\\AI科研训练营\\第04讲_知识库工作区\"\n\nNew-Item -ItemType Directory -Force -Path configs,zotero_exports,obsidian_vault,matrices,graphs,reusable_insights,qc,output,skills\nNew-Item -ItemType Directory -Force -Path obsidian_vault\\00_Inbox,obsidian_vault\\01_Papers,obsidian_vault\\02_Theories,obsidian_vault\\03_Variables,obsidian_vault\\04_Methods,obsidian_vault\\05_Projects,obsidian_vault\\06_Matrices,obsidian_vault\\07_Graphs,obsidian_vault\\08_Templates,obsidian_vault\\99_Archive\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md", 0.65, 1.45, 12.2, 4.1, size=9.5)
    add_card(slide, "检查", "打开文件夹后应看到 configs、zotero_exports、obsidian_vault、matrices、graphs、reusable_insights、qc。", 0.9, 6.0, 11.4, 0.75, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)

    slide = new_slide("Step 5：复制本讲 Skills")
    add_code(slide, "cd \"D:\\desk\\AI agent科研资料\\本地Skills功能分类库\"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName zotero-library-sync,obsidian-paper-card,theory-variable-method-matrix,research-question-knowledge-graph,paper-reusable-insight-bank,citation-management,pdf,markitdown,xlsx `\n  -Workspace \"D:\\AI科研训练营\\第04讲_知识库工作区\"", 0.85, 1.55, 11.8, 2.25, size=11.5)
    add_table(slide, [
        ["检查命令", "Get-ChildItem \"D:\\AI科研训练营\\第04讲_知识库工作区\\skills\" -Directory"],
        ["合格结果", "能看到 zotero-library-sync、obsidian-paper-card、matrix、graph、insight-bank"],
    ], 0.85, 4.25, 11.8, 1.45, col_widths=[2.3, 9.5], font_size=12)

    slide = new_slide("Step 6：Zotero Collection 结构")
    rows = [
        ["Collection", "用途"],
        ["00_inbox", "刚收进来的文献，未清理"],
        ["01_to_read", "未来要读"],
        ["02_reading", "正在读"],
        ["03_core_theory", "理论源文献"],
        ["04_methods", "方法源文献"],
        ["05_measures", "量表、测量和变量操作化"],
        ["06_datasets", "数据集、语料、平台"],
        ["07_reviews_meta", "综述和 meta-analysis"],
        ["08_manuscript_citations", "当前论文要引用的文献"],
        ["09_daily_literature", "每日文献日报导入内容"],
        ["99_archive", "已归档"],
    ]
    add_table(slide, rows, 0.65, 1.25, 12.05, 5.8, col_widths=[3.0, 9.05], font_size=10.8)

    slide = new_slide("Step 7：Zotero 标签体系")
    add_card(slide, "字段", "field:management\nfield:psychology\nfield:neuroscience\nfield:bci\nfield:hci", 0.75, 1.45, 3.0, 2.2, fill=COLORS["white"], accent=COLORS["teal"], body_size=12)
    add_card(slide, "方法", "method:eeg\nmethod:erp\nmethod:scenario_experiment\nmethod:text_mining\nmethod:machine_learning\nmethod:causal_inference", 3.95, 1.45, 3.0, 2.2, fill=COLORS["white"], accent=COLORS["blue"], body_size=12)
    add_card(slide, "角色", "role:classic\nrole:must_read\nrole:method_source\nrole:measure_source\nrole:theory_source", 7.15, 1.45, 2.8, 2.2, fill=COLORS["white"], accent=COLORS["amber"], body_size=12)
    add_card(slide, "状态", "status:to_read\nstatus:read\nstatus:verified\nstatus:needs_check", 10.15, 1.45, 2.45, 2.2, fill=COLORS["white"], accent=COLORS["red"], body_size=12)
    add_text(slide, "每篇文献至少 3 个标签：field + method + status。重要文献再加 role。", 0.9, 4.55, 11.4, 0.45, size=18, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, "不要做", "不要一会儿写 method:eeg，一会儿写 EEG，一会儿写 脑电。标签必须稳定。", 1.0, 5.55, 11.15, 0.85, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("Step 8：导入 10 篇文献")
    rows = [
        ["方式", "操作", "适合场景"],
        ["浏览器 Connector", "打开论文网页，点击 Zotero 保存", "出版社页面、数据库页面"],
        ["拖入 PDF", "PDF 拖进 Zotero，右键 Retrieve Metadata", "已有 PDF"],
        ["DOI 魔法棒", "点击魔法棒，输入 DOI", "知道 DOI"],
        ["BibTeX 导入", "File > Import", "第三讲生成 BibTeX"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 3.35, col_widths=[2.4, 5.8, 3.65], font_size=12)
    add_bullets(slide, [
        "至少 10 篇文献",
        "至少 5 篇有 PDF 附件",
        "每篇检查题名、作者、年份、期刊、DOI",
        "缺 DOI 的标 status:needs_check",
    ], 1.0, 5.35, 11.1, 1.1, size=17)

    slide = new_slide("Step 9：导出 Zotero")
    add_bullets(slide, [
        "在 Zotero 中选中 Collection",
        "右键 Export Collection",
        "格式选择 BibTeX",
        "保存到 zotero_exports/references.bib",
        "如果有 CSV/JSON，也保存到 zotero_exports/",
    ], 0.95, 1.65, 5.65, 2.4, size=18)
    add_code(slide, "目标路径：\nD:\\AI科研训练营\\第04讲_知识库工作区\\zotero_exports\\references.bib", 7.05, 1.75, 5.25, 0.9, size=13)
    add_card(slide, "检查点", "references.bib 存在；citekey 不重复；重要文献 DOI 没丢；中文题名没有乱码。", 7.05, 3.25, 5.25, 1.1, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)
    add_card(slide, "下一步", "用 zotero-library-sync 审计缺 DOI、重复、缺 PDF、未打标签。", 7.05, 4.85, 5.25, 0.95, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=13)

    slide = new_slide("Step 10：Zotero API 与安全")
    add_code(slide, "API Key 页面：\nhttps://www.zotero.org/settings/keys\n\n环境变量：\nZOTERO_LIBRARY_ID=your_user_or_group_id\nZOTERO_LIBRARY_TYPE=user\nZOTERO_API_KEY=your_private_key", 0.85, 1.55, 5.8, 2.3, size=13)
    add_bullets(slide, [
        "没有 API Key 也可以手动导出 BibTeX/CSV",
        "API Key 只放环境变量或私有 .env",
        "不要写进 Obsidian vault",
        "不要提交到作业、飞书或 Git",
        "课堂默认先读，不写回 Zotero",
    ], 7.15, 1.7, 4.9, 2.5, size=17)
    add_card(slide, "安全底线", "任何 API Key 泄露，本讲作业直接不合格。", 1.0, 5.25, 11.15, 0.85, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=15)

    slide = new_slide("Step 11：审计 Zotero 文献库")
    add_code(slide, "Use zotero-library-sync to audit my Zotero library export.\nDo not display API keys.\nUse files in zotero_exports/.\nExport zotero_items.csv, references.bib, metadata_audit.xlsx, and obsidian_paper_index.md.\nCheck missing DOI, duplicate titles, missing PDF attachments, untagged items, and inconsistent venue names.\nDo not write changes back to Zotero.", 0.85, 1.45, 11.8, 2.35, size=12)
    add_table(slide, [
        ["输出", "用途"],
        ["zotero_items.csv", "全库文献元数据"],
        ["metadata_audit.xlsx", "缺 DOI、缺 PDF、重复、未打标签"],
        ["obsidian_paper_index.md", "给 Obsidian 的文献索引"],
    ], 0.85, 4.25, 11.8, 1.6, col_widths=[3.2, 8.6], font_size=12)

    slide = new_slide("Step 12：打开 Obsidian Vault")
    add_bullets(slide, [
        "打开 Obsidian",
        "选择 Open folder as vault",
        "选择第04讲工作区下的 obsidian_vault",
        "确认左侧出现 00_Inbox 到 99_Archive",
        "把模板放入 08_Templates",
    ], 0.95, 1.65, 5.7, 2.6, size=18)
    add_code(slide, "Vault 路径：\nD:\\AI科研训练营\\第04讲_知识库工作区\\obsidian_vault\n\n模板路径：\nobsidian_vault\\08_Templates\\obsidian_paper_note_template.md", 7.05, 1.75, 5.3, 1.65, size=12.5)
    add_card(slide, "检查点", "在 Obsidian 中能看到 01_Papers、02_Theories、03_Variables、04_Methods、08_Templates。", 7.05, 4.05, 5.3, 1.0, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)

    slide = new_slide("Step 13：论文卡片命名与 YAML")
    add_code(slide, "@FirstAuthorYear_ShortTitle.md\n\n示例：\n@Smith2024_AITrust.md", 0.85, 1.55, 4.8, 1.25, size=15)
    add_code(slide, "---\nzotero_key:\ncitation_key:\ntitle:\nauthors:\nyear:\nvenue:\ndoi:\nfields: []\nmethods: []\ntheories: []\nvariables: []\nstatus: to_read\nverification_status: unverified\n---", 6.15, 1.25, 5.95, 4.6, size=11)
    add_bullets(slide, [
        "一篇论文一张 canonical note",
        "YAML 放可检索字段",
        "卡片主体放理解和启发",
        "未核验内容不要写成 verified",
    ], 0.95, 3.35, 4.65, 2.05, size=17)

    slide = new_slide("Step 14：生成 3 张论文卡片")
    add_code(slide, "Use obsidian-paper-card to create Obsidian notes for three core papers.\nUse obsidian_paper_note_template.md.\nEach note must include YAML frontmatter, citation, one-sentence takeaway, research question, theory mechanism, variables, method, data/sample, key results, limitations, reusable ideas, and links.\nSeparate paper evidence from my interpretation.\nOutput notes to obsidian_vault/01_Papers/.", 0.85, 1.45, 11.8, 2.25, size=12)
    add_table(slide, [
        ["合格卡片", "不合格卡片"],
        ["有理论机制、变量、方法、结果、启发", "只有摘要或翻译"],
        ["区分论文原意和个人解释", "把自己的猜测写成作者结论"],
        ["有 Zotero / PDF / 相关 note 链接", "没有来源链接"],
    ], 0.85, 4.15, 11.8, 1.85, col_widths=[5.9, 5.9], font_size=12)

    slide = new_slide("Step 15：建立理论、变量、方法笔记")
    add_code(slide, "02_Theories/Technology_Acceptance.md\n03_Variables/Trust_in_AI.md\n04_Methods/Scenario_Experiment.md\n\n论文卡片中链接：\n[[Technology_Acceptance]]\n[[Trust_in_AI]]\n[[Scenario_Experiment]]", 0.85, 1.55, 5.85, 2.5, size=13)
    add_bullets(slide, [
        "理论笔记：理论解释什么机制",
        "变量笔记：定义、测量、量表、操作化",
        "方法笔记：设计、样本、分析模型、风险",
        "项目笔记：自己的论文如何复用这些材料",
    ], 7.15, 1.75, 4.9, 2.3, size=18)
    add_card(slide, "检查点", "点击 [[...]] 链接可以跳转；核心概念不再散落在多篇论文卡片里。", 1.0, 5.1, 11.15, 0.9, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=14)

    slide = new_slide("Step 16：理论-变量-方法矩阵")
    rows = [
        ["字段", "说明"],
        ["paper_id / zotero_key / citation_key", "追溯来源"],
        ["research_question", "论文解决的问题"],
        ["theory / mechanism", "理论和解释机制分开"],
        ["construct / variable_role / variable_name", "构念、角色、变量名"],
        ["measurement", "量表、题项、特征、操作化"],
        ["method_design / sample_data", "研究设计和数据"],
        ["analysis_model / key_result", "模型和结果"],
        ["reuse_for_my_project", "如何复用"],
        ["verification_status / note_link", "核验状态和 Obsidian 链接"],
    ]
    add_table(slide, rows, 0.65, 1.25, 12.05, 5.8, col_widths=[4.0, 8.05], font_size=10.8)

    slide = new_slide("Step 17：生成矩阵的 Agent 提示词")
    add_code(slide, "Use theory-variable-method-matrix to build a matrix from my Obsidian paper cards.\nExtract theory, mechanism, constructs, variable roles, measurement, method design, sample, analysis model, key result, limitations, and reuse_for_my_project.\nExport matrices/theory_variable_method_matrix.xlsx and matrices/theory_variable_method_matrix.md.\nDo not treat theories as decorative labels.", 0.85, 1.45, 11.8, 2.0, size=12.5)
    add_bullets(slide, [
        "至少 10 行",
        "一篇论文可以多行",
        "理论必须写机制",
        "变量必须写测量",
        "每行有 note_link",
    ], 1.0, 4.1, 11.1, 1.5, size=18)

    slide = new_slide("Step 18：研究问题知识图谱")
    add_table(slide, [
        ["节点类型", "paper, author, venue, theory, mechanism, construct, variable, method, dataset, measure, research_question, hypothesis, finding, limitation, project, idea"],
        ["边类型", "cites, supports, contradicts, uses_theory, tests_mechanism, measures, uses_method, has_gap, informs_hypothesis"],
        ["输出", "graphs/research_question_graph.md 和 graphs/research_question_edges.xlsx"],
    ], 0.75, 1.55, 11.85, 2.25, col_widths=[2.0, 9.85], font_size=11.5)
    add_code(slide, "graph TD\n  RQ[Research Question]\n  T[Theory]\n  M[Mechanism]\n  V[Variable]\n  P[Paper]\n  Method[Method]\n  RQ -->|uses_theory| T\n  T -->|explains| M\n  M -->|predicts| V\n  P -->|supports| M\n  P -->|uses_method| Method", 0.95, 4.15, 11.4, 2.1, size=12)

    slide = new_slide("Step 19：生成知识图谱的提示词")
    add_code(slide, "Use research-question-knowledge-graph to build a knowledge graph for my research question.\nUse Zotero items and Obsidian paper cards as sources.\nReturn node tables, typed edges, a Mermaid graph, and gap candidates requiring verification.\nOutput graphs/research_question_graph.md and graphs/research_question_edges.xlsx.", 0.85, 1.55, 11.8, 1.75, size=12.5)
    add_card(slide, "人工核验", "如果一条边不能被论文或笔记支持，标为 idea edge；不能直接当作事实。", 1.0, 3.9, 11.15, 0.9, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)
    add_bullets(slide, [
        "gap 不是“我没读到”，而是检索后仍缺少证据",
        "图谱太密时拆成理论层、方法层、证据层",
        "每个关键节点都应能回到 Zotero 或 Obsidian note",
    ], 1.0, 5.25, 11.1, 1.1, size=17)

    slide = new_slide("Step 20：可复用启发库")
    add_code(slide, "Use paper-reusable-insight-bank to extract reusable insights from the three paper cards and the matrix.\nSeparate reusable theory, variables, methods, experiment designs, EEG/ERP/BCI ideas, text-mining ideas, and writing/rebuttal ideas.\nOutput reusable_insights/paper_reusable_insight_bank.md.", 0.85, 1.55, 11.8, 1.65, size=12.5)
    add_table(slide, [
        ["启发类型", "例子"],
        ["理论", "某理论适合解释什么机制"],
        ["变量", "某变量如何测量或操纵"],
        ["方法", "某设计如何复用到自己项目"],
        ["EEG/BCI", "某特征、范式或分类评估方式"],
        ["写作", "某篇论文如何处理局限或贡献"],
    ], 0.85, 3.75, 11.8, 2.15, col_widths=[2.4, 9.4], font_size=12)

    slide = new_slide("Step 21：每日维护流程")
    add_flow(slide, ["收文献", "补元数据", "打标签", "制卡", "更新矩阵", "更新图谱", "写日志"], 0.7, 1.55, 12.0, box_h=0.7, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=11)
    add_bullets(slide, [
        "新文献先进 Zotero 00_inbox",
        "补 DOI、作者、年份、期刊、PDF",
        "核心文献生成 Obsidian 卡片",
        "把理论、变量、方法写入矩阵",
        "新增关系写入知识图谱",
        "真正可复用的内容进入 insight bank",
        "最后更新 research_log.md",
    ], 1.0, 2.8, 11.2, 2.7, size=18)
    add_card(slide, "时间", "每天 20-40 分钟维护一次，比月底整理 300 篇文献更可靠。", 1.0, 6.0, 11.15, 0.7, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("知识库审计清单")
    rows = [
        ["审计项", "检查内容"],
        ["Zotero 元数据", "题名、作者、年份、期刊、DOI"],
        ["PDF 附件", "是否与条目匹配"],
        ["citekey", "唯一、稳定、可读"],
        ["Obsidian 卡片", "是否混淆论文原意和个人推断"],
        ["理论机制", "是否只是贴理论标签"],
        ["变量测量", "构念和测量是否分开"],
        ["知识图谱边", "每条边是否有证据来源"],
        ["API Key", "是否泄露到 vault、作业、共享文件"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.2, col_widths=[3.0, 8.85], font_size=12)

    slide = new_slide("课堂验收：学生必须完成")
    rows = [
        ["提交物", "最低要求"],
        ["Zotero 文献库", "10 篇文献，5 篇有 PDF，每篇 3 个标签"],
        ["references.bib", "BibTeX 导出，citekey 不重复"],
        ["metadata_audit.xlsx", "缺 DOI、缺 PDF、重复、未打标签"],
        ["Obsidian Vault", "标准文件夹结构"],
        ["论文卡片", "至少 3 张，含 YAML、链接和核验项"],
        ["理论-变量-方法矩阵", "至少 10 行"],
        ["研究问题知识图谱", "节点、边、Mermaid 图"],
        ["可复用启发库", "至少 5 条具体启发"],
    ]
    add_table(slide, rows, 0.75, 1.25, 11.85, 5.75, col_widths=[3.1, 8.75], font_size=11)

    slide = new_slide("常见错误")
    rows = [
        ["错误", "后果", "修正"],
        ["把 PDF 只存在本地文件夹", "引用和元数据不可管理", "进入 Zotero"],
        ["Obsidian 只摘抄摘要", "无法服务综述和选题", "按理论/变量/方法制卡"],
        ["理论只写名称", "机制不成立", "写清理论如何解释变量关系"],
        ["矩阵没有 note_link", "无法追溯证据", "每行链接论文卡片"],
        ["图谱边无证据", "伪知识图谱", "标注 evidence 或 idea edge"],
        ["API Key 写进 vault", "密钥泄露", "用环境变量或私有配置"],
    ]
    add_table(slide, rows, 0.6, 1.45, 12.15, 4.85, col_widths=[3.15, 4.25, 4.75], font_size=11)

    slide = new_slide("课后作业")
    add_bullets(slide, [
        "提交：姓名_第04讲_Zotero_Obsidian长期知识库",
        "Zotero：至少 10 篇文献，5 篇 PDF，Collection 和 Tags 完整",
        "导出：zotero_items.csv、references.bib、metadata_audit.xlsx",
        "Obsidian：至少 3 张论文卡片",
        "矩阵：theory_variable_method_matrix.xlsx 至少 10 行",
        "图谱：research_question_graph.md 含 Mermaid 图",
        "启发库：至少 5 条可复用启发",
        "反思：Zotero 与 Obsidian 为什么不能混用",
    ], 0.9, 1.5, 11.4, 4.2, size=17)
    add_card(slide, "直接不合格", "泄露 API Key；只有截图没有文件；论文卡片只有摘要；矩阵无机制；图谱边无证据。", 1.0, 6.0, 11.15, 0.8, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=14)

    slide = new_slide("本讲一句话总结")
    add_text(slide, "长期知识库的价值，\n不是把文献存起来，\n而是让每篇文献都能进入理论、变量、方法、问题和项目的复用网络。", 1.15, 1.95, 11.0, 1.8, size=28, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
    add_flow(slide, ["Zotero 事实", "Obsidian 理解", "矩阵比较", "图谱关联", "启发复用"], 1.35, 4.75, 10.55, box_h=0.75, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=13)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
