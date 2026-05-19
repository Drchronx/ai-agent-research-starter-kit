from pathlib import Path
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第05讲_实操版课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(87, 99, 112),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "teal_dark": RGBColor(8, 78, 74),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "green": RGBColor(22, 120, 75),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "pale_green": RGBColor(226, 246, 236),
    "panel": RGBColor(255, 255, 255),
}

FONT = "Microsoft YaHei"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=10.5, fill_color=COLORS["dark"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.13)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 05 讲 · 实操版", 0.45, 0.055, 2.2, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12.5, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])
    return rect


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.2, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248)
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
    return shape


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=10):
    gap = 0.13
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(bx - gap + 0.01),
                Inches(y + h / 2),
                Inches(bx - 0.02),
                Inches(y + h / 2),
            )
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def section_label(slide, label, x=0.58, y=1.36, color=COLORS["teal"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.35), Inches(0.32))
    fill(rect, color)
    no_line(rect)
    textbox(slide, label, x + 0.10, y + 0.065, 1.15, 0.18, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 05 讲\n实操版", 0.52, 0.90, 2.05, 0.95, size=26, color=COLORS["white"], bold=True)
    textbox(slide, "文献综述\n理论建构\n研究问题生成", 0.50, 2.15, 2.30, 1.65, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "从材料到选题", 0.55, 6.18, 1.8, 0.30, size=12, color=COLORS["white"], bold=True)
    textbox(slide, "学生照着做的 4 小时课堂实操", 3.55, 1.10, 8.8, 0.58, size=29, color=COLORS["ink"], bold=True)
    textbox(slide, "真实文献 → 证据表 → 文献树 → 缺口审计 → 理论机制 → 假设组 → 一页方案", 3.58, 1.92, 8.8, 0.35, size=15, color=COLORS["muted"])
    flow(slide, ["输入材料", "证据表", "文献树", "缺口审计", "机制链", "假设", "方案"], 3.60, 3.05, 8.85, h=0.62, size=10.5)
    card(slide, "本讲交付", "11 个可检查文件：evidence_table、literature_tree、streams、gap_candidates、gap_audit、synthesis_outline、mechanism_chain、hypothesis_set、one_page_proposal、citation_check、final_audit。", 3.60, 4.42, 8.72, 1.25, accent=COLORS["amber"], body_size=12.5)
    card(slide, "课堂原则", "先结构化证据，再写综述；先审计缺口，再讲贡献；先解释机制，再写假设。", 3.60, 5.95, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def add_step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    section_label(slide, f"STEP {idx}")
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.15, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.15, accent=COLORS["amber"], body_size=11.5, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.70, size=9.4)
    textbox(slide, "人工核验", 4.65, 4.80, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.8)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么科研问题")
    bullets(slide, [
        "文献综述只堆材料，没有结构推进",
        "研究缺口写成“已有研究较少”，但没有证据",
        "理论只是贴标签，不能解释变量关系",
        "假设只是统计关系，不是理论预测",
        "研究问题太大，无法被数据或实验检验",
        "不知道如何从论文卡片压缩出一页研究方案",
    ], 0.85, 1.60, 5.85, 3.0, size=16)
    card(slide, "本讲不做", "不直接让 AI 写一篇看起来流畅的综述。流畅文字会掩盖文献不足、缺口不成立和理论错位。", 7.10, 1.75, 5.00, 1.25, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)
    card(slide, "本讲要做", "把文献变成问题结构：证据、流派、争议、缺口、机制、假设、设计。每一步都有文件和人工核验。", 7.10, 3.45, 5.00, 1.25, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.5)
    flow(slide, ["材料", "结构", "缺口", "理论", "假设", "方案"], 1.00, 5.70, 11.35, h=0.62, size=12)

    slide = new("本讲最终产出文件")
    rows = [
        ["阶段", "文件", "作用"],
        ["证据整理", "review/evidence_table.md", "把论文卡片转成可追溯证据"],
        ["领域结构", "review/literature_tree.md", "建立 Novelty Tree 与 Challenge-Insight Tree"],
        ["争议识别", "review/research_streams_and_controversies.md", "找流派、争议和可能缺口"],
        ["缺口判断", "review/research_gap_candidates.md / research_gap_audit.md", "生成、删减和修改研究缺口"],
        ["理论建构", "theory/theory_mechanism_chain.md", "说明变量关系背后的机制"],
        ["假设形成", "hypotheses/hypothesis_set.md", "把机制转成可检验假设"],
        ["方案与审计", "proposal/one_page_research_proposal.md / qc/gap_and_theory_audit.md", "形成导师可评估的一页方案"],
    ]
    table(slide, rows, 0.72, 1.55, 11.95, 4.75, col_widths=[1.6, 4.6, 5.7], font_size=9.5)
    card(slide, "提交底线", "任何进入正文或方案的引用，都必须标注已核验或待核验。不能用模型记忆替代文献检索。", 0.82, 6.48, 11.70, 0.58, accent=COLORS["red"], body_size=11.8, fill_color=COLORS["pale_red"])

    slide = new("需要哪些 Skills")
    rows = [
        ["功能", "Skills", "本讲用途"],
        ["文献综述", "literature-review / ssci-literature-review", "组织真实文献和综述结构"],
        ["平台检索", "ai4scholar-research / ai4scholar-citation-network", "补充真实文献与引用网络"],
        ["引用核验", "citation-management", "核验 DOI、BibTeX、作者、年份"],
        ["假设生成", "hypothesis-generation", "把机制转成可检验假设"],
        ["批判审计", "scientific-critical-thinking", "模拟导师和审稿人挑问题"],
        ["方案写作", "research-proposal", "压缩成一页研究方案"],
        ["知识库", "theory-variable-method-matrix / research-question-knowledge-graph", "使用第 04 讲矩阵和图谱"],
    ]
    table(slide, rows, 0.70, 1.48, 12.0, 4.85, col_widths=[1.55, 4.4, 6.05], font_size=9.4)
    card(slide, "放置原则", "Skills 复制到当前项目工作区的 skills/ 文件夹。后续需要什么再复制什么，不污染 Codex 根目录。", 0.85, 6.45, 11.55, 0.56, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=11.5)

    slide = new("课前输入材料检查")
    rows = [
        ["输入", "来源", "最低要求"],
        ["candidate_papers.xlsx", "第 03 讲", "至少 20 篇真实候选文献"],
        ["verified_references.bib", "第 03 讲", "至少 10 条已核验 BibTeX"],
        ["paper_cards/*.md", "第 02/04 讲", "至少 5 张论文卡片"],
        ["theory_variable_method_matrix.xlsx", "第 04 讲", "至少 10 篇文献的理论-变量-方法记录"],
        ["research_question_graph.md", "第 04 讲", "至少 5 个研究问题节点"],
    ]
    table(slide, rows, 0.82, 1.55, 11.60, 3.55, col_widths=[3.4, 2.2, 6.0], font_size=10)
    bullets(slide, [
        "材料不完整可以跟做流程，但不能把结果当作真实选题。",
        "缺失文献时要补检索；缺失 BibTeX 时要先核验。",
        "任何无法确认来源的条目都标注“待核验”。",
    ], 1.00, 5.55, 11.0, 1.0, size=15, color=COLORS["red"])

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第05讲_综述理论与选题工作区"\ncd "D:\\AI科研训练营\\第05讲_综述理论与选题工作区"\n\nNew-Item -ItemType Directory -Force -Path input,review,theory,hypotheses,proposal,figures,qc,output,skills\nNew-Item -ItemType Directory -Force -Path input\\paper_cards,input\\matrices,input\\graphs,input\\bibliography\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md', 0.85, 1.55, 11.65, 2.40, size=11)
    card(slide, "检查标准", "工作区中必须有 input、review、theory、hypotheses、proposal、qc、skills 文件夹，以及 AGENTS.md 和 research_log.md。", 0.95, 4.35, 5.45, 1.10, accent=COLORS["teal"], body_size=12.5)
    card(slide, "路径提醒", "如果没有 D 盘，可以换成 C 盘或自己的项目盘。路径中有中文没有问题，但尽量不要混用多个项目目录。", 6.80, 4.35, 5.45, 1.10, accent=COLORS["amber"], body_size=12.5, fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName literature-review,ssci-literature-review,ai4scholar-research,ai4scholar-citation-network,citation-management,hypothesis-generation,scientific-critical-thinking,research-proposal,theory-variable-method-matrix,research-question-knowledge-graph `\n  -Workspace "D:\\AI科研训练营\\第05讲_综述理论与选题工作区"', 0.80, 1.50, 11.75, 2.72, size=10)
    bullets(slide, [
        "复制后检查工作区 skills/ 文件夹。",
        "如果某个 Skill 缺失，在 research_log.md 记录，不要中断课程。",
        "不要安装到 Codex 根目录；本课按项目复制、按功能使用。",
    ], 0.95, 4.70, 11.1, 1.2, size=15)
    card(slide, "缺失处理", "缺少 research-ideation 或 idea-tournament 时，可以先用本讲提示词完成流程；课后再补 Skill。", 0.95, 6.12, 11.1, 0.65, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.5)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第05讲工作区 Agent 指令\n\n你是科研选题、文献综述与理论建构助手。\n\n原则：\n- 只使用真实文献、已核验 BibTeX、论文卡片和可核验检索结果。\n- 禁止编造引用、作者、期刊、DOI、数据结果。\n- 文献综述围绕问题、理论、机制、变量、方法和证据推进。\n- 理论必须解释机制，禁止只贴理论标签。\n- 假设必须可检验，并说明设计、测量和分析方法。", 0.85, 1.55, 6.25, 3.35, size=10.5)
    card(slide, "为什么要写", "AGENTS.md 是项目级约束。它让 Agent 每次处理本项目时都遵守真实文献、理论机制和人工核验要求。", 7.45, 1.70, 4.65, 1.25, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.5)
    card(slide, "课堂要求", "学生可以直接使用 lesson05_workspace_AGENTS_template.md，不需要从零写。", 7.45, 3.45, 4.65, 0.95, accent=COLORS["blue"], body_size=12.5)
    card(slide, "不能省略", "没有项目指令时，Agent 更容易把未核验内容写成确定事实。", 7.45, 4.92, 4.65, 0.95, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    slide = new("实操 4：写 research_log.md")
    codebox(slide, "# 第05讲研究日志\n\n## 我的研究方向\n- 中文方向：\n- 英文方向：\n- 学科归属：\n- 目标期刊或期刊层级：\n\n## 我目前关注的现象\n- 现象是什么：\n- 发生在什么场景：\n- 重要性在哪里：\n- 目前我不理解的地方：\n\n## 我可能使用的方法\n- 情景实验：\n- EEG/ERP：\n- 文本挖掘：\n- 机器学习/深度学习：\n- 因果推断：", 0.85, 1.48, 6.35, 4.55, size=10.2)
    card(slide, "不合格写法", "我想研究 AI Agent 对科研效率的影响。问题太大，没有对象、场景、机制和方法。", 7.55, 1.70, 4.55, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.2)
    card(slide, "合格方向", "在文献筛选任务中，AI Agent 的解释透明度是否通过降低不确定性感知并提升信任，影响博士生采纳 AI 推荐文献的意愿？", 7.55, 3.35, 4.55, 1.45, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)
    card(slide, "教师检查", "必须写清研究对象、任务场景、核心机制、结果变量和可能方法。", 7.55, 5.28, 4.55, 0.90, accent=COLORS["amber"], body_size=12.2)

    add_step_slide(
        prs, 5, "实操 5：生成证据表",
        "把论文卡片和矩阵压成可追溯证据，防止 Agent 直接写流畅但无证据的综述。",
        "请读取 input/paper_cards、input/candidate_papers.xlsx 和 input/matrices 中的材料，生成 review/evidence_table.md。\n\n不要写综述正文。每篇文献提取 citation_key、研究问题、理论、变量、机制、方法、发现、局限和可复用启发。\n没有依据的字段写“未提供”，不要猜。",
        "review/evidence_table.md",
        ["每行能追溯到真实材料", "没有凭空出现的新文献", "理论列不是变量名", "方法列能看出实验/问卷/文本/EEG/建模"],
        "证据表是本讲最重要的底层文件。后续所有判断都要回到这里。"
    )

    slide = new("证据表应该长什么样")
    codebox(slide, "| ID | 文献 | 问题 | 理论 | IV/操纵 | 机制 | DV | 方法 | 发现 | 局限 | 启发 |\n|---|---|---|---|---|---|---|---|---|---|---|\n| P01 | 已核验 citation_key |  |  |  |  |  |  |  |  |  |\n| P02 | 待核验 |  |  |  |  |  |  |  |  |  |", 0.80, 1.50, 11.75, 1.65, size=10.5)
    card(slide, "合格证据表", "能看出每篇论文的研究问题、理论、变量、方法、发现和局限。空白处写“未提供”，不是让 Agent 猜。", 0.95, 3.65, 5.45, 1.10, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "不合格证据表", "只写摘要、没有变量和方法、没有 citation_key、把未核验文献写成确定事实。", 6.85, 3.65, 5.45, 1.10, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "课堂动作", "教师随机点 1-2 个学生打开 evidence_table，检查是否存在“看起来合理但找不到来源”的内容。", 0.95, 5.35, 11.35, 0.75, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    add_step_slide(
        prs, 6, "实操 6：生成文献树",
        "把文献从逐篇摘要变成领域结构，找出贡献类型、挑战和未解决问题。",
        "基于 review/evidence_table.md，生成 review/literature_tree.md。\n\n建立两棵树：\n1. Novelty Tree：理论解释、新变量/新构念、新情境/新对象、新方法/新数据、新机制/边界条件。\n2. Challenge-Insight Tree：核心挑战、已有解决方式、未解决问题、对我选题的启发。\n\n不要按年份排列，不要逐篇摘要。",
        "review/literature_tree.md",
        ["有 Novelty Tree", "有 Challenge-Insight Tree", "每个分支有代表文献", "能看出已有解释的边界"],
        "文献树不是时间线。只要按年份写，基本就不合格。"
    )

    slide = new("文献树的两个核心结构")
    card(slide, "Novelty Tree", "按贡献类型整理：理论解释、新变量/新构念、新情境/新对象、新方法/新数据、新机制/边界条件。它回答“已有研究贡献了什么”。", 0.85, 1.55, 5.60, 1.75, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "Challenge-Insight Tree", "按挑战组织：领域面对什么问题，已有研究如何解决，哪些解释仍不足。它回答“还有什么没被解决”。", 6.85, 1.55, 5.60, 1.75, accent=COLORS["blue"], fill_color=COLORS["pale_blue"])
    flow(slide, ["核心问题", "已有解释", "证据类型", "争议", "未解决点", "我的切入"], 1.00, 4.10, 11.35, h=0.66, size=11.5)
    card(slide, "判断标准", "读完文献树后，应能说清：这个领域不是缺文献，而是在哪个机制、边界、测量或方法上还解释不够。", 1.00, 5.55, 11.35, 0.90, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=13)

    add_step_slide(
        prs, 7, "实操 7：识别研究流派与争议",
        "看清领域内部有哪些解释路线，以及已有结论为什么不一致。",
        "基于 review/evidence_table.md 和 review/literature_tree.md，生成 review/research_streams_and_controversies.md。\n\n列出 3-5 个研究流派和 2-4 个主要争议。\n每个争议说明来源：理论、样本、测量、方法、数据或情境。\n不要把所有差异都写成“研究视角不同”。",
        "review/research_streams_and_controversies.md",
        ["每个流派有核心问题", "每个争议有具体来源", "争议能否转化为缺口被说明", "没有万能化地写“视角不同”"],
        "争议不等于缺口。只有能导向机制、边界或方法改进时，才可能形成选题。"
    )

    slide = new("争议来源判断")
    rows = [
        ["争议来源", "判断问题", "可能形成的缺口"],
        ["理论不同", "用不同理论解释同一关系？", "理论整合或机制比较"],
        ["样本/情境不同", "关系只在特定群体或任务中成立？", "边界条件或情境机制"],
        ["测量不同", "构念被不同量表或指标测量？", "测量有效性或构念区分"],
        ["方法不同", "实验、问卷、二手数据结论不一致？", "识别策略或方法改进"],
        ["数据不同", "文本、行为、脑电、问卷看到不同过程？", "多源数据机制检验"],
    ]
    table(slide, rows, 0.85, 1.55, 11.65, 4.15, col_widths=[2.0, 4.6, 5.05], font_size=10)
    card(slide, "教师提醒", "不要把所有争议都写成“研究视角不同”。要明确冲突来自理论、样本、测量、方法、数据还是情境。", 0.95, 6.05, 11.35, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    add_step_slide(
        prs, 8, "实操 8：生成研究缺口候选",
        "从文献树和争议中生成 5-8 个可能选题，但先不急着保留。",
        "基于文献树、流派争议和证据表，生成 5-8 个研究缺口候选，输出 review/research_gap_candidates.md。\n\n每个缺口包含：类型、已知内容、未知内容、重要性、支持证据、可能研究问题、可行方法、审稿风险、补充检索需求。\n禁止空话。",
        "review/research_gap_candidates.md",
        ["每个 gap 有类型", "写清已知与未知", "有真实文献或待核验文献支撑", "至少有一个审稿风险"],
        "这一阶段允许多生成，但不能直接相信。下一步要严格审计。"
    )

    slide = new("六类常见研究缺口")
    rows = [
        ["类型", "说明", "常见错误"],
        ["理论缺口", "现有理论不能解释新现象或关键关系", "只贴一个新理论名"],
        ["机制缺口", "知道 X 影响 Y，但不知道为什么", "硬塞中介"],
        ["边界缺口", "关系在不同条件下强弱不同", "为了复杂而加调节"],
        ["测量缺口", "现有测量不能捕捉关键构念", "只是换量表"],
        ["方法缺口", "现有方法不能回答关键问题", "只炫技，不服务问题"],
        ["情境缺口", "新场景改变了原有机制", "只是换样本"],
    ]
    table(slide, rows, 0.82, 1.55, 11.70, 4.65, col_widths=[1.75, 5.25, 4.70], font_size=10)
    card(slide, "判断句式", "不是“已有研究较少”，而是“已有解释在某个机制、情境、测量或方法上不能回答关键问题”。", 0.95, 6.45, 11.35, 0.55, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)

    slide = new("弱缺口与强缺口示范")
    card(slide, "弱缺口", "已有研究较少关注 AI Agent 对科研效率的影响。", 0.95, 1.60, 5.25, 1.05, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=14)
    card(slide, "问题", "研究对象太泛；结果变量太泛；没有机制；没有任务场景；没有说明已有研究哪里解释不了。", 0.95, 3.00, 5.25, 1.30, accent=COLORS["red"], body_size=12.5)
    card(slide, "较强缺口", "已有研究虽讨论 AI 建议采纳和自动化信任，但在科研文献筛选这类高不确定、高专业判断任务中，解释透明度是否通过降低不确定性感知影响采纳，仍需要结合具体任务行为进行检验。", 6.70, 1.60, 5.65, 1.75, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)
    card(slide, "仍需核验", "已有研究是否已经讨论解释透明度—信任—采纳；科研任务是否真的改变机制；采纳意愿和实际行为是否一致。", 6.70, 3.82, 5.65, 1.30, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    add_step_slide(
        prs, 9, "实操 9：缺口审计",
        "删掉不成立的选题，避免把弱想法包装成贡献。",
        "请严格审计 review/research_gap_candidates.md，输出 review/research_gap_audit.md。\n\n对每个 gap 判断：是否已被解决、是否只是换情境、需要的新贡献、可行性、审稿风险、最终决策 keep/revise/drop。\n不要迎合我，弱 gap 必须指出。",
        "review/research_gap_audit.md",
        ["至少 1 个 gap 被 drop 或 revise", "保留 gap 有理由", "每个保留 gap 有反证检索建议", "没有“所有想法都很好”"],
        "缺口审计比生成缺口更重要。弱选题越早删越省时间。"
    )

    slide = new("缺口评分表")
    rows = [
        ["gap_id", "importance", "novelty", "feasibility", "theory_value", "method_fit", "evidence", "decision"],
        ["G1", "1-5", "1-5", "1-5", "1-5", "1-5", "1-5", "keep/revise/drop"],
        ["G2", "1-5", "1-5", "1-5", "1-5", "1-5", "1-5", "keep/revise/drop"],
        ["G3", "1-5", "1-5", "1-5", "1-5", "1-5", "1-5", "keep/revise/drop"],
    ]
    table(slide, rows, 0.78, 1.60, 11.75, 2.20, col_widths=[1.15, 1.35, 1.25, 1.45, 1.55, 1.35, 1.25, 2.45], font_size=9.2)
    bullets(slide, [
        "importance：问题是否重要",
        "novelty：是否真有新意",
        "feasibility：数据、实验、时间、能力是否可行",
        "theory_value：是否能推进理论解释",
        "method_fit：是否适合你的方法能力",
        "evidence：是否有足够真实文献支持",
    ], 0.95, 4.25, 5.6, 1.75, size=13.5)
    card(slide, "主选题公式", "在 [具体情境] 中，[核心自变量/刺激/技术特征] 如何通过 [心理/认知/神经/组织机制] 影响 [结果变量]，并受到 [边界条件] 的影响？", 6.95, 4.30, 5.35, 1.50, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.5)

    add_step_slide(
        prs, 10, "实操 10：建立理论机制链",
        "把最强研究问题转成理论解释，而不是只写变量关系。",
        "基于最强研究问题，生成 theory/theory_mechanism_chain.md。\n\n请区分情境/刺激、理论、构念、变量、测量、机制和结果。\n说明理论为什么适配，也说明理论不能解释的边界。\n给出适合的研究设计：情景实验、问卷、EEG/ERP、文本挖掘、机器学习或因果推断。",
        "theory/theory_mechanism_chain.md",
        ["理论不是只写名称", "构念和变量被区分", "每个变量有测量或操纵", "理论边界被写出来"],
        "理论要解释为什么 X 会影响 Y，以及什么时候关系更强或更弱。"
    )

    slide = new("理论机制链模板")
    flow(slide, ["情境/刺激", "心理/认知/神经机制", "行为/态度/绩效结果"], 1.10, 1.65, 11.0, h=0.78, size=14)
    rows = [
        ["层次", "含义", "示例"],
        ["情境/刺激", "研究中发生变化的场景或操纵", "AI 解释透明度高/低"],
        ["构念", "理论层面的抽象概念", "不确定性感知、信任"],
        ["变量", "可测量或可操纵指标", "量表得分、行为选择"],
        ["机制", "为什么 X 会影响 Y", "信息可解释性降低不确定性"],
        ["结果", "最终被解释的行为或绩效", "采纳意愿、任务表现"],
    ]
    table(slide, rows, 0.90, 3.05, 11.55, 3.15, col_widths=[1.55, 4.5, 5.5], font_size=10)
    card(slide, "底线", "如果理论不能解释机制，就不要把它写进论文。", 0.98, 6.48, 11.35, 0.55, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)

    add_step_slide(
        prs, 11, "实操 11：生成假设组",
        "把理论机制链转成可检验假设，并标明测量、设计和分析方法。",
        "基于 theory/theory_mechanism_chain.md，生成 hypotheses/hypothesis_set.md。\n\n每个假设包括：假设文本、机制解释、文献支持、测量/操纵、分析方法、因果语言风险。\n只写理论必要的中介和调节，不要为了复杂而添加。",
        "hypotheses/hypothesis_set.md",
        ["假设不是变量关系口号", "每个假设有机制解释", "每个假设有检验方法", "没有滥用因果语言"],
        "中介和调节不是越多越好。没有理论依据的复杂模型会增加审稿风险。"
    )

    slide = new("假设质量检查")
    rows = [
        ["问题", "不合格写法", "修改方向"],
        ["只有关系", "X 正向影响 Y", "写清为什么影响"],
        ["中介硬塞", "X 通过 M 影响 Y", "说明 M 为什么是过程机制"],
        ["调节硬塞", "W 调节 X 与 Y", "说明 W 为什么改变关系强度"],
        ["因果过度", "问卷证明 X 导致 Y", "改成关联，或补实验/识别策略"],
    ]
    table(slide, rows, 0.85, 1.60, 11.65, 3.45, col_widths=[1.6, 4.1, 5.95], font_size=10.5)
    codebox(slide, "较弱：AI 透明度正向影响用户信任。\n\n较好：AI 透明度通过降低用户对算法决策过程的不确定性感知，提高用户对 AI 建议的信任。", 0.95, 5.45, 11.35, 0.95, size=11)

    add_step_slide(
        prs, 12, "实操 12：生成综述提纲",
        "把前面的文献树、缺口和机制链转成可以写论文的综述结构。",
        "基于 review/literature_tree.md、review/research_gap_audit.md 和 theory/theory_mechanism_chain.md，生成 review/synthesis_outline.md。\n\n综述提纲按以下逻辑组织：研究现象与核心问题、已有理论解释、关键变量与机制、方法和证据类型、争议与不足、我的研究如何进入这个缺口。\n不要按年份写，不要逐篇摘要。",
        "review/synthesis_outline.md",
        ["提纲有问题推进", "每一节服务于研究问题", "不是按作者年份堆材料", "最后能自然引出自己的研究"],
        "综述提纲是后续写 Introduction 和 Literature Review 的骨架。"
    )

    add_step_slide(
        prs, 13, "实操 13：生成一页研究方案",
        "把选题压缩成导师能快速判断的一页方案。",
        "基于前面所有输出，生成 proposal/one_page_research_proposal.md。\n\n必须包含：中文英文题目、一句话研究问题、重要性、文献缺口、理论机制、研究模型、假设组、研究设计、数据与分析方法、预期贡献、最大风险、下一步计划。\n控制在 1-2 页，所有引用标注已核验或待核验。",
        "proposal/one_page_research_proposal.md",
        ["导师 3 分钟能判断", "缺口具体", "理论机制支撑假设", "方法可执行", "风险没有被隐藏"],
        "一页方案不是摘要，而是决策文件。"
    )

    slide = new("一页研究方案结构")
    rows = [
        ["模块", "必须回答的问题"],
        ["题目", "研究对象、机制、场景是否清楚"],
        ["研究问题", "一句话能否说明要解释什么"],
        ["重要性", "为什么值得研究，不只是“有意义”"],
        ["文献缺口", "现有研究哪里解释不足"],
        ["理论机制", "为什么 X 影响 Y"],
        ["研究模型/假设", "关系是否可检验"],
        ["研究设计", "实验、问卷、EEG/ERP、文本、建模或因果推断怎么做"],
        ["风险", "最大审稿风险和下一步补救"],
    ]
    table(slide, rows, 0.85, 1.48, 11.65, 5.25, col_widths=[2.1, 9.55], font_size=10.2)

    add_step_slide(
        prs, 14, "实操 14：导师前审计",
        "在交给导师前，让 Agent 用严格审稿人标准挑问题。",
        "以严格审稿人和导师视角，基于 proposal/one_page_research_proposal.md，生成 qc/gap_and_theory_audit.md。\n\n检查：研究问题是否过大、缺口是否成立、是否只是换情境、理论是否解释机制、构念变量测量是否一致、假设是否可检验、方法能否支持因果语言、文献是否真实、贡献是否夸大。",
        "qc/gap_and_theory_audit.md",
        ["列出致命问题", "列出重要问题", "列出可修改问题", "给出下一步行动清单"],
        "不要害怕审计结果难看。现在暴露问题比投稿后暴露问题成本低。"
    )

    slide = new("引用真实性核验")
    bullets(slide, [
        "任何引用进入正文前必须核验：题名、作者、年份、期刊、DOI 或出版页面。",
        "可用链路：Ai4Scholar、AMiner、OpenAlex、Zotero、期刊官网、DOI 页面。",
        "BibTeX 只说明格式，不自动证明文献真实。",
        "无法核验时，标注“待核验”，不能写成确定事实。",
    ], 0.90, 1.55, 6.0, 2.3, size=15.5)
    codebox(slide, "# qc/citation_reality_check.md\n\n| citation_key | title | DOI | platform_checked | status | issue |\n|---|---|---|---|---|---|\n|  |  |  | Ai4Scholar / AMiner / OpenAlex / DOI | verified/pending/error |  |", 7.15, 1.62, 5.10, 2.25, size=9.5)
    card(slide, "课堂底线", "模型不能凭记忆生成参考文献。凡是查不到来源的引用，都不能作为真实文献使用。", 0.95, 4.60, 11.30, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=13)
    card(slide, "学生动作", "把所有 proposal 中出现的引用复制到 qc/citation_reality_check.md，逐条核验。", 0.95, 5.75, 11.30, 0.70, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.5)

    slide = new("课堂示范：模糊想法如何压缩")
    card(slide, "原始想法", "我想研究 AI Agent 对科研效率的影响。", 0.90, 1.55, 4.20, 0.95, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=14)
    card(slide, "主要问题", "对象泛、场景泛、结果变量泛、没有机制、没有可检验设计。", 0.90, 2.90, 4.20, 1.10, accent=COLORS["red"], body_size=12.5)
    card(slide, "压缩后", "在文献筛选任务中，AI Agent 的解释透明度是否通过降低不确定性感知并提升信任，增加博士生采纳 AI 推荐文献的意愿？", 5.55, 1.55, 6.65, 1.45, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=13)
    card(slide, "方法化", "高解释透明度 vs. 低解释透明度；测量不确定性感知、信任、采纳意愿和实际采纳行为；使用中介模型检验机制。", 5.55, 3.45, 6.65, 1.35, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.5)
    card(slide, "仍需核验", "已有研究是否已经覆盖解释透明度—信任—采纳；科研任务是否真正改变机制；量表和行为指标是否有效。", 0.95, 5.40, 11.25, 0.85, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    slide = new("如果加入 EEG/ERP 怎么处理")
    card(slide, "可以做", "把解释透明度、冲突信息、风险提示等设计成刺激材料；记录判断过程中的认知负荷、冲突加工或注意分配相关指标。", 0.90, 1.55, 5.55, 1.35, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.5)
    card(slide, "不能做", "不能把脑电指标当作装饰。必须说明 EEG/ERP 指标与心理机制之间的理论对应。", 6.85, 1.55, 5.55, 1.10, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)
    rows = [
        ["问题", "需要补充"],
        ["机制是什么", "认知负荷、冲突监测、注意分配、奖赏/风险加工等"],
        ["指标为什么相关", "用真实神经科学文献说明，不凭感觉"],
        ["任务是否可行", "刺激时长、试次数、伪迹、反应方式"],
        ["统计如何报告", "ERP/时频/行为指标与假设一一对应"],
    ]
    table(slide, rows, 0.95, 3.35, 11.35, 2.85, col_widths=[2.5, 8.85], font_size=10.5)

    slide = new("常见错误与修正")
    rows = [
        ["错误", "后果", "修正"],
        ["直接写综述正文", "掩盖证据不足", "先证据表和文献树"],
        ["缺口写“研究较少”", "审稿人不认可", "说明现有解释哪里不足"],
        ["理论只写名称", "无法支撑假设", "写理论命题和机制链"],
        ["硬加中介调节", "模型复杂但不可信", "只保留理论必要关系"],
        ["编造文献", "学术风险严重", "逐条核验引用"],
        ["问卷写因果", "方法与结论不匹配", "改语言或补实验/识别策略"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.3, 4.0, 5.45], font_size=10)

    slide = new("课后提交要求")
    bullets(slide, [
        "提交文件夹：姓名_第05讲_综述理论与研究问题生成",
        "必须包含 review/、theory/、hypotheses/、proposal/、qc/ 和 research_log.md",
        "必须附 300 字反思：原始想法哪里被文献树、缺口审计或理论机制链修正了",
        "未核验引用必须保留“待核验”标记，不能改成确定引用",
    ], 0.90, 1.55, 11.3, 1.75, size=16)
    rows = [
        ["评分项", "占比"],
        ["证据表与引用真实性", "25%"],
        ["文献树与流派争议", "20%"],
        ["缺口审计质量", "20%"],
        ["理论机制与假设", "20%"],
        ["一页方案与反思", "15%"],
    ]
    table(slide, rows, 2.15, 3.80, 8.85, 2.20, col_widths=[6.6, 2.25], font_size=11.5)
    card(slide, "最低合格线", "能清楚说明：我的研究问题是什么，已有文献哪里不足，我用什么理论解释，准备用什么方法检验。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.5)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是让 AI 替你写综述，而是让 Agent 帮你把模糊想法压缩成可讨论、可核验、可检验的研究方案。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["读文献", "建结构", "找缺口", "审缺口", "讲机制", "写假设", "成方案"], 1.00, 3.50, 11.20, h=0.70, size=12)
    card(slide, "下一步", "补真实文献核验，修改被审计指出的问题，把一页方案交给导师或同伴做第二轮评审。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
