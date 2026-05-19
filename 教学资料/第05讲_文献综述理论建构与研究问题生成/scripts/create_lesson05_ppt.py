from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第05讲_课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(247, 248, 245),
    "ink": RGBColor(31, 41, 51),
    "muted": RGBColor(91, 103, 112),
    "teal": RGBColor(15, 118, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(180, 83, 9),
    "red": RGBColor(185, 28, 28),
    "line": RGBColor(210, 215, 220),
    "white": RGBColor(255, 255, 255),
    "pale_teal": RGBColor(219, 245, 241),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(254, 243, 199),
    "pale_red": RGBColor(254, 226, 226),
    "dark": RGBColor(33, 43, 54),
}

FONT = "Microsoft YaHei"


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color=COLORS["line"], width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def clear_line(shape):
    shape.line.fill.background()


def add_text(slide, text, x, y, w, h, size=20, color=COLORS["ink"], bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=COLORS["ink"], gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def add_code(slide, text, x, y, w, h, size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, COLORS["dark"])
    clear_line(shape)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(235, 239, 245)
    return shape


def add_header(slide, title, section="第 05 讲", num=None):
    set_fill(slide.background, COLORS["bg"])
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    set_fill(bar, COLORS["teal"])
    clear_line(bar)
    add_text(slide, section, 0.45, 0.05, 1.6, 0.22, size=8, color=COLORS["white"], bold=True)
    if num is not None:
        add_text(slide, f"{num:02d}", 12.05, 0.04, 0.7, 0.25, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.62, 11.8, 0.55, size=25, color=COLORS["ink"], bold=True)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.21), Inches(1.1), Inches(0.06))
    set_fill(line, COLORS["amber"])
    clear_line(line)


def add_card(slide, title, body, x, y, w, h, fill=COLORS["white"], accent=COLORS["teal"], title_size=15, body_size=12):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill)
    set_line(shape, COLORS["line"], 0.8)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    set_fill(stripe, accent)
    clear_line(stripe)
    add_text(slide, title, x + 0.18, y + 0.15, w - 0.28, 0.35, size=title_size, color=accent, bold=True)
    add_text(slide, body, x + 0.18, y + 0.62, w - 0.32, h - 0.72, size=body_size, color=COLORS["ink"])
    return shape


def add_table(slide, rows, x, y, w, h, col_widths=None, font_size=10, header_fill=COLORS["teal"]):
    nrows, ncols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size if r else font_size + 0.5)
                p.font.bold = bool(r == 0)
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(242, 245, 247))
    return table_shape


def add_flow(slide, labels, x, y, w, box_h=0.58, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11):
    gap = 0.18
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(box_h))
        set_fill(shape, color)
        set_line(shape, accent, 1)
        add_text(slide, label, bx + 0.08, y + 0.14, box_w - 0.16, box_h - 0.2, size=font_size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.02), Inches(y + box_h / 2), Inches(bx - 0.02), Inches(y + box_h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.2)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_fill(slide.background, COLORS["bg"])
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.0), Inches(7.5))
    set_fill(left, COLORS["teal"])
    clear_line(left)
    add_text(slide, "文献综述\n理论建构\n研究问题生成", 0.45, 0.9, 2.35, 2.2, size=23, color=COLORS["white"], bold=True)
    add_text(slide, "第 05 讲", 0.6, 5.95, 1.4, 0.35, size=14, color=COLORS["white"], bold=True)
    add_text(slide, "从论文卡片到研究缺口、机制、假设与一页方案", 3.45, 1.22, 9.0, 0.7, size=27, color=COLORS["ink"], bold=True)
    add_text(slide, "Literature Tree · Gap Audit · Theory Mechanism · Hypothesis Set · One-Page Proposal", 3.5, 2.1, 8.7, 0.45, size=15, color=COLORS["muted"])
    add_flow(slide, ["文献卡片", "文献树", "缺口审计", "理论机制", "假设组", "一页方案"], 3.55, 3.25, 8.8, box_h=0.66, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_card(slide, "本讲交付", "literature_tree、research_gap_audit、synthesis_outline、theory_mechanism_chain、hypothesis_set、one_page_proposal", 3.55, 4.65, 8.75, 1.35, fill=COLORS["white"], accent=COLORS["amber"], title_size=14, body_size=13)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new_slide(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_header(slide, title, num=len(prs.slides))
        return slide

    slide = new_slide("这节课解决什么问题")
    add_bullets(slide, [
        "文献综述只堆材料，没有结构推进",
        "研究缺口写成“已有研究较少”，但没有证据",
        "理论只是贴标签，不能解释机制",
        "变量、构念、测量和假设之间不匹配",
        "假设只是统计关系，不是理论预测",
        "研究问题太大，无法用数据或实验检验",
        "不知道如何从文献卡片生成选题",
    ], 0.9, 1.55, 6.1, 3.8, size=17)
    add_card(slide, "本讲核心", "把文献从“材料”转成“问题结构”：问题、理论、机制、变量、方法、证据、缺口。", 7.25, 1.85, 4.95, 1.55, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)
    add_card(slide, "底线", "不能把“我没读到”写成“文献没有”；不能把理论名写成理论解释。", 7.25, 4.05, 4.95, 1.35, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=14)

    slide = new_slide("本讲操作路线")
    add_flow(slide, ["输入材料", "文献树", "缺口候选", "缺口审计", "综述提纲", "机制链", "假设组", "一页方案"], 0.55, 1.55, 12.25, box_h=0.7, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=10)
    rows = [
        ["阶段", "输出"],
        ["输入", "论文卡片、矩阵、图谱、BibTeX、候选文献"],
        ["结构化综述", "literature_tree.md, synthesis_outline.md"],
        ["选题", "research_gap_candidates.md, research_gap_audit.md"],
        ["理论", "theory_mechanism_chain.md"],
        ["假设", "hypothesis_set.md"],
        ["提案", "one_page_research_proposal.md"],
        ["审计", "gap_and_theory_audit.md"],
    ]
    add_table(slide, rows, 0.8, 2.75, 11.85, 3.8, col_widths=[2.2, 9.65], font_size=12)

    slide = new_slide("需要哪些 Skills")
    rows = [
        ["功能", "Skills"],
        ["文献综述", "literature-review, ssci-literature-review, literature-reviewer-skill"],
        ["研究构思", "research-ideation"],
        ["想法比较", "idea-tournament"],
        ["假设生成", "hypothesis-generation"],
        ["批判审计", "scientific-critical-thinking"],
        ["论文规划", "paper-planning, research-proposal"],
        ["文献核验", "ai4scholar-research, ai4scholar-citation-network, citation-management"],
        ["知识库", "theory-variable-method-matrix, research-question-knowledge-graph"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 5.25, col_widths=[3.0, 9.05], font_size=11.5)

    slide = new_slide("输入材料来自前几讲")
    rows = [
        ["输入", "来源", "最低要求"],
        ["candidate_papers.xlsx", "第 03 讲", "至少 20 篇真实候选文献"],
        ["verified_references.bib", "第 03 讲", "至少 10 条已核验 BibTeX"],
        ["paper_cards/*.md", "第 02/04 讲", "至少 5 张论文卡片"],
        ["theory_variable_method_matrix.xlsx", "第 04 讲", "理论、变量、方法矩阵"],
        ["research_question_graph.md", "第 04 讲", "初步研究问题图谱"],
        ["研究方向/目标期刊", "学生自选", "不能太大，要有方法可行性"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 4.55, col_widths=[3.4, 2.4, 6.25], font_size=11.5)
    add_text(slide, "没有真实输入，就不要让 Agent 生成研究缺口。", 0.9, 6.35, 11.2, 0.35, size=16, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("综述不是摘要：反例与修正")
    add_table(slide, [
        ["错误写法", "为什么不行", "修正方向"],
        ["A 研究发现... B 研究发现...", "逐篇堆材料，没有问题推进", "按问题、机制、证据组织"],
        ["已有研究较少关注 X", "没有证据，像主观感受", "说明已有研究做了什么、没解释什么"],
        ["基于某理论，本文认为...", "理论只是标签", "写清理论如何解释变量关系"],
        ["加入中介/调节变量", "统计模型复杂但理论弱", "先说明机制和边界条件必要性"],
    ], 0.65, 1.55, 12.05, 4.25, col_widths=[4.0, 4.0, 4.05], font_size=11)
    add_card(slide, "好综述", "这个领域围绕什么问题形成了哪些解释、争议、证据和缺口？", 1.0, 6.15, 11.15, 0.75, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)

    slide = new_slide("Step 1：创建工作区")
    add_code(slide, "New-Item -ItemType Directory -Force -Path \"D:\\AI科研训练营\\第05讲_综述与选题工作区\"\ncd \"D:\\AI科研训练营\\第05讲_综述与选题工作区\"\n\nNew-Item -ItemType Directory -Force -Path input,review,theory,hypotheses,proposal,figures,qc,output,skills\nNew-Item -ItemType Directory -Force -Path input\\paper_cards,input\\matrices,input\\graphs,input\\bibliography\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md", 0.75, 1.55, 12.0, 2.65, size=11.5)
    add_bullets(slide, [
        "把第 03 讲候选文献和 BibTeX 放入 input",
        "把第 04 讲论文卡片、矩阵、图谱放入 input",
        "本讲输出全部放 review/theory/hypotheses/proposal/qc",
    ], 1.0, 4.65, 11.1, 1.3, size=18)

    slide = new_slide("Step 2：复制 Skills")
    add_code(slide, "cd \"D:\\desk\\AI agent科研资料\\本地Skills功能分类库\"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName literature-review,ssci-literature-review,research-ideation,idea-tournament,hypothesis-generation,scientific-critical-thinking,paper-planning,research-proposal,ai4scholar-research,ai4scholar-citation-network,citation-management,theory-variable-method-matrix,research-question-knowledge-graph `\n  -Workspace \"D:\\AI科研训练营\\第05讲_综述与选题工作区\"", 0.75, 1.45, 12.0, 2.6, size=10.5)
    add_card(slide, "检查点", "skills 文件夹中至少有 literature/research ideation/hypothesis/critical thinking/citation 相关能力。", 1.0, 4.7, 11.15, 0.9, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=14)

    slide = new_slide("Step 3：明确研究方向")
    add_code(slide, "## 第05讲研究方向\n\n### 我的研究方向\n\n### 长期目标\n\n### 目标学科/期刊\n\n### 可用方法\n- 情景实验\n- EEG/ERP\n- 文本挖掘\n- 机器学习/深度学习\n- 因果推断\n\n### 本讲希望生成的问题", 0.85, 1.45, 5.65, 4.9, size=12)
    add_bullets(slide, [
        "不能只写“AI agent”",
        "要写研究对象、情境、机制或方法",
        "要说明目标期刊或学科标准",
        "要明确自己能做的数据或实验",
    ], 7.0, 1.8, 5.1, 2.2, size=18)
    add_card(slide, "例子", "AI agent 透明度如何通过降低不确定性感知影响用户对科研推荐系统的信任。", 7.0, 4.7, 5.1, 1.0, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)

    slide = new_slide("Step 4：生成文献树")
    add_code(slide, "Use research-ideation and literature-review to build a literature tree from my paper cards and theory-variable-method matrix.\nCreate both:\n1. Novelty Tree\n2. Challenge-Insight Tree\nOutput review/literature_tree.md.\nDo not summarize papers one by one.\nOrganize by problems, mechanisms, variables, methods, and unresolved challenges.", 0.85, 1.45, 11.8, 2.25, size=12)
    add_table(slide, [
        ["Novelty Tree", "理论新解释、新变量/构念、新情境、新方法/数据、新机制/边界"],
        ["Challenge-Insight Tree", "挑战、已有解决方式、代表文献、是否充分、新机会"],
        ["验收", "不是按年份，不是逐篇摘要，每个分支有代表文献"],
    ], 0.85, 4.15, 11.8, 1.65, col_widths=[3.1, 8.7], font_size=12)

    slide = new_slide("文献树模板")
    rows = [
        ["模块", "要填什么"],
        ["研究方向", "方向、长期目标、目标学科/期刊、主要方法"],
        ["Novelty Tree", "已有研究按贡献类型分类"],
        ["Challenge-Insight Tree", "挑战、已有解决方式、未解决问题"],
        ["研究流派图", "流派、核心问题、理论、变量、方法、局限"],
        ["争议与矛盾", "观点差异、差异来源、可能解释"],
        ["缺口池", "gap_id、缺口表述、支持证据、仍需核验"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.15, col_widths=[3.0, 8.85], font_size=12)

    slide = new_slide("Step 5：生成研究缺口候选池")
    add_code(slide, "Based on review/literature_tree.md and input/matrices/theory_variable_method_matrix.xlsx, generate 5-8 candidate research gaps.\nFor each gap, state:\ngap type, supporting papers, what is known, what remains unknown, why it matters, and what evidence is still missing.\nOutput review/research_gap_candidates.md.", 0.85, 1.45, 11.8, 1.75, size=12.5)
    add_table(slide, [
        ["缺口类型", "说明"],
        ["理论缺口", "现有理论不能解释某机制"],
        ["机制缺口", "变量关系存在，但中间过程不清楚"],
        ["边界条件缺口", "关系在什么条件下成立不清楚"],
        ["测量缺口", "构念测量不稳定或不适配新情境"],
        ["方法缺口", "现有方法不能捕捉关键现象"],
        ["因果识别缺口", "现有研究只能相关，不能支持因果"],
    ], 0.85, 3.65, 11.8, 2.75, col_widths=[3.0, 8.8], font_size=12)

    slide = new_slide("好缺口与坏缺口")
    add_card(slide, "坏缺口", "已有研究较少关注 AI agent 在科研中的应用。", 0.85, 1.55, 5.6, 1.15, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=15)
    add_card(slide, "好缺口", "已有研究解释了推荐系统透明度对用户信任的影响，但在 AI agent 主动生成科研建议的情境中，用户面对的是持续自主行动者而非单次算法输出，原有透明度理论尚不能解释用户如何形成对代理行为的过程性信任。", 6.85, 1.55, 5.7, 2.35, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=12)
    add_bullets(slide, [
        "好缺口必须说明已有研究解释了什么",
        "说明新情境为什么改变机制",
        "说明缺口为什么重要",
        "说明可以如何检验",
    ], 1.0, 4.35, 11.1, 1.4, size=18)

    slide = new_slide("Step 6：缺口审计")
    add_code(slide, "Use research_gap_audit_template.md to audit each candidate gap.\nFor each gap, perform a well-established solution check:\nHas existing literature already solved it?\nIs it merely a new context/sample?\nDoes it require new theory, new mechanism, new measurement, or new method?\nOutput review/research_gap_audit.md.", 0.85, 1.45, 11.8, 1.85, size=12.5)
    add_table(slide, [
        ["审计问题", "判断"],
        ["是否已有直接解决方案", "有则修改或放弃"],
        ["是否只是换样本/情境", "没有机制变化则弱"],
        ["是否需要新理论或新机制", "有才更可能贡献"],
        ["是否能被数据检验", "不能检验则暂缓"],
        ["是否适合目标期刊", "不适合则调整定位"],
    ], 0.85, 3.8, 11.8, 2.3, col_widths=[4.0, 7.8], font_size=12)

    slide = new_slide("Step 7：生成综述提纲")
    add_code(slide, "Create a literature review synthesis outline.\nDo not organize by year or by individual paper.\nOrganize by:\n1. core problem,\n2. theoretical explanations,\n3. mechanisms and variables,\n4. methods and evidence,\n5. unresolved gaps,\n6. my study's position.\nOutput review/synthesis_outline.md.", 0.85, 1.45, 6.0, 4.0, size=12)
    add_bullets(slide, [
        "综述提纲必须能推出自己的研究",
        "每一节都有问题推进",
        "理论、机制、方法和证据要互相连接",
        "不要把文献综述写成参考文献说明书",
    ], 7.3, 1.75, 4.7, 2.4, size=18)
    add_card(slide, "合格标准", "读完综述提纲，别人能明白为什么你的研究问题必须出现。", 7.3, 4.8, 4.7, 1.0, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("Step 8：选择最强研究缺口")
    add_table(slide, [
        ["gap_id", "importance", "novelty", "feasibility", "theory_value", "method_fit", "decision"],
        ["G1", "1-5", "1-5", "1-5", "1-5", "1-5", "keep/revise/drop"],
        ["G2", "1-5", "1-5", "1-5", "1-5", "1-5", "keep/revise/drop"],
        ["G3", "1-5", "1-5", "1-5", "1-5", "1-5", "keep/revise/drop"],
    ], 0.7, 1.55, 11.95, 2.5, col_widths=[1.6, 1.6, 1.6, 1.7, 1.9, 1.8, 1.75], font_size=11)
    add_bullets(slide, [
        "保留 1 个最强 gap，最多 1 个备选",
        "不要同时推进 5 个方向",
        "理论价值和方法可行性必须同时成立",
        "低可行性但高新颖性的方向先放入备选库",
    ], 1.0, 4.55, 11.1, 1.5, size=18)

    slide = new_slide("Step 9：理论机制链")
    add_code(slide, "Use theory_mechanism_chain_template.md to build a theory mechanism chain for the strongest research gap.\nConnect:\ncontext/stimulus/technology change → psychological/cognitive mechanism → behavioral/attitudinal/performance outcome.\nSeparate constructs from measured variables.\nOutput theory/theory_mechanism_chain.md.", 0.85, 1.45, 11.8, 1.8, size=12.5)
    add_flow(slide, ["情境/刺激/技术变化", "心理/认知机制", "行为/态度/绩效结果", "边界条件"], 1.0, 3.85, 11.2, box_h=0.75, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=13)
    add_card(slide, "检查", "理论必须解释机制；构念、变量和测量必须分开；每个环节有文献证据。", 1.0, 5.55, 11.15, 0.85, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("理论机制链模板")
    rows = [
        ["模块", "内容"],
        ["理论选择", "理论名、来源文献、原始领域、适用原因、边界条件"],
        ["理论核心命题", "该理论真正解释什么"],
        ["机制链", "前因、机制、结果、边界条件"],
        ["构念与变量", "理论含义、变量名、测量/操纵方式"],
        ["假设草案", "H1/H2/H3"],
        ["可检验预测", "如果成立/不成立，应观察到什么"],
        ["错误检查", "理论标签化、变量错位、不可检验"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.15, col_widths=[3.1, 8.75], font_size=12)

    slide = new_slide("Step 10：生成假设组")
    add_code(slide, "Use hypothesis-generation and hypothesis_set_template.md to create a testable hypothesis set.\nInclude main effect, mediation, moderation, and conditional indirect effect only if theoretically justified.\nFor each hypothesis, include mechanism explanation, literature support, measurement/design, analysis method, and causal language warning.\nOutput hypotheses/hypothesis_set.md.", 0.85, 1.45, 11.8, 2.0, size=12.5)
    add_table(slide, [
        ["假设类型", "什么时候写"],
        ["主效应", "理论预测 IV 影响 DV"],
        ["中介", "理论解释为什么会影响"],
        ["调节", "理论解释何时更强/更弱"],
        ["调节中介", "边界条件改变间接机制"],
    ], 0.85, 4.05, 11.8, 2.0, col_widths=[2.8, 9.0], font_size=12)

    slide = new_slide("假设写法：弱 vs 强")
    add_card(slide, "弱", "H1: AI 透明度正向影响用户信任。", 0.85, 1.55, 5.65, 1.0, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=15)
    add_card(slide, "强", "H1: AI agent 的过程透明度通过降低用户对其自主决策过程的不确定性感知，提高用户对科研建议的信任。", 6.85, 1.55, 5.65, 1.35, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)
    add_bullets(slide, [
        "强假设说明机制",
        "强假设暗含可测中介",
        "强假设能指导实验操纵",
        "强假设能对应统计模型",
    ], 1.0, 3.6, 11.1, 1.5, size=19)
    add_card(slide, "注意", "不是每个研究都需要复杂模型；理论不支持时不要硬加中介/调节。", 1.0, 5.7, 11.15, 0.8, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("Step 11：一页研究方案")
    rows = [
        ["模块", "内容"],
        ["标题", "清楚表达研究对象和机制"],
        ["研究问题", "一句话说明关注什么"],
        ["为什么重要", "实践重要性 + 理论重要性"],
        ["文献缺口", "已有研究做了什么，仍缺什么"],
        ["理论机制", "理论 → 机制 → 变量关系"],
        ["研究模型/假设", "IV, mediator, moderator, DV"],
        ["研究设计", "数据、样本、方法、测量、模型"],
        ["贡献与风险", "预期贡献、最大风险、下一步"],
    ]
    add_table(slide, rows, 0.75, 1.35, 11.85, 5.35, col_widths=[3.1, 8.75], font_size=11.5)

    slide = new_slide("生成一页方案的提示词")
    add_code(slide, "Use one_page_research_proposal_template.md to write a one-page research proposal.\nThe proposal must include:\ntitle, research question, importance, literature gap, theory mechanism, research model, hypotheses, design, expected contribution, risks, and next steps.\nOutput proposal/one_page_research_proposal.md.", 0.85, 1.55, 11.8, 1.75, size=12.5)
    add_bullets(slide, [
        "导师应能 3 分钟读懂",
        "不要写成 3000 字综述",
        "不要隐藏最大风险",
        "下一步必须可执行",
    ], 1.0, 4.1, 11.1, 1.5, size=19)

    slide = new_slide("Step 12：导师前审计")
    add_code(slide, "Use scientific-critical-thinking to audit my one-page research proposal.\nCheck:\ntheory fit, gap validity, variable-measure alignment, method fit, causal language, contribution overclaim, feasibility, and target journal fit.\nOutput qc/gap_and_theory_audit.md.", 0.85, 1.45, 11.8, 1.75, size=12.5)
    add_table(slide, [
        ["审计项", "要问的问题"],
        ["理论匹配", "理论是否真的解释机制"],
        ["gap 有效性", "是否已有文献解决"],
        ["变量测量", "构念和测量是否对应"],
        ["方法匹配", "设计能否检验假设"],
        ["因果语言", "是否过度声称因果"],
        ["贡献", "是否夸大为理论贡献"],
        ["可行性", "数据、样本、时间是否现实"],
    ], 0.85, 3.75, 11.8, 2.8, col_widths=[3.0, 8.8], font_size=11.5)

    slide = new_slide("必须人工核验的地方")
    rows = [
        ["风险", "人工核验"],
        ["文献缺口", "不能把“我没读到”写成“文献没有”"],
        ["真实文献", "引用必须可核验，不能靠模型记忆"],
        ["理论机制", "理论必须解释变量关系，不能只写理论名"],
        ["变量测量", "构念、变量和测量不能混为一谈"],
        ["假设", "每个假设必须能被数据或实验检验"],
        ["因果语言", "没有实验/准实验/识别策略时不能说因果"],
        ["贡献", "不要把换情境、换样本包装成理论贡献"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 5.25, col_widths=[2.7, 9.35], font_size=12)

    slide = new_slide("本讲提交物")
    rows = [
        ["提交物", "最低要求"],
        ["literature_tree.md", "Novelty Tree + Challenge-Insight Tree"],
        ["research_gap_candidates.md", "至少 5 个候选 gap"],
        ["research_gap_audit.md", "每个 gap 有反证搜索和保留/放弃理由"],
        ["synthesis_outline.md", "按问题、理论、机制、方法、缺口组织"],
        ["theory_mechanism_chain.md", "理论、构念、变量、测量对应"],
        ["hypothesis_set.md", "至少 2 个可检验假设"],
        ["one_page_research_proposal.md", "导师可快速判断"],
        ["gap_and_theory_audit.md", "具体指出风险和修改建议"],
    ]
    add_table(slide, rows, 0.65, 1.35, 12.05, 5.5, col_widths=[4.0, 8.05], font_size=11.3)

    slide = new_slide("评分标准")
    rows = [
        ["项目", "分值", "标准"],
        ["输入材料", "10", "来自前几讲，真实可追溯"],
        ["文献树", "20", "有结构、有流派、有挑战和争议"],
        ["缺口审计", "20", "有证据、有反证搜索、有保留/放弃理由"],
        ["理论机制链", "15", "理论、构念、变量、测量对应"],
        ["假设组", "15", "可检验，有理论解释和方法对应"],
        ["一页研究方案", "10", "简洁清楚，导师可判断"],
        ["审计与反思", "10", "能指出真实风险和修正"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 5.05, col_widths=[2.7, 1.05, 8.3], font_size=11)

    slide = new_slide("直接不合格的情况")
    add_bullets(slide, [
        "编造文献或引用",
        "研究缺口只写“已有研究较少”",
        "文献综述按年份堆文献",
        "理论只是标签，没有机制",
        "假设无法用数据或实验检验",
        "把相关研究设计写成因果",
        "没有人工审计",
        "只提交聊天截图，没有本地文件",
    ], 1.0, 1.55, 11.1, 4.0, size=21, color=COLORS["red"])
    add_card(slide, "底线", "第五讲产出的不是漂亮想法，而是经文献、理论和方法审计后的可讨论研究方案。", 1.0, 6.0, 11.15, 0.8, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("本讲一句话总结")
    add_text(slide, "文献综述的目标，\n不是证明你读了很多论文，\n而是证明你的研究问题从已有知识结构中自然长出来。", 1.15, 1.95, 11.0, 1.8, size=28, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
    add_flow(slide, ["文献证据", "结构综述", "缺口审计", "理论机制", "可检验假设"], 1.35, 4.75, 10.55, box_h=0.75, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=13)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
