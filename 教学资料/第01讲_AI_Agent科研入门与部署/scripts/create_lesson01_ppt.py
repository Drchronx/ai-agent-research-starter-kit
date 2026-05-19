from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第01讲_课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(247, 248, 245),
    "ink": RGBColor(31, 41, 51),
    "muted": RGBColor(91, 103, 112),
    "teal": RGBColor(15, 118, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(180, 83, 9),
    "red": RGBColor(185, 28, 28),
    "line": RGBColor(210, 215, 220),
    "white": RGBColor(255, 255, 255),
    "pale_teal": RGBColor(219, 245, 241),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(254, 243, 199),
    "pale_red": RGBColor(254, 226, 226),
    "soft": RGBColor(236, 239, 241),
}

FONT = "Microsoft YaHei"


def rgb_hex(color):
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color=COLORS["line"], width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def clear_line(shape):
    shape.line.fill.background()


def add_text(slide, text, x, y, w, h, size=20, color=COLORS["ink"], bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=COLORS["ink"], gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.1
    return box


def add_code(slide, text, x, y, w, h, size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, RGBColor(33, 43, 54))
    clear_line(shape)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(235, 239, 245)
    return shape


def add_header(slide, title, section="第 01 讲", num=None):
    set_fill(slide.background, COLORS["bg"])
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34)
    )
    bar = slide.shapes[-1]
    set_fill(bar, COLORS["teal"])
    clear_line(bar)
    add_text(slide, section, 0.45, 0.05, 1.6, 0.22, size=8, color=COLORS["white"], bold=True)
    if num is not None:
        add_text(slide, f"{num:02d}", 12.05, 0.04, 0.7, 0.25, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.62, 11.8, 0.55, size=25, color=COLORS["ink"], bold=True)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.21), Inches(1.1), Inches(0.06)
    )
    set_fill(line, COLORS["amber"])
    clear_line(line)


def add_card(slide, title, body, x, y, w, h, fill=COLORS["white"], accent=COLORS["teal"],
             title_size=15, body_size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, COLORS["line"], 0.8)
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    set_fill(stripe, accent)
    clear_line(stripe)
    add_text(slide, title, x + 0.18, y + 0.15, w - 0.28, 0.35, size=title_size, color=accent, bold=True)
    add_text(slide, body, x + 0.18, y + 0.62, w - 0.32, h - 0.72, size=body_size, color=COLORS["ink"])
    return shape


def add_table(slide, rows, x, y, w, h, col_widths=None, font_size=10, header_fill=COLORS["teal"]):
    nrows, ncols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size if r else font_size + 0.5)
                p.font.bold = bool(r == 0)
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(242, 245, 247))
    return table_shape


def add_flow(slide, labels, x, y, w, box_h=0.58, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11):
    gap = 0.18
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    boxes = []
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(box_h)
        )
        set_fill(shape, color)
        set_line(shape, accent, 1)
        add_text(slide, label, bx + 0.08, y + 0.14, box_w - 0.16, box_h - 0.2,
                 size=font_size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        boxes.append(shape)
        if i > 0:
            x1 = Inches(bx - gap + 0.02)
            x2 = Inches(bx - 0.02)
            ymid = Inches(y + box_h / 2)
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, ymid, x2, ymid)
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.2)
    return boxes


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_fill(slide.background, COLORS["bg"])
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.0), Inches(7.5))
    set_fill(left, COLORS["teal"])
    clear_line(left)
    add_text(slide, "AI Agent\n科研入门与部署", 0.55, 1.0, 2.1, 1.7, size=26, color=COLORS["white"], bold=True)
    add_text(slide, "第 01 讲", 0.6, 5.95, 1.4, 0.35, size=14, color=COLORS["white"], bold=True)
    add_text(slide, "从零搭建可复用的科研 Agent 工作区", 3.55, 1.28, 8.7, 0.65, size=28, color=COLORS["ink"], bold=True)
    add_text(slide, "Codex / Claude Code / OpenClaw / Hermes · Skills · MCP · Ai4Scholar · AMiner · 飞书/IMA", 3.58, 2.08, 8.6, 0.45, size=15, color=COLORS["muted"])
    add_flow(slide, ["选 Agent", "选模型", "配记忆", "装 Skills", "接 MCP", "跑任务"], 3.55, 3.25, 8.8, box_h=0.66, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_card(slide, "本讲交付", "一个项目级科研工作区：AGENTS.md、skills/、mcp.example.json、research_log.md、first_agent_test.md", 3.55, 4.65, 8.75, 1.35, fill=COLORS["white"], accent=COLORS["amber"], title_size=14, body_size=13)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    title_slide(prs)

    slides = []

    def new_slide(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_header(slide, title, num=len(prs.slides))
        return slide

    slide = new_slide("本节课解决什么科研问题")
    add_bullets(slide, [
        "工具混用：不知道任务该交给哪个 Agent",
        "模型乱选：只看“强不强”，忽视费用、合规、工具调用",
        "配置散落：API Key、MCP、Skills、记忆文件没有统一管理",
        "输出不可复现：今天做出的流程，明天学生自己跑不出来",
        "文献不可信：模型可能生成看似合理但不可核验的引用",
        "渠道割裂：飞书、IMA、文献平台和本地工作区没有打通",
    ], 0.9, 1.65, 5.7, 4.5, size=18)
    add_card(slide, "本讲核心目标", "把“会聊天”升级为“会部署科研工作流”：每个学生都建立一个可复制、可核验、可交作业的 Agent 工作区。", 7.0, 1.82, 5.2, 2.05, fill=COLORS["pale_teal"], accent=COLORS["teal"], title_size=16, body_size=15)
    add_card(slide, "课堂底线", "文献、数据、API Key 和组织权限必须人工核验。Agent 可以执行流程，但不能替代科研责任。", 7.0, 4.25, 5.2, 1.6, fill=COLORS["pale_amber"], accent=COLORS["amber"], title_size=16, body_size=15)

    slide = new_slide("每一讲固定教学模板")
    rows = [
        ["模块", "课堂必须回答的问题"],
        ["科研问题和需求", "这节课到底解决什么真实科研痛点"],
        ["需要哪些 Skills", "哪些能力要复制到项目工作区"],
        ["输入材料", "学生要准备什么文件、账号和数据"],
        ["Agent 怎么执行", "任务卡、步骤、命令、日志怎么写"],
        ["输出文件", "最终交付物路径、格式和合格标准"],
        ["人工核验", "哪些地方不能交给模型自动决定"],
        ["课后作业", "如何证明学生真的能独立完成"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.9, col_widths=[2.6, 9.25], font_size=13)
    add_text(slide, "课程不是工具清单，而是“输入 - 执行 - 输出 - 核验”的科研生产线。", 0.9, 6.55, 10.8, 0.35, size=14, color=COLORS["amber"], bold=True)

    slide = new_slide("本讲结束时学生必须产出")
    add_flow(slide, ["项目文件夹", "记忆文件", "skills/", "MCP 示例", "测试输出", "日志"], 0.8, 1.55, 11.8, box_h=0.7, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    rows = [
        ["文件", "合格标准"],
        ["AGENTS.md", "学科背景、引用规则、理论规则、写作风格、安全规则"],
        ["skills/", "至少包含 Ai4Scholar / AMiner / 项目管理相关 Skills"],
        ["configs/mcp.example.json", "密钥全部用环境变量占位，不出现真实 Token"],
        ["project_status.md", "写清目标、输入、下一步任务"],
        ["research_log.md", "记录一次 Agent 执行过程和人工核验结果"],
        ["output/first_agent_test.md", "列出已发现 Skills、配置状态、未核验项"],
    ]
    add_table(slide, rows, 0.8, 2.7, 11.8, 3.65, col_widths=[3.3, 8.5], font_size=12)

    slide = new_slide("科研 Agent 基础设施架构")
    add_flow(slide, ["学生需求", "Agent 客户端", "模型/中转站", "Skills", "MCP/外部平台", "本地输出"], 0.75, 1.7, 11.95, box_h=0.74, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11)
    add_card(slide, "Agent 客户端", "Codex、Claude Code、OpenClaw、Hermes。负责理解任务、调用工具、读写文件。", 0.85, 3.15, 3.75, 1.35, fill=COLORS["white"], accent=COLORS["teal"])
    add_card(slide, "模型/中转站", "按任务选择国内/国外模型，评估推理、代码、中文、长上下文、工具调用和成本。", 4.85, 3.15, 3.75, 1.35, fill=COLORS["white"], accent=COLORS["blue"])
    add_card(slide, "MCP/平台", "Ai4Scholar、AMiner、Zotero、飞书、浏览器、数据库。用真实工具替代模型记忆。", 8.85, 3.15, 3.75, 1.35, fill=COLORS["white"], accent=COLORS["amber"])
    add_text(slide, "关键原则：模型负责推理，Skills 负责流程，MCP 负责外部真实能力，本地日志负责复现。", 1.0, 5.55, 11.3, 0.45, size=18, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("四个 Agent 的区别")
    rows = [
        ["Agent", "核心定位", "科研中优先使用的场景"],
        ["Codex", "本地工作区、文件、代码与自动化", "资料整理、生成 docx/pptx/xlsx、批处理数据、维护 Skills"],
        ["Claude Code", "终端型长上下文代码/文档协作", "长论文结构修改、大项目理解、CLAUDE.md 记忆、MCP 调用"],
        ["OpenClaw", "Agent 网关、插件和多工具连接", "AI4Scholar 插件、MCP、浏览器、飞书/渠道工具"],
        ["Hermes", "可配置记忆、人格和工具的 Agent 客户端", "演示 SOUL/MEMORY/USER、工具装配、本地优先 Agent 行为"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 4.1, col_widths=[1.55, 3.35, 7.15], font_size=11)
    add_card(slide, "教学判断", "不是哪个 Agent 永远最好，而是“任务 - 工具 - 数据 - 输出格式”决定 Agent 选择。", 0.85, 6.0, 11.65, 0.75, fill=COLORS["pale_amber"], accent=COLORS["amber"], title_size=12, body_size=13)

    slide = new_slide("Codex：本地科研工作区自动化")
    add_bullets(slide, [
        "适合：资料整理、脚本、批量文件处理、Word/PPT/Excel 生成",
        "项目记忆：优先使用项目根目录 AGENTS.md",
        "Skills：课堂默认放在项目工作区 skills/，不塞根目录",
        "常用输出：docs/、output/、analysis/、research_log.md",
    ], 0.85, 1.55, 5.7, 2.5, size=18)
    add_code(slide, "npm install -g @openai/codex\ncd \"D:\\AI科研训练营\\第01讲_我的Agent工作区\"\ncodex", 7.0, 1.65, 5.35, 1.3, size=14)
    add_card(slide, "课堂提醒", "安装命令以 OpenAI 官方或当前课程版本为准。Codex 的价值不是替代文献数据库，而是把本地科研流程做成可复现文件。", 7.0, 3.35, 5.35, 1.65, fill=COLORS["pale_blue"], accent=COLORS["blue"])
    add_text(slide, "典型任务：把 30 篇 PDF 读取为论文卡片 + 文献筛选表 + 未核验项清单。", 0.9, 5.7, 11.4, 0.45, size=18, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("Claude Code：长上下文项目协作")
    add_code(slide, "npm install -g @anthropic-ai/claude-code\ncd \"D:\\AI科研训练营\\第01讲_我的Agent工作区\"\nclaude", 0.85, 1.6, 5.65, 1.25, size=14)
    add_card(slide, "适合", "长代码库理解、长篇论文结构调整、项目级记忆、MCP 工具调用、复杂文档重构。", 0.85, 3.2, 5.65, 1.35, fill=COLORS["white"], accent=COLORS["teal"])
    add_bullets(slide, [
        "CLAUDE.md 写项目规则，不写密钥",
        "稳定偏好可以放用户级记忆",
        "长文档任务要让 Agent 输出修改依据和核验项",
        "未配置文献工具时，不允许声称已完成真实检索",
    ], 7.0, 1.55, 5.3, 3.2, size=17)
    add_card(slide, "课堂提醒", "Claude Code 的记忆能提升协作一致性，但不能替代事实核验。", 7.0, 5.35, 5.3, 0.9, fill=COLORS["pale_amber"], accent=COLORS["amber"], title_size=12, body_size=13)

    slide = new_slide("OpenClaw：插件、MCP 与渠道网关")
    add_bullets(slide, [
        "适合：统一接入 AI4Scholar、浏览器、MCP、插件和渠道工具",
        "课堂重点：OpenClaw gateway + 插件 + MCP 配置",
        "Windows 安装失败时，优先切换轻量 MCP 模式",
        "所有 API Key 使用环境变量或私有配置文件",
    ], 0.85, 1.55, 5.7, 3.25, size=17)
    add_code(slide, "openclaw plugins install ai4scholar\n# 或按版本使用：openclaw plugins install clawhub:ai4scholar\nopenclaw gateway stop\nopenclaw gateway start\nopenclaw plugins list", 6.95, 1.55, 5.55, 2.05, size=12)
    add_card(slide, "科研应用", "AI4Scholar 插件/工具可用于论文搜索、全文、引用网络、自动引用、科学绘图；具体可用工具以当前插件返回为准。", 6.95, 4.05, 5.55, 1.45, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=12)

    slide = new_slide("Hermes：记忆、灵魂与工具装配")
    add_card(slide, "课堂定位", "Hermes 用来讲清 Agent 的“配置层”：SOUL、USER、MEMORY、TOOLS 如何影响行为。", 0.85, 1.55, 5.7, 1.45, fill=COLORS["white"], accent=COLORS["teal"], body_size=13)
    rows = [
        ["文件", "用途"],
        ["SOUL.md", "原则、边界、科研伦理、沟通方式"],
        ["USER.md", "使用者背景、研究方向、写作偏好"],
        ["MEMORY.md", "长期项目记忆、已排除方向、可复用经验"],
        ["TOOLS.md", "可用工具、MCP、Skills、脚本入口"],
    ]
    add_table(slide, rows, 0.85, 3.4, 5.7, 2.55, col_widths=[1.55, 4.15], font_size=11)
    add_bullets(slide, [
        "版本差异较大，安装命令必须按官方或课程包确认",
        "Windows 学生可优先 WSL / Git Bash",
        "不要把人格设定当成理论解释",
        "工具未配置时，不要让 Agent 假装能检索真实文献",
    ], 7.05, 1.75, 5.1, 2.7, size=17)
    add_card(slide, "教学价值", "让学生理解：Agent 能力不是来自“口号”，而是来自明确的记忆、工具、数据和核验流程。", 7.05, 4.95, 5.1, 1.05, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=13)

    slide = new_slide("如何按科研任务选择 Agent")
    rows = [
        ["任务", "优先 Agent", "原因"],
        ["生成 Word/PPT/Excel", "Codex", "本地文件读写和脚本执行强"],
        ["长论文结构修改", "Claude Code", "长上下文和项目记忆适合复杂文本"],
        ["接入 AI4Scholar/MCP", "OpenClaw", "插件和网关型工具更合适"],
        ["演示记忆/人格/工具装配", "Hermes", "适合讲配置层和本地 Agent 行为"],
        ["飞书/IMA 协作", "OpenClaw 或 Lark 工具 Agent", "关键是渠道权限和工具配置"],
    ]
    add_table(slide, rows, 0.7, 1.55, 11.95, 4.35, col_widths=[3.25, 2.55, 6.15], font_size=11)
    add_text(slide, "判断句式：这个任务的输入在哪里？要调用什么工具？输出文件是什么？哪些必须人工核验？", 0.9, 6.35, 11.2, 0.4, size=16, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("模型与中转站选择")
    add_text(slide, "课堂参考入口：", 0.85, 1.48, 2.0, 0.35, size=16, color=COLORS["muted"])
    add_code(slide, "https://hvoy.ai/model-select", 2.55, 1.35, 5.2, 0.62, size=15)
    rows = [
        ["任务", "更适合模型", "选择依据"],
        ["脚本/数据/文档生成", "强代码模型", "工具调用、代码可靠性、成本"],
        ["英文论文写作", "强推理/长上下文模型", "篇章逻辑、语气、审稿式修改"],
        ["中文资料/协作", "国内模型或中文强模型", "访问稳定、中文理解、合规"],
        ["文献核验", "模型 + MCP/数据库", "不能只靠模型记忆"],
        ["EEG/BCI 代码", "强代码模型 + 本地环境", "可复现、可调试"],
    ]
    add_table(slide, rows, 0.75, 2.35, 11.85, 3.75, col_widths=[3.25, 3.3, 5.3], font_size=11)

    slide = new_slide("国内模型与国外模型的取舍")
    add_card(slide, "国内模型优势", "国内访问稳定；中文任务和本土平台集成方便；成本可能更低；组织合规沟通更容易。", 0.85, 1.55, 5.7, 1.55, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)
    add_card(slide, "国外模型优势", "复杂推理、英文论文写作、代码和长上下文通常更强；工具生态成熟；适合跨学科英文材料。", 6.95, 1.55, 5.7, 1.55, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=13)
    add_bullets(slide, [
        "不是按国家标签判断，而是按任务能力、成本、合规和工具支持判断",
        "未公开数据、企业数据、被试隐私数据不能随意上传",
        "中转站要评估日志留存、稳定性、价格、限速和模型版本",
        "文献检索必须接数据库或 MCP，不能只问模型",
    ], 1.0, 3.65, 11.3, 2.25, size=18)

    slide = new_slide("API Key 与安全边界")
    add_card(slide, "允许", "使用环境变量、私有配置文件、课堂占位符。示例中统一写 ${API_KEY_NAME}。", 0.85, 1.55, 5.7, 1.4, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=13)
    add_card(slide, "禁止", "把真实 Key 写进 PPT、讲义、聊天记录、Git、飞书共享文档或学生作业。", 6.95, 1.55, 5.7, 1.4, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=13)
    add_code(slide, "$env:AI4SCHOLAR_API_KEY=\"你的真实key\"\n$env:AMINER_MCP_TOKEN=\"你的真实token\"\n# 提交作业时只保留 ${AI4SCHOLAR_API_KEY}", 1.15, 3.45, 11.0, 1.25, size=14)
    add_text(slide, "科研自动化的第一条合规规则：能不上传的原始数据不要上传，能不共享的密钥绝不共享。", 1.0, 5.65, 11.3, 0.45, size=18, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("记忆、灵魂、用户画像文件")
    rows = [
        ["文件", "写什么", "不写什么"],
        ["AGENTS.md", "项目规则、引用标准、理论规则、写作风格", "密钥、密码"],
        ["CLAUDE.md", "Claude Code 项目记忆", "不可公开数据"],
        ["SOUL.md", "Agent 原则、边界、科研伦理", "虚构事实"],
        ["USER.md", "学生背景、研究方向、输出偏好", "私人敏感信息"],
        ["MEMORY.md", "项目进展、可复用经验、待核验项", "未经授权的原始数据"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 4.4, col_widths=[1.65, 5.6, 4.8], font_size=11)
    add_text(slide, "记忆文件越具体，Agent 越稳定；记忆文件越混乱，Agent 越容易把偏好当事实。", 0.9, 6.35, 11.4, 0.35, size=16, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("AGENTS.md 最小模板")
    add_code(slide, "# 项目协作规则\n\n## 文献规则\n- 必须引用真实文献\n- 无法核验时标注“未核验”\n- 禁止编造 DOI、作者、期刊、年份\n\n## 理论规则\n- 理论必须解释机制\n- 禁止标签化套用理论\n\n## 写作风格\n- 简洁、准确、有推进感\n- 避免空泛表达\n\n## 安全规则\n- 不输出 API Key\n- 不上传敏感数据到不合规模型", 0.85, 1.45, 6.05, 5.25, size=11)
    add_bullets(slide, [
        "项目级规则写在项目根目录",
        "稳定规则可以跨 Agent 复用",
        "学生每次新项目先改 AGENTS.md",
        "引用和数据安全规则必须显式写出",
    ], 7.25, 1.7, 4.95, 2.25, size=18)
    add_card(slide, "课堂要求", "每个学生都要提交自己的 AGENTS.md，不能只复制老师模板。", 7.25, 4.5, 4.95, 1.05, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("Skills 是什么")
    add_text(slide, "Skill = 本地指令 + 可选脚本/模板/参考资料", 0.9, 1.45, 11.4, 0.5, size=22, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)
    add_code(slide, "skill-name/\n├─ SKILL.md\n├─ scripts/\n├─ references/\n├─ assets/\n└─ templates/", 0.95, 2.3, 4.05, 2.2, size=15)
    add_bullets(slide, [
        "告诉 Agent 什么时候用这个能力",
        "规定输入、步骤、输出和安全边界",
        "把课程经验封装成可复用科研流程",
        "不是模型本身，也不是插件市场宣传语",
        "好的 Skill 必须能产出可检查文件",
    ], 5.65, 2.1, 6.4, 2.8, size=18)
    add_card(slide, "本地库位置", "D:\\desk\\AI agent科研资料\\本地Skills功能分类库", 0.95, 5.25, 11.4, 0.9, fill=COLORS["pale_blue"], accent=COLORS["blue"], title_size=12, body_size=13)

    slide = new_slide("Skills 放根目录还是工作区")
    rows = [
        ["位置", "优点", "风险", "课堂建议"],
        ["Agent 根目录", "全局可用、省事", "污染所有项目；触发原因不透明", "固定成熟流程再用"],
        ["项目 skills/", "可复现、可交作业、可隔离", "每个项目要复制", "本课程默认"],
        ["临时复制", "演示快", "容易丢失、不可追踪", "只用于课堂演示"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 3.1, col_widths=[2.2, 2.9, 3.9, 3.05], font_size=11)
    add_code(slide, "powershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName ai4scholar-research,aminer-mcp-research `\n  -Workspace \"D:\\AI科研训练营\\第01讲_我的Agent工作区\"", 0.9, 5.15, 11.6, 1.15, size=12)

    slide = new_slide("本地 Skills 功能分类库")
    rows = [
        ["功能区", "课堂用途"],
        ["02 文献检索综述引用", "Ai4Scholar、AMiner、OpenAlex、引用核验"],
        ["07 自动化 MCP 浏览器与 Agent", "OpenClaw、MCP、Skill 创建、浏览器自动化"],
        ["10 情景实验与行为研究", "顶刊情景实验设计、操纵检验、结果汇报"],
        ["12 EEG/ERP 神经科学实验", "预处理、ERP、频域、连接、脑电写作"],
        ["15 BCI 脑电智能建模", "机器学习、深度学习、跨被试、防泄漏评估"],
        ["17 Zotero/Obsidian 知识库", "论文卡片、理论变量方法矩阵、知识图谱"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 4.9, col_widths=[4.0, 7.85], font_size=12)
    add_text(slide, "本讲只复制最小组合；后续课程按功能区逐步复制。", 1.0, 6.55, 11.2, 0.35, size=15, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("MCP：Agent 连接外部能力的标准接口")
    add_flow(slide, ["Agent", "MCP Client", "MCP Server", "工具/数据库", "真实结果"], 1.1, 1.55, 11.1, box_h=0.72, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_card(slide, "为什么科研需要 MCP", "文献、学者、Zotero、飞书、浏览器、数据库不能只靠模型记忆。MCP 让 Agent 能调用可核验工具。", 0.9, 3.0, 5.45, 1.65, fill=COLORS["white"], accent=COLORS["blue"], body_size=13)
    add_card(slide, "MCP 不等于万能", "工具返回什么、权限是否够、数据是否真实、结果是否过期，仍然需要人工核验。", 6.9, 3.0, 5.45, 1.65, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)
    add_text(slide, "课堂判断：凡是涉及外部事实、文献、学者、机构、文件系统，都优先考虑 MCP 或真实数据库。", 0.9, 5.65, 11.45, 0.45, size=17, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("Ai4Scholar-MCP：真实文献能力")
    add_bullets(slide, [
        "论文搜索、论文详情、PDF/全文阅读",
        "DOI/PMID/arXiv 等元数据核验",
        "引用网络、推荐、作者信息",
        "自动引用、BibTeX、科学绘图等能力按当前工具暴露为准",
    ], 0.85, 1.55, 5.75, 2.45, size=18)
    add_code(slide, "{\n  \"mcpServers\": {\n    \"ai4scholar\": {\n      \"url\": \"https://mcp.ai4scholar.net/sse\",\n      \"headers\": {\n        \"Authorization\": \"Bearer ${AI4SCHOLAR_API_KEY}\"\n      }\n    }\n  }\n}", 6.95, 1.45, 5.55, 3.25, size=11)
    add_card(slide, "人工核验", "检索结果不等于可引用结果。必须检查 DOI、期刊、作者、年份、主题匹配和是否为真实来源。", 0.9, 4.95, 11.5, 1.05, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("AMiner-MCP：学者、机构与知识图谱")
    add_bullets(slide, [
        "学者画像：姓名、机构、研究领域、代表作",
        "作者同名消歧：机构、合作网络、代表论文交叉判断",
        "机构/团队分析：研究方向、合作关系、影响力",
        "专利和应用成果：适合管理、HCI、神经营销、BCI 交叉研究",
    ], 0.85, 1.55, 5.85, 2.8, size=17)
    add_code(slide, "{\n  \"mcpServers\": {\n    \"aminer\": {\n      \"url\": \"https://mcp.aminer.cn/sse\",\n      \"headers\": {\n        \"Authorization\": \"Bearer ${AMINER_MCP_TOKEN}\"\n      }\n    }\n  }\n}", 7.0, 1.45, 5.5, 3.05, size=11)
    add_card(slide, "人工核验", "中文姓名同名非常常见，不能把第一条结果直接当作目标学者。必须用机构、领域、代表作和主页交叉确认。", 0.9, 5.0, 11.5, 1.0, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=13)

    slide = new_slide("Ai4Scholar 与 AMiner 怎么分工")
    rows = [
        ["问题", "优先工具", "原因"],
        ["这篇论文是否真实、DOI 是否对", "Ai4Scholar / OpenAlex / Crossref", "论文级元数据核验"],
        ["某个学者是谁、属于哪个机构", "AMiner", "学者画像和知识图谱"],
        ["一个方向有哪些核心论文", "Ai4Scholar + AMiner", "论文检索 + 专家/团队补充"],
        ["文献池如何进入 Zotero", "Zotero / pyzotero", "长期知识库管理"],
        ["结果如何写进论文", "paper-writing + citation-management", "写作和引用格式"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 4.35, col_widths=[4.0, 3.45, 4.6], font_size=11)
    add_text(slide, "科研检索不是一个工具通吃，而是多平台交叉核验。", 0.9, 6.3, 11.4, 0.4, size=17, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("飞书 / IMA 渠道连接")
    add_flow(slide, ["飞书/IMA", "Agent 网关", "模型", "Skills/MCP", "本地/云端输出"], 1.0, 1.55, 11.3, box_h=0.72, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=12)
    rows = [
        ["连接对象", "科研用途"],
        ["飞书文档", "讲义、论文草稿、会议纪要"],
        ["飞书知识库", "课程资料、论文卡片、方法手册"],
        ["飞书多维表格", "文献筛选、实验进度、任务看板"],
        ["飞书消息", "触发 Agent、接收摘要、推送提醒"],
        ["IMA / 其他渠道", "作为消息或知识库入口，按 Webhook/Bot/API 能力接入"],
    ]
    add_table(slide, rows, 0.85, 2.7, 5.95, 3.35, col_widths=[2.0, 3.95], font_size=11)
    add_bullets(slide, [
        "确认组织权限和 scope",
        "密钥只放环境变量",
        "先用低风险测试文档",
        "不要把未公开数据直接发到外部渠道",
        "输出要回写到 research_log.md",
    ], 7.35, 2.85, 4.8, 2.5, size=17)

    slide = new_slide("标准科研工作区结构")
    add_code(slide, "第01讲_我的Agent工作区/\n├─ AGENTS.md\n├─ CLAUDE.md\n├─ MEMORY.md\n├─ SOUL.md\n├─ USER.md\n├─ project_status.md\n├─ research_log.md\n├─ skills/\n├─ input/\n├─ literature/\n├─ data/\n│  ├─ raw/\n│  └─ processed/\n├─ output/\n└─ configs/\n   ├─ mcp.example.json\n   └─ env.example", 0.85, 1.45, 5.65, 5.3, size=12)
    add_bullets(slide, [
        "input：原始输入，不随意覆盖",
        "literature：文献检索、筛选表、BibTeX",
        "data/raw：原始数据；data/processed：清洗后数据",
        "output：Agent 输出文件",
        "configs：只放示例和占位，不放真实密钥",
        "research_log：每次执行都记录",
    ], 7.05, 1.65, 5.25, 3.3, size=17)
    add_card(slide, "课堂标准", "学生提交的是文件夹，不是聊天截图。", 7.05, 5.55, 5.25, 0.85, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("Agent 日常任务部署模板")
    add_code(slide, "# 今日 Agent 任务\n\n## 目标\n形成 20 篇候选文献表。\n\n## 输入\n关键词、时间范围、数据库。\n\n## 可用 Skills\nai4scholar-research, aminer-mcp-research。\n\n## 执行要求\n每篇论文必须有题名、作者、年份、来源、DOI/URL。\n无法核验标注“未核验”。\n\n## 输出\nliterature/candidate_papers.xlsx\nliterature/search_log.md\nliterature/unverified_items.md", 0.85, 1.4, 6.0, 5.35, size=11)
    add_bullets(slide, [
        "任务卡把需求转成可执行流程",
        "先定义输入和输出，再让 Agent 执行",
        "所有未核验项单独成文件",
        "执行后更新 project_status.md 和 research_log.md",
    ], 7.25, 1.85, 4.95, 2.35, size=18)
    add_card(slide, "课堂要求", "以后每次用 Agent 做科研，都先写任务卡。", 7.25, 4.85, 4.95, 0.95, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)

    slide = new_slide("课堂实操 1：建立工作区")
    add_code(slide, "New-Item -ItemType Directory -Force -Path \"D:\\AI科研训练营\\第01讲_我的Agent工作区\"\ncd \"D:\\AI科研训练营\\第01讲_我的Agent工作区\"\n\nNew-Item -ItemType Directory -Force -Path skills,input,literature,output,configs,data,docs\nNew-Item -ItemType Directory -Force -Path data\\raw,data\\processed\nNew-Item -ItemType File -Force -Path AGENTS.md,CLAUDE.md,SOUL.md,USER.md,MEMORY.md,project_status.md,research_log.md", 0.8, 1.55, 12.0, 2.5, size=12)
    add_bullets(slide, [
        "先建结构，再让 Agent 工作",
        "不要把所有输出堆在桌面",
        "不要把 Skills 直接装到根目录",
        "每个学生用自己的姓名/主题命名项目",
    ], 1.05, 4.55, 11.2, 1.6, size=18)

    slide = new_slide("课堂实操 2：复制本讲 Skills")
    add_code(slide, "cd \"D:\\desk\\AI agent科研资料\\本地Skills功能分类库\"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName ai4scholar-mcp-openclaw-setup,aminer-mcp-research,find-skills,skill-creator `\n  -Workspace \"D:\\AI科研训练营\\第01讲_我的Agent工作区\"\n\nGet-ChildItem \"D:\\AI科研训练营\\第01讲_我的Agent工作区\\skills\" -Directory", 0.8, 1.55, 12.0, 2.6, size=12)
    add_card(slide, "检查点", "看到 skills 文件夹下出现对应 Skill 目录，且每个目录中至少有 SKILL.md。", 1.0, 4.7, 11.15, 0.9, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=14)
    add_text(slide, "如果失败，先检查路径空格、PowerShell 执行策略和 Skill 名称。", 1.0, 6.0, 11.2, 0.35, size=15, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("课堂实操 3：写 MCP 示例配置")
    add_code(slide, "{\n  \"mcpServers\": {\n    \"ai4scholar\": {\n      \"url\": \"https://mcp.ai4scholar.net/sse\",\n      \"headers\": {\"Authorization\": \"Bearer ${AI4SCHOLAR_API_KEY}\"}\n    },\n    \"aminer\": {\n      \"url\": \"https://mcp.aminer.cn/sse\",\n      \"headers\": {\"Authorization\": \"Bearer ${AMINER_MCP_TOKEN}\"}\n    }\n  }\n}", 0.85, 1.45, 6.25, 4.55, size=12)
    add_bullets(slide, [
        "文件名：configs/mcp.example.json",
        "真实 Key 只在本机环境变量中配置",
        "配置后重启 Agent 或 MCP gateway",
        "第一次测试只跑低成本查询",
        "失败时输出排查清单，不编造结果",
    ], 7.45, 1.75, 4.6, 2.8, size=17)

    slide = new_slide("课堂实操 4：第一次低风险测试")
    add_code(slide, "请检查当前工作区的 skills 文件夹，确认是否有 ai4scholar-mcp-openclaw-setup 和 aminer-mcp-research。\n不要读取或输出任何真实 API Key。\n根据 configs/mcp.example.json 生成 MCP 配置核验表。\n如果无法连接 MCP，请不要编造检索结果，只输出下一步排查清单。", 0.85, 1.55, 11.8, 2.35, size=13)
    add_table(slide, [
        ["输出文件", "必须包含"],
        ["output/first_agent_test.md", "已发现 Skills、MCP 配置状态、需要人工补充的 Token、下一步测试提示词、未核验项列表"],
        ["research_log.md", "执行时间、Agent、输入、输出路径、人工核验结果"],
    ], 0.85, 4.35, 11.8, 1.6, col_widths=[3.0, 8.8], font_size=11)
    add_text(slide, "课堂验收只看可复现文件，不看“我问过 AI 了”。", 0.9, 6.35, 11.4, 0.35, size=16, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("必须人工核验的地方")
    rows = [
        ["类型", "人工核验点"],
        ["密钥", "是否泄露；权限是否过大；是否写入共享文件"],
        ["模型", "费用、隐私、长上下文、工具调用是否满足任务"],
        ["文献", "标题、作者、年份、期刊、DOI/URL 是否真实匹配"],
        ["AMiner", "同名作者是否误配；机构和代表作是否一致"],
        ["飞书/IMA", "组织权限、文档权限、消息范围、日志留存"],
        ["结论", "是否把相关写成因果；是否夸大贡献"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.55, col_widths=[2.2, 9.65], font_size=12)
    add_text(slide, "Agent 能提高效率，但不能承担作者责任。", 1.0, 6.45, 11.2, 0.35, size=17, color=COLORS["amber"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("课后作业")
    add_bullets(slide, [
        "提交：姓名_第01讲_Agent科研工作区",
        "必须包含：AGENTS.md、CLAUDE.md、SOUL.md、USER.md、MEMORY.md",
        "必须包含：configs/mcp.example.json、skills/、research_log.md",
        "完成一次 Ai4Scholar 或 AMiner 低风险测试",
        "提交 300 字反思：你更适合先用哪个 Agent，为什么",
    ], 0.9, 1.55, 6.0, 3.1, size=18)
    add_table(slide, [
        ["评分点", "分值"],
        ["工作区结构", "20"],
        ["记忆文件质量", "20"],
        ["Skills 部署", "15"],
        ["MCP 配置", "15"],
        ["测试记录", "20"],
        ["反思", "10"],
    ], 7.35, 1.65, 4.55, 3.25, col_widths=[3.2, 1.35], font_size=12)
    add_card(slide, "不合格情况", "出现真实 API Key、编造文献、只交聊天截图、没有 research_log.md。", 1.0, 5.65, 11.15, 0.9, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=14)

    slide = new_slide("课程总路线：按功能区逐讲推进")
    rows = [
        ["阶段", "讲次", "核心产出"],
        ["基础设施", "01", "Agent 工作区、Skills、MCP、渠道连接"],
        ["文献与知识库", "02-05", "论文卡片、文献池、Zotero/Obsidian、研究问题"],
        ["方法设计", "06-09", "CNKI 选题、情景实验、统计因果、文本挖掘"],
        ["脑电与 BCI", "10-11", "EEG/ERP 分析、BCI 智能建模、防泄漏评估"],
        ["论文与传播", "12-13", "论文写作审稿排版、PPT/海报/图表"],
        ["综合项目", "14", "可复现科研 Agent 项目包"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.55, col_widths=[2.4, 1.5, 7.95], font_size=12)
    add_text(slide, "第一讲是地基；后面每一讲都在同一个工作区规范上增加能力。", 1.0, 6.45, 11.2, 0.35, size=17, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("外部资料入口")
    add_bullets(slide, [
        "Claude Code：https://docs.anthropic.com/en/docs/claude-code/overview",
        "Claude Code memory：https://docs.anthropic.com/en/docs/claude-code/memory",
        "OpenAI Codex：https://github.com/openai/codex",
        "MCP 官方文档：https://modelcontextprotocol.io/docs",
        "Ai4Scholar：https://ai4scholar.net；MCP：https://mcp.ai4scholar.net/sse",
        "AMiner 开放文档：https://www.aminer.cn/open/docs；MCP：https://mcp.aminer.cn/sse",
        "模型选择参考：https://hvoy.ai/model-select",
    ], 0.85, 1.55, 11.8, 3.4, size=15)
    add_card(slide, "版本原则", "部署命令会随工具版本变化。课件中涉及外部工具的安装步骤，学生必须以官方当前文档或课程包当前版本再次确认。", 0.9, 5.35, 11.45, 1.0, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("本讲一句话总结")
    add_text(slide, "科研 Agent 不是一个聊天窗口，\n而是一套可追踪的工作区、记忆、Skills、MCP、模型和人工核验流程。", 1.2, 2.35, 10.9, 1.4, size=30, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
    add_flow(slide, ["输入明确", "工具明确", "输出明确", "核验明确", "日志明确"], 1.35, 4.55, 10.55, box_h=0.75, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=13)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
