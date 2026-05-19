from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第08讲_课件.pptx"
W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}
FONT = "Microsoft YaHei"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, level = item if isinstance(item, tuple) else (item, 0)
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=9.3):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 08 讲 · 高级统计与因果推断", 0.45, 0.055, 3.2, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12.1, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])
    return rect


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.2, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
    return shape


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=10):
    gap = 0.12
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.01), Inches(y + h / 2), Inches(bx - 0.02), Inches(y + h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"])
    no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.16, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.18, accent=COLORS["amber"], body_size=11.2, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.72, size=8.7)
    textbox(slide, "人工核验", 4.65, 4.82, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.4)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 08 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "数据分析\n高级统计\n因果推断", 0.52, 1.70, 2.25, 1.65, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "从结果到可发表解释", 0.55, 6.10, 2.2, 0.35, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "变量字典、模型选择、机制检验、因果识别与结果汇报", 3.55, 1.10, 8.75, 0.72, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "Statistics · PROCESS · SEM/CFA · Multilevel · DID/PSM/IV/RDD/DML · Tables · Claims Audit", 3.58, 1.95, 8.75, 0.35, size=13.5, color=COLORS["muted"])
    flow(slide, ["数据", "变量", "假设", "模型", "诊断", "结果表", "论文解释", "因果边界"], 3.60, 3.10, 8.85, h=0.62, size=9.4)
    card(slide, "本讲交付", "variable_dictionary、hypothesis_model_map、model_selection、causal_identification_audit、analysis_plan、results_tables、result_writeup、statistical_claims_audit。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=12)
    card(slide, "底线", "不改数值、不伪造显著性、不隐藏非显著结果、不把关联写成因果。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, [
        "会跑回归，但不知道模型是否匹配研究问题",
        "会做中介/调节，但不知道什么能说成机制",
        "看到显著结果就想写因果，忽略设计边界",
        "不知道 SEM/CFA、多层模型、DID、IV、RDD、DML 的适用条件",
        "结果表、结果段落和附录无法追踪到原始模型输出",
        "稳健性检验、模型诊断、因果识别被混在一起",
    ], 0.85, 1.60, 5.95, 3.12, size=15.5)
    card(slide, "本讲不做", "不教学生追求“高级模型感”；不根据显著性临时换模型；不让 Agent 猜统计结果。", 7.05, 1.72, 5.10, 1.22, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "建立从数据、假设、模型、诊断、结果表到论文解释的一条可审计链路。", 7.05, 3.35, 5.10, 1.22, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["变量字典", "模型选择", "机制检验", "识别审计", "结果表", "结果段落", "声称边界"], 0.95, 5.65, 11.35, h=0.62, size=10.2)

    slide = new("本讲最终产出文件")
    rows = [
        ["阶段", "文件", "作用"],
        ["变量", "analysis/variable_dictionary.xlsx", "变量名、类型、角色、编码、来源、缺失和转换规则"],
        ["数据审计", "analysis/data_audit.md", "样本量、缺失、异常、重复、排除和版本记录"],
        ["假设映射", "analysis/hypothesis_model_map.xlsx", "每个假设对应 X/Y/M/W、模型、诊断和表格"],
        ["模型选择", "analysis/model_selection.md", "说明为什么用回归、PROCESS、SEM、多层或因果模型"],
        ["因果审计", "analysis/causal_identification_audit.md", "DAG、混淆、坏控制、识别假设和可说边界"],
        ["结果交付", "analysis/results_tables.xlsx / result_writeup.md", "结果表、图形计划和论文段落"],
        ["质量控制", "qc/statistical_claims_audit.md", "检查数值一致、显著性、效应量、因果语言和遗漏结果"],
    ]
    table(slide, rows, 0.70, 1.46, 11.95, 4.98, col_widths=[1.55, 4.95, 5.45], font_size=8.9)
    card(slide, "提交底线", "任何数值都必须来自真实模型输出；任何因果表述都必须有设计或识别假设支撑。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [
        ["功能", "Skills", "用途"],
        ["基础统计", "statistical-analysis", "t 检验、ANOVA、回归、假设检验、效应量和 APA 报告"],
        ["机制/边界", "process-mediation-moderation", "中介、调节、调节中介、bootstrap、simple slopes"],
        ["SEM/CFA", "sem-cfa-path-latent", "潜变量、测量模型、结构路径、拟合指标"],
        ["多层纵向", "multilevel-longitudinal-modeling", "嵌套、重复测量、随机效应、中心化"],
        ["因果推断", "did-psm-iv-rdd-dml-event-study", "DID、event study、PSM、IV、RDD、DML"],
        ["因果审计", "causal-inference-design-audit", "DAG、混淆、坏控制、识别假设和声称边界"],
        ["结果表", "statistical-results-tables", "主效应、机制、稳健性、附录表和结果段落"],
        ["脚本/可视化", "empirical-analysis-skill-python / data-visualization-analysis", "分析脚本、诊断图、交互图和事件研究图"],
    ]
    table(slide, rows, 0.68, 1.42, 12.0, 5.05, col_widths=[1.50, 4.9, 5.60], font_size=8.7)
    card(slide, "配套模板", "本讲提供 5 个 Excel 模板：变量字典、假设模型映射、结果表、诊断清单、分析运行日志。", 0.85, 6.50, 11.55, 0.52, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=11.5)

    slide = new("实操 1：创建分析工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第08讲_数据分析高级统计与因果推断工作区"\ncd "D:\\AI科研训练营\\第08讲_数据分析高级统计与因果推断工作区"\n\nNew-Item -ItemType Directory -Force -Path data,analysis,analysis\\scripts,analysis\\figures,qc,output,skills\nNew-Item -ItemType Directory -Force -Path input\\lesson05_outputs,input\\lesson07_outputs,input\\target_journal_tables\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md,analysis\\model_selection.md,analysis\\analysis_plan.md', 0.85, 1.55, 11.65, 2.48, size=9.4)
    card(slide, "检查标准", "工作区至少包含 data、analysis、qc、output、skills；原始数据进入 data/raw，只读备份。", 0.95, 4.42, 5.60, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "输入来源", "第 05 讲提供假设和理论机制；第 07 讲提供实验设计、量表和数据字典草稿。", 6.85, 4.42, 5.60, 1.05, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName statistical-analysis,process-mediation-moderation,sem-cfa-path-latent,multilevel-longitudinal-modeling,did-psm-iv-rdd-dml-event-study,causal-inference-design-audit,statistical-results-tables,empirical-analysis-skill-python,data-visualization-analysis `\n  -Workspace "D:\\AI科研训练营\\第08讲_数据分析高级统计与因果推断工作区"', 0.80, 1.48, 11.75, 2.82, size=8.7)
    bullets(slide, ["复制后检查 skills/ 文件夹。", "Skills 放工作区，不放 Codex 根目录。", "缺失 Skill 写入 research_log.md，并用现有 Skills 临时替代。"], 0.95, 4.75, 11.1, 1.0, size=15)
    card(slide, "课堂提醒", "统计和因果任务可以交给 Agent 起草流程，但模型假设、变量编码和因果声称必须人工核验。", 0.95, 6.05, 11.1, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第08讲工作区 Agent 指令\n\n你是数据分析、高级统计与因果推断助手。\n\n原则：\n- 禁止编造、修改或平滑任何统计数值。\n- 禁止隐藏非显著结果。\n- 不得把横截面相关写成因果。\n- 所有模型选择必须说明数据结构、变量角色和诊断要求。\n- 原始数据只读保存，清洗必须生成新文件。\n- 结果段落必须能追踪到结果表和脚本输出。", 0.85, 1.55, 6.45, 3.50, size=9.6)
    card(slide, "为什么要写", "第8讲最容易发生数值错配、模型错配和因果过度声称。项目指令用来先设底线。", 7.62, 1.70, 4.45, 1.10, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson08_workspace_AGENTS_template.md。", 7.62, 3.20, 4.45, 0.80, accent=COLORS["blue"])
    card(slide, "不能省略", "没有约束时，Agent 容易为了写得好看而替你扩大结论。", 7.62, 4.45, 4.45, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    step_slide(
        prs, 4, "实操 4：建立变量字典",
        "把每个变量的含义、类型、角色、编码、来源、缺失规则和分析用途写清楚。",
        "请读取 data/raw、input/lesson07_outputs 和量表文件，填写 analysis/variable_dictionary.xlsx。\n\n每个变量必须包含 variable_name、label、type、role、coding、source、expected_range、reverse_code、missing_rule、transformation、analysis_use。不要覆盖原始数据；反向题和重编码必须生成新列。",
        "analysis/variable_dictionary.xlsx",
        ["变量名与数据列一致", "IV/DV/M/W/controls 角色清楚", "反向题明确", "缺失和异常规则先验", "原始数据未被覆盖"],
        "变量字典是后面所有模型和结果表的地基。"
    )

    step_slide(
        prs, 5, "实操 5：做数据审计",
        "在建模前检查样本量、缺失、重复、异常、编码、随机化和平衡性。",
        "请基于 variable_dictionary.xlsx 和 data/raw 数据生成 analysis/data_audit.md。\n\n按数据版本、样本量、变量范围、缺失、重复、异常值、反向题、条件编码、随机化平衡、排除规则和待人工核验项写。不得删除原始数据；清洗输出 data/data_clean_v1.csv。",
        "analysis/data_audit.md\n data/data_clean_v1.csv",
        ["样本量流向清楚", "缺失与排除可复核", "条件编码无误", "清洗脚本可复现", "不根据显著性排除样本"],
        "先审计数据，再跑模型。顺序不能反。"
    )

    step_slide(
        prs, 6, "实操 6：假设到模型映射",
        "把每个研究假设映射到变量、模型、诊断、结果表和可说结论。",
        "请读取 hypotheses、causal_model 和 variable_dictionary.xlsx，填写 analysis/hypothesis_model_map.xlsx。\n\n每个假设写明 X、Y、mediator、moderator、controls、design_type、recommended_model、model_reason、causal_language_allowed、required_diagnostics、output_table。",
        "analysis/hypothesis_model_map.xlsx",
        ["每个假设都有模型", "模型理由不是“高级”", "控制变量有理论理由", "诊断要求清楚", "因果语言边界明确"],
        "没有模型映射，结果段落很容易乱写。"
    )

    slide = new("模型选择总表")
    rows = [
        ["问题", "数据结构", "建议模型", "最低核验"],
        ["二组差异", "随机实验 + 连续 DV", "t-test / OLS", "随机化、操纵检验、效应量"],
        ["2×2 或连续调节", "实验/问卷", "Interaction + simple slopes", "交互图、简单斜率、理论边界"],
        ["中介/调节中介", "理论机制明确", "PROCESS-style bootstrap", "时间顺序、bootstrap CI、不能过度因果"],
        ["潜变量路径", "多题项构念", "CFA + SEM", "先测量模型再结构模型"],
        ["嵌套/重复测量", "学生-班级/日记/纵向", "Multilevel / mixed model", "ICC、随机效应、中心化"],
        ["准实验政策冲击", "处理组+对照组+时间", "DID / event study", "平行趋势、聚类 SE、处理时点"],
        ["观察数据因果", "有识别策略", "PSM / IV / RDD / DML", "识别假设、失败边界、稳健性"],
    ]
    table(slide, rows, 0.72, 1.46, 11.95, 4.96, col_widths=[2.25, 3.0, 3.0, 3.7], font_size=8.8)
    card(slide, "判断句", "先问“研究问题和数据结构是什么”，再问“哪个模型最适合”。不要倒过来。", 0.85, 6.50, 11.55, 0.52, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=11.8)

    slide = new("基础统计与回归：不要跳过地基")
    bullets(slide, [
        "描述统计：N、均值、标准差、范围、缺失率。",
        "相关矩阵：标注样本量、显著性规则和变量来源。",
        "t-test / ANOVA：适合实验条件差异和组间比较。",
        "OLS / Logistic：根据 DV 类型选择，报告效应量和置信区间。",
        "稳健标准误或聚类标准误：根据数据结构决定，不是万能修复。",
    ], 0.90, 1.55, 5.85, 2.65, size=14.6)
    card(slide, "Agent 任务", "生成 analysis/scripts/01_descriptives.py、02_main_effects.py 和 results_tables.xlsx 的 descriptives、correlations、main_effects 表。", 7.05, 1.58, 5.10, 1.55, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)
    card(slide, "人工核验", "检查变量类型、编码、异常值、标准误类型、N 是否和数据审计一致。", 7.05, 3.60, 5.10, 1.15, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.2)
    flow(slide, ["描述", "相关", "主效应", "诊断", "效应量", "表格"], 1.00, 5.55, 11.2, h=0.62, size=11)

    slide = new("PROCESS 风格：中介、调节、调节中介")
    rows = [
        ["模型", "回答的问题", "必须报告"],
        ["中介", "X 是否通过 M 与 Y 相关/产生作用", "a、b、c'、间接效应、bootstrap CI"],
        ["调节", "W 是否改变 X-Y 关系", "交互项、simple slopes、交互图"],
        ["调节中介", "间接效应是否随 W 改变", "conditional indirect effects、index of moderated mediation"],
        ["Johnson-Neyman", "连续调节下哪些 W 区间显著", "显著区间和图形"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.25, col_widths=[2.25, 5.10, 4.40], font_size=10)
    card(slide, "边界", "横截面自陈问卷的中介结果只能写“与理论机制一致的间接效应证据”，不能直接写因果机制。", 0.95, 5.10, 11.35, 0.82, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)
    card(slide, "实操输出", "analysis/results_tables.xlsx 的 mediation_moderation 表；analysis/figures/interaction_plot.png；result_writeup.md 中的机制段落。", 0.95, 6.15, 11.35, 0.70, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=11.8)

    slide = new("SEM / CFA：先测量，再结构")
    bullets(slide, [
        "CFA：先确认题项是否测量预期构念。",
        "信效度：loadings、alpha/omega、CR、AVE、HTMT 或区分效度证据。",
        "结构模型：在测量模型可接受后再估计路径。",
        "拟合指标：CFI、TLI、RMSEA、SRMR 需要结合模型和样本解释。",
        "禁止为了拟合好看而无理论理由地加 correlated errors。",
    ], 0.90, 1.55, 5.95, 2.72, size=14.5)
    card(slide, "适合", "多题项潜变量、测量误差重要、需要同时估计多个路径。", 7.05, 1.65, 5.10, 1.00, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "不适合", "样本量不足、构念不是潜变量、普通回归已经能回答问题。", 7.05, 3.02, 5.10, 1.00, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])
    card(slide, "审稿风险", "好拟合不等于因果；CFA 和 SEM 不能弥补差的理论和设计。", 7.05, 4.40, 5.10, 1.00, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    flow(slide, ["题项", "CFA", "信效度", "结构路径", "拟合", "解释"], 1.00, 5.85, 11.2, h=0.62, size=11)

    slide = new("多层 / 纵向模型：处理非独立数据")
    rows = [
        ["数据", "为什么不能普通 OLS", "建模要点"],
        ["学生嵌套班级", "同班学生并非独立", "ICC、随机截距、班级层变量"],
        ["员工嵌套组织", "组织文化/制度产生共享方差", "跨层主效应与跨层交互"],
        ["日记研究", "同一人多天重复测量", "within-person 与 between-person 分开"],
        ["面板数据", "单位随时间重复观测", "固定效应/随机效应、时间趋势"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.35, col_widths=[2.45, 4.75, 4.55], font_size=10)
    card(slide, "关键概念", "中心化不是技术细节。group-mean centering 与 grand-mean centering 对解释不同。", 0.95, 5.20, 5.45, 0.88, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)
    card(slide, "人工核验", "聚类变量、时间变量、重复测量次数、缺失模式、随机斜率是否有理论必要。", 6.75, 5.20, 5.45, 0.88, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)

    slide = new("因果推断地图：先识别，再估计")
    rows = [
        ["方法", "适合场景", "核心假设"],
        ["DID / event study", "处理组与对照组在某时点受到冲击", "平行趋势、无预期效应、处理定义清楚"],
        ["PSM / weighting", "观察数据中处理组和对照组可按可观测变量平衡", "无未观测混淆、common support"],
        ["IV", "有影响处理但不直接影响结果的外生工具", "相关性、排除限制、单调性"],
        ["RDD", "处理由阈值规则决定", "截点附近不可操纵、协变量连续"],
        ["DML", "高维控制或非线性 nuisance model", "识别假设仍成立、样本拆分和交叉拟合"],
    ]
    table(slide, rows, 0.68, 1.45, 12.0, 4.45, col_widths=[2.1, 5.1, 4.8], font_size=9.2)
    card(slide, "判断句", "因果推断不是更复杂的回归，而是回答“为什么这个比较可以代表反事实”。", 0.85, 6.25, 11.55, 0.65, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.3)

    slide = new("DID / Event Study 操作手册")
    flow(slide, ["定义处理", "定义对照", "构造时间", "画事件图", "检验趋势", "估计模型", "做稳健性"], 0.82, 1.45, 11.7, h=0.62, size=9.6)
    codebox(slide, "请基于面板数据和政策/平台冲击说明，生成 DID/event study 分析计划。\n\n必须包含：处理组定义、对照组定义、处理时间、样本窗口、固定效应、聚类标准误、平行趋势检验、事件研究图、错位处理风险、稳健性和可说/不可说结论。", 0.92, 2.45, 6.15, 2.25, size=9.1)
    card(slide, "人工核验", "平行趋势不是写一句话；需要处理前趋势图或估计证据。错位处理时不能盲用传统 TWFE。", 7.35, 2.45, 4.95, 1.15, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)
    card(slide, "输出文件", "analysis/causal_identification_audit.md、analysis/figures/event_study.png、results_tables.xlsx 的 causal_models 表。", 7.35, 4.05, 4.95, 1.15, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)
    card(slide, "结果写法", "“在平行趋势等识别假设成立下，处理与结果变化存在一致证据。”", 0.92, 5.45, 11.38, 0.72, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.2)

    slide = new("PSM / IV / RDD / DML 的使用边界")
    rows = [
        ["方法", "能帮什么", "不能替你解决什么"],
        ["PSM", "让处理组和对照组在可观测变量上更平衡", "不能处理未观测混淆，匹配后仍需结果模型和敏感性分析"],
        ["IV", "在有效工具下识别局部平均处理效应", "不能只因为一阶段强就证明排除限制成立"],
        ["RDD", "在阈值附近识别局部处理效应", "不能推广为全样本平均效应"],
        ["DML", "用机器学习处理高维控制和非线性 nuisance", "不能绕过处理变量的识别假设"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.80, col_widths=[1.9, 4.55, 5.30], font_size=9.7)
    card(slide, "课堂判断", "只要不能说清反事实比较来自哪里，就先降级为关联分析或探索性证据。", 0.95, 5.80, 11.35, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.8)

    slide = new("DAG、混淆与坏控制")
    bullets(slide, [
        "混淆变量：同时影响 X 和 Y，通常需要在设计或模型中处理。",
        "中介变量：在估计总效应时不能随意控制。",
        "Collider：由两个变量共同影响，控制它可能制造虚假关联。",
        "Post-treatment variable：处理发生后的变量，控制它可能改变估计目标。",
        "控制变量不是越多越好，关键看它在因果图中的位置。",
    ], 0.90, 1.55, 5.95, 2.85, size=14.4)
    card(slide, "Agent 任务", "画出 X、Y、M、W、controls、潜在混淆的 DAG 草图，并标注哪些变量能控制、不能控制、需要稳健性。", 7.05, 1.62, 5.10, 1.35, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)
    card(slide, "人工核验", "坏控制通常不是统计软件能发现的，它来自理论和时间顺序判断。", 7.05, 3.42, 5.10, 1.05, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.2)
    flow(slide, ["理论", "时间顺序", "DAG", "控制集", "估计目标", "声称边界"], 1.00, 5.68, 11.2, h=0.62, size=10.6)

    step_slide(
        prs, 7, "实操 7：生成因果识别审计",
        "把研究设计、DAG、混淆、坏控制和识别假设写成可审计文件。",
        "请读取 hypothesis_model_map.xlsx、variable_dictionary.xlsx 和研究设计说明，生成 analysis/causal_identification_audit.md。\n\n必须包括 estimand、DAG 文本说明、混淆变量、坏控制、识别假设、诊断方法、稳健性、哪些结论可说因果、哪些只能说关联。",
        "analysis/causal_identification_audit.md",
        ["识别目标清楚", "不控制 post-treatment mediator", "列出不可检验假设", "失败时降级结论", "不把稳健性当识别"],
        "这是本讲最关键的人工核验环节。"
    )

    step_slide(
        prs, 8, "实操 8：生成分析计划和脚本骨架",
        "把数据处理、模型、诊断、稳健性、图表和输出文件写成可执行计划。",
        "请基于 variable_dictionary、hypothesis_model_map 和 causal_identification_audit，生成 analysis/analysis_plan.md 和 analysis/scripts/analysis_template.py。\n\n脚本只写骨架和可复现流程：读取 data_clean_v1、生成描述统计、主模型、机制模型、诊断图、结果表。禁止写入虚构结果。",
        "analysis/analysis_plan.md\nanalysis/scripts/analysis_template.py",
        ["脚本不覆盖原始数据", "所有输出路径固定", "模型和假设一一对应", "诊断和稳健性写清", "随机种子和数据版本记录"],
        "先让脚本可复现，再谈结果是否好看。"
    )

    slide = new("模型诊断：结果可信度检查")
    rows = [
        ["模型", "诊断", "常见修正"],
        ["OLS", "残差图、异方差、异常点、多重共线性", "稳健 SE、转换、敏感性分析"],
        ["Logistic", "类别不平衡、分离、校准", "报告 OR/AME、使用合适指标"],
        ["PROCESS", "bootstrap CI、交互图、simple slopes", "避免只报显著交互项"],
        ["SEM/CFA", "拟合指标、载荷、CR/AVE/HTMT", "理论支持下修改，不为拟合乱改"],
        ["Multilevel", "ICC、随机效应、中心化、收敛", "简化随机结构或解释中心化"],
        ["Causal", "识别假设诊断、安慰剂、敏感性", "降级因果语言或换设计"],
    ]
    table(slide, rows, 0.72, 1.46, 11.95, 4.72, col_widths=[2.0, 5.15, 4.8], font_size=9.4)
    card(slide, "输出", "model_diagnostics_checklist_template.xlsx 可直接复制到 analysis/ 下作为诊断清单。", 0.85, 6.45, 11.55, 0.58, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)

    step_slide(
        prs, 9, "实操 9：生成结果表",
        "把模型输出整理成描述统计、相关、主效应、机制、SEM、多层、因果和稳健性表。",
        "请使用 statistical-results-tables，基于真实模型输出填写 analysis/results_tables.xlsx。\n\n表格至少包含 descriptives、correlations、main_effects、mediation_moderation、causal_models、robustness、figure_plan、writeup_trace。禁止修改 b、SE、p、N 或隐藏非显著结果。",
        "analysis/results_tables.xlsx",
        ["数值来自脚本输出", "N 与数据审计一致", "非显著结果保留", "CI 和效应量完整", "表格可追踪到模型"],
        "表格不是美化结果的地方，是审计链路。"
    )

    step_slide(
        prs, 10, "实操 10：写结果段落",
        "把结果表转成 APA 或目标期刊风格段落，并解释效应量和边界。",
        "请基于 analysis/results_tables.xlsx 和 result_writeup_template.md 生成 analysis/result_writeup.md。\n\n每段必须写明模型、核心估计值、SE、p、CI、N、效应量解释、是否支持假设、因果语言边界。不要只写“显著”。不支持的假设也要报告。",
        "analysis/result_writeup.md",
        ["数值和表格一致", "报告效应量/CI", "非显著结果透明", "结论不超出设计", "能追踪到表格行"],
        "论文结果段落不是报喜，是准确描述证据。"
    )

    slide = new("示范：AI Agent 解释透明度研究")
    rows = [
        ["模块", "示范内容"],
        ["研究问题", "解释透明度是否提升博士生采纳 AI 推荐文献的意愿"],
        ["主效应", "condition -> adoption_intention；二组实验可用 OLS / t-test"],
        ["机制", "condition -> trust / uncertainty -> adoption_intention；bootstrap indirect effect"],
        ["边界", "任务复杂度或 AI 使用经验作为 moderator；simple slopes"],
        ["因果语言", "若随机分配且操纵有效，可说操纵解释透明度对 DV 的影响"],
        ["不能说", "不能声称所有真实科研场景都成立；不能把未操纵的中介写成强因果机制"],
    ]
    table(slide, rows, 0.78, 1.45, 11.75, 4.95, col_widths=[2.0, 9.75], font_size=9.7)
    card(slide, "最大风险", "解释透明度可能同时提升感知能力、友好度和可靠性，必须在情景材料和混淆检验中处理。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("常见错误与修正")
    rows = [
        ["错误", "后果", "修正"],
        ["变量编码没核验", "方向反了，结论相反", "变量字典 + coding_rules"],
        ["模型为高级而高级", "审稿人质疑必要性", "回到假设和数据结构"],
        ["控制 post-treatment 变量", "估计目标被改变", "DAG 审计控制集"],
        ["只报 p 值", "效应大小不可判断", "补 b/SE/CI/效应量/N"],
        ["隐藏非显著结果", "学术诚信风险", "完整结果表和附录"],
        ["稳健性当因果识别", "因果声称站不住", "区分 robustness 与 identification"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.55, 4.10, 5.10], font_size=9.5)

    step_slide(
        prs, 11, "实操 11：统计声称审计",
        "用审稿人视角检查数值一致性、模型假设、效应量、非显著结果和因果语言。",
        "请读取 analysis 下所有输出，生成 qc/statistical_claims_audit.md。\n\n逐条检查：正文数值是否与表格一致；N 是否一致；是否报告效应量和 CI；是否隐藏非显著结果；是否把关联写成因果；是否把稳健性写成识别；哪些句子必须改写。",
        "qc/statistical_claims_audit.md",
        ["指出具体句子", "列出数值错配", "标注因果过度声称", "给出改写建议", "明确 pass/revise/fail"],
        "这是提交前最后一道闸门。"
    )

    slide = new("课堂 90 分钟带做安排")
    rows = [
        ["时间", "教师带做", "学生产出"],
        ["0-10 分钟", "讲清模型选择和因果边界", "理解本讲输出文件"],
        ["10-25 分钟", "创建工作区、复制 Skills、写 AGENTS.md", "可运行工作区"],
        ["25-45 分钟", "变量字典和数据审计", "variable_dictionary / data_audit"],
        ["45-60 分钟", "假设-模型映射和因果审计", "hypothesis_model_map / causal_audit"],
        ["60-75 分钟", "分析计划和结果表结构", "analysis_plan / results_tables"],
        ["75-90 分钟", "结果段落和声称审计", "result_writeup / claims_audit"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.85, col_widths=[1.8, 5.0, 4.95], font_size=9.5)
    card(slide, "教师重点", "现场纠正三件事：模型错配、坏控制、因果过度声称。", 0.95, 6.50, 11.35, 0.50, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12)

    slide = new("课后提交要求")
    bullets(slide, [
        "提交文件夹：姓名_第08讲_数据分析高级统计与因果推断",
        "必须包含 variable_dictionary、data_audit、hypothesis_model_map、model_selection、causal_identification_audit、analysis_plan、analysis_template、results_tables、result_writeup、statistical_claims_audit",
        "附 300 字反思：你的结果中哪些能说因果，哪些只能说关联，为什么",
        "所有未核验模型假设和数据问题必须保留待核验标记",
    ], 0.90, 1.55, 11.3, 1.75, size=15.0)
    rows = [
        ["评分项", "占比"],
        ["变量字典与数据审计", "20%"],
        ["模型选择与假设映射", "20%"],
        ["机制/因果识别审计", "25%"],
        ["结果表与结果段落", "20%"],
        ["统计声称审计与反思", "15%"],
    ]
    table(slide, rows, 2.15, 3.80, 8.85, 2.20, col_widths=[6.6, 2.25], font_size=11.5)
    card(slide, "最低合格线", "能说清：变量怎么来的、模型为什么选、结果如何解释、哪些话不能说成因果。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.3)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是让学生多会几个模型，而是建立一个不会乱改数值、不会乱选模型、不会乱写因果的分析系统。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["数据", "变量", "假设", "模型", "诊断", "识别", "表格", "写作", "审计"], 1.00, 3.50, 11.20, h=0.70, size=9.5)
    card(slide, "下一步", "学生把第 07 讲实验方案或自己的数据带入第 08 讲流程，先做变量字典和模型映射，再决定是否跑高级模型。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
