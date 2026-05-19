from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"
TEAL = "0D746E"
BLUE = "2563EB"
AMBER = "B55D0E"
RED = "B91C1C"
PALE = "F4F7F8"
WHITE = "FFFFFF"
INK = "202A36"
LINE = "D2D8DE"


def style_sheet(ws, widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    side = Side(style="thin", color=LINE)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
    ws.row_dimensions[1].height = 34
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    else:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18


def make_wb(sheets, path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows, widths in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_sheet(ws, widths)
    wb.save(path)
    load_workbook(path).close()


def add_validations(path):
    wb = load_workbook(path)
    if "variable_dictionary" in wb.sheetnames:
        ws = wb["variable_dictionary"]
        type_dv = DataValidation(type="list", formula1='"continuous,binary,categorical,ordinal,text,date,id"', allow_blank=True)
        role_dv = DataValidation(type="list", formula1='"IV,DV,mediator,moderator,control,cluster,time,treatment,outcome,id,weight,diagnostic"', allow_blank=True)
        yn_dv = DataValidation(type="list", formula1='"yes,no,not_applicable,to_check"', allow_blank=True)
        ws.add_data_validation(type_dv)
        ws.add_data_validation(role_dv)
        ws.add_data_validation(yn_dv)
        type_dv.add("C2:C500")
        role_dv.add("D2:D500")
        yn_dv.add("H2:H500")
    if "hypothesis_model_map" in wb.sheetnames:
        ws = wb["hypothesis_model_map"]
        dv = DataValidation(type="list", formula1='"association,experimental_causal,quasi_experimental_causal,mechanism_only,exploratory,to_check"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("K2:K500")
    if "diagnostics_checklist" in wb.sheetnames:
        ws = wb["diagnostics_checklist"]
        dv = DataValidation(type="list", formula1='"pass,warning,fail,not_checked,not_applicable"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("E2:E500")
    wb.save(path)
    load_workbook(path).close()


variable_sheets = [
    (
        "variable_dictionary",
        [
            "variable_name", "variable_label", "type", "role", "coding_or_anchor", "source",
            "expected_range", "reverse_code", "missing_rule", "transformation", "analysis_use", "notes",
        ],
        [
            ["participant_id", "被试编号", "id", "id", "匿名编号", "问卷平台导出", "unique", "not_applicable", "不得缺失", "none", "merge/check only", "不能作为预测变量"],
            ["condition", "实验条件", "categorical", "IV", "0=低解释透明度; 1=高解释透明度", "随机分配", "0/1", "no", "不得缺失", "factor/dummy", "main effect", "检查随机化是否平衡"],
            ["trust_mean", "信任均值", "continuous", "mediator", "1-7 Likert", "真实量表或课程样例", "1-7", "to_check", "按量表规则处理", "mean scale", "mediation", "反向题需先重编码"],
            ["adoption_intention", "采纳意愿", "continuous", "DV", "1-7 Likert", "问卷题项", "1-7", "no", "按先验规则处理", "mean or single item", "outcome", "报告信度或说明单题测量"],
        ],
        [20, 24, 16, 18, 28, 24, 18, 16, 26, 22, 24, 34],
    ),
    (
        "coding_rules",
        ["variable_name", "raw_code", "recoded_value", "rationale", "audit_status", "notes"],
        [
            ["condition", "低解释透明度", "0", "对照条件", "to_check", "保持原始列，只新增 recoded 列"],
            ["condition", "高解释透明度", "1", "处理条件", "to_check", "不要覆盖原始数据"],
            ["reverse_item_example", "1-7", "8 - raw_value", "反向题重编码", "to_check", "只有确认反向题后才能执行"],
        ],
        [22, 22, 22, 30, 18, 34],
    ),
    (
        "missing_and_outlier_rules",
        ["variable_name", "missing_values", "missing_rate", "outlier_rule", "winsorize_or_exclude", "decision", "reviewer_risk"],
        [
            ["duration_seconds", "空值/极端短时长", "待计算", "低于先验阈值", "exclude if preregistered", "to_check", "不能根据显著性临时调整排除规则"],
            ["attention_check", "答错", "待计算", "答错关键事实题", "exclude if preregistered", "to_check", "排除规则必须先于结果确定"],
            ["continuous_scale", "空值", "待计算", "箱线图/IQR/理论范围", "prefer sensitivity analysis", "to_check", "winsorize 不能替代稳健性解释"],
        ],
        [22, 26, 18, 30, 26, 18, 42],
    ),
    (
        "README",
        ["section", "instruction"],
        [
            ["用途", "把所有变量的角色、编码、来源和处理规则写清楚，避免模型解释时变量含义混乱。"],
            ["底线", "不得覆盖原始数据；所有重编码都应生成新列，并在 coding_rules 中记录。"],
            ["人工核验", "反向题、条件编码、处理组、时间变量、聚类变量、缺失值和排除规则。"],
        ],
        [20, 90],
    ),
]

hypothesis_sheets = [
    (
        "hypothesis_model_map",
        [
            "hypothesis_id", "hypothesis_text", "X", "Y", "mediator", "moderator", "controls",
            "design_type", "recommended_model", "model_reason", "causal_language_allowed",
            "required_diagnostics", "output_table",
        ],
        [
            ["H1", "AI Agent 解释透明度与采纳意愿正相关/存在处理效应", "condition", "adoption_intention", "", "", "age; gender; prior_ai_use", "experiment", "OLS / t-test / ANOVA", "二组随机实验且 DV 连续", "experimental_causal", "随机化平衡; 操纵检验; 残差诊断", "main_effects"],
            ["H2", "信任在解释透明度与采纳意愿之间起中介作用", "condition", "adoption_intention", "trust_mean", "", "age; gender; prior_ai_use", "experiment", "bootstrap mediation", "检验间接效应", "mechanism_only", "时间顺序; 中介与操纵检验区分", "mediation_moderation"],
            ["H3", "任务复杂度调节解释透明度对信任的影响", "condition", "trust_mean", "", "task_complexity", "age; gender; prior_ai_use", "2x2 experiment", "interaction + simple slopes", "检验交互和简单斜率", "experimental_causal", "simple slopes; Johnson-Neyman if continuous W", "mediation_moderation"],
        ],
        [18, 42, 20, 22, 22, 22, 32, 22, 28, 34, 26, 42, 24],
    ),
    (
        "model_selection_decision",
        ["research_question_type", "data_structure", "outcome_type", "recommended_model", "minimum_checks", "what_not_to_claim"],
        [
            ["二组差异", "随机实验", "连续", "t-test / OLS", "随机化平衡; 操纵检验; 效应量", "不要忽略操纵失败或混淆检验"],
            ["多组/2x2", "随机实验", "连续", "ANOVA / OLS with interactions", "planned contrasts; interaction plot", "不要只解释主效应而忽略交互"],
            ["中介/调节", "问卷或实验", "连续", "PROCESS-style bootstrap / OLS", "理论顺序; bootstrap CI; simple slopes", "横截面自陈问卷不能直接说因果机制"],
            ["潜变量模型", "多题项量表", "连续/有序", "CFA + SEM", "测量模型先行; fit indices; CR/AVE/HTMT", "好拟合不证明因果"],
            ["嵌套/重复测量", "学生-班级/日记/面板", "连续/二元", "multilevel / mixed model", "ICC; random effects; centering", "不要把非独立观测当独立样本"],
            ["政策/平台冲击", "处理组+对照组+时间", "连续/二元", "DID / event study", "平行趋势; 处理时点; 聚类 SE", "不要盲用 TWFE，尤其错位处理"],
            ["选择偏差", "观察数据", "连续/二元", "PSM + outcome model / weighting", "匹配平衡; common support", "PSM 本身不等于强因果识别"],
            ["工具变量", "观察数据", "连续/二元", "2SLS / IV", "相关性; 排除限制; 弱工具检验", "不能只因工具显著就认为有效"],
            ["阈值规则", "截点附近样本", "连续/二元", "RDD", "带宽; 密度操纵; 协变量连续", "RDD 结论是局部因果效应"],
            ["高维控制", "观察数据", "连续/二元", "DML", "样本拆分; nuisance model; 识别假设", "DML 不是绕过因果设计的黑箱捷径"],
        ],
        [28, 28, 18, 34, 46, 46],
    ),
    (
        "causal_language_boundary",
        ["evidence_type", "allowed_language", "prohibited_language", "required_caveat"],
        [
            ["横截面问卷", "X 与 Y 显著相关；结果与理论机制一致", "X 导致 Y；M 解释了因果机制", "同源、同时间点测量限制因果解释"],
            ["随机情景实验", "操纵 X 对 Y 产生影响", "所有现实场景都成立", "结论受情景真实性、样本和操纵强度限制"],
            ["DID/event study", "在平行趋势等假设下，处理与结果变化相关/可解释为处理效应", "不检验趋势就声称政策因果效应", "说明平行趋势、错位处理和聚类标准误"],
            ["IV", "在相关性和排除限制成立下识别局部平均处理效应", "工具变量天然有效", "讨论排除限制不可完全检验"],
            ["RDD", "截点附近的局部处理效应", "总体平均因果效应", "强调局部性和带宽敏感性"],
        ],
        [28, 46, 42, 52],
    ),
]

results_sheets = [
    (
        "descriptives",
        ["variable", "N", "mean", "SD", "min", "max", "scale_or_unit", "notes"],
        [["", "", "", "", "", "", "", "填入真实统计量；不要让 Agent 猜测数值"]],
        [24, 12, 12, 12, 12, 12, 24, 52],
    ),
    (
        "correlations",
        ["variable", "mean", "SD", "1", "2", "3", "4", "5", "notes"],
        [["1. condition", "", "", "-", "", "", "", "", "相关矩阵需要标注显著性规则和样本量"]],
        [24, 12, 12, 12, 12, 12, 12, 12, 52],
    ),
    (
        "main_effects",
        ["model", "DV", "predictor", "b_or_beta", "SE", "t_or_z", "p", "CI_low", "CI_high", "N", "controls", "interpretation_boundary"],
        [["M1", "adoption_intention", "condition", "", "", "", "", "", "", "", "age; gender; prior_ai_use", "填真实结果；解释效应量和因果边界"]],
        [14, 24, 24, 14, 12, 12, 12, 14, 14, 12, 32, 46],
    ),
    (
        "mediation_moderation",
        ["hypothesis", "model_type", "path_or_effect", "estimate", "SE", "p", "CI_low", "CI_high", "bootstrap_n", "simple_slope_level", "claim"],
        [["H2", "mediation", "indirect effect X->M->Y", "", "", "", "", "", "5000/10000", "", "间接效应证据；横截面数据需谨慎"]],
        [16, 22, 32, 14, 12, 12, 14, 14, 18, 24, 50],
    ),
    (
        "sem_cfa",
        ["model", "construct_or_path", "loading_or_coef", "SE", "p", "CR", "AVE", "HTMT_or_fit", "decision", "notes"],
        [["CFA", "trust item loadings", "", "", "", "", "", "CFI/TLI/RMSEA/SRMR", "to_check", "先测量模型，后结构模型"]],
        [16, 32, 18, 12, 12, 12, 12, 26, 18, 52],
    ),
    (
        "causal_models",
        ["design", "estimand", "model", "key_assumption", "diagnostic", "estimate", "SE", "p", "CI", "robustness", "claim_boundary"],
        [["DID", "ATT", "event study / DID", "parallel trends", "pre-trend test; event plot", "", "", "", "", "alternative windows; clustered SE", "在假设成立下解释"]],
        [16, 24, 28, 34, 36, 14, 12, 12, 18, 34, 46],
    ),
    (
        "robustness",
        ["check_id", "purpose", "model_change", "sample_change", "result_pattern", "decision", "cannot_solve"],
        [["R1", "检验主结果是否受控制变量影响", "替换控制变量组", "不变", "", "to_check", "稳健性不能替代识别假设"]],
        [14, 34, 34, 28, 28, 18, 44],
    ),
    (
        "figure_plan",
        ["figure_id", "figure_type", "data_source", "message", "variables", "must_show", "risk"],
        [["Fig1", "interaction plot / event study plot", "analysis output", "展示效应方向和不确定性", "X; Y; W/time", "CI/error bars; sample info", "图不能美化或隐藏非显著结果"]],
        [14, 28, 24, 40, 26, 34, 44],
    ),
    (
        "writeup_trace",
        ["paragraph", "statistic_source", "exact_values_used", "claim", "causal_boundary", "human_checked"],
        [["Results H1", "main_effects!row", "b/SE/p/CI/N", "支持/不支持 H1", "因果/关联边界", "no"]],
        [24, 30, 32, 38, 38, 18],
    ),
]

diagnostic_sheets = [
    (
        "diagnostics_checklist",
        ["model", "assumption", "diagnostic", "threshold_or_rule", "result", "action", "notes"],
        [
            ["OLS", "线性/同方差/残差异常", "residual plot; robust SE", "按领域和样本判断", "not_checked", "to_check", "不要只看 p 值"],
            ["Logistic", "事件数/分离/拟合", "class balance; separation check", "避免完全分离", "not_checked", "to_check", "报告 OR 或 AME"],
            ["Mediation", "间接效应", "bootstrap CI", "CI 不跨 0", "not_checked", "to_check", "理论顺序比显著性更重要"],
            ["Moderation", "交互解释", "simple slopes / Johnson-Neyman", "报告条件效应", "not_checked", "to_check", "不要只报交互项"],
            ["SEM/CFA", "测量模型", "fit indices; loading; CR/AVE/HTMT", "按期刊规范", "not_checked", "to_check", "不要乱加相关误差"],
            ["Multilevel", "非独立观测", "ICC; random effects", "ICC 和理论共同判断", "not_checked", "to_check", "注意中心化和跨层解释"],
        ],
        [22, 30, 36, 34, 18, 22, 46],
    ),
    (
        "causal_identification",
        ["design", "core_assumption", "diagnostic_or_test", "pass_criteria", "failure_action", "claim_boundary"],
        [
            ["DID", "平行趋势", "event study pre-trends", "处理前趋势接近", "改设计或降级为关联", "假设成立下的处理效应"],
            ["PSM", "可观测变量平衡", "standardized mean differences", "匹配后显著改善", "调整匹配/权重或不做因果声称", "减少可观测选择偏差"],
            ["IV", "相关性和排除限制", "first-stage F; theory argument", "弱工具风险低", "更换工具或放弃 IV", "LATE 且依赖排除限制"],
            ["RDD", "截点附近随机性", "density test; covariate continuity", "无明显操纵", "改变带宽/重新定义设计", "截点附近局部效应"],
            ["DML", "识别假设 + 样本拆分", "cross-fitting; sensitivity", "流程可复现", "回到因果设计审计", "不能绕过混淆问题"],
        ],
        [18, 30, 34, 28, 36, 42],
    ),
    (
        "reporting_checklist",
        ["item", "question", "status", "evidence_file", "fix"],
        [
            ["数值一致", "正文、表格、图中的 b/SE/p/N 是否一致？", "not_checked", "", "统一从 results_tables.xlsx 读取"],
            ["效应量", "是否报告效应大小和置信区间？", "not_checked", "", "补充效应量和 CI"],
            ["因果边界", "是否把关联结果写成因果？", "not_checked", "", "改写为关联或增加识别假设"],
            ["非显著结果", "是否隐藏了不显著结果？", "not_checked", "", "保留并解释"],
            ["稳健性", "是否把稳健性误写为识别？", "not_checked", "", "区分 robustness 与 identification"],
        ],
        [22, 58, 18, 32, 48],
    ),
]

run_log_sheets = [
    (
        "analysis_run_log",
        ["run_id", "date", "script", "data_version", "model", "purpose", "key_output", "decision", "next_step"],
        [
            ["run_001", "YYYY-MM-DD", "analysis/scripts/01_descriptives.py", "data_clean_v1.csv", "descriptives", "检查样本和变量范围", "待填写", "to_check", "完成变量字典后运行"],
            ["run_002", "YYYY-MM-DD", "analysis/scripts/02_main_effects.py", "data_clean_v1.csv", "OLS", "检验 H1", "待填写", "to_check", "生成结果表"],
        ],
        [16, 16, 36, 26, 24, 34, 34, 18, 36],
    ),
    (
        "file_lineage",
        ["file", "source", "created_by", "created_at", "modified_from", "do_not_overwrite", "notes"],
        [
            ["data/raw_data.csv", "问卷平台/数据库原始导出", "human", "YYYY-MM-DD", "none", "yes", "原始数据只读保存"],
            ["data/data_clean_v1.csv", "data/raw_data.csv", "script", "YYYY-MM-DD", "清洗脚本", "no", "保留清洗日志"],
            ["analysis/results_tables.xlsx", "model outputs", "agent+human", "YYYY-MM-DD", "真实模型输出", "no", "数值需人工核验"],
        ],
        [34, 34, 18, 18, 34, 18, 42],
    ),
]


outputs = [
    ("variable_dictionary_template.xlsx", variable_sheets),
    ("hypothesis_model_map_template.xlsx", hypothesis_sheets),
    ("results_tables_template.xlsx", results_sheets),
    ("model_diagnostics_checklist_template.xlsx", diagnostic_sheets),
    ("analysis_run_log_template.xlsx", run_log_sheets),
]

for filename, sheets in outputs:
    path = ROOT / filename
    make_wb(sheets, path)
    add_validations(path)
    print(path)
