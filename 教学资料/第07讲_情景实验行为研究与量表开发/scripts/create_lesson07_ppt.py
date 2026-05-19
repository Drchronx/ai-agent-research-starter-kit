from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第07讲_课件.pptx"
W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}
FONT = "Microsoft YaHei"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=9.4):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 07 讲 · 情景实验与量表", 0.45, 0.055, 2.6, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12.2, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])
    return rect


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.3, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
    return shape


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=10):
    gap = 0.13
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.01), Inches(y + h / 2), Inches(bx - 0.02), Inches(y + h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"]); no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.15, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.15, accent=COLORS["amber"], body_size=11.4, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.70, size=9.0)
    textbox(slide, "人工核验", 4.65, 4.80, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.5)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"]); no_line(left)
    textbox(slide, "第 07 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "情景实验\n行为研究\n量表开发", 0.52, 1.70, 2.25, 1.60, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "从假设到可执行实验", 0.55, 6.10, 2.0, 0.35, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "从顶刊实验范式到 2 组 / 2×2 情景实验方案", 3.55, 1.12, 8.75, 0.66, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "Benchmark · Stimuli · Checks · Scales · Pretest · Analysis · Reporting", 3.58, 1.95, 8.75, 0.35, size=14.5, color=COLORS["muted"])
    flow(slide, ["因果结构", "顶刊范式", "情景材料", "检查项", "量表", "预测试", "报告"], 3.60, 3.10, 8.85, h=0.62, size=10.2)
    card(slide, "本讲交付", "causal_model、benchmark_matrix、stimuli、condition_difference_table、checks、scales、pretest_plan、procedure、analysis_plan、method_results_template、design_audit。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=12.2)
    card(slide, "底线", "不编造顶刊论文和量表来源；不把问卷包装成实验；不让操纵同时改变多个构念。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, [
        "研究假设有了，但不知道如何实验操纵",
        "情景材料像作文，不能形成干净因果对比",
        "操纵同时改变多个构念，导致混淆",
        "不会设计操纵、现实感、注意力和混淆检验",
        "量表来源不清楚，翻译后误当成验证",
        "不会规划预测试、主实验、分析和报告",
    ], 0.85, 1.60, 5.90, 3.10, size=16)
    card(slide, "本讲不做", "不把问卷包装成实验；不追求复杂模型；不凭空引用顶刊和量表。", 7.05, 1.72, 5.10, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "把研究问题转成可操纵、可测量、可预测试、可分析、可报告的实验系统。", 7.05, 3.35, 5.10, 1.20, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["假设", "操纵", "检查", "量表", "流程", "分析", "报告"], 0.95, 5.65, 11.35, h=0.62, size=11.5)

    slide = new("本讲最终产出文件")
    rows = [
        ["阶段", "文件", "作用"],
        ["因果结构", "experiment/causal_model.md", "IV、DV、中介、调节、竞争解释"],
        ["顶刊范式", "experiment/benchmark_matrix.xlsx", "抽取真实论文的实验结构"],
        ["情景材料", "experiment/stimuli.md", "写各条件材料和操纵句"],
        ["差异控制", "experiment/condition_difference_table.xlsx", "检查各条件只改变目标操纵"],
        ["检查项", "experiment/checks_and_attention.md", "操纵、现实感、注意力、混淆、怀疑检验"],
        ["量表", "experiment/scales.xlsx", "真实量表来源、题项、锚点和适配"],
        ["分析报告", "analysis/analysis_plan.md / reporting/method_results_template.md", "模型、效应量、结果报告和附录"],
    ]
    table(slide, rows, 0.72, 1.48, 11.95, 4.85, col_widths=[1.55, 4.7, 5.7], font_size=9.2)
    card(slide, "提交底线", "任何量表和顶刊样例都必须可核验；情景实验必须有预测试计划和设计审计。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [
        ["功能", "Skills", "用途"],
        ["顶刊范式", "scenario-experiment-benchmark-mining", "抽取真实顶刊实验结构"],
        ["实验设计", "scenario-experiment-design", "设计 2 组、2×2、中介、调节实验"],
        ["数据分析", "scenario-experiment-analysis", "操纵检验、主效应、中介/调节"],
        ["结果汇报", "scenario-experiment-reporting", "Method、Results、表格、附录"],
        ["量表选择", "scale-selection-adaptation", "真实量表、翻译回译、题项适配"],
        ["信效度", "scale-reliability-validity / efa-cfa-measurement-model", "alpha、omega、CR、AVE、CFA"],
        ["偏差审计", "common-method-bias-invariance", "共同方法偏差和测量不变性"],
    ]
    table(slide, rows, 0.70, 1.48, 12.0, 4.85, col_widths=[1.55, 4.8, 5.65], font_size=9.1)
    card(slide, "配套模板", "本讲已提供 benchmark、condition difference、scales、data dictionary 四个 Excel 模板。", 0.85, 6.45, 11.55, 0.56, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=11.5)

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第07讲_情景实验行为研究与量表开发工作区"\ncd "D:\\AI科研训练营\\第07讲_情景实验行为研究与量表开发工作区"\n\nNew-Item -ItemType Directory -Force -Path input,experiment,analysis,reporting,qc,output,skills\nNew-Item -ItemType Directory -Force -Path input\\lesson05_outputs,input\\lesson06_outputs,input\\target_journal_examples\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md,input\\initial_experiment_idea.md', 0.85, 1.55, 11.65, 2.50, size=10)
    card(slide, "检查标准", "有 input、experiment、analysis、reporting、qc、skills 文件夹；有 AGENTS.md、research_log.md 和初始实验想法。", 0.95, 4.45, 5.60, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "输入来源", "第 05 讲提供理论机制和假设；第 06 讲提供目标语境和投稿版本。", 6.85, 4.45, 5.60, 1.05, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName scenario-experiment-benchmark-mining,scenario-experiment-design,scenario-experiment-analysis,scenario-experiment-reporting,scale-selection-adaptation,scale-reliability-validity,efa-cfa-measurement-model,common-method-bias-invariance,questionnaire-reporting-template,statistical-analysis,causal-inference-design-audit `\n  -Workspace "D:\\AI科研训练营\\第07讲_情景实验行为研究与量表开发工作区"', 0.80, 1.50, 11.75, 2.80, size=9.1)
    bullets(slide, ["复制后检查 skills/ 文件夹。", "缺失 Skill 写入 research_log.md。", "Skills 放当前工作区，不放 Codex 根目录。"], 0.95, 4.75, 11.1, 1.0, size=15)
    card(slide, "课堂提醒", "实验设计和量表来源需要真实文献支持，缺失时先标待核验，不让 Agent 补假来源。", 0.95, 6.05, 11.1, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第07讲工作区 Agent 指令\n\n你是情景实验、行为研究与量表开发助手。\n\n原则：\n- 禁止编造顶刊论文、样本量、DOI、量表来源、信度或效度。\n- 情景操纵必须对应理论构念。\n- 各条件只改变目标操纵，不引入能力、风险、成本、隐私等混淆。\n- 操纵检验、现实感检验、注意力检验、混淆检验和怀疑检验要分开。\n- 翻译量表不等于验证量表。\n- 所有伦理风险都要说明。", 0.85, 1.55, 6.45, 3.50, size=9.8)
    card(slide, "为什么要写", "本讲最容易发生“情景不干净”和“量表来源不实”。项目指令用来提前设边界。", 7.62, 1.70, 4.45, 1.10, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson07_workspace_AGENTS_template.md。", 7.62, 3.20, 4.45, 0.80, accent=COLORS["blue"])
    card(slide, "不能省略", "没有约束时，Agent 容易把普通问卷写成实验。", 7.62, 4.45, 4.45, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    step_slide(
        prs, 4, "实操 4：生成因果结构",
        "从研究问题中抽出可实验化的 IV、DV、中介、调节和竞争解释。",
        "请读取 input/lesson05_outputs、input/lesson06_outputs 和 input/initial_experiment_idea.md，生成 experiment/causal_model.md。\n\n提取研究问题、IV、DV、中介、调节、竞争解释、目标被试、单位随机化和理论机制。指出哪些关系能用情景实验检验，哪些不能。",
        "experiment/causal_model.md",
        ["IV 能被操纵", "DV 能被测量", "中介不是操纵检验重复", "有竞争解释", "标注不能实验化的关系"],
        "实验不是从写情景开始，而是从因果结构开始。"
    )

    slide = new("情景实验不是问卷包装")
    card(slide, "弱设计", "让被试看一个 AI Agent 的介绍，然后问信任、满意度、使用意愿。", 0.90, 1.55, 5.35, 1.00, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=13)
    card(slide, "问题", "没有明确操纵；不知道哪部分导致结果变化；可能同时改变透明度、能力、风险、亲和力和新颖性。", 0.90, 2.95, 5.35, 1.25, accent=COLORS["red"], body_size=12.3)
    card(slide, "较好设计", "随机分配到高解释透明度 vs 低解释透明度。两组保持任务、角色、成本、结果、系统能力一致，只改变 AI 是否解释推荐依据。", 6.70, 1.55, 5.65, 1.55, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.2)
    card(slide, "关键判断", "能否说清：操纵改变了哪个构念？哪些替代解释被控制？", 6.70, 3.55, 5.65, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])
    flow(slide, ["操纵", "构念", "机制", "测量", "分析"], 1.00, 5.45, 11.20, h=0.65, size=12)

    step_slide(
        prs, 5, "实操 5：抽取顶刊 benchmark",
        "从真实顶刊论文中学习情景结构、操纵、预测试、检查项和报告方式。",
        "请基于 input/target_journal_examples 中的真实论文，使用 scenario-experiment-benchmark-mining，生成 experiment/benchmark_matrix.xlsx 的内容草稿。\n\n提取论文信息、实验类型、操纵、样本、情景结构、操纵检验、现实感检验、注意力检验、混淆检验、排除规则、分析模型、报告风格和可复用设计经验。无法确认字段写 not verified。",
        "experiment/benchmark_matrix.xlsx",
        ["每篇论文真实可查", "确实包含情景实验", "没有编造 DOI/N/结果", "提取的是可复用设计"],
        "不要凭想象设计实验。先看真实论文怎么做。"
    )

    slide = new("选择实验类型")
    rows = [
        ["类型", "适用场景", "例子"],
        ["两组 between-subjects", "一个清晰因果对比", "AI 解释高 vs 低"],
        ["2×2 between-subjects", "两个 IV 或 IV×边界", "透明度 × 任务复杂度"],
        ["中介实验", "机制是核心贡献", "解释透明度 → 信任 → 采纳"],
        ["调节中介", "机制和边界都有理论必要", "透明度经信任影响采纳，受专业经验调节"],
        ["选择实验", "DV 是选择、偏好或权衡", "选择人类专家 vs AI Agent"],
        ["多研究包", "贡献需要三角验证", "预测试 + 主实验 + 行为跟进"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.75, col_widths=[2.3, 4.65, 4.80], font_size=9.6)
    card(slide, "原则", "选择最简洁、最能检验理论的设计。不要为了复杂而做 2×2 或调节中介。", 0.95, 6.45, 11.35, 0.55, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)

    step_slide(
        prs, 6, "实操 6：写情景材料",
        "生成各实验条件材料，并确保条件之间只改变目标操纵。",
        "请基于 experiment/causal_model.md，生成 experiment/stimuli.md。\n\n写出所有条件情景材料；各条件长度尽量一致；保持角色、任务、背景、风险、成本、结果和语气一致；只改变目标操纵；标注每个操纵句对应的构念；不要让材料直接暴露假设。",
        "experiment/stimuli.md",
        ["条件只改变目标操纵", "非操纵元素一致", "不额外改变能力/风险/隐私", "语言符合被试经验", "不暴露假设"],
        "情景材料写得漂亮不够，关键是操纵干净。"
    )

    slide = new("条件差异表")
    rows = [
        ["字段", "作用"],
        ["condition_id / name", "区分实验条件"],
        ["manipulated_construct", "说明操纵构念"],
        ["exact_manipulated_text", "记录具体操纵句"],
        ["constant_elements", "哪些元素保持一致"],
        ["potential_confound", "可能引入的混淆"],
        ["confound_control", "如何控制混淆"],
        ["expected_check_direction", "操纵检验预期方向"],
    ]
    table(slide, rows, 0.80, 1.55, 11.75, 4.45, col_widths=[3.25, 8.50], font_size=10)
    card(slide, "模板文件", "condition_difference_table_template.xlsx 已生成。学生复制到 experiment/ 下填写。", 0.95, 6.35, 11.35, 0.60, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.2)

    step_slide(
        prs, 7, "实操 7：设计检查项",
        "设计操纵检验、现实感、清晰度、注意力、混淆和怀疑检验。",
        "请基于 experiment/stimuli.md，生成 experiment/checks_and_attention.md。\n\n设计操纵检验、现实感检验、清晰度检验、注意力检验、混淆检验、怀疑检验和排除规则。操纵检验测目标操纵，不要和中介或因变量措辞高度重叠。混淆检验覆盖能力、风险、成本、友好度、任务难度、隐私等替代解释。",
        "experiment/checks_and_attention.md",
        ["操纵检验不等于中介", "有现实感和清晰度", "有混淆检验", "排除规则提前确定"],
        "操纵检验显著不能补救混淆操纵。"
    )

    slide = new("检查项的分工")
    rows = [
        ["检查项", "回答的问题", "常见错误"],
        ["操纵检验", "被试是否感知到目标操纵", "题项和中介/DV 高度重复"],
        ["现实感检验", "情景是否可信", "只问“是否认真阅读”"],
        ["注意力检验", "是否理解关键事实", "设置过难或诱导性题目"],
        ["混淆检验", "替代解释是否被改变", "遗漏能力、风险、隐私、成本"],
        ["怀疑检验", "是否猜到研究目的", "只在结果不好时才看"],
    ]
    table(slide, rows, 0.78, 1.55, 11.75, 4.20, col_widths=[2.0, 4.7, 5.05], font_size=10)
    card(slide, "底线", "检查项要分开设计。操纵检验不是中介测量，也不是因变量测量。", 0.95, 6.25, 11.35, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(
        prs, 8, "实操 8：选择和适配量表",
        "为每个构念选择真实量表，记录来源、题项、锚点、翻译和风险。",
        "请基于 experiment/causal_model.md，使用 scale-selection-adaptation，生成 experiment/scales.xlsx 的内容草稿。\n\n每个构念包括定义、量表来源、来源核验状态、原始题项、翻译题项、锚点、反向题、适配说明、版权说明、既有信度、预测试计划和测量风险。不要编造量表来源或信度。",
        "experiment/scales.xlsx",
        ["量表来源真实", "有原始题项和锚点", "翻译回译有计划", "反向题编码明确", "不把翻译当验证"],
        "找题项不是量表开发。翻译也不是验证。"
    )

    slide = new("量表开发与心理测量底线")
    bullets(slide, [
        "不要编造 scale source 或 reliability。",
        "不要混用不同构念题项，除非有理论理由。",
        "不要因为 alpha 提升就随意删题。",
        "不要用同一批数据把 EFA 和 CFA 当成独立验证。",
        "不要说 Harman 单因子检验完全解决共同方法偏差。",
    ], 0.90, 1.60, 5.85, 2.6, size=15.5)
    card(slide, "量表表必须有", "构念定义、量表来源、原始题项、中文翻译、锚点、反向题、适配说明、版权或使用限制、既有信度、预测试计划。", 7.05, 1.70, 5.10, 1.55, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "后续分析", "alpha / omega / CR / AVE / HTMT / CFA / 共同方法偏差 / 测量不变性根据数据和研究设计选择。", 7.05, 3.80, 5.10, 1.25, accent=COLORS["blue"], fill_color=COLORS["pale_blue"])

    step_slide(
        prs, 9, "实操 9：规划预测试",
        "在主实验前检验操纵强度、现实感、清晰度、混淆和题项质量。",
        "请基于 experiment/stimuli.md、experiment/checks_and_attention.md 和 experiment/scales.xlsx，生成 experiment/pretest_plan.md。\n\n包含预测试目的、样本来源和建议 N、随机分配方式、操纵检验判据、现实感和清晰度判据、混淆检验判据、修改规则和进入主实验的标准。",
        "experiment/pretest_plan.md",
        ["有明确预测试目的", "有判据", "有修改规则", "先预测试再主实验", "不看主结果后才改规则"],
        "预测试不是形式，是避免主实验失败的低成本环节。"
    )

    step_slide(
        prs, 10, "实操 10：规划主实验流程",
        "写出可直接执行的招募、知情同意、随机分配、情景阅读、测量和排除流程。",
        "请生成 experiment/procedure.md。\n\n包括招募平台和目标被试、知情同意、随机分配、情景阅读、操纵检验、中介和因变量测量、注意力/现实感/混淆检验、人口统计变量、退出权和伦理说明、排除规则、数据保存和匿名化。",
        "experiment/procedure.md",
        ["招募和目标被试清楚", "随机分配清楚", "排除规则透明", "有伦理说明", "有数据匿名化说明"],
        "流程要能直接交给问卷平台或研究助理执行。"
    )

    step_slide(
        prs, 11, "实操 11：生成分析计划",
        "提前确定数据审计、排除、操纵检验、主模型、机制模型、效应量和因果边界。",
        "请基于 experiment/causal_model.md、experiment/procedure.md 和 hypotheses，生成 analysis/analysis_plan.md。\n\n按数据审计、缺失和重复、排除规则、随机化检验、操纵检验、信度、主效应、中介或调节、效应量、稳健性、因果语言边界、表格和图形计划的顺序写。",
        "analysis/analysis_plan.md",
        ["模型匹配设计", "排除规则先于结果", "报告效应量", "区分确认性和探索性", "标注因果语言边界"],
        "不能根据显著性临时改模型、改排除或换 DV。"
    )

    slide = new("分析顺序")
    flow(slide, ["数据审计", "排除", "随机化", "操纵检验", "信度", "主效应", "机制/边界", "稳健性"], 0.80, 1.55, 11.75, h=0.62, size=10)
    rows = [
        ["设计", "默认分析"],
        ["两组 + 连续 DV", "t-test + OLS"],
        ["多组条件", "ANOVA + planned contrasts"],
        ["2×2", "OLS/ANOVA interaction"],
        ["二元选择 DV", "Logistic regression"],
        ["中介", "Bootstrap indirect effect"],
        ["调节", "Interaction + simple slopes"],
        ["调节中介", "Conditional indirect effects"],
    ]
    table(slide, rows, 1.00, 2.75, 11.25, 3.45, col_widths=[3.2, 8.05], font_size=10.5)

    step_slide(
        prs, 12, "实操 12：生成 Method / Results 模板",
        "把实验设计转成论文方法、结果、表格、图和附录的写作模板。",
        "请使用 scenario-experiment-reporting 和 questionnaire-reporting-template，生成 reporting/method_results_template.md。\n\n包括 Method section draft、Design and participants、Procedure、Stimuli and manipulation、Measures、Manipulation checks、Analysis plan、Results reporting template、Table plan、Figure plan、Appendix materials、Transparency statement。",
        "reporting/method_results_template.md",
        ["报告设计和条件", "报告样本与排除", "报告操纵检验和信度", "报告效应量/CI/p值", "有附录和透明度说明"],
        "顶刊式报告要能让读者复现实验。"
    )

    step_slide(
        prs, 13, "实操 13：实验设计审计",
        "用严格审稿人视角检查情景、操纵、量表、预测试、分析和伦理风险。",
        "请以严格审稿人视角，审计本工作区所有实验设计文件，生成 qc/experiment_design_audit.md。\n\n检查情景真实、操纵是否只影响目标构念、是否存在能力/风险/成本/隐私/情绪混淆、操纵检验是否与中介/DV 重叠、量表来源是否真实、预测试判据、排除规则、分析计划、因果语言和伦理风险。",
        "qc/experiment_design_audit.md",
        ["列出致命问题", "列出重要问题", "列出可修改问题", "给出下一步行动", "明确 keep/revise/drop"],
        "设计审计越严格，后期数据越不容易白收。"
    )

    slide = new("课堂示范：AI Agent 解释透明度")
    rows = [
        ["模块", "示范内容"],
        ["研究问题", "解释透明度是否通过降低不确定性感知并提升信任，增加博士生采纳 AI 推荐文献的意愿"],
        ["条件 A", "AI 只给出推荐列表，不说明理由"],
        ["条件 B", "AI 给出推荐理由、证据来源、匹配逻辑和不确定性说明"],
        ["保持一致", "AI 能力、推荐数量、任务背景、时间压力、推荐质量、隐私描述"],
        ["检查项", "解释充分性感知；能力/友好度/隐私风险/任务难度混淆检验"],
        ["中介/DV", "不确定性感知、信任、采纳意愿或实际选择"],
    ]
    table(slide, rows, 0.78, 1.45, 11.75, 4.95, col_widths=[2.0, 9.75], font_size=9.7)
    card(slide, "最大风险", "解释透明度可能同时提升感知能力。必须用材料控制和能力混淆检验处理。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("常见错误与修正")
    rows = [
        ["错误", "后果", "修正"],
        ["问卷包装成实验", "没有因果识别", "明确随机分配和操纵"],
        ["操纵改变多个构念", "结果无法解释", "条件差异表和混淆检验"],
        ["操纵检验等同中介", "机制证据污染", "分开测量，题项不重叠"],
        ["量表来源不实", "学术风险", "逐条核验来源和题项"],
        ["预测试省略", "主实验失败成本高", "先小样本检操纵和混淆"],
        ["只报 p 值", "结果不可审计", "报告效应量、CI、表格和图"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.4, 4.05, 5.30], font_size=9.8)

    slide = new("课后提交要求")
    bullets(slide, [
        "提交文件夹：姓名_第07讲_情景实验行为研究与量表开发",
        "必须包含 causal_model、benchmark_matrix、stimuli、condition_difference_table、checks、scales、pretest_plan、procedure、analysis_plan、method_results_template、design_audit",
        "附 300 字反思：你的实验最大混淆风险是什么，如何控制",
        "未核验论文和量表必须保留待核验标记",
    ], 0.90, 1.55, 11.3, 1.75, size=15.5)
    rows = [
        ["评分项", "占比"],
        ["因果结构与实验类型", "15%"],
        ["情景材料与条件差异控制", "25%"],
        ["检查项和预测试设计", "20%"],
        ["量表来源与适配", "20%"],
        ["分析计划、报告和审计", "20%"],
    ]
    table(slide, rows, 2.15, 3.80, 8.85, 2.20, col_widths=[6.6, 2.25], font_size=11.5)
    card(slide, "最低合格线", "能说清：操纵什么、为什么操纵、如何避免混淆、怎么测量、怎么预测试、怎么分析和怎么报告。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.3)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是写了一段情景材料，而是建立了一个能被预测试、主实验、统计分析和审稿人审计的实验系统。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["因果结构", "顶刊范式", "操纵", "检查", "量表", "预测试", "分析", "报告"], 1.00, 3.50, 11.20, h=0.70, size=10.5)
    card(slide, "下一步", "补真实量表来源和顶刊样例，修改混淆风险，先做预测试，再决定是否进入主实验。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
