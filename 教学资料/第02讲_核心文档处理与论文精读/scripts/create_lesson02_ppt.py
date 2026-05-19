from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第02讲_课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(247, 248, 245),
    "ink": RGBColor(31, 41, 51),
    "muted": RGBColor(91, 103, 112),
    "teal": RGBColor(15, 118, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(180, 83, 9),
    "red": RGBColor(185, 28, 28),
    "line": RGBColor(210, 215, 220),
    "white": RGBColor(255, 255, 255),
    "pale_teal": RGBColor(219, 245, 241),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(254, 243, 199),
    "pale_red": RGBColor(254, 226, 226),
    "soft": RGBColor(236, 239, 241),
    "dark": RGBColor(33, 43, 54),
}

FONT = "Microsoft YaHei"


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color=COLORS["line"], width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def clear_line(shape):
    shape.line.fill.background()


def add_text(slide, text, x, y, w, h, size=20, color=COLORS["ink"], bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=COLORS["ink"], gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def add_code(slide, text, x, y, w, h, size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, COLORS["dark"])
    clear_line(shape)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(235, 239, 245)
    return shape


def add_header(slide, title, section="第 02 讲", num=None):
    set_fill(slide.background, COLORS["bg"])
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34)
    )
    set_fill(bar, COLORS["teal"])
    clear_line(bar)
    add_text(slide, section, 0.45, 0.05, 1.6, 0.22, size=8, color=COLORS["white"], bold=True)
    if num is not None:
        add_text(slide, f"{num:02d}", 12.05, 0.04, 0.7, 0.25, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.62, 11.8, 0.55, size=25, color=COLORS["ink"], bold=True)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.21), Inches(1.1), Inches(0.06)
    )
    set_fill(line, COLORS["amber"])
    clear_line(line)


def add_card(slide, title, body, x, y, w, h, fill=COLORS["white"], accent=COLORS["teal"],
             title_size=15, body_size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, COLORS["line"], 0.8)
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    set_fill(stripe, accent)
    clear_line(stripe)
    add_text(slide, title, x + 0.18, y + 0.15, w - 0.28, 0.35, size=title_size, color=accent, bold=True)
    add_text(slide, body, x + 0.18, y + 0.62, w - 0.32, h - 0.72, size=body_size, color=COLORS["ink"])
    return shape


def add_table(slide, rows, x, y, w, h, col_widths=None, font_size=10, header_fill=COLORS["teal"]):
    nrows, ncols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size if r else font_size + 0.5)
                p.font.bold = bool(r == 0)
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(242, 245, 247))
    return table_shape


def add_flow(slide, labels, x, y, w, box_h=0.58, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11):
    gap = 0.18
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(box_h)
        )
        set_fill(shape, color)
        set_line(shape, accent, 1)
        add_text(slide, label, bx + 0.08, y + 0.14, box_w - 0.16, box_h - 0.2,
                 size=font_size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            x1 = Inches(bx - gap + 0.02)
            x2 = Inches(bx - 0.02)
            ymid = Inches(y + box_h / 2)
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, ymid, x2, ymid)
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.2)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_fill(slide.background, COLORS["bg"])
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.0), Inches(7.5))
    set_fill(left, COLORS["teal"])
    clear_line(left)
    add_text(slide, "核心文档处理\n与论文精读", 0.5, 1.0, 2.2, 1.7, size=26, color=COLORS["white"], bold=True)
    add_text(slide, "第 02 讲", 0.6, 5.95, 1.4, 0.35, size=14, color=COLORS["white"], bold=True)
    add_text(slide, "把混乱文件变成可复用科研资产", 3.55, 1.25, 8.75, 0.65, size=29, color=COLORS["ink"], bold=True)
    add_text(slide, "PDF / DOCX / XLSX / PPTX / Markdown · 文件盘点 · 论文卡片 · 表格抽取 · 人工核验", 3.58, 2.08, 8.6, 0.45, size=15, color=COLORS["muted"])
    add_flow(slide, ["盘点", "预检", "转换", "抽取", "卡片", "核验"], 3.55, 3.25, 8.8, box_h=0.66, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=12)
    add_card(slide, "本讲交付", "file_inventory.xlsx、converted/*.md、paper_cards/*.md、extracted/tables/*.xlsx、manual_verification_checklist.md", 3.55, 4.65, 8.75, 1.35, fill=COLORS["white"], accent=COLORS["amber"], title_size=14, body_size=13)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    title_slide(prs)

    def new_slide(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_header(slide, title, num=len(prs.slides))
        return slide

    slide = new_slide("这节课解决什么科研问题")
    add_bullets(slide, [
        "文件混乱：PDF、Word、PPT、Excel、扫描件散落在不同位置",
        "抽取不可靠：PDF 断行、OCR 错字、表格错位、公式丢失",
        "读完不沉淀：只有“看过了”，没有论文卡片和阅读矩阵",
        "来源不可追踪：总结内容没有页码、表号、图号和原始文件",
        "后续污染：错误抽取进入综述、理论、变量和写作流程",
    ], 0.9, 1.65, 5.9, 3.4, size=18)
    add_card(slide, "本讲目标", "建立一条标准生产线：文件盘点 → 质量预检 → Markdown 转换 → 表格/图注抽取 → 论文卡片 → 人工核验。", 7.1, 1.75, 5.1, 1.75, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)
    add_card(slide, "课堂底线", "不要一上来就让 Agent 总结。先盘点、再抽取、再核验，最后才概括和沉淀。", 7.1, 4.05, 5.1, 1.45, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("本讲固定模板")
    rows = [
        ["模块", "第二讲怎么落实"],
        ["科研问题和需求", "把混乱文件转成可复用科研资产"],
        ["需要哪些 Skills", "pdf、docx、xlsx、pptx、markitdown"],
        ["输入材料", "PDF 论文、Word、Excel、PPT、扫描件"],
        ["Agent 怎么执行", "先盘点文件，再转换、抽取、制卡、核验"],
        ["输出文件", "file_inventory、converted、paper_cards、tables、qc"],
        ["人工核验", "OCR、表格、公式、图注、引用、隐私"],
        ["课后作业", "处理 5 个文件，提交卡片、矩阵、核验表和日志"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.9, col_widths=[2.65, 9.2], font_size=13)
    add_text(slide, "这节课的评价标准是“可追踪、可复核、可复用”，不是总结文字越多越好。", 0.9, 6.55, 10.9, 0.35, size=14, color=COLORS["amber"], bold=True)

    slide = new_slide("文档处理生产线")
    add_flow(slide, ["文件盘点", "质量预检", "格式转换", "内容抽取", "论文卡片", "人工核验", "知识沉淀"], 0.7, 1.55, 12.0, box_h=0.7, color=COLORS["pale_blue"], accent=COLORS["blue"], font_size=11)
    add_card(slide, "输入", "PDF 论文、Word 草稿、Excel 数据、PPT 课件、扫描件。", 0.85, 2.95, 3.7, 1.3, fill=COLORS["white"], accent=COLORS["teal"])
    add_card(slide, "中间格式", "Markdown 文本、抽取表格、图注索引、元数据。", 4.85, 2.95, 3.7, 1.3, fill=COLORS["white"], accent=COLORS["blue"])
    add_card(slide, "输出资产", "论文卡片、阅读矩阵、核验清单、研究日志。", 8.85, 2.95, 3.7, 1.3, fill=COLORS["white"], accent=COLORS["amber"])
    add_text(slide, "从第二讲开始，所有资料都必须进入工作区结构，而不是堆在桌面或聊天窗口。", 0.9, 5.65, 11.5, 0.45, size=18, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("需要哪些 Skills")
    rows = [
        ["Skill", "主要能力", "课堂任务"],
        ["pdf", "文本、表格、图片、元数据、OCR、拆分合并", "论文正文、表格、图注、页码来源"],
        ["docx", "Word 读取、生成、编辑、样式、批注、修订", "草稿、读书报告、卡片导出"],
        ["xlsx", "Excel/CSV 读取、清洗、公式、格式化", "文件盘点表、阅读矩阵、抽取表格"],
        ["pptx", "PPT 读取、生成、检查、转 Markdown", "课程课件、讲者备注、汇报结构"],
        ["markitdown", "多格式转 Markdown", "统一中间格式，便于 Agent 处理"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 4.35, col_widths=[1.45, 5.15, 5.45], font_size=11)
    add_code(slide, "copy_skill_to_workspace.ps1 -SkillName pdf,docx,xlsx,pptx,markitdown -Workspace \"D:\\AI科研训练营\\第02讲_文档处理工作区\"", 0.9, 6.2, 11.5, 0.65, size=12)

    slide = new_slide("输入材料怎么准备")
    rows = [
        ["输入", "数量", "要求"],
        ["PDF 论文", "3 篇", "至少 1 篇含表格，最好 1 篇含复杂公式或图"],
        ["Word 文档", "1 个", "论文草稿、读书笔记或课程资料"],
        ["Excel 表格", "1 个", "文献表、实验数据或课程名单"],
        ["PPT 课件", "1 个", "用于抽取结构和讲者备注"],
        ["扫描件/图片", "可选", "用于演示 OCR 风险"],
        ["研究主题", "1 个", "决定哪些信息值得抽取"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.45, col_widths=[2.2, 1.35, 8.3], font_size=12)
    add_text(slide, "输入文件命名要服务追踪：年份_第一作者_关键词.pdf，而不是“下载.pdf”。", 0.9, 6.35, 11.2, 0.35, size=15, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("标准工作区结构")
    add_code(slide, "第02讲_文档处理工作区/\n├─ AGENTS.md\n├─ skills/\n├─ input/\n│  ├─ papers/\n│  ├─ word/\n│  ├─ slides/\n│  ├─ spreadsheets/\n│  └─ scans/\n├─ converted/\n├─ extracted/\n│  ├─ tables/\n│  ├─ figures/\n│  └─ metadata/\n├─ paper_cards/\n├─ reading_matrix/\n├─ qc/\n└─ output/", 0.85, 1.45, 5.45, 5.4, size=12)
    add_bullets(slide, [
        "input：原始材料，只读不覆盖",
        "converted：Markdown 中间结果",
        "extracted：表格、图注、元数据",
        "paper_cards：单篇论文精读卡片",
        "reading_matrix：多篇论文比较表",
        "qc：人工核验清单",
    ], 7.0, 1.7, 5.3, 3.1, size=18)
    add_card(slide, "关键原则", "每个输出都能追溯到原始文件、页码、表号或幻灯片编号。", 7.0, 5.35, 5.3, 0.9, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("第一步：文件盘点表")
    rows = [
        ["字段", "作用"],
        ["file_id", "给每个文件一个唯一编号"],
        ["file_name / file_type", "记录原始文件名和格式"],
        ["source / research_use", "记录来源和科研用途"],
        ["need_ocr", "判断是否扫描件或图片"],
        ["has_tables / has_figures", "判断是否需要抽表或图注索引"],
        ["process_status", "未处理 / 已转换 / 已抽取 / 已核验"],
        ["verification_notes", "人工核验备注"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.9, col_widths=[3.0, 8.85], font_size=12)
    add_text(slide, "没有 file_inventory，就没有可追踪的文档处理流程。", 0.9, 6.55, 11.2, 0.35, size=16, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("第二步：质量预检")
    add_card(slide, "PDF", "是否扫描件、是否加密、是否断行、是否包含跨页表、是否有图注和公式。", 0.85, 1.55, 3.75, 1.55, fill=COLORS["white"], accent=COLORS["teal"], body_size=12)
    add_card(slide, "Word", "是否有修订、批注、样式混乱、图片嵌入、表格和脚注。", 4.8, 1.55, 3.75, 1.55, fill=COLORS["white"], accent=COLORS["blue"], body_size=12)
    add_card(slide, "Excel", "是否有隐藏列、合并单元格、公式错误、单位混乱、表头错位。", 8.75, 1.55, 3.75, 1.55, fill=COLORS["white"], accent=COLORS["amber"], body_size=12)
    add_card(slide, "PPT", "是否有讲者备注、图片文字、复杂图表、嵌入对象、页码结构。", 0.85, 3.75, 3.75, 1.45, fill=COLORS["white"], accent=COLORS["blue"], body_size=12)
    add_card(slide, "扫描件", "OCR 语言、分辨率、倾斜、表格边框、手写批注。", 4.8, 3.75, 3.75, 1.45, fill=COLORS["white"], accent=COLORS["red"], body_size=12)
    add_card(slide, "安全", "未公开数据、被试隐私、企业数据不能直接上传外部模型。", 8.75, 3.75, 3.75, 1.45, fill=COLORS["pale_red"], accent=COLORS["red"], body_size=12)

    slide = new_slide("第三步：转换为 Markdown")
    add_bullets(slide, [
        "Markdown 是 Agent 友好的中间格式，不是最终事实",
        "转换后必须保留原始文件名、工具、时间、核验状态",
        "标题层级、页码、图注和表格位置要尽量保留",
        "不要在转换阶段改写论文主张",
    ], 0.9, 1.55, 5.7, 2.65, size=18)
    add_code(slide, "# 每个 Markdown 文件开头\n---\nsource_file: 2024_Zhang_AI_agent.pdf\nconversion_tool: markitdown/pdf\nconversion_time: 2026-05-16\nverification_status: 未核验\n---", 7.0, 1.65, 5.4, 2.0, size=13)
    add_card(slide, "课堂判断", "转换结果越干净，后续卡片越稳定；转换结果越脏，后续综述越容易被污染。", 1.0, 4.85, 11.2, 0.9, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=14)

    slide = new_slide("第四步：论文精读卡片")
    rows = [
        ["模块", "必须抽取什么"],
        ["基本信息", "题名、作者、年份、期刊、DOI/URL"],
        ["研究问题", "问题、重要性、与自己研究的关系"],
        ["理论机制", "理论如何解释变量关系，不只列理论名"],
        ["变量与构念", "自变量、因变量、中介、调节、控制变量"],
        ["样本与数据", "样本量、来源、情境、时间范围"],
        ["方法与结果", "模型、检验、主要发现、证据位置"],
        ["启发与局限", "可复用点、边界条件、必须核验项"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.05, col_widths=[2.4, 9.45], font_size=12)
    add_text(slide, "论文卡片不是摘要，而是后续综述、选题、变量设计和写作的结构化原料。", 0.9, 6.65, 11.2, 0.35, size=15, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("第五步：表格抽取")
    add_bullets(slide, [
        "PDF 表格抽取最容易出现错位",
        "跨页表、合并单元格、多层表头需要特别标注",
        "统计显著性星号和表注不能丢",
        "不要自动补全缺失数值",
        "每张表记录来源文件、页码、表号、表题",
    ], 0.85, 1.55, 5.65, 3.1, size=18)
    rows = [
        ["风险", "处理方式"],
        ["表头错位", "标注需人工核验"],
        ["跨页表", "拆分后回原文核验"],
        ["星号/表注丢失", "保留原始注释"],
        ["变量名换行", "不要自动猜测"],
    ]
    add_table(slide, rows, 6.95, 1.75, 5.35, 2.5, col_widths=[2.0, 3.35], font_size=12)
    add_card(slide, "原则", "宁可留下“需人工核验”，不要让 Agent 猜出一个看似整齐但错误的表。", 6.95, 4.75, 5.35, 1.0, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=13)

    slide = new_slide("第六步：图注和图表索引")
    add_table(slide, [
        ["字段", "说明"],
        ["source_file", "来源文件"],
        ["page", "页码或幻灯片编号"],
        ["figure_or_table_id", "Figure 1 / Table 2 等编号"],
        ["caption", "图题或表题"],
        ["related_section", "相关正文部分"],
        ["verification_status", "未核验 / 部分核验 / 已核验"],
    ], 0.8, 1.55, 5.8, 4.1, col_widths=[2.2, 3.6], font_size=12)
    add_bullets(slide, [
        "图表索引不是替代原图",
        "索引用于后续写作、复盘和快速定位",
        "图号、表号和正文引用必须一致",
        "不要把正文普通句子误判成图注",
    ], 7.0, 1.85, 5.1, 2.4, size=18)
    add_card(slide, "后续用途", "第 13 讲做科研 PPT、论文图和海报时，图表索引会直接复用。", 7.0, 4.8, 5.1, 0.95, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=13)

    slide = new_slide("第七步：文献阅读矩阵")
    rows = [
        ["字段", "用途"],
        ["title/authors/year/journal", "文献基本信息"],
        ["research_question", "综述中的问题线索"],
        ["theory", "理论建构素材"],
        ["variables", "变量设计和假设来源"],
        ["method/sample", "方法对比和可行性判断"],
        ["key_findings", "结果综合"],
        ["limitations", "研究缺口来源"],
        ["reusable_insights", "自己的选题启发"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.15, col_widths=[3.2, 8.65], font_size=12)
    add_text(slide, "阅读矩阵是第 05 讲文献综述和研究问题生成的直接输入。", 0.9, 6.65, 11.2, 0.35, size=15, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("第八步：人工核验清单")
    rows = [
        ["风险点", "核验方法"],
        ["OCR 错字", "回到原 PDF 页面或图片"],
        ["PDF 断行", "检查摘要、理论段落、结果段落"],
        ["表格错位", "对照原始表格行列和表注"],
        ["公式/统计符号", "检查希腊字母、上下标、p、N、t、F、β"],
        ["图注与图号", "检查编号和正文引用是否一致"],
        ["参考文献", "下一讲用 DOI/数据库核验"],
        ["Excel 公式", "检查 #REF!、#DIV/0!、隐藏列、合并单元格"],
        ["数据隐私", "未公开数据不能上传外部模型"],
    ]
    add_table(slide, rows, 0.65, 1.45, 12.05, 5.25, col_widths=[2.6, 9.45], font_size=11)

    slide = new_slide("Agent 执行任务卡")
    add_code(slide, "# 今日 Agent 任务：核心文档处理\n\n## 目标\n把 input/ 中的论文和课程材料转成 Markdown、表格和论文卡片。\n\n## 可用 Skills\npdf, docx, xlsx, pptx, markitdown\n\n## 执行要求\n- 先生成 file_inventory.xlsx，不要直接总结\n- 所有抽取内容记录来源文件和页码/幻灯片编号\n- 表格无法确认时标注“需人工核验”\n- 不要补全缺失 DOI 或改写论文主张\n\n## 输出\nreading_matrix/file_inventory.xlsx\nconverted/*.md\nextracted/tables/*.xlsx\npaper_cards/*.md\nqc/manual_verification_checklist.md", 0.85, 1.45, 6.35, 5.45, size=10.5)
    add_bullets(slide, [
        "任务卡先规定输入、工具、输出和核验",
        "每一步都要落到文件路径",
        "失败项必须进入 qc 或 research_log",
        "输出越结构化，后续课程越省力",
    ], 7.65, 1.9, 4.5, 2.35, size=18)
    add_card(slide, "课堂口令", "先盘点，不总结；先核验，不写作。", 7.65, 4.9, 4.5, 0.9, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=15)

    slide = new_slide("课堂实操 1：创建工作区")
    add_code(slide, "New-Item -ItemType Directory -Force -Path \"D:\\AI科研训练营\\第02讲_文档处理工作区\"\ncd \"D:\\AI科研训练营\\第02讲_文档处理工作区\"\n\nNew-Item -ItemType Directory -Force -Path skills,input,converted,extracted,paper_cards,reading_matrix,qc,output\nNew-Item -ItemType Directory -Force -Path input\\papers,input\\word,input\\slides,input\\spreadsheets,input\\scans\nNew-Item -ItemType Directory -Force -Path extracted\\tables,extracted\\figures,extracted\\metadata\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md,project_status.md", 0.8, 1.55, 12.0, 2.7, size=12)
    add_bullets(slide, [
        "把 PDF 放入 input/papers",
        "把 Word 放入 input/word",
        "把 PPT 放入 input/slides",
        "把 Excel 放入 input/spreadsheets",
    ], 1.0, 4.85, 11.2, 1.2, size=18)

    slide = new_slide("课堂实操 2：复制 Skills")
    add_code(slide, "cd \"D:\\desk\\AI agent科研资料\\本地Skills功能分类库\"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName pdf,docx,xlsx,pptx,markitdown `\n  -Workspace \"D:\\AI科研训练营\\第02讲_文档处理工作区\"\n\nGet-ChildItem \"D:\\AI科研训练营\\第02讲_文档处理工作区\\skills\" -Directory", 0.85, 1.55, 11.8, 2.45, size=12)
    add_card(slide, "检查点", "skills 文件夹下应出现 pdf、docx、xlsx、pptx、markitdown，每个目录至少包含 SKILL.md。", 1.0, 4.65, 11.1, 0.95, fill=COLORS["pale_blue"], accent=COLORS["blue"], body_size=14)

    slide = new_slide("课堂实操 3：生成 file_inventory")
    add_code(slide, "请扫描 input/ 下所有文件，先不要总结内容。\n请生成 reading_matrix/file_inventory.xlsx，字段包括：\nfile_id, file_name, file_type, source, research_use, need_ocr, has_tables, has_figures, process_status, output_path, verification_notes。\n如果无法判断某项，请写“需人工核验”。", 0.85, 1.55, 11.8, 1.75, size=13)
    add_table(slide, [
        ["人工检查", "标准"],
        ["每个输入文件都有一行", "不能漏文件"],
        ["文件类型正确", "PDF/DOCX/PPTX/XLSX/Image"],
        ["OCR/表格/图标注", "不确定写需人工核验"],
        ["输出路径", "指向 converted/extracted/paper_cards"],
    ], 0.85, 3.85, 11.8, 2.0, col_widths=[3.6, 8.2], font_size=12)

    slide = new_slide("课堂实操 4：转换 Markdown")
    add_code(slide, "请把 input/papers、input/word、input/slides 中的文件转换为 Markdown。\n输出到 converted/。\n每个 Markdown 文件开头保留：原始文件名、转换时间、转换工具、人工核验状态。\n不要改写原文主张。", 0.85, 1.55, 11.8, 1.55, size=13)
    add_bullets(slide, [
        "检查标题层级",
        "检查摘要完整性",
        "检查表格是否被误拆",
        "检查图注是否丢失",
        "检查中文乱码和英文断行",
    ], 1.0, 3.65, 5.3, 1.85, size=18)
    add_card(slide, "输出", "converted/*.md", 7.0, 3.8, 5.0, 0.85, fill=COLORS["pale_teal"], accent=COLORS["teal"], body_size=15)

    slide = new_slide("课堂实操 5：生成论文卡片")
    add_code(slide, "请基于 converted/ 中的论文 Markdown，按 paper_card_template.md 生成论文精读卡片。\n每张卡片输出到 paper_cards/。\n每个结论都要保留来源位置或段落线索。\n无法确认的内容写“需人工核验”，不要猜。", 0.85, 1.55, 11.8, 1.55, size=13)
    add_table(slide, [
        ["卡片模块", "必须有"],
        ["理论", "不能只列理论名，要写机制"],
        ["变量", "类型、操作化、测量来源"],
        ["方法", "研究设计、模型、样本、数据"],
        ["结果", "证据位置和支持情况"],
        ["启发", "对自己研究可复用的具体点"],
    ], 0.85, 3.65, 11.8, 2.45, col_widths=[2.8, 9.0], font_size=12)

    slide = new_slide("课堂实操 6：抽取表格和图注")
    add_code(slide, "请从 PDF 或 Markdown 中识别表格和图注。\n表格输出为 extracted/tables/*.xlsx。\n图注索引输出为 extracted/figures/figure_index.md。\n每个条目必须记录来源文件和页码；无法确认页码时标注“需人工核验”。", 0.85, 1.55, 11.8, 1.65, size=13)
    add_bullets(slide, [
        "表格：保留表号、表题、表注、页码",
        "图注：保留 Figure/Table 编号和相关章节",
        "跨页表：优先标注风险，再人工核验",
        "不自动补全任何数值",
    ], 1.0, 3.75, 11.0, 1.8, size=18)

    slide = new_slide("课堂实操 7：生成核验清单和日志")
    add_code(slide, "请生成 qc/manual_verification_checklist.md。\n按文件列出 OCR、表格、公式、统计符号、参考文献、页码、图注、Excel 公式的核验任务。\n不要把未核验内容写成已确认。\n\n随后更新 research_log.md，写清输入、Skills、执行步骤、输出路径、未核验项和人工核验结果。", 0.85, 1.55, 11.8, 2.0, size=13)
    add_card(slide, "提交标准", "能复现的科研流程一定有日志；没有日志的自动化输出不能进入论文写作。", 1.0, 4.45, 11.1, 1.0, fill=COLORS["pale_amber"], accent=COLORS["amber"], body_size=14)

    slide = new_slide("本讲提交物")
    rows = [
        ["提交物", "最低要求"],
        ["file_inventory.xlsx", "每个输入文件一行"],
        ["converted/*.md", "至少 3 个转换结果"],
        ["paper_cards/*.md", "至少 3 张论文卡片"],
        ["extracted/tables/*.xlsx", "至少 1 个抽取表格"],
        ["figure_index.md", "图表标题、页码、核验状态"],
        ["literature_reading_matrix.xlsx", "每篇论文一行"],
        ["manual_verification_checklist.md", "具体列出风险和核验任务"],
        ["research_log.md", "执行过程和未核验项"],
    ]
    add_table(slide, rows, 0.75, 1.45, 11.85, 5.25, col_widths=[3.6, 8.25], font_size=12)

    slide = new_slide("评分标准")
    rows = [
        ["项目", "分值", "标准"],
        ["工作区结构", "10", "输入、转换、抽取、卡片、核验目录清楚"],
        ["文件盘点表", "15", "每个文件有记录，状态和风险标注清楚"],
        ["Markdown 转换", "15", "转换完整，保留来源和核验状态"],
        ["论文卡片", "20", "结构完整，能服务综述和选题"],
        ["表格/图注抽取", "15", "有来源、页码和人工核验标注"],
        ["核验清单", "15", "能指出具体风险，不空泛"],
        ["日志与反思", "10", "过程可追踪，反思具体"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 4.75, col_widths=[2.5, 1.05, 8.5], font_size=11)

    slide = new_slide("常见错误")
    rows = [
        ["错误", "后果", "修正"],
        ["直接总结 PDF", "来源不可追踪", "先做 file_inventory"],
        ["OCR 不核验", "错字进入综述", "回原文页面核验"],
        ["表格抽取后直接使用", "行列错位导致数据错", "对照原表核验"],
        ["卡片只写摘要", "不能服务选题和综述", "按理论/变量/方法/结果结构化"],
        ["引用信息不核验", "可能产生假引用", "第三讲用 DOI/数据库核验"],
    ]
    add_table(slide, rows, 0.65, 1.55, 12.05, 4.35, col_widths=[3.1, 4.05, 4.9], font_size=11)
    add_text(slide, "第二讲最重要的能力：把 Agent 输出当作待核验材料，而不是最终事实。", 0.9, 6.3, 11.4, 0.4, size=17, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("与后续课程的连接")
    rows = [
        ["后续课程", "会复用第二讲什么产出"],
        ["第 03 讲 文献检索与引用核验", "file_inventory、论文基本信息、未核验引用"],
        ["第 04 讲 Zotero/Obsidian", "paper_cards、reading_matrix"],
        ["第 05 讲 文献综述与研究问题", "理论、变量、方法、局限、启发"],
        ["第 08 讲 数据分析", "抽取表格、变量字典"],
        ["第 12 讲 论文写作与排版", "论文卡片、图表索引、参考文献信息"],
        ["第 13 讲 PPT/海报", "图表索引、课程 PPT 结构"],
    ]
    add_table(slide, rows, 0.75, 1.55, 11.85, 4.65, col_widths=[3.6, 8.25], font_size=12)
    add_text(slide, "第二讲的输出质量，直接决定后续综述、选题和写作的质量。", 0.9, 6.45, 11.2, 0.35, size=16, color=COLORS["teal"], bold=True, align=PP_ALIGN.CENTER)

    slide = new_slide("本讲一句话总结")
    add_text(slide, "文档处理不是“帮我总结”，\n而是把每个文件变成有来源、有结构、有核验状态的科研资产。", 1.15, 2.25, 11.0, 1.45, size=30, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
    add_flow(slide, ["文件清楚", "来源清楚", "结构清楚", "风险清楚", "用途清楚"], 1.35, 4.55, 10.55, box_h=0.75, color=COLORS["pale_teal"], accent=COLORS["teal"], font_size=13)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
