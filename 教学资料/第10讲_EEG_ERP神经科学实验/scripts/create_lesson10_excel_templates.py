from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"
TEAL = "0D746E"
PALE = "F4F7F8"
WHITE = "FFFFFF"
INK = "202A36"
LINE = "D2D8DE"


def style_sheet(ws, widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    side = Side(style="thin", color=LINE)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
    ws.row_dimensions[1].height = 36
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    else:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18


def make_wb(sheets, path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows, widths in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_sheet(ws, widths)
    wb.save(path)
    load_workbook(path).close()


def add_validations(path):
    wb = load_workbook(path)
    status = DataValidation(type="list", formula1='"pass,revise,fail,to_check,not_applicable"', allow_blank=True)
    yn = DataValidation(type="list", formula1='"yes,no,to_check,not_applicable"', allow_blank=True)
    role = DataValidation(type="list", formula1='"stimulus,response,feedback,fixation,instruction,baseline,artifact,other"', allow_blank=True)
    decision = DataValidation(type="list", formula1='"include,exclude,revise,to_check"', allow_blank=True)
    for ws in wb.worksheets:
        ws.add_data_validation(status)
        ws.add_data_validation(yn)
        ws.add_data_validation(role)
        ws.add_data_validation(decision)
    if "event_marker_table" in wb.sheetnames:
        ws = wb["event_marker_table"]
        role.add("D2:D500")
        yn.add("I2:J500")
        status.add("L2:L500")
    if "participant_qc" in wb.sheetnames:
        ws = wb["participant_qc"]
        decision.add("N2:N500")
    if "analysis_audit" in wb.sheetnames:
        ws = wb["analysis_audit"]
        status.add("D2:D500")
    wb.save(path)
    load_workbook(path).close()


event_sheets = [
    (
        "event_marker_table",
        [
            "marker_code", "event_name", "condition", "time_locking_role", "trial_phase",
            "stimulus_type", "response_mapping", "expected_count", "included_in_erp",
            "included_in_frequency", "behavior_link", "audit_status", "notes",
        ],
        [
            ["101", "low_transparency_stimulus", "low_transparency", "stimulus", "stimulus_onset", "AI recommendation text", "N/A", "to_check", "yes", "yes", "trial_id", "to_check", "示例：低解释透明度刺激出现"],
            ["102", "high_transparency_stimulus", "high_transparency", "stimulus", "stimulus_onset", "AI recommendation text", "N/A", "to_check", "yes", "yes", "trial_id", "to_check", "示例：高解释透明度刺激出现"],
            ["201", "accept_response", "all", "response", "response", "button press", "accept", "to_check", "response-locked if needed", "no", "trial_id+rt", "to_check", "反应事件"],
            ["301", "positive_feedback", "feedback_positive", "feedback", "feedback", "feedback screen", "N/A", "to_check", "yes", "yes", "trial_id", "to_check", "反馈锁定可用于 FRN/P3"],
        ],
        [16, 30, 24, 22, 20, 28, 26, 18, 18, 22, 22, 18, 42],
    ),
    (
        "condition_count_check",
        ["condition", "expected_trials", "actual_events", "valid_behavior_trials", "difference", "action", "notes"],
        [
            ["low_transparency", "to_check", "", "", "", "to_check", "事件数应与行为日志基本一致"],
            ["high_transparency", "to_check", "", "", "", "to_check", "检查漏发 trigger"],
        ],
        [26, 18, 18, 22, 16, 22, 46],
    ),
    (
        "trial_timing",
        ["phase", "duration_ms", "jitter", "marker", "purpose", "artifact_risk", "notes"],
        [
            ["fixation", "500", "yes/no", "", "baseline/attention", "blink before stimulus", ""],
            ["stimulus", "to_check", "yes/no", "101/102", "time-locking", "eye movement if reading", ""],
            ["response", "max/to_check", "yes/no", "201/202", "behavior", "motor artifacts", ""],
            ["feedback", "to_check", "yes/no", "301/302", "feedback-locked ERP", "late blinks", ""],
        ],
        [22, 18, 16, 18, 34, 34, 42],
    ),
]

qc_sheets = [
    (
        "participant_qc",
        [
            "subject_id", "raw_file", "behavior_file", "sampling_rate", "n_channels",
            "bad_channels", "interpolated_channels", "ica_components_removed",
            "total_events", "valid_events", "total_epochs", "retained_epochs",
            "retention_rate", "qc_decision", "exclusion_reason", "notes",
        ],
        [
            ["sub-001", "data/raw/sub-001_task-demo_raw.vhdr", "data/behavior/sub-001_behavior.csv", "to_check", "", "", "", "", "", "", "", "", "", "to_check", "", "示例行"],
        ],
        [16, 36, 36, 18, 14, 24, 26, 26, 16, 16, 16, 18, 18, 18, 34, 42],
    ),
    (
        "ica_component_log",
        ["subject_id", "component_id", "artifact_type", "evidence", "removed", "reviewer", "notes"],
        [
            ["sub-001", "ICA000", "eye blink/to_check", "topography/time series/correlation", "to_check", "student", "删除组件必须有证据"],
        ],
        [16, 16, 24, 40, 16, 18, 46],
    ),
    (
        "exclusion_flow",
        ["stage", "n_participants", "n_trials", "reason", "decision_rule", "notes"],
        [
            ["recorded", "", "", "all recorded", "", ""],
            ["after_event_check", "", "", "missing markers", "predefined", ""],
            ["after_preprocessing", "", "", "bad channels/artifacts", "predefined", ""],
            ["final_analysis", "", "", "retained", "", ""],
        ],
        [26, 18, 18, 34, 28, 46],
    ),
]

erp_sheets = [
    (
        "erp_component_table",
        [
            "component", "time_locking_event", "epoch_window", "baseline_window", "time_window",
            "roi_electrodes", "metric", "conditions", "statistical_model",
            "literature_or_rationale", "confirmatory_status", "notes",
        ],
        [
            ["P3", "stimulus_onset", "-200 to 800 ms", "-200 to 0 ms", "300-600 ms", "Pz; CPz; Cz", "mean amplitude", "low vs high transparency", "repeated-measures ANOVA / mixed model", "需要结合任务和文献核验", "to_check", "课堂示例"],
            ["FRN", "feedback_onset", "-200 to 600 ms", "-200 to 0 ms", "200-350 ms", "FCz; Cz", "mean amplitude", "positive vs negative feedback", "condition x feedback", "反馈评价相关示例", "to_check", "课堂示例"],
        ],
        [18, 24, 20, 22, 20, 28, 22, 34, 36, 42, 22, 42],
    ),
    (
        "trial_retention_by_condition",
        ["subject_id", "condition", "expected_epochs", "valid_epochs", "rejected_epochs", "retention_rate", "action", "notes"],
        [
            ["sub-001", "low_transparency", "", "", "", "", "to_check", "条件间 trial 数差异过大需处理"],
            ["sub-001", "high_transparency", "", "", "", "", "to_check", ""],
        ],
        [16, 28, 20, 18, 18, 18, 22, 46],
    ),
    (
        "erp_figure_plan",
        ["figure_id", "figure_type", "component", "conditions", "roi", "time_window", "must_show", "notes"],
        [
            ["Fig_ERP_01", "waveform", "P3", "low/high", "Pz/CPz/Cz", "300-600 ms", "time zero; CI/SE; condition labels", ""],
            ["Fig_ERP_02", "scalp map", "P3 difference", "high-low", "whole scalp", "300-600 ms", "color scale; window; condition", ""],
        ],
        [18, 24, 20, 28, 22, 20, 42, 42],
    ),
]

freq_sheets = [
    (
        "frequency_plan",
        [
            "analysis_type", "band_or_metric", "frequency_range", "method", "baseline",
            "normalization", "roi", "time_window", "statistical_model", "multiple_comparison", "interpretation_limit", "notes",
        ],
        [
            ["band_power", "alpha", "8-12 Hz/to_check", "PSD or wavelet", "pre-stimulus", "dB or percent change", "parietal/occipital", "to_check", "condition model", "planned ROI or correction", "alpha 不自动等于注意", "示例"],
            ["time_frequency", "theta", "4-7 Hz/to_check", "wavelet/ERSP", "pre-stimulus", "dB", "frontal-midline", "to_check", "condition x time", "cluster/permutation if many tests", "theta 解释需结合任务", "示例"],
        ],
        [22, 22, 22, 24, 22, 22, 26, 22, 34, 34, 44, 34],
    ),
    (
        "connectivity_plan",
        ["metric", "frequency_band", "node_or_roi_1", "node_or_roi_2", "time_window", "baseline", "statistical_model", "claim_boundary", "notes"],
        [
            ["PLV/wPLI/coherence", "theta/alpha/to_check", "frontal", "parietal", "to_check", "to_check", "condition model", "functional coupling, not causal direction", "连接解释必须克制"],
        ],
        [24, 24, 24, 24, 22, 20, 34, 42, 44],
    ),
    (
        "frequency_figure_plan",
        ["figure_id", "figure_type", "metric", "conditions", "frequency", "time_window", "must_show", "notes"],
        [
            ["Fig_TF_01", "ERSP", "power", "low/high", "theta/alpha", "to_check", "baseline; color scale; ROI", ""],
            ["Fig_CONN_01", "connectivity matrix", "wPLI", "low/high", "theta", "to_check", "metric; threshold; nodes", ""],
        ],
        [18, 26, 20, 26, 20, 20, 42, 42],
    ),
]

ml_sheets = [
    (
        "eeg_ml_plan",
        [
            "target", "label_source", "feature_type", "unit", "split_strategy", "leakage_risk",
            "baseline_model", "advanced_model", "metrics", "chance_level", "validation", "interpretation_limit",
        ],
        [
            ["condition classification", "event condition", "ERP amplitudes / band power", "epoch or subject", "subject-wise split if cross-subject", "same subject epochs across train/test", "logistic regression / SVM", "CNN/RNN/Transformer only if justified", "balanced accuracy; AUC; F1; confusion matrix", "to_check", "CV / LOSO / permutation", "prediction is not mechanism"],
        ],
        [28, 24, 34, 18, 34, 42, 32, 36, 38, 20, 28, 42],
    ),
    (
        "leakage_audit",
        ["risk", "example", "detection", "required_fix", "status", "notes"],
        [
            ["epoch_leakage", "same participant epochs in train and test", "group split by subject", "use subject-wise split", "to_check", ""],
            ["normalization_leakage", "scaler fit on all data", "inspect pipeline", "fit scaler inside CV", "to_check", ""],
            ["test_tuning", "hyperparameters tuned on test", "review logs", "nested CV or validation set", "to_check", ""],
            ["label_leakage", "feature contains label/event code", "feature audit", "remove leaking feature", "to_check", ""],
        ],
        [28, 42, 32, 36, 18, 42],
    ),
    (
        "metrics_report",
        ["model", "split", "accuracy", "balanced_accuracy", "auc", "f1", "ci_or_sd", "chance_level", "decision", "notes"],
        [
            ["baseline_svm", "LOSO/to_check", "", "", "", "", "", "", "to_check", "不要只报告单一 accuracy"],
        ],
        [22, 22, 14, 20, 14, 14, 18, 18, 18, 42],
    ),
]

audit_sheets = [
    (
        "analysis_audit",
        ["audit_item", "question", "evidence_file", "status", "fix_required", "notes"],
        [
            ["theory_indicator_match", "神经指标是否匹配研究问题和任务时序？", "study_design.md", "to_check", "", ""],
            ["event_integrity", "事件码是否完整、唯一、可 time-lock？", "event_marker_table.xlsx", "to_check", "", ""],
            ["behavior_sync", "行为数据和 EEG 是否能按 trial 对齐？", "behavior/events", "to_check", "", ""],
            ["preprocessing_preregistered", "滤波、参考、坏道、ICA、epoch 是否先验？", "preprocessing_plan.md", "to_check", "", ""],
            ["qc_transparency", "被试、通道、ICA、trial 剔除是否透明？", "preprocessing_qc.xlsx", "to_check", "", ""],
            ["erp_window_roi", "ERP 窗口和 ROI 是否有理论/文献依据？", "erp_component_table.xlsx", "to_check", "", ""],
            ["multiple_comparison", "多电极、多窗口、多频段是否控制多重比较？", "analysis plans", "to_check", "", ""],
            ["connectivity_claim", "连接分析是否避免因果方向解释？", "frequency_connectivity_plan.md", "to_check", "", ""],
            ["ml_leakage", "机器学习是否存在训练测试泄漏？", "eeg_ml_plan.xlsx", "to_check", "", ""],
            ["writing_boundary", "论文解释是否克制并报告非显著计划检验？", "eeg_results_template.md", "to_check", "", ""],
        ],
        [28, 60, 34, 18, 38, 42],
    ),
    (
        "reviewer_risk_register",
        ["risk", "severity", "why_it_matters", "mitigation", "owner", "status"],
        [
            ["low trial retention", "high", "ERP averages may be unstable", "increase trials or report limitation", "student", "to_check"],
            ["post-hoc ERP window", "high", "circular analysis", "label exploratory or use preregistered window", "student", "to_check"],
            ["overinterpreted component", "medium/high", "component is not construct", "rewrite as neural evidence consistent with theory", "student", "to_check"],
        ],
        [30, 18, 42, 42, 18, 18],
    ),
]

run_sheets = [
    (
        "pipeline_run_log",
        ["run_id", "date", "step", "script_or_tool", "input_file", "output_file", "parameters", "result", "next_step"],
        [
            ["run_001", "YYYY-MM-DD", "event_check", "MNE/EEGLAB/manual", "data/raw + data/events", "qc/event_check.md", "", "to_check", "preprocessing"],
            ["run_002", "YYYY-MM-DD", "preprocessing", "MNE/EEGLAB", "data/raw", "data/derivatives", "filter/reference/ICA", "to_check", "QC"],
            ["run_003", "YYYY-MM-DD", "ERP extraction", "MNE/EEGLAB", "clean epochs", "erp_component_values.csv", "window/ROI", "to_check", "statistics"],
        ],
        [16, 16, 24, 30, 34, 34, 34, 18, 34],
    ),
    (
        "file_lineage",
        ["file", "source", "created_by", "created_at", "modified_from", "do_not_overwrite", "notes"],
        [
            ["data/raw/sub-001_task-demo_raw.vhdr", "acquisition system", "human", "YYYY-MM-DD", "none", "yes", "原始数据只读保存"],
            ["data/derivatives/sub-001_clean-epo.fif", "raw EEG", "script", "YYYY-MM-DD", "preprocessing script", "no", "衍生文件"],
            ["eeg/erp_component_values.csv", "clean epochs", "script", "YYYY-MM-DD", "ERP extraction", "no", "进入统计分析"],
        ],
        [36, 34, 18, 18, 34, 18, 42],
    ),
]


outputs = [
    ("event_marker_table_template.xlsx", event_sheets),
    ("preprocessing_qc_template.xlsx", qc_sheets),
    ("erp_component_table_template.xlsx", erp_sheets),
    ("frequency_connectivity_template.xlsx", freq_sheets),
    ("eeg_ml_plan_template.xlsx", ml_sheets),
    ("eeg_analysis_audit_template.xlsx", audit_sheets),
    ("eeg_pipeline_log_template.xlsx", run_sheets),
]

for filename, sheets in outputs:
    path = ROOT / filename
    make_wb(sheets, path)
    add_validations(path)
    print(path)
