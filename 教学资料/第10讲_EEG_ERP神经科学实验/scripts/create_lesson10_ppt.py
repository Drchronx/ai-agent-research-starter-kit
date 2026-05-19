from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第10讲_课件.pptx"
W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}
FONT = "Microsoft YaHei"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, level = item if isinstance(item, tuple) else (item, 0)
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=8.9):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 10 讲 · EEG/ERP 神经科学实验", 0.45, 0.055, 3.3, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])
    return rect


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.2, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
    return shape


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=10):
    gap = 0.12
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.01), Inches(y + h / 2), Inches(bx - 0.02), Inches(y + h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"])
    no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.16, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.18, accent=COLORS["amber"], body_size=11.2, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.72, size=8.55)
    textbox(slide, "人工核验", 4.65, 4.82, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.3)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 10 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "EEG / ERP\n神经科学\n实验", 0.52, 1.70, 2.25, 1.65, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "从范式到论文报告", 0.55, 6.10, 2.2, 0.35, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "实验设计、事件码、预处理、ERP、频域连接与 QC", 3.55, 1.10, 8.75, 0.72, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "Paradigm · Markers · Preprocessing · ERP · ERSP · Connectivity · ML · Reporting", 3.58, 1.95, 8.75, 0.35, size=14, color=COLORS["muted"])
    flow(slide, ["问题", "范式", "事件码", "预处理", "ERP", "频域", "ML", "报告"], 3.60, 3.10, 8.85, h=0.62, size=9.5)
    card(slide, "本讲交付", "study_design、event_marker_table、data_structure、preprocessing_plan、preprocessing_qc、erp_analysis_plan、erp_component_table、frequency_connectivity_plan、eeg_ml_plan、results_template、analysis_audit。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=11.6)
    card(slide, "底线", "不覆盖原始数据、不看结果改窗口、不把成分直接等同心理构念、不隐藏 QC。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, [
        "不知道如何把理论问题转成可记录的 EEG/ERP 范式",
        "事件码设计混乱，后期无法分条件和 time-lock",
        "预处理参数随结果调整，滤波、ICA、窗口缺少先验规则",
        "ERP 时间窗、ROI、频段和连接解释过度",
        "QC 不透明，无法说明剔除被试、通道和 trial",
        "EEG 机器学习训练测试泄漏，准确率虚高",
    ], 0.85, 1.60, 5.95, 3.05, size=15.2)
    card(slide, "本讲不做", "不把 EEG 分析变成跑脚本；不追求漂亮波形；不把成分标签当心理构念。", 7.05, 1.72, 5.10, 1.22, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "建立从范式、事件码、预处理、分析、QC 到论文报告的可审计流程。", 7.05, 3.35, 5.10, 1.22, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["范式", "事件码", "数据结构", "预处理", "ERP", "频域连接", "ML", "报告"], 0.95, 5.65, 11.35, h=0.62, size=9.8)

    slide = new("本讲最终产出文件")
    rows = [
        ["阶段", "文件", "作用"],
        ["实验设计", "eeg/study_design.md", "研究问题、神经假设、范式、trial 时序和解释边界"],
        ["事件码", "eeg/event_marker_table.xlsx", "trigger、条件、反应、反馈和 time-locking"],
        ["数据结构", "eeg/data_structure.md", "raw、behavior、events、derivatives、scripts、qc"],
        ["预处理", "preprocessing_plan.md / preprocessing_qc.xlsx", "参数、坏道、ICA、trial 保留率和剔除原因"],
        ["ERP", "erp_analysis_plan.md / erp_component_table.xlsx", "epoch、baseline、窗口、ROI、指标和统计模型"],
        ["频域/连接", "frequency_connectivity_plan.md", "频段、ERSP、ITC、PLV/wPLI、图表和解释边界"],
        ["ML/报告", "eeg_ml_plan.md / eeg_results_template.md / eeg_analysis_audit.md", "分类、防泄漏、图表、写作和审稿审计"],
    ]
    table(slide, rows, 0.70, 1.46, 11.95, 4.98, col_widths=[1.55, 4.95, 5.45], font_size=8.6)
    card(slide, "提交底线", "每个结果都必须能追溯到事件码、预处理参数、QC 和先验分析计划。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [
        ["功能", "Skills", "用途"],
        ["实验规划", "eeg-experiment-planning", "研究问题、范式、事件码、采样率、电极、伦理"],
        ["预处理", "eeg-preprocessing-pipeline", "导入、滤波、重参考、坏道、ICA、QC"],
        ["ERP", "erp-segmentation-analysis", "epoch、baseline、trial 剔除、成分窗口、ROI、统计"],
        ["频域连接", "eeg-frequency-connectivity", "PSD、ERSP、ITC、coherence、PLV、wPLI"],
        ["机器学习", "eeg-ml-classification", "特征、交叉验证、防泄漏、指标和解释边界"],
        ["结果写作", "eeg-results-writing-figures", "Methods、Results、波形图、头皮图、QC 表"],
        ["项目规范", "lv-naodianshujuchul", "复现实验脚本、路径追踪、正式运行入口管理"],
    ]
    table(slide, rows, 0.68, 1.42, 12.0, 4.72, col_widths=[1.55, 4.6, 5.85], font_size=8.9)
    card(slide, "配套模板", "本讲提供事件码、预处理 QC、ERP 成分、频域连接、ML、防泄漏和审计 Excel 模板。", 0.85, 6.38, 11.55, 0.62, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=11.8)

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第10讲_EEG_ERP神经科学实验工作区"\ncd "D:\\AI科研训练营\\第10讲_EEG_ERP神经科学实验工作区"\n\nNew-Item -ItemType Directory -Force -Path data,data\\raw,data\\behavior,data\\events,data\\derivatives,eeg,eeg\\scripts,eeg\\figures,qc,output,skills,input\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md', 0.85, 1.55, 11.65, 2.48, size=9.4)
    card(slide, "检查标准", "原始 EEG 放 data/raw；行为和事件分开；所有衍生文件进入 derivatives。", 0.95, 4.42, 5.60, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "输入示例", "BrainVision、Neuroscan、EDF、BDF、set/fdt、fif；行为 CSV/XLSX；events.tsv。", 6.85, 4.42, 5.60, 1.05, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName eeg-experiment-planning,eeg-preprocessing-pipeline,erp-segmentation-analysis,eeg-frequency-connectivity,eeg-ml-classification,eeg-results-writing-figures,statistical-analysis `\n  -Workspace "D:\\AI科研训练营\\第10讲_EEG_ERP神经科学实验工作区"', 0.80, 1.48, 11.75, 2.82, size=8.7)
    bullets(slide, ["复制后检查 skills/ 文件夹。", "原始数据不要复制到 Skills 目录。", "若有真实脑电项目，记录代码根目录、数据根目录、Python/MATLAB 环境和结果根目录。"], 0.95, 4.75, 11.1, 1.0, size=14.2)
    card(slide, "课堂提醒", "本讲先做设计和审计，不强制跑真实 EEG 数据。", 0.95, 6.05, 11.1, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第10讲工作区 Agent 指令\n\n你是 EEG/ERP 神经科学实验助手。\n\n原则：\n- 禁止覆盖原始 EEG 数据。\n- 禁止看到结果后临时改滤波、窗口、ROI 或剔除规则。\n- 禁止把 ERP 成分、频段或连接指标直接等同心理构念。\n- ICA、坏道、被试和 trial 剔除必须写理由。\n- ML 必须先做泄漏审计。\n- 报告必须包含 QC 和解释边界。", 0.85, 1.55, 6.45, 3.50, size=9.4)
    card(slide, "为什么要写", "EEG 分析自由度很高，项目指令先把不能碰的边界固定下来。", 7.62, 1.70, 4.45, 1.10, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson10_workspace_AGENTS_template.md。", 7.62, 3.20, 4.45, 0.80, accent=COLORS["blue"])
    card(slide, "不能省略", "没有约束时，Agent 容易按“显著效果”倒推参数。", 7.62, 4.45, 4.45, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    slide = new("EEG/ERP 全流程地图")
    flow(slide, ["研究问题", "范式", "事件码", "采集", "预处理", "QC", "ERP/频域", "统计", "报告"], 0.78, 1.50, 11.8, h=0.62, size=9.2)
    card(slide, "关键逻辑", "先确定 time-locking event，再设计 trigger；先写预处理和窗口规则，再看结果；先通过 QC，再解释神经效应。", 0.95, 2.70, 11.35, 0.95, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=13.2)
    rows = [
        ["如果缺失", "后果"],
        ["事件码不清", "无法可靠分条件和锁时"],
        ["行为日志不同步", "无法解释反应和 trial"],
        ["预处理无记录", "审稿人无法复核"],
        ["窗口后验选择", "循环分析风险"],
        ["QC 不透明", "显著结果可信度低"],
    ]
    table(slide, rows, 2.05, 4.15, 9.25, 2.15, col_widths=[3.5, 5.75], font_size=10.5)

    step_slide(
        prs, 4, "实操 4：生成实验设计",
        "把理论问题转成可记录、可分析、可解释的 EEG/ERP 实验。",
        "请基于我的研究主题，生成 eeg/study_design.md。\n\n包括研究问题、理论机制、行为假设、神经假设、目标 ERP 成分或频段指标、实验范式、trial 时序、事件码设计原则、EEG 记录设置、预处理路线、分析路线、伦理和风险。不要把 ERP 成分直接等同心理构念。",
        "eeg/study_design.md",
        ["理论和范式匹配", "神经指标有任务时序", "行为指标清楚", "解释边界清楚", "伦理风险已列出"],
        "先定义神经假设，再选择成分或频段。"
    )

    step_slide(
        prs, 5, "实操 5：设计事件码表",
        "让每个 trial 能被条件、刺激、反应和反馈准确 time-lock。",
        "请生成 eeg/event_marker_table.xlsx 的内容草稿。\n\n字段包括 marker_code、event_name、condition、trial_phase、time_locking_role、expected_count、behavior_link、included_in_erp、included_in_frequency、notes。检查事件码重复、条件不可区分、无法 time-lock 和行为日志无法对齐的问题。",
        "eeg/event_marker_table.xlsx",
        ["marker 唯一", "条件可区分", "能锁定刺激/反应/反馈", "与 trial_id 对齐", "expected_count 可核验"],
        "事件码设计错了，后面分析很难补救。"
    )

    slide = new("事件码表应长什么样")
    rows = [
        ["marker", "event", "condition", "role", "phase", "behavior_link"],
        ["101", "low transparency stimulus", "low", "stimulus", "stimulus_onset", "trial_id"],
        ["102", "high transparency stimulus", "high", "stimulus", "stimulus_onset", "trial_id"],
        ["201", "accept response", "all", "response", "response", "trial_id + RT"],
        ["301", "positive feedback", "feedback_positive", "feedback", "feedback", "trial_id"],
        ["302", "negative feedback", "feedback_negative", "feedback", "feedback", "trial_id"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.10, col_widths=[1.4, 3.0, 2.4, 1.8, 2.0, 1.15], font_size=9.6)
    card(slide, "核验问题", "同一个 trigger 是否出现在多个条件？反馈是否能和具体 trial 对齐？反应码是否区分正确/错误或接受/拒绝？", 0.95, 5.00, 11.35, 0.85, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    step_slide(
        prs, 6, "实操 6：建立数据结构",
        "把 raw、behavior、events、derivatives、scripts、figures 和 qc 分开管理。",
        "请生成 eeg/data_structure.md。\n\n设计本项目的文件夹和命名规范。要求原始数据只读保存，所有衍生文件有版本号，subject/session/task/run 命名一致。示例：sub-001_task-aiTrust_raw.vhdr、sub-001_task-aiTrust_events.tsv。",
        "eeg/data_structure.md",
        ["原始数据只读", "行为和事件分开", "衍生数据可追踪", "命名含 subject/task/run", "脚本和图表有固定路径"],
        "文件结构混乱会直接破坏复现。"
    )

    slide = new("环境检查")
    codebox(slide, 'python -c "import sys; print(sys.version)"\npython -c "import mne; print(mne.__version__)"\npython -c "import numpy, pandas, scipy, sklearn; print(\'basic ok\')"\n\n# MATLAB / EEGLAB：记录 MATLAB 版本、EEGLAB 版本、插件版本、系统和数据格式。', 0.85, 1.55, 6.25, 2.20, size=9.2)
    card(slide, "Agent 任务", "生成 eeg/environment_check.md，记录 Python/MATLAB、MNE/EEGLAB、关键包版本、数据格式支持、缺失依赖和运行入口。", 7.35, 1.65, 4.85, 1.25, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "人工核验", "版本差异会影响滤波、ICA、数据导入和图表输出。真实项目必须记录环境。", 7.35, 3.38, 4.85, 1.05, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    step_slide(
        prs, 7, "实操 7：生成预处理计划",
        "固定导入、通道、滤波、参考、坏道、ICA、事件码和 epoch 前检查规则。",
        "请基于 study_design.md、event_marker_table.xlsx 和数据格式，生成 eeg/preprocessing_plan.md。\n\n包括原始数据保护、导入、通道名、montage、采样率、滤波、重参考、坏道、ICA、事件码完整性、epoch 前检查、剔除规则、输出命名和 reviewer-facing 预处理段落。",
        "eeg/preprocessing_plan.md",
        ["参数先验", "不覆盖 raw", "坏道规则清楚", "ICA 删除有证据", "事件码先检查"],
        "预处理计划写在看结果之前。"
    )

    slide = new("预处理顺序")
    flow(slide, ["备份 raw", "导入", "通道/事件", "滤波", "坏道", "重参考", "ICA", "epoch", "baseline", "QC"], 0.74, 1.48, 11.9, h=0.62, size=8.6)
    rows = [
        ["环节", "风险"],
        ["滤波", "参数不适合 ERP/频域，可能改变波形或边界"],
        ["坏道", "ROI 通道被插值过多会影响解释"],
        ["ICA", "删除成分无依据会被质疑"],
        ["epoch", "事件码没核验就分段会全错"],
        ["baseline", "窗口不合适会影响成分幅度"],
    ]
    table(slide, rows, 1.05, 2.80, 11.10, 2.85, col_widths=[2.2, 8.9], font_size=10.2)
    card(slide, "原则", "不同条件不能用不同预处理标准，除非有清楚理由并透明报告。", 1.05, 6.10, 11.10, 0.58, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)

    step_slide(
        prs, 8, "实操 8：生成预处理 QC 表",
        "记录每个被试的坏道、ICA、事件、epoch、保留率和剔除原因。",
        "请生成 eeg/preprocessing_qc.xlsx 的内容草稿。\n\n每个被试记录 raw_file、sampling_rate、bad_channels、interpolated_channels、rejected_ica_components、total_events、valid_events、total_epochs、retained_epochs、retention_rate、exclusion_reason、qc_decision。",
        "eeg/preprocessing_qc.xlsx",
        ["每个被试一行", "坏道/ICA 有记录", "事件和 epoch 数清楚", "保留率可计算", "剔除理由透明"],
        "不报告 QC，就不要报告漂亮 ERP 图。"
    )

    step_slide(
        prs, 9, "实操 9：生成 ERP 分析计划",
        "固定 epoch、baseline、成分、ROI、时间窗、指标和统计模型。",
        "请生成 eeg/erp_analysis_plan.md。\n\n包括 time-locking event、epoch window、baseline window、artifact rejection、condition averaging、ERP component、ROI electrodes、time window、amplitude/latency metric、statistical model、figure plan、confirmatory/exploratory 标记。",
        "eeg/erp_analysis_plan.md",
        ["time-locking 明确", "窗口先验", "ROI 有依据", "指标清楚", "探索/确认区分"],
        "看波形后选最大差异窗口，只能标探索性。"
    )

    slide = new("ERP 成分表")
    rows = [
        ["成分", "常见窗口", "常见 ROI", "典型用途"],
        ["P1/N1", "80-180 ms", "枕区/顶枕区", "早期感觉/注意线索"],
        ["N2", "200-350 ms", "额中区", "冲突或抑制相关线索"],
        ["P3", "300-600 ms", "顶区/中央顶区", "注意资源或任务相关加工线索"],
        ["N400", "300-500 ms", "中央-顶区", "语义整合线索"],
        ["LPP", "400-800 ms 或更晚", "中央-顶区", "情绪/动机相关持续加工线索"],
        ["FRN/ERN", "200-350 ms / 0-100 ms", "额中区", "反馈或错误监控线索"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 4.05, col_widths=[1.8, 2.8, 2.4, 4.75], font_size=9.2)
    card(slide, "课堂提醒", "这些是教学示例。真实研究必须用任务时序和真实文献确定窗口与 ROI。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    slide = new("ERP 图表和统计")
    rows = [
        ["输出", "必须包含"],
        ["ERP waveform", "time zero、条件、ROI、CI/SE、时间窗标记"],
        ["Scalp map", "成分窗口、统一色标、条件或差异图"],
        ["Trial retention table", "每条件保留 trial 数和剔除比例"],
        ["Component table", "被试 × 条件 × 成分指标"],
        ["Statistics", "组内/组间模型、效应量、CI、多重比较控制"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.70, col_widths=[3.0, 8.75], font_size=10.3)
    card(slide, "底线", "不要只展示显著电极、显著窗口或好看的波形。planned tests 即使不显著也要报告。", 0.95, 5.75, 11.35, 0.78, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(
        prs, 10, "实操 10：频域 / 时频 / 连接计划",
        "定义频段、方法、baseline、ROI、时间窗、统计模型和连接解释边界。",
        "请生成 eeg/frequency_connectivity_plan.md。\n\n包括目标频段、PSD/STFT/wavelet/ERSP/ITC/coherence/PLV/wPLI、baseline、归一化、ROI、时间窗、频段、多重比较、图表计划和解释边界。强调连接不是因果方向。",
        "eeg/frequency_connectivity_plan.md",
        ["频段有依据", "baseline 清楚", "ROI/窗口先验", "多重比较处理", "连接不写因果方向"],
        "频段和连接解释比 ERP 更容易过度。"
    )

    slide = new("连接分析的解释边界")
    rows = [
        ["指标", "能说什么", "不能说什么"],
        ["coherence", "频域同步或相关线索", "不能自动说明信息流方向"],
        ["PLV", "相位一致性线索", "不能说明因果驱动"],
        ["wPLI", "降低零相位体积传导影响的连接线索", "不能完全解决源混叠"],
        ["graph metrics", "网络拓扑描述", "不能直接等同认知机制"],
        ["source-space connectivity", "可能更接近源空间", "依赖头模型和源定位假设"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.85, col_widths=[2.4, 4.55, 4.80], font_size=9.6)
    card(slide, "审稿风险", "把功能连接写成因果方向，通常会被直接质疑。", 0.95, 5.80, 11.35, 0.75, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.8)

    step_slide(
        prs, 11, "实操 11：EEG 机器学习计划",
        "构建最小可行 EEG 分类方案，同时防止训练测试泄漏。",
        "请生成 eeg/eeg_ml_plan.md。\n\n包括 prediction target、label 来源、feature 类型、被试内/跨被试/留一被试验证、train/validation/test 划分、防泄漏规则、baseline model、deep model 是否必要、metrics、permutation test/nested CV 和解释边界。",
        "eeg/eeg_ml_plan.md",
        ["标签来源清楚", "按被试隔离", "标准化在 CV 内", "有 baseline", "指标不只 accuracy"],
        "预测准确率不是神经机制证据。"
    )

    slide = new("EEG 机器学习防泄漏")
    rows = [
        ["泄漏类型", "例子", "修正"],
        ["epoch 泄漏", "同一被试 epoch 同时进 train/test", "按 subject 分组划分"],
        ["标准化泄漏", "全数据 fit scaler 后再划分", "scaler 放进 CV pipeline"],
        ["调参泄漏", "在 test set 上调超参", "validation 或 nested CV"],
        ["标签泄漏", "特征中包含 marker/condition code", "特征审计并删除"],
        ["重复 trial 泄漏", "增强数据跨集合重复", "增强只在训练集内做"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.80, col_widths=[2.4, 4.95, 4.40], font_size=9.7)
    card(slide, "课堂判断", "如果声称跨被试泛化，就不能把同一被试的 epoch 放入测试集。", 0.95, 5.80, 11.35, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.8)

    step_slide(
        prs, 12, "实操 12：结果写作与图表模板",
        "把记录、预处理、QC、ERP、频域/连接、ML 和限制写成论文可用模板。",
        "请生成 eeg/eeg_results_template.md。\n\n包括 Participants and exclusions、EEG recording、Preprocessing、ERP analysis、Frequency/time-frequency/connectivity analysis、ML analysis、QC disclosure、Figure plan、Table plan、Results paragraph template、Limitations and reviewer risks。",
        "eeg/eeg_results_template.md",
        ["报告剔除", "报告 retained trials", "图表与计划一致", "非显著也报告", "解释不超出指标"],
        "写作不是美化结果，而是透明报告证据。"
    )

    slide = new("课堂示范：AI Agent 解释透明度")
    rows = [
        ["模块", "示范内容"],
        ["研究问题", "解释透明度是否影响用户对 AI 推荐结果的信任和采纳"],
        ["范式", "阅读 AI 推荐与解释 → 采纳/拒绝 → 反馈"],
        ["事件码", "低/高透明度刺激、接受/拒绝反应、正/负反馈"],
        ["ERP", "刺激锁定 P3/LPP；反馈锁定 FRN/P3"],
        ["频域", "额中 theta 或 alpha 变化作为探索性指标"],
        ["边界", "P3/LPP 不能直接写成信任；只能写与注意或动机加工线索一致"],
    ]
    table(slide, rows, 0.78, 1.45, 11.75, 4.95, col_widths=[2.0, 9.75], font_size=9.7)
    card(slide, "最大风险", "阅读材料会引入眼动和眨眼伪迹，trial 时序与刺激长度必须控制。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("常见错误与修正")
    rows = [
        ["错误", "后果", "修正"],
        ["事件码不区分条件", "无法做 ERP 条件比较", "重设计 trigger 表"],
        ["看结果选窗口", "循环分析风险", "先验窗口或标探索性"],
        ["ICA 删除无记录", "预处理不可复核", "ICA log + 证据"],
        ["只报显著电极", "选择性报告", "planned ROI + 多重比较"],
        ["QC 不透明", "结果可信度低", "报告 retained trials 和 attrition"],
        ["ML epoch 泄漏", "准确率虚高", "按被试分组划分"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.55, 4.10, 5.10], font_size=9.5)

    step_slide(
        prs, 13, "实操 13：审稿人式审计",
        "用严格审稿人视角检查设计、事件码、预处理、窗口、连接、ML 和写作边界。",
        "请生成 qc/eeg_analysis_audit.md。\n\n检查研究问题与 EEG 指标匹配、事件码完整性、行为和 EEG 对齐、预处理先验、ICA/坏道依据、ERP 窗口和 ROI、频域/连接解释、ML 泄漏、QC 透明和写作是否过度。输出 pass/revise/fail。",
        "qc/eeg_analysis_audit.md",
        ["指出致命问题", "列出修改项", "标注探索性分析", "检查泄漏", "给出 pass/revise/fail"],
        "审计不通过，不建议进入正式论文结果。"
    )

    slide = new("课堂 135 分钟带做安排")
    rows = [
        ["时间", "教师带做", "学生产出"],
        ["0-18 分钟", "讲清 EEG/ERP 可信结果链路", "理解交付物"],
        ["18-30 分钟", "创建工作区、复制 Skills、写 AGENTS.md", "可运行工作区"],
        ["30-45 分钟", "研究问题、范式和神经假设", "study_design"],
        ["45-60 分钟", "事件码与 trial 时序", "event_marker_table"],
        ["60-78 分钟", "预处理计划", "preprocessing_plan"],
        ["78-92 分钟", "ERP 分析计划和成分表", "erp_analysis_plan / component_table"],
        ["92-105 分钟", "频域/连接计划", "frequency_connectivity_plan"],
        ["105-125 分钟", "ML 防泄漏和结果写作", "eeg_ml_plan / results_template"],
        ["125-135 分钟", "审稿人式审计", "eeg_analysis_audit"],
    ]
    table(slide, rows, 0.78, 1.45, 11.75, 5.05, col_widths=[1.8, 5.0, 4.95], font_size=8.7)
    card(slide, "教师重点", "现场纠正三件事：事件码不完整、窗口后验选择、神经解释过度。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12)

    slide = new("课后提交要求")
    bullets(slide, [
        "提交文件夹：姓名_第10讲_EEG_ERP神经科学实验",
        "必须包含 study_design、event_marker_table、data_structure、preprocessing_plan、preprocessing_qc、erp_analysis_plan、erp_component_table、frequency_connectivity_plan、eeg_ml_plan、eeg_results_template、eeg_analysis_audit",
        "至少完成一种主分析路线：ERP、频域/时频、连接或 EEG 机器学习",
        "附 300 字反思：最大审稿风险是什么，准备用什么 QC 或稳健性证据回应",
    ], 0.90, 1.55, 11.3, 1.75, size=14.4)
    rows = [
        ["评分项", "占比"],
        ["实验设计与神经假设", "20%"],
        ["事件码和数据结构", "20%"],
        ["预处理计划与 QC", "25%"],
        ["ERP/频域/连接/ML 分析计划", "25%"],
        ["写作模板与审稿审计", "10%"],
    ]
    table(slide, rows, 2.15, 3.80, 8.85, 2.20, col_widths=[6.6, 2.25], font_size=11.5)
    card(slide, "一票否决", "覆盖原始数据、事件码无法区分条件、看结果选窗口、ICA/QC 无记录、ML 泄漏。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.1)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是跑出一张脑电图，而是建立一套能让审稿人复核的 EEG/ERP 实验与分析系统。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["理论", "范式", "事件码", "预处理", "QC", "ERP", "频域", "ML", "报告"], 1.00, 3.50, 11.20, h=0.70, size=9.2)
    card(slide, "下一步", "学生把第10讲的 EEG/ERP 变量和结果接入第8讲统计分析，或进入第11讲 BCI 脑电智能建模。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
