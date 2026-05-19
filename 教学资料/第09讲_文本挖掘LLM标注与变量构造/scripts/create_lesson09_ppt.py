from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第09讲_课件.pptx"
W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}
FONT = "Microsoft YaHei"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, level = item if isinstance(item, tuple) else (item, 0)
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=9.0):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 09 讲 · 文本挖掘与变量构造", 0.45, 0.055, 3.3, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])
    return rect


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.2, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
    return shape


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=10):
    gap = 0.12
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.01), Inches(y + h / 2), Inches(bx - 0.02), Inches(y + h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"])
    no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.16, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.18, accent=COLORS["amber"], body_size=11.2, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.72, size=8.55)
    textbox(slide, "人工核验", 4.65, 4.82, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.3)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 09 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "文本挖掘\nLLM 标注\n变量构造", 0.52, 1.70, 2.25, 1.65, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "从文本到可分析变量", 0.55, 6.10, 2.2, 0.35, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "让评论、公告、访谈、社媒和摘要进入严谨实证研究", 3.55, 1.10, 8.75, 0.72, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "Cleaning · Segmentation · Codebook · LLM Labeling · Reliability · Features · Audit", 3.58, 1.95, 8.75, 0.35, size=14, color=COLORS["muted"])
    flow(slide, ["文本", "清洗", "构念", "标注", "可靠性", "特征", "变量表", "实证模型"], 3.60, 3.10, 8.85, h=0.62, size=9.5)
    card(slide, "本讲交付", "text_data_inventory、text_cleaning_plan、text_variable_definition、annotation_codebook、annotation_sample、feature_dictionary、features、reliability_report、variable_construction_report、text_variable_audit。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=11.6)
    card(slide, "底线", "不编造标签、不泄露隐私、不把词云当变量、不让 LLM 替代人工核验。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, [
        "有评论、公告、访谈、社媒或摘要文本，但不知道如何变成变量",
        "会做词云、词频和主题，但解释不了理论构念",
        "想用 LLM 标注，但没有 codebook、样本复核和一致性指标",
        "情感分数、主题概率和关键词频率直接进入回归，缺少有效性说明",
        "训练、测试和文本合并过程中出现数据泄漏",
        "最终结果不能和第 08 讲统计分析衔接",
    ], 0.85, 1.60, 5.95, 3.05, size=15.2)
    card(slide, "本讲不做", "不把 NLP 做成炫技；不把 LLM 标签当成客观事实；不只交词云和主题词。", 7.05, 1.72, 5.10, 1.22, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "把文本转成可解释、可复核、可合并、可进入实证模型的变量。", 7.05, 3.35, 5.10, 1.22, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["盘点", "清洗", "构念", "路线", "标注", "审计", "变量"], 0.95, 5.65, 11.35, h=0.62, size=11)

    slide = new("本讲最终产出文件")
    rows = [
        ["阶段", "文件", "作用"],
        ["数据盘点", "textmining/text_data_inventory.xlsx", "文档 ID、文本列、来源、时间、风险和分析单位"],
        ["清洗", "textmining/text_cleaning_plan.md", "去噪、分词、停用词、自定义词典和输出路径"],
        ["构念", "textmining/text_variable_definition.md", "说明文本变量代表什么理论构念"],
        ["标注", "annotation_codebook.md / annotation_sample.xlsx", "标签定义、正例、反例、边界例和样本标注"],
        ["可靠性", "textmining/reliability_report.md", "人工-人工、人工-LLM、模型-人工一致性和错误类型"],
        ["变量", "feature_dictionary.xlsx / features.xlsx", "最终可合并文本变量和构造说明"],
        ["审计", "qc/text_variable_audit.md", "隐私、泄漏、构念错配、过度解释和可进入模型性"],
    ]
    table(slide, rows, 0.70, 1.46, 11.95, 4.98, col_widths=[1.55, 4.95, 5.45], font_size=8.8)
    card(slide, "提交底线", "最终变量必须能追踪到文本来源、清洗规则、构念定义、标注准则和可靠性审计。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [
        ["功能", "Skills", "用途"],
        ["基础处理", "text-analysis-basic", "PDF/文本抽取、分词、词频、TF-IDF、embedding、相似度"],
        ["主题模型", "topic-modeling", "LDA、DTM、BERTopic、主题词、文档主题概率"],
        ["情感分析", "sentiment-analysis", "词典法、机器学习、LLM 情感标注"],
        ["标注与变量", "big-data-labeling-variable-construction", "LDA、sklearn、预训练模型、OpenAI 兼容 LLM 标注路由"],
        ["深度学习", "deep-learning-nlp", "DNN/RNN/GRU 文本分类训练和预测"],
        ["衔接分析", "empirical-analysis-skill-python", "变量合并、分析数据表和第 08 讲模型衔接"],
        ["可视化", "data-visualization-analysis", "词频图、主题图、情感趋势和变量分布"],
    ]
    table(slide, rows, 0.68, 1.42, 12.0, 4.72, col_widths=[1.55, 4.8, 5.65], font_size=8.8)
    card(slide, "配套模板", "本讲提供 6 个 Excel 模板：数据盘点、标注样本、特征字典、结果汇总、变量审计、运行日志。", 0.85, 6.38, 11.55, 0.62, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=11.8)

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第09讲_文本挖掘LLM标注与变量构造工作区"\ncd "D:\\AI科研训练营\\第09讲_文本挖掘LLM标注与变量构造工作区"\n\nNew-Item -ItemType Directory -Force -Path data,data\\raw,data\\clean,data\\samples,textmining,textmining\\scripts,textmining\\figures,qc,output,skills,input\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md', 0.85, 1.55, 11.65, 2.48, size=9.4)
    card(slide, "检查标准", "原始文本放 data/raw；清洗后文本放 data/clean；标注、变量和审计都放 textmining 与 qc。", 0.95, 4.42, 5.60, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "输入示例", "评论、年报、公告、新闻、社媒、访谈、开放题回答、论文摘要。", 6.85, 4.42, 5.60, 1.05, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName text-analysis-basic,topic-modeling,sentiment-analysis,big-data-labeling-variable-construction,deep-learning-nlp,empirical-analysis-skill-python,data-visualization-analysis,citation-management,literature-review `\n  -Workspace "D:\\AI科研训练营\\第09讲_文本挖掘LLM标注与变量构造工作区"', 0.80, 1.48, 11.75, 2.82, size=8.7)
    bullets(slide, ["复制后检查 skills/ 文件夹。", "API key 通过环境变量读取，不写进文件。", "缺失 Skill 写入 research_log.md，并用现有路线替代。"], 0.95, 4.75, 11.1, 1.0, size=15)
    card(slide, "课堂提醒", "文本挖掘任务容易涉及隐私和平台协议，先审计文本来源，再批量处理。", 0.95, 6.05, 11.1, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第09讲工作区 Agent 指令\n\n你是文本挖掘、LLM 标注与变量构造助手。\n\n原则：\n- 禁止编造文本、标签、准确率、一致性指标。\n- 禁止把词频、主题或情感分数直接等同理论构念。\n- 原始文本只读保存，清洗输出新文件。\n- LLM 标注必须遵守 codebook。\n- API key 只通过环境变量读取。\n- 最终变量必须可追踪到输入、清洗、标注和审计。", 0.85, 1.55, 6.45, 3.50, size=9.45)
    card(slide, "为什么要写", "第9讲最容易出现 LLM 越界推断、隐私泄露、变量解释过度和数据泄漏。", 7.62, 1.70, 4.45, 1.10, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson09_workspace_AGENTS_template.md。", 7.62, 3.20, 4.45, 0.80, accent=COLORS["blue"])
    card(slide, "不能省略", "没有边界时，Agent 会把容易计算的文本结果包装成研究变量。", 7.62, 4.45, 4.45, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    step_slide(
        prs, 4, "实操 4：文本数据盘点",
        "确认文本文件、文档 ID、文本列、来源、时间、风险和分析单位。",
        "请读取 data/raw/ 下的文本数据，生成 textmining/text_data_inventory.xlsx 的内容草稿。\n\n检查文件类型、编码、行数、doc_id、text、source、date、重复、空文本、异常长短文本、个人信息、敏感信息、版权风险和最小分析单位。不要修改原始数据。",
        "textmining/text_data_inventory.xlsx",
        ["doc_id 唯一", "文本列明确", "来源和时间可追踪", "风险已标注", "原始文件未覆盖"],
        "没有 ID 和来源的文本，不建议进入实证变量构造。"
    )

    slide = new("文本伦理、隐私与版权")
    rows = [
        ["风险", "常见场景", "处理方式"],
        ["个人信息", "访谈、社媒、评论含姓名/账号/邮箱", "脱敏、最小化使用、限制共享原文"],
        ["敏感信息", "健康、政治、宗教、身份属性", "伦理审查、权限控制、必要性说明"],
        ["版权/协议", "平台文本、付费报告、论文全文", "检查授权、只输出变量和摘要"],
        ["API 传输", "把原文发给 LLM 服务", "脱敏、确认服务条款、控制日志"],
        ["再识别风险", "少量文本可识别个体", "聚合变量、删除原文片段"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 4.15, col_widths=[2.0, 4.85, 4.90], font_size=9.7)
    card(slide, "课堂底线", "不能为了方便标注把未脱敏访谈或社媒个人信息直接发给外部模型。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(
        prs, 5, "实操 5：生成清洗计划",
        "明确去噪、分句、分词、停用词、自定义词典、去重和输出规则。",
        "请基于 text_data_inventory.xlsx，生成 textmining/text_cleaning_plan.md。\n\n写清保留列、去除噪声、分句/分段、中文分词、自定义词典、停用词、专业词保护、去重、异常文本处理、脱敏规则和输出路径 data/clean/text_clean_v1.csv。",
        "textmining/text_cleaning_plan.md",
        ["清洗步骤可复现", "专业词不被拆错", "停用词不误删构念词", "去重规则清楚", "输出新文件"],
        "清洗规则会直接改变后续变量，不能随手处理。"
    )

    slide = new("中文分词与自定义词典")
    codebox(slide, 'python .\\skills\\text-analysis-basic\\scripts\\jieba_tokenize.py `\n  .\\data\\clean\\text_clean_v1.csv `\n  --text-column text `\n  --output .\\textmining\\tokenized_text.csv\n\npython .\\skills\\text-analysis-basic\\scripts\\frequencies.py `\n  .\\textmining\\tokenized_text.csv `\n  --text-column tokens `\n  --output .\\textmining\\word_frequency.xlsx', 0.85, 1.55, 6.25, 2.60, size=9.2)
    card(slide, "自定义词典", "AI Agent、ChatGPT、脑机接口、神经营销、Zotero、Obsidian、ERP、LLM 标注等专业词要保护。", 7.35, 1.70, 4.85, 1.20, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "人工核验", "抽看 50 条分词结果。错分的专业词、机构名、产品名和理论构念词要进入自定义词典。", 7.35, 3.35, 4.85, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "输出", "tokenized_text.csv、word_frequency.xlsx、清洗日志。", 0.95, 5.25, 11.35, 0.72, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)

    step_slide(
        prs, 6, "实操 6：定义文本变量",
        "把目标变量从“好算的指标”转成“有理论解释的文本构念”。",
        "请基于研究问题和文本数据，生成 textmining/text_variable_definition.md。\n\n每个变量包含变量名、理论构念、构念定义、文本证据、正例、反例、边界例、推荐技术路线、最终分析单位、可进入模型和不能支持的解释。",
        "textmining/text_variable_definition.md",
        ["构念定义明确", "有文本证据", "有反例和边界例", "单位可合并", "解释不过度"],
        "文本变量首先是测量问题，其次才是模型问题。"
    )

    slide = new("技术路线选择表")
    rows = [
        ["研究需求", "推荐路线", "主要风险"],
        ["看高频词和专业词", "分词 + 词频 / TF-IDF", "词频不能直接代表理论变量"],
        ["已有理论词典", "词典法", "词典领域适配不足"],
        ["态度、情绪、风险语气", "情感分析 / 词典 / LLM", "讽刺、否定、领域词误判"],
        ["探索文本主题", "LDA / BERTopic", "主题命名硬贴理论"],
        ["已有标签样本", "sklearn / BERT / ERNIE", "训练测试泄漏、类别不平衡"],
        ["复杂语义判断", "LLM 标注 + 人工复核", "越界推断、标签漂移、成本"],
        ["比较相似度", "TF-IDF / embedding similarity", "相似度不等于同一构念"],
    ]
    table(slide, rows, 0.72, 1.46, 11.95, 4.80, col_widths=[3.0, 4.2, 4.75], font_size=9.1)
    card(slide, "判断句", "技术路线由研究问题、构念和数据决定，不由模型新旧决定。", 0.85, 6.45, 11.55, 0.58, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.2)

    slide = new("词典法与 TF-IDF：可解释但要克制")
    bullets(slide, [
        "词典法适合已有清晰词表的构念，例如风险、不确定性、前瞻性、创新关注。",
        "TF-IDF 适合关键词特征、相似度和轻量分类基线。",
        "词频比例通常要除以文本长度或有效词数。",
        "词典要记录来源、增删理由和人工核验。",
        "TF-IDF 高权重词是特征，不自动等于理论构念。",
    ], 0.90, 1.55, 5.95, 2.72, size=14.5)
    codebox(slide, 'python .\\skills\\text-analysis-basic\\scripts\\vectorize.py `\n  .\\data\\clean\\text_clean_v1.csv `\n  --method tfidf `\n  --text-column text `\n  --output .\\textmining\\tfidf_features.xlsx', 7.05, 1.62, 5.10, 1.45, size=9)
    card(slide, "输出", "tfidf_features.xlsx、feature_dictionary.xlsx、变量构造报告中的词典/TF-IDF 说明。", 7.05, 3.55, 5.10, 1.00, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)
    card(slide, "人工核验", "抽看高低分文本，确认变量真的捕捉目标构念。", 7.05, 4.88, 5.10, 0.82, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)

    slide = new("情感分析：先验证领域适配")
    rows = [
        ["路线", "适合", "核验"],
        ["词典法", "词语情感方向清晰、需解释", "领域词、否定词、程度副词、讽刺"],
        ["机器学习", "已有人工标签样本", "训练/测试隔离、类别平衡、F1"],
        ["LLM 情感标注", "语义复杂、样本量中等", "抽样复核、一致性、成本"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 2.50, col_widths=[2.3, 4.6, 4.85], font_size=10.5)
    codebox(slide, 'python .\\skills\\sentiment-analysis\\scripts\\dictionary_sentiment.py `\n  .\\data\\clean\\text_clean_v1.csv `\n  --text-column text `\n  --output .\\textmining\\sentiment_scores.xlsx', 0.92, 4.35, 6.10, 1.35, size=9)
    card(slide, "常见问题", "“低风险”“不差”“模型不会保证正确”等句子容易被简单词典误判。", 7.35, 4.18, 4.95, 1.05, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)

    slide = new("主题模型：探索主题，不自动生成理论")
    flow(slide, ["清洗", "分词", "诊断 K", "看主题词", "看代表文档", "命名主题", "导出概率"], 0.82, 1.45, 11.7, h=0.62, size=9.8)
    codebox(slide, 'python .\\skills\\topic-modeling\\scripts\\lda_model.py `\n  .\\data\\clean\\text_clean_v1.csv `\n  --text-column text `\n  --topic-range 2-10 `\n  --output-dir .\\textmining\\lda_diagnostics', 0.92, 2.45, 6.15, 1.70, size=9)
    card(slide, "关键规则", "若未明确 K，先跑诊断和候选主题，人工选择 K 后再训练最终模型。", 7.35, 2.35, 4.95, 1.05, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12)
    card(slide, "人工核验", "主题命名必须看代表文档，不只看 top words。主题概率可作为变量，但解释要克制。", 7.35, 3.85, 4.95, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)
    card(slide, "输出", "topic_summary、topic_doc_probs、可视化图和 feature_dictionary。", 0.92, 5.35, 11.38, 0.70, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)

    slide = new("监督分类与深度学习")
    rows = [
        ["路线", "输入", "适用场景", "核心核验"],
        ["sklearn", "文本 + 标签", "轻量基线、可解释、速度快", "train/test split、F1、错误类型"],
        ["BERT/ERNIE", "文本 + 标签 + 模型", "语义复杂、标签样本较多", "依赖环境、GPU、泄漏、过拟合"],
        ["DNN/RNN/GRU", "课程文本分类样例", "教学演示和基线扩展", "保存模型、可复现日志"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.15, col_widths=[2.0, 2.7, 3.6, 3.45], font_size=9.5)
    card(slide, "防泄漏", "同一用户、企业、事件或重复模板文本不能跨训练集和测试集。需要 group split 或 time split。", 0.95, 5.05, 5.45, 0.95, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)
    card(slide, "输出", "模型指标、混淆矩阵、错误分析、预测标签和特征字典。", 6.75, 5.05, 5.45, 0.95, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)

    slide = new("LLM 标注流程：先计划，后执行")
    flow(slide, ["Codebook", "抽样", "试标注", "人工复核", "修订", "批量标注", "可靠性"], 0.82, 1.45, 11.7, h=0.62, size=9.8)
    codebox(slide, 'python .\\skills\\big-data-labeling-variable-construction\\scripts\\prepare_labeling_plan.py `\n  .\\data\\clean\\text_clean_v1.csv `\n  --category openai_llm `\n  --output-dir .\\textmining\\llm_labeling_plan\n\n# 展示计划后，学生确认：确认无误，请执行', 0.92, 2.45, 6.15, 2.05, size=8.9)
    card(slide, "安全", "API key 用环境变量；脱敏后再调用外部模型；输出中不保留敏感原文。", 7.35, 2.42, 4.95, 1.05, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12)
    card(slide, "Skill 规则", "该工作流要求先检查前 10 行并展示推断计划，只有学生回复“确认无误，请执行”后才运行。", 7.35, 3.88, 4.95, 1.20, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=11.8)

    step_slide(
        prs, 7, "实操 7：生成 Codebook",
        "让不同标注者能按同一标准判断文本标签。",
        "请基于 text_variable_definition.md，生成 textmining/annotation_codebook.md。\n\n每个标签包含标签名、定义、标注单位、标 1 条件、标 0 条件、正例、反例、边界例、冲突规则、不确定规则和禁止文本外推断规则。",
        "textmining/annotation_codebook.md",
        ["定义具体", "有正反例", "有边界例", "允许 uncertain", "禁止文本外推断"],
        "Codebook 写不清，后面的准确率没有意义。"
    )

    step_slide(
        prs, 8, "实操 8：抽样人工标注",
        "用小样本校准标签定义，避免批量标注后才发现规则错。",
        "请生成 textmining/annotation_sampling_plan.md，并填写 annotation_sample.xlsx。\n\n说明抽样数量、随机种子、分层变量、覆盖来源/时间/评分、纳入排除规则。样本表包含 doc_id、text、human_label_1、human_label_2、llm_label、final_label、disagreement_reason。",
        "textmining/annotation_sample.xlsx",
        ["样本覆盖不同来源", "人工标签独立", "差异原因记录", "边界例保留", "最终标签可追踪"],
        "先标 50-200 条校准，再决定是否批量。"
    )

    step_slide(
        prs, 9, "实操 9：可靠性审计",
        "判断人工标注、LLM 标注或模型标签是否足以用于变量构造。",
        "请读取 annotation_sample.xlsx，生成 textmining/reliability_report.md。\n\n报告人工-人工一致率、Cohen's Kappa 或 Krippendorff's Alpha、LLM-人工一致率、precision/recall/F1、错误类型、边界例、系统性偏差、codebook 修改建议和是否可批量标注。",
        "textmining/reliability_report.md",
        ["报告一致性指标", "分析错误类型", "修订 codebook", "判断是否批量", "低可靠性变量不进入主模型"],
        "可靠性不足时，不要硬把标签放进回归。"
    )

    step_slide(
        prs, 10, "实操 10：生成特征字典与变量表",
        "把分词、词典、情感、主题、标注和分类结果整理成最终文本变量。",
        "请基于清洗文本和模型/标注输出，生成 textmining/feature_dictionary.xlsx 与 textmining/features.xlsx。\n\n每个变量说明 feature_name、construct、method、input_file、parameters、unit、range、aggregation_rule、interpretation、limitations、audit_status。",
        "feature_dictionary.xlsx\nfeatures.xlsx",
        ["变量有合并键", "方法和参数清楚", "取值范围清楚", "解释边界清楚", "可进入第8讲模型"],
        "最终表要像实证数据表，而不是 NLP 输出碎片。"
    )

    slide = new("数据泄漏审计")
    rows = [
        ["泄漏类型", "例子", "处理"],
        ["重复文本泄漏", "同一评论同时在训练集和测试集", "去重后划分"],
        ["作者泄漏", "同一用户或企业跨训练测试", "按 user_id / firm_id 分组划分"],
        ["时间泄漏", "用未来文本预测过去结果", "time split / lag features"],
        ["标签泄漏", "文本中直接包含评分或结果", "删除泄漏字段或改任务"],
        ["模板泄漏", "公告/年报模板句重复出现", "模板识别和去除"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.80, col_widths=[2.3, 5.0, 4.45], font_size=9.7)
    card(slide, "课堂判断", "如果模型效果异常高，第一反应不是高兴，而是检查泄漏。", 0.95, 5.80, 11.35, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.8)

    slide = new("与第 08 讲统计分析衔接")
    rows = [
        ["分析单位", "合并键", "文本变量示例"],
        ["评论级", "doc_id", "sentiment_score、uncertainty_label、topic_prob"],
        ["用户级", "user_id", "risk_expression_ratio、avg_sentiment"],
        ["企业-年级", "firm_id + year", "ai_attention_score、innovation_topic_prob"],
        ["事件级", "event_id/date", "policy_similarity、media_tone"],
        ["论文级", "paper_id/year", "method_topic_prob、novelty_signal"],
    ]
    table(slide, rows, 0.78, 1.48, 11.75, 3.80, col_widths=[2.3, 3.1, 6.35], font_size=9.8)
    card(slide, "输出给第8讲", "features.xlsx + feature_dictionary.xlsx + variable_construction_report.md。", 0.95, 5.80, 11.35, 0.80, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.8)

    slide = new("课堂示范：AI Agent 工具评论")
    rows = [
        ["模块", "示范内容"],
        ["文本", "学生或研究者对 AI Agent 工具的使用评论"],
        ["变量 1", "efficiency_benefit：节省时间、减少重复劳动、提高检索效率"],
        ["变量 2", "uncertainty_expression：可能出错、需核验、不确定来源"],
        ["变量 3", "trust_signal：可信、可靠、愿意采纳、可作为辅助依据"],
        ["路线", "Codebook + LLM 标注 + 人工复核；同时构造情感分数和文本长度控制变量"],
        ["进入模型", "用文本变量解释采纳意愿、满意度或后续使用行为"],
    ]
    table(slide, rows, 0.78, 1.45, 11.75, 4.95, col_widths=[2.0, 9.75], font_size=9.7)
    card(slide, "最大风险", "文本中的信任表达不能直接等同真实信任心理状态，需要和问卷或行为指标区分。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("常见错误与修正")
    rows = [
        ["错误", "后果", "修正"],
        ["只有词云", "不能进入实证模型", "构造文档级/单位级变量"],
        ["主题硬贴理论", "理论解释站不住", "看代表文档和文献依据"],
        ["LLM 自创标签", "标签体系漂移", "固定 codebook 和 JSON 输出"],
        ["无人工复核", "可靠性不可判断", "抽样标注和一致性审计"],
        ["训练测试泄漏", "指标虚高", "group split / time split"],
        ["把文本变量当事实", "解释过度", "写清测量边界"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.55, 4.10, 5.10], font_size=9.5)

    step_slide(
        prs, 11, "实操 11：变量构造报告与审计",
        "把文本变量生产线写成论文可用的 Method 说明，并做审稿人式审计。",
        "请生成 textmining/variable_construction_report.md 和 qc/text_variable_audit.md。\n\n报告文本来源、清洗流程、构念定义、技术路线、标注/模型、可靠性结果、最终变量、合并方式、有效性证据、稳健性和解释边界。审计 pass/revise/fail。",
        "variable_construction_report.md\nqc/text_variable_audit.md",
        ["来源合规", "流程可复现", "可靠性达标", "无泄漏", "解释边界清楚"],
        "没有审计通过的文本变量，不建议进入主模型。"
    )

    slide = new("课堂 120 分钟带做安排")
    rows = [
        ["时间", "教师带做", "学生产出"],
        ["0-15 分钟", "讲清文本变量不是词云", "理解最终交付物"],
        ["15-30 分钟", "创建工作区、复制 Skills、写 AGENTS.md", "可运行工作区"],
        ["30-45 分钟", "数据盘点、隐私和来源审计", "text_data_inventory"],
        ["45-60 分钟", "清洗计划和分词检查", "text_cleaning_plan"],
        ["60-75 分钟", "变量定义和路线选择", "text_variable_definition / route_selection"],
        ["75-95 分钟", "Codebook、抽样和 LLM 标注计划", "annotation_codebook / sample"],
        ["95-110 分钟", "可靠性审计和特征字典", "reliability_report / feature_dictionary"],
        ["110-120 分钟", "变量表、衔接第8讲和审计", "features / text_variable_audit"],
    ]
    table(slide, rows, 0.78, 1.45, 11.75, 5.05, col_widths=[1.8, 5.0, 4.95], font_size=8.9)
    card(slide, "教师重点", "现场纠正三件事：构念错配、LLM 越界推断、数据泄漏。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12)

    slide = new("课后提交要求")
    bullets(slide, [
        "提交文件夹：姓名_第09讲_文本挖掘LLM标注与变量构造",
        "必须包含 text_data_inventory、text_cleaning_plan、text_variable_definition、route_selection、annotation_codebook、annotation_sample、feature_dictionary、features、reliability_report、variable_construction_report、text_variable_audit",
        "至少构造 1 个可进入实证模型的文本变量",
        "附 300 字反思：变量最可能被审稿人质疑什么，准备用什么证据回应",
    ], 0.90, 1.55, 11.3, 1.75, size=14.7)
    rows = [
        ["评分项", "占比"],
        ["文本盘点与清洗计划", "20%"],
        ["构念定义与路线匹配", "25%"],
        ["Codebook 与标注样本", "20%"],
        ["可靠性与偏差审计", "20%"],
        ["最终变量表和论文可用性", "15%"],
    ]
    table(slide, rows, 2.15, 3.80, 8.85, 2.20, col_widths=[6.6, 2.25], font_size=11.5)
    card(slide, "一票否决", "编造标签/指标、泄露 API key、未脱敏敏感文本、只有词云没有变量。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是完成一个 NLP 分析，而是建立一条能被复核、能进实证模型、能写进论文方法部分的文本变量生产线。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["文本", "清洗", "构念", "Codebook", "标注", "可靠性", "变量", "模型", "审计"], 1.00, 3.50, 11.20, h=0.70, size=9.2)
    card(slide, "下一步", "学生把 features.xlsx 接入第 08 讲流程，建立变量字典、模型映射和结果解释。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
