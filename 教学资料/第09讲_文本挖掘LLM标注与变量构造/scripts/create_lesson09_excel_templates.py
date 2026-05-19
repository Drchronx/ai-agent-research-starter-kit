from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"
TEAL = "0D746E"
PALE = "F4F7F8"
WHITE = "FFFFFF"
INK = "202A36"
LINE = "D2D8DE"


def style_sheet(ws, widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    side = Side(style="thin", color=LINE)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
    ws.row_dimensions[1].height = 36
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    else:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18


def make_wb(sheets, path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows, widths in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_sheet(ws, widths)
    wb.save(path)
    load_workbook(path).close()


def add_common_validations(path):
    wb = load_workbook(path)
    status = DataValidation(type="list", formula1='"pass,revise,fail,to_check,not_applicable"', allow_blank=True)
    yes_no = DataValidation(type="list", formula1='"yes,no,to_check,not_applicable"', allow_blank=True)
    label_status = DataValidation(type="list", formula1='"labeled,unlabeled,partially_labeled,to_check"', allow_blank=True)
    route = DataValidation(type="list", formula1='"dictionary,tfidf,sentiment,lda,bertopic,sklearn,llm_labeling,embedding,deep_learning,to_check"', allow_blank=True)
    for ws in wb.worksheets:
        ws.add_data_validation(status)
        ws.add_data_validation(yes_no)
        ws.add_data_validation(label_status)
        ws.add_data_validation(route)
    if "text_data_inventory" in wb.sheetnames:
        ws = wb["text_data_inventory"]
        yes_no.add("J2:L500")
        status.add("M2:M500")
    if "annotation_sample" in wb.sheetnames:
        ws = wb["annotation_sample"]
        label_status.add("J2:J500")
        status.add("M2:M500")
    if "feature_dictionary" in wb.sheetnames:
        ws = wb["feature_dictionary"]
        route.add("D2:D500")
        status.add("L2:L500")
    if "text_variable_audit" in wb.sheetnames:
        ws = wb["text_variable_audit"]
        status.add("D2:D500")
    wb.save(path)
    load_workbook(path).close()


inventory_sheets = [
    (
        "text_data_inventory",
        [
            "file_name", "file_path", "file_type", "encoding", "row_count", "doc_id_column",
            "text_column", "time_column", "source_column", "has_personal_info", "has_sensitive_info",
            "copyright_or_terms_risk", "audit_status", "minimum_analysis_unit", "notes",
        ],
        [
            ["raw_text_data.xlsx", "data/raw/raw_text_data.xlsx", "xlsx", "utf-8/to_check", "", "doc_id", "text", "date", "source", "to_check", "to_check", "to_check", "to_check", "document/user/firm-year", "原始文件只读保存"],
            ["sample_comments.csv", "data/raw/sample_comments.csv", "csv", "utf-8-sig", "", "doc_id", "comment", "created_at", "platform", "to_check", "to_check", "to_check", "to_check", "comment", "课堂样例"],
        ],
        [24, 38, 14, 16, 14, 18, 18, 18, 18, 18, 18, 24, 18, 24, 42],
    ),
    (
        "text_quality_checks",
        ["check_item", "method", "threshold_or_rule", "result", "action", "notes"],
        [
            ["empty_text", "count empty/NA text", "0 preferred", "to_check", "remove or recover source", "不得删除前先记录"],
            ["duplicate_text", "exact duplicate and near duplicate", "记录重复比例", "to_check", "deduplicate or cluster", "重复模板文本可能造成泄漏"],
            ["too_short_text", "length below threshold", "按任务设定", "to_check", "exclude or label uncertain", "短文本可能无法判断构念"],
            ["too_long_text", "length above threshold", "按分位数/领域判断", "to_check", "split into sections", "长文档需要分段或分句"],
            ["template_noise", "regular expressions / manual scan", "模板词占比", "to_check", "remove template blocks", "年报免责声明常见"],
        ],
        [26, 34, 26, 18, 34, 46],
    ),
    (
        "privacy_ethics_audit",
        ["risk_type", "detected_examples", "risk_level", "action", "owner", "status", "notes"],
        [
            ["personal_identifier", "姓名/手机号/邮箱/账号", "high", "脱敏或删除", "student", "to_check", "访谈和社媒文本优先检查"],
            ["sensitive_attribute", "健康/政治/宗教等", "high", "伦理审查和最小化使用", "student", "to_check", "根据数据来源处理"],
            ["copyright_terms", "平台协议/爬虫限制", "medium/high", "检查授权和引用边界", "student", "to_check", "不要共享受限原文"],
        ],
        [26, 38, 18, 34, 16, 18, 48],
    ),
]

annotation_sheets = [
    (
        "annotation_sample",
        [
            "doc_id", "text", "source", "date", "human_label_1", "human_label_2",
            "llm_label", "final_label", "confidence", "label_status", "evidence_span",
            "disagreement_reason", "audit_status", "notes",
        ],
        [
            ["doc_001", "AI 推荐结果有帮助，但我仍需要核验来源。", "course_sample", "YYYY-MM-DD", "uncertainty_expression", "", "", "", "", "partially_labeled", "仍需要核验来源", "", "to_check", "课堂样例"],
            ["doc_002", "这个 Agent 节省了很多查文献时间。", "course_sample", "YYYY-MM-DD", "efficiency_benefit", "", "", "", "", "partially_labeled", "节省了很多查文献时间", "", "to_check", "课堂样例"],
        ],
        [16, 58, 18, 16, 20, 20, 20, 20, 14, 18, 32, 36, 18, 36],
    ),
    (
        "label_set",
        ["label", "definition", "positive_rule", "negative_rule", "boundary_rule", "examples", "notes"],
        [
            ["uncertainty_expression", "文本中表达对结果、来源或系统可靠性的不确定", "明确出现可能错误、需核验、不确定等线索", "只是描述一般风险但无不确定表达", "讽刺或引用句需人工复核", "可能出错；需要人工核验", "不能等同真实风险"],
            ["efficiency_benefit", "文本中表达节省时间、提升效率或降低工作量", "明确出现节省时间、提高效率、减少重复劳动", "只是说体验好但没有效率线索", "泛泛夸奖标 uncertain", "节省查文献时间", "不能等同真实效率提升"],
        ],
        [24, 42, 42, 42, 38, 34, 40],
    ),
    (
        "sampling_plan",
        ["sampling_goal", "sample_size", "stratification", "random_seed", "inclusion_rule", "exclusion_rule", "notes"],
        [
            ["codebook calibration", "50-200", "source/date/rating if available", "20260516", "有足够文本内容", "空文本/明显乱码", "先小样本校准，再批量标注"],
        ],
        [28, 16, 34, 18, 34, 34, 46],
    ),
]

feature_sheets = [
    (
        "feature_dictionary",
        [
            "feature_name", "construct", "feature_type", "method", "input_file", "parameters",
            "unit", "range", "aggregation_rule", "interpretation", "limitations", "audit_status",
        ],
        [
            ["text_length", "文本长度", "control", "rule_based", "data/clean/text_clean_v1.csv", "char count", "doc", "0+", "mean/sum by unit", "控制文本长度差异", "不能代表内容质量", "to_check"],
            ["uncertainty_expression_label", "不确定性表达", "main_variable", "llm_labeling", "annotation output", "codebook v1", "doc", "0/1/uncertain", "ratio by unit", "文本中不确定线索", "不能代表真实风险", "to_check"],
            ["positive_sentiment_score", "积极情绪", "main/control", "sentiment", "data/clean/text_clean_v1.csv", "dictionary or model", "doc", "0-1 or score", "mean by unit", "文本积极态度", "需验证领域适配", "to_check"],
            ["topic_ai_governance_prob", "AI 治理主题", "topic_variable", "lda/bertopic", "tokenized text", "K/to_check", "doc", "0-1", "mean by unit", "文档属于某主题概率", "主题命名需人工解释", "to_check"],
        ],
        [28, 28, 18, 20, 34, 34, 16, 18, 28, 38, 42, 18],
    ),
    (
        "merge_plan",
        ["target_unit", "merge_key", "time_window", "aggregation", "main_dataset", "text_feature_file", "risk", "status"],
        [
            ["user-level", "user_id", "same survey wave", "mean/ratio", "analysis/main_data.csv", "textmining/features.xlsx", "ID 不一致或重复用户", "to_check"],
            ["firm-year", "firm_id + year", "annual", "sum/mean/ratio", "analysis/panel_data.csv", "textmining/features.xlsx", "时间错配和滞后处理", "to_check"],
        ],
        [20, 24, 24, 24, 30, 30, 40, 18],
    ),
    (
        "validity_checks",
        ["check_type", "question", "evidence", "result", "action"],
        [
            ["face_validity", "变量是否看起来测量了目标构念？", "人工阅读高低分文本", "to_check", "展示代表性文本片段"],
            ["convergent_validity", "是否与相关变量方向一致？", "与问卷/评分/外部指标相关", "to_check", "进入第8讲相关和回归"],
            ["discriminant_validity", "是否区别于相邻构念？", "与相近标签低到中等相关", "to_check", "合并或重写 codebook"],
            ["robustness", "换词典/模型/阈值是否稳健？", "替代构造", "to_check", "报告稳健性和边界"],
        ],
        [24, 48, 38, 18, 42],
    ),
]

results_sheets = [
    (
        "method_results",
        ["route", "input", "output", "key_metric", "metric_value", "human_review", "decision", "notes"],
        [
            ["jieba+tfidf", "data/clean/text_clean_v1.csv", "tfidf_features.xlsx", "top terms reviewed", "", "yes", "to_check", "不要只看词频解释构念"],
            ["llm_labeling", "annotation_sample.xlsx", "llm_labels.xlsx", "LLM-human agreement", "", "yes", "to_check", "必须抽样复核"],
            ["lda", "tokenized_text.csv", "topic_doc_probs.xlsx", "coherence/perplexity", "", "yes", "to_check", "K 需要人工选择"],
        ],
        [22, 32, 32, 28, 18, 18, 18, 46],
    ),
    (
        "topic_summary",
        ["topic_id", "top_words", "representative_docs", "proposed_name", "theoretical_link", "human_decision", "notes"],
        [
            ["topic_01", "", "", "", "", "to_check", "主题命名必须看代表文本，不只看关键词"],
        ],
        [16, 42, 42, 28, 38, 18, 46],
    ),
    (
        "sentiment_summary",
        ["method", "lexicon_or_model", "score_range", "validation_sample", "known_bias", "decision", "notes"],
        [
            ["dictionary", "to_check", "to_check", "manual sample", "讽刺、否定、领域词", "to_check", "情感分数需领域适配"],
        ],
        [20, 30, 20, 28, 38, 18, 46],
    ),
]

audit_sheets = [
    (
        "text_variable_audit",
        ["audit_item", "question", "evidence_file", "status", "fix_required", "notes"],
        [
            ["source_compliance", "文本来源是否合法合规？", "text_data_inventory.xlsx", "to_check", "", ""],
            ["privacy", "是否含个人信息且已脱敏？", "privacy_ethics_audit", "to_check", "", ""],
            ["cleaning_reproducible", "清洗规则和脚本是否可复现？", "text_cleaning_plan.md", "to_check", "", ""],
            ["construct_alignment", "文本变量是否对应理论构念？", "text_variable_definition.md", "to_check", "", ""],
            ["codebook_quality", "标签定义是否足够具体？", "annotation_codebook.md", "to_check", "", ""],
            ["llm_boundary", "LLM 是否有越界推断？", "annotation_sample.xlsx", "to_check", "", ""],
            ["reliability", "一致性是否足以批量使用？", "reliability_report.md", "to_check", "", ""],
            ["leakage", "训练测试或合并是否泄漏？", "feature_dictionary.xlsx", "to_check", "", ""],
            ["analysis_ready", "最终变量能否进入第8讲模型？", "features.xlsx", "to_check", "", ""],
        ],
        [24, 58, 34, 18, 38, 44],
    ),
    (
        "leakage_checks",
        ["risk", "example", "detection", "decision", "notes"],
        [
            ["duplicate_leakage", "同一文本出现在训练和测试", "hash/near duplicate", "to_check", "去重后再划分"],
            ["author_leakage", "同一用户/企业跨训练测试", "group split", "to_check", "按用户或企业分组划分"],
            ["time_leakage", "用未来文本预测过去结果", "time split", "to_check", "时间序列任务必须按时间切分"],
            ["label_leakage", "文本中直接包含结果标签", "manual scan", "to_check", "删除泄漏字段或改任务"],
        ],
        [28, 42, 28, 18, 46],
    ),
]

run_sheets = [
    (
        "pipeline_run_log",
        ["run_id", "date", "step", "script_or_skill", "input_file", "output_file", "parameters", "result", "next_step"],
        [
            ["run_001", "YYYY-MM-DD", "inventory", "manual/agent", "data/raw/raw_text_data.xlsx", "text_data_inventory.xlsx", "", "to_check", "清洗计划"],
            ["run_002", "YYYY-MM-DD", "tokenize", "text-analysis-basic", "data/clean/text_clean_v1.csv", "tokenized_text.csv", "jieba + custom_dict", "to_check", "词频/TF-IDF"],
            ["run_003", "YYYY-MM-DD", "llm_labeling", "big-data-labeling-variable-construction", "data/clean/text_clean_v1.csv", "llm_labels.xlsx", "codebook v1", "to_check", "可靠性审计"],
        ],
        [16, 16, 22, 34, 34, 34, 34, 18, 36],
    ),
    (
        "file_lineage",
        ["file", "source", "created_by", "created_at", "modified_from", "do_not_overwrite", "notes"],
        [
            ["data/raw/raw_text_data.xlsx", "原始下载/导出", "human", "YYYY-MM-DD", "none", "yes", "原始文本只读保存"],
            ["data/clean/text_clean_v1.csv", "data/raw/raw_text_data.xlsx", "script", "YYYY-MM-DD", "cleaning script", "no", "清洗日志需保留"],
            ["textmining/features.xlsx", "model/label outputs", "agent+human", "YYYY-MM-DD", "feature construction", "no", "最终变量表"],
        ],
        [34, 34, 18, 18, 34, 18, 42],
    ),
]


outputs = [
    ("text_data_inventory_template.xlsx", inventory_sheets),
    ("annotation_sample_template.xlsx", annotation_sheets),
    ("feature_dictionary_template.xlsx", feature_sheets),
    ("text_mining_results_template.xlsx", results_sheets),
    ("text_variable_audit_template.xlsx", audit_sheets),
    ("text_mining_pipeline_log_template.xlsx", run_sheets),
]

for filename, sheets in outputs:
    path = ROOT / filename
    make_wb(sheets, path)
    add_common_validations(path)
    print(path)
