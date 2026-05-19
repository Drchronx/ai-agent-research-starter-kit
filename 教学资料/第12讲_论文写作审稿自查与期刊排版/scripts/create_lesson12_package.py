from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"


def write_texts():
    files = {
        "第12讲_教学大纲.md": """# 第12讲 论文写作、审稿自查与期刊排版

## 这节课解决什么科研问题和需求

第1-11讲已经把 AI Agent 部署、文献、知识库、综述、选题、实验、统计、文本挖掘、EEG/ERP 和 BCI 建模串起来。本讲解决最后一个高耗时环节：如何把研究材料变成可投稿论文，并在投稿前完成审稿人式自查和期刊格式转换。

本讲不是让 AI 替你编论文，而是让 AI 做三件确定性工作：

1. 把研究证据组织成清晰论文故事线。
2. 按章节写出可核验、可追溯、不过度声称的正文。
3. 按目标期刊规则完成格式、引用、图表、统计报告和盲审匿名化审计。

## 需要哪些 Skills

建议复制到第12讲工作区：

```text
academic-paper-composer
empirical-paper-writer
scientific-writing
peer-review
scientific-critical-thinking
venue-templates
journal-title-page-metadata
journal-abstract-keywords
journal-heading-hierarchy
journal-intext-citation-style
journal-reference-list-format
journal-table-figure-caption
journal-statistics-units-style
journal-layout-spacing
journal-footnote-endnote
journal-blind-review-anonymizer
citation-management
docx
```

其中 `journal-*` 是论文排版功能区的 10 个核心 Skills。

## 输入材料是什么

最低输入材料：

1. 研究问题、理论机制链、假设组和一页方案。
2. 文献矩阵、BibTeX、真实 DOI 或数据库核验记录。
3. 实验/数据/文本/EEG/BCI 的方法计划。
4. 结果表、图表计划、统计输出和审计文件。
5. 目标期刊或目标格式：APA 7、SSCI/SCI/C刊、自定义期刊指南。
6. 作者信息、基金、致谢、数据可用性、伦理声明和利益冲突声明。

## Agent 应该怎么执行

固定顺序：

1. 先整理 manuscript evidence pack。
2. 再生成论文故事线。
3. 再做 reverse outline。
4. 再逐章节写作，不直接一次性写整篇。
5. 每章写完后做证据、引用、数值、理论和因果语言审计。
6. 完成全文后做 peer-review 自查。
7. 再进入期刊排版流水线。
8. 最后执行盲审匿名化和提交包检查。

## 输出文件应该长什么样

```text
第12讲_论文写作审稿自查与期刊排版工作区/
├─ AGENTS.md
├─ manuscript/
│  ├─ manuscript_storyline.md
│  ├─ reverse_outline.md
│  ├─ title_abstract_keywords.md
│  ├─ introduction.md
│  ├─ literature_review.md
│  ├─ theory_hypotheses.md
│  ├─ methods.md
│  ├─ results.md
│  ├─ discussion.md
│  ├─ conclusion.md
│  └─ full_manuscript.md
├─ references/
│  ├─ references.bib
│  ├─ reference_audit.xlsx
│  └─ citation_support_matrix.xlsx
├─ formatting/
│  ├─ formatting_pipeline.md
│  ├─ journal_formatting_audit.xlsx
│  ├─ blind_review_checklist.xlsx
│  └─ submission_package_tracker.xlsx
├─ qc/
│  ├─ reviewer_risk_audit.md
│  ├─ claims_audit.md
│  ├─ statistics_consistency_audit.md
│  └─ final_submission_checklist.md
└─ output/
   ├─ manuscript_author_version.docx
   ├─ manuscript_blind_version.docx
   └─ cover_letter.md
```

## 哪些地方必须人工核验

1. 每条引用是否真实、是否支持对应句子。
2. 统计数值、p 值、置信区间和表格是否一致。
3. 结果是否支持因果语言。
4. 理论是否真正解释机制，而不是贴标签。
5. 期刊格式是否符合最新 author guidelines。
6. 盲审稿件是否残留作者姓名、单位、自引线索和文件元数据。
7. 图表标题、脚注、单位、统计符号是否正确。
8. 摘要和结论是否夸大贡献。

## 学生课后作业

提交一份小论文写作包：

1. `manuscript_storyline.md`
2. `reverse_outline.md`
3. 至少 3 个章节草稿。
4. `reference_audit.xlsx`
5. `reviewer_risk_audit.md`
6. `formatting_pipeline.md`
7. `journal_formatting_audit.xlsx`
8. `blind_review_checklist.xlsx`
9. `cover_letter.md`
""",
        "第12讲_零基础实操教程.md": """# 第12讲 零基础实操教程

## 0. 你今天要学会什么

你要学会把前面课程产出的材料变成投稿论文。注意：论文写作不是把材料堆在一起。你要先确定故事线，再写章节，再做审稿自查，最后排版。

本讲有四个关键判断：

1. 论文每个主张都要有证据。
2. 每个引用都要真实且支持句子。
3. 统计结果不能被文字改写。
4. 排版自动化只能改格式，不能改学术内容。

## 1. 创建工作区

```powershell
New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第12讲_论文写作审稿自查与期刊排版工作区"
cd "D:\\AI科研训练营\\第12讲_论文写作审稿自查与期刊排版工作区"

New-Item -ItemType Directory -Force -Path `
  input,manuscript,references,figures,tables,formatting,qc,output,skills

New-Item -ItemType File -Force -Path AGENTS.md,research_log.md
```

## 2. 复制 Skills

```powershell
cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"

powershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `
  -SkillName academic-paper-composer,empirical-paper-writer,scientific-writing,peer-review,scientific-critical-thinking,venue-templates,journal-title-page-metadata,journal-abstract-keywords,journal-heading-hierarchy,journal-intext-citation-style,journal-reference-list-format,journal-table-figure-caption,journal-statistics-units-style,journal-layout-spacing,journal-footnote-endnote,journal-blind-review-anonymizer,citation-management,docx `
  -Workspace "D:\\AI科研训练营\\第12讲_论文写作审稿自查与期刊排版工作区"
```

## 3. 写 AGENTS.md

复制 `lesson12_workspace_AGENTS_template.md` 到工作区 `AGENTS.md`。

必须保留这些规则：

```text
- 禁止编造引用。
- 禁止改动统计数值。
- 禁止把非因果设计写成因果证明。
- 禁止排版时改变学术内容。
- 参考文献缺字段只能标注，不得补造。
- 盲审前必须清除身份线索和文件元数据。
```

## 4. 整理输入材料

把前面课程材料放到 `input/`：

```text
input/研究问题与假设.md
input/文献矩阵.xlsx
input/references.bib
input/方法计划.md
input/结果表.xlsx
input/图表计划.md
input/统计输出.md
input/目标期刊指南.pdf 或 .md
```

如果没有目标期刊，先用 APA 7 作为训练格式，但要标注“投稿前需按目标期刊更新”。

## 5. 生成论文证据包

复制给 Agent：

```text
请读取 input/ 下的研究问题、假设、文献、方法计划、结果表和图表计划，生成 manuscript/evidence_pack.md。

每条证据必须包含：
- claim
- evidence source
- result/table/figure
- citation needed
- support strength
- risk
- manuscript section

禁止编造不存在的数据、引用或结果。
```

人工核验：

- 结果表是否真实存在。
- 统计值是否与原输出一致。
- 引用是否来自可核验文献。

## 6. 生成论文故事线

复制给 Agent：

```text
请基于 evidence_pack.md，生成 manuscript/manuscript_storyline.md。

结构：
1. research problem
2. practical context
3. theoretical tension
4. gap
5. mechanism
6. hypotheses
7. method fit
8. main findings
9. contribution
10. limitations

要求：不要写空泛贡献，不要把方法优势包装成理论贡献。
```

人工核验：

- gap 是否真实来自文献，不是“很少有人研究”。
- contribution 是否能被结果支持。
- 理论机制是否解释变量关系。

## 7. 反向提纲 Reverse Outline

复制给 Agent：

```text
请生成 manuscript/reverse_outline.md。

按论文章节列出：
- paragraph purpose
- topic sentence
- evidence
- citation
- link to next paragraph
- risk

每个段落只服务一个核心功能。
```

人工核验：

- Introduction 是否从问题推进到 gap。
- Literature review 是否按主题组织，而不是按作者罗列。
- Methods 是否足以复现。
- Results 是否先主结果后补充分析。
- Discussion 是否先解释发现，再讲贡献和限制。

## 8. 分章节写作

不要一次性让 Agent 写整篇。按顺序写：

1. Methods：最容易从计划和数据写起。
2. Results：只报告结果，不解释过多。
3. Introduction：写问题、gap、贡献。
4. Literature review / theory：写机制和假设。
5. Discussion：解释结果、贡献、限制。
6. Abstract：最后写。

Prompt 示例：

```text
请基于 reverse_outline.md、evidence_pack.md 和 input/结果表.xlsx，写 manuscript/results.md。

要求：
- 用完整段落，不用 bullet。
- 所有统计值必须来自结果表，不得改数值。
- 不显著结果也要报告。
- 不把相关或回归结果写成因果证明。
- 每段后附“证据核验清单”。
```

## 9. 真实引用核验

复制给 Agent：

```text
请使用 citation-management。读取 references/references.bib 和 manuscript/ 下的草稿，生成 references/reference_audit.xlsx。

检查：
- in-text citation 是否在 bibliography 中。
- bibliography 是否被正文引用。
- DOI、year、journal、volume、issue、pages 是否缺失。
- 每条引用支持哪一句话。
- 缺失字段标注待人工核验，不得编造。
```

人工核验：

- DOI 是否真实。
- 文献结论是否真的支持该句。
- AI 不确定的引用必须人工查数据库。

## 10. 审稿人式自查

复制给 Agent：

```text
请使用 peer-review 和 scientific-critical-thinking，生成 qc/reviewer_risk_audit.md。

按严重程度列出：
- contribution risk
- theory-mechanism mismatch
- hypothesis unsupported
- method-design mismatch
- statistics/reporting issue
- causal overclaim
- citation weakness
- figure/table weakness
- reproducibility issue
- ethics/data availability issue

每项给 severity、why it matters、fix、must-fix before submission。
```

人工核验：

- 不要只保留轻微问题。
- 若有致命设计问题，不能靠润色解决。

## 11. 期刊排版流水线

复制给 Agent：

```text
请生成 formatting/formatting_pipeline.md。

目标格式：[APA 7 / 目标期刊名 / 自定义指南]

请依次调用或模拟以下 journal Skills：
1. journal-title-page-metadata
2. journal-abstract-keywords
3. journal-heading-hierarchy
4. journal-intext-citation-style
5. journal-reference-list-format
6. journal-table-figure-caption
7. journal-statistics-units-style
8. journal-layout-spacing
9. journal-footnote-endnote
10. journal-blind-review-anonymizer

每一步输出：输入、处理规则、输出、人工核验项。
```

人工核验：

- 目标期刊指南可能更新，投稿前必须看官网。
- 排版只能改格式，不改数据和学术主张。

## 12. 盲审匿名化

复制给 Agent：

```text
请使用 journal-blind-review-anonymizer，生成 formatting/blind_review_checklist.md。

检查：
- title page
- author names
- affiliations
- acknowledgments
- funding
- self-citations
- file properties
- supplementary files
- data/code repository
- figures and appendices

输出 blind version 和 author version 的差异清单。
```

人工核验：

- Word 文件属性是否仍有作者姓名。
- 自引是否需要匿名处理。
- 数据仓库、OSF、GitHub 链接是否暴露身份。

## 13. 投稿包

最终输出：

```text
output/manuscript_author_version.docx
output/manuscript_blind_version.docx
output/cover_letter.md
output/submission_checklist.md
formatting/journal_formatting_audit.xlsx
formatting/blind_review_checklist.xlsx
references/reference_audit.xlsx
qc/reviewer_risk_audit.md
```

没有目标期刊时，不要声称“已经完全符合投稿要求”，只能说“已按 APA 7 训练格式整理，目标期刊投稿前需二次核验”。
""",
        "第12讲_课堂任务单.md": """# 第12讲 课堂任务单

## 任务 1：工作区和 Skills

- [ ] 已创建第12讲工作区。
- [ ] 已复制写作、审稿、排版和引用 Skills。
- [ ] 已写入 AGENTS.md。
- [ ] 已准备 input/ 材料。

## 任务 2：论文证据包

输出：`manuscript/evidence_pack.md`

- [ ] 每条 claim 有 evidence source。
- [ ] 统计结果有表格或输出来源。
- [ ] 引用需求已标注。
- [ ] 风险已标注。

## 任务 3：论文故事线

输出：`manuscript/manuscript_storyline.md`

- [ ] research problem 清楚。
- [ ] theoretical tension 不是空泛背景。
- [ ] gap 由文献支持。
- [ ] mechanism 连接变量。
- [ ] contribution 能被结果支持。

## 任务 4：Reverse Outline

输出：`manuscript/reverse_outline.md`

- [ ] 每段只有一个目的。
- [ ] topic sentence 明确。
- [ ] evidence 和 citation 对应。
- [ ] 段落之间有推进。

## 任务 5：分章节写作

至少完成 3 个章节：

- [ ] Methods
- [ ] Results
- [ ] Introduction
- [ ] Literature review / theory
- [ ] Discussion
- [ ] Abstract

检查：

- [ ] 引用真实。
- [ ] 统计值未改。
- [ ] 因果语言不过度。
- [ ] 结果与表图一致。

## 任务 6：引用核验

输出：`references/reference_audit.xlsx`

- [ ] 正文引用在参考文献中。
- [ ] 参考文献被正文引用。
- [ ] DOI、year、journal、volume、issue、pages 缺失项已标注。
- [ ] 每条引用支持对应句子。
- [ ] 未编造 DOI 或页码。

## 任务 7：审稿自查

输出：`qc/reviewer_risk_audit.md`

- [ ] contribution risk。
- [ ] theory risk。
- [ ] method risk。
- [ ] statistics risk。
- [ ] causality risk。
- [ ] citation risk。
- [ ] reproducibility risk。

## 任务 8：期刊排版

输出：

- `formatting/formatting_pipeline.md`
- `formatting/journal_formatting_audit.xlsx`

检查 10 个模块：

- [ ] 扉页与元数据。
- [ ] 摘要与关键词。
- [ ] 标题层级。
- [ ] 正文引用。
- [ ] 参考文献列表。
- [ ] 图表标题。
- [ ] 统计与单位。
- [ ] 布局与间距。
- [ ] 脚注/尾注。
- [ ] 盲审匿名化。

## 任务 9：投稿包

- [ ] 作者版稿件。
- [ ] 盲审版稿件。
- [ ] cover letter。
- [ ] submission checklist。
- [ ] 人工核验清单。
""",
        "第12讲_教师带做讲稿.md": """# 第12讲 教师带做讲稿

## 开场 0-10 分钟

告诉学生：前面 11 讲解决的是“研究材料如何生产”，第12讲解决的是“材料如何变成可投稿论文”。

核心句：

> 写论文不是让 AI 拼接材料，而是让 AI 在证据边界内组织论证。

强调三条底线：

1. 引用不能编。
2. 数值不能改。
3. 排版不能改变学术内容。

## 10-25 分钟：工作区与 AGENTS.md

教师带做创建工作区、复制 Skills、写 AGENTS.md。

重点解释：

- `academic-paper-composer` 和 `scientific-writing` 用于章节写作。
- `peer-review` 和 `scientific-critical-thinking` 用于审稿自查。
- `journal-*` 用于格式转换。
- `citation-management` 用于引用核验。
- `docx` 用于最终 Word 文件处理。

## 25-45 分钟：证据包

让学生把前几讲材料放进 `input/`。

教师讲解 evidence pack 的作用：

- 论文每个 claim 必须能追溯到证据。
- claim、evidence、citation、risk 必须绑定。
- 没有证据的贡献不能写。

现场提问：

“如果某句写‘本研究首次证明’，需要什么证据？”

标准回答：

- 文献检索证明 gap。
- 研究设计支持证明程度。
- 结果支持 claim。
- 不能只凭作者感觉。

## 45-65 分钟：论文故事线

教师展示故事线结构：

```text
problem -> tension -> gap -> mechanism -> hypothesis -> method -> finding -> contribution
```

纠正常见错误：

- 把现实背景写成学术问题。
- 把“研究少”写成 gap。
- 把方法复杂写成理论贡献。
- 把结果显著写成理论成立。

## 65-85 分钟：Reverse Outline

教师解释：反向提纲不是目录，是每一段在论文中承担的功能。

每段要有：

1. 段落目的。
2. 主题句。
3. 证据。
4. 引用。
5. 与下一段的连接。
6. 风险。

让学生现场检查 Introduction 是否真正推进。

## 85-110 分钟：分章节写作

教师建议写作顺序：

1. Methods
2. Results
3. Introduction
4. Literature review / theory
5. Discussion
6. Abstract

解释原因：

- Methods 和 Results 最依赖事实，先写能防止故事线漂移。
- Abstract 最后写，因为它必须总结全文。

强调：

- 最终论文正文要用完整段落，不用 bullet。
- 统计值必须来自表格。
- 非显著结果也要报告。

## 110-125 分钟：引用核验

教师讲：

AI 生成 citation 最危险。学生必须把引用从“看起来像真的”变成“可以核验的”。

引用核验表至少包含：

- citation key
- title
- DOI
- journal
- year
- in-text sentence
- support status
- missing metadata
- manual verification

若缺 DOI 或页码，Agent 只能标注，不能补造。

## 125-145 分钟：审稿人式自查

教师用 peer-review 视角讲 7 类问题：

1. 贡献不足。
2. 理论机制错位。
3. 方法无法回答问题。
4. 统计报告不完整。
5. 因果语言过度。
6. 引用不支持主张。
7. 图表和复现不足。

要求学生把每一项写成 must-fix 或 optional-fix。

## 145-170 分钟：期刊排版流水线

逐个讲 10 个排版 Skills：

1. 扉页与元数据。
2. 摘要长度与关键词结构。
3. 标题层级。
4. 正文引用格式。
5. 参考文献列表。
6. 图表标题。
7. 统计报告与单位。
8. 布局与间距。
9. 脚注转尾注。
10. 盲审匿名化。

强调：这些 Skills 是句法和结构处理器，不负责改变论文结论。

## 170-180 分钟：收尾

总结：

> 一篇能投稿的论文，不是写得漂亮，而是每个主张、引用、数值、图表和格式都能被追踪和核验。
""",
        "第12讲_课后作业.md": """# 第12讲 课后作业

## 作业目标

把一个研究项目整理成可投稿论文写作包。

## 必交材料

```text
manuscript/evidence_pack.md
manuscript/manuscript_storyline.md
manuscript/reverse_outline.md
manuscript/methods.md
manuscript/results.md
manuscript/introduction.md
references/reference_audit.xlsx
qc/reviewer_risk_audit.md
formatting/formatting_pipeline.md
formatting/journal_formatting_audit.xlsx
formatting/blind_review_checklist.xlsx
output/cover_letter.md
```

## 评分标准

| 项目 | 占比 |
|---|---:|
| 证据包与故事线 | 20% |
| 章节写作质量 | 25% |
| 引用真实性核验 | 20% |
| 审稿自查深度 | 20% |
| 期刊排版与盲审检查 | 15% |

## 一票否决

1. 编造引用。
2. 编造统计结果。
3. 排版时改变原始学术数据。
4. 非因果设计写成因果证明。
5. 盲审版残留明显作者身份。

## 反思题

用 400-600 字回答：

1. 你的论文最大拒稿风险是什么？
2. 哪个主张最需要补文献或补分析？
3. 哪个引用最需要人工核验？
4. 目标期刊格式中最容易出错的部分是什么？
""",
        "lesson12_workspace_AGENTS_template.md": """# 第12讲 论文写作、审稿自查与期刊排版工作区 Agent 指令

你是论文写作、审稿自查与期刊排版助手。

## 学术底线

- 禁止编造引用、DOI、页码、期号、样本量、p 值或模型结果。
- 禁止修改统计数值、方向、置信区间、样本量和表格数据。
- 禁止把非因果研究写成因果证明。
- 禁止把理论当标签堆砌，必须解释机制。
- 若无法验证文献真实性，必须标注“待人工核验”。

## 写作规则

- 最终论文正文使用完整段落，不使用 bullet 代替正文。
- 每个段落必须有清晰功能：问题、gap、机制、方法、结果或贡献。
- 所有主张必须能追溯到 evidence pack。
- Results 只报告结果，Discussion 再解释。
- 不显著结果也要透明报告。

## 排版规则

- 排版 Skills 只处理格式、结构、标点、层级和匿名化。
- 排版过程不得改变学术内容和数据。
- 参考文献缺字段只能标注，不能补造。
- 盲审版必须检查正文、参考文献、自引、致谢、补充材料和文件元数据。

## 输出目录

```text
manuscript/
references/
formatting/
qc/
output/
```
""",
        "lesson12_agent_prompt_bank.md": """# 第12讲 Agent Prompt Bank

## Prompt 1：生成证据包

```text
请读取 input/ 下的研究问题、假设、文献、方法计划、结果表和图表计划，生成 manuscript/evidence_pack.md。每条证据包含 claim、evidence source、result/table/figure、citation needed、support strength、risk、manuscript section。禁止编造不存在的数据、引用或结果。
```

## Prompt 2：生成论文故事线

```text
请基于 evidence_pack.md，生成 manuscript/manuscript_storyline.md。结构包括 problem、context、theoretical tension、gap、mechanism、hypotheses、method fit、main findings、contribution、limitations。不要写空泛贡献。
```

## Prompt 3：生成 Reverse Outline

```text
请生成 manuscript/reverse_outline.md。按章节列出 paragraph purpose、topic sentence、evidence、citation、link to next paragraph、risk。每段只服务一个核心功能。
```

## Prompt 4：写 Methods

```text
请基于 input/方法计划.md、input/数据说明.md 和 evidence_pack.md，写 manuscript/methods.md。要求可复现，包含样本、程序、变量、测量、统计方法、伦理和排除规则。禁止编造未提供的信息，不确定处标待人工核验。
```

## Prompt 5：写 Results

```text
请基于 input/结果表.xlsx 和 input/统计输出.md，写 manuscript/results.md。所有数值必须来自输入表，不得改动。报告主结果、机制/调节、稳健性、非显著结果、效应量和置信区间。不要把相关结果写成因果证明。
```

## Prompt 6：写 Introduction

```text
请基于 manuscript_storyline.md、reverse_outline.md 和 references/reference_audit.xlsx，写 manuscript/introduction.md。按问题、文献张力、gap、研究方案、贡献推进。引用必须真实且支持句子，不确定引用标待人工核验。
```

## Prompt 7：引用核验

```text
请使用 citation-management。读取 manuscript/ 草稿和 references/references.bib，生成 references/reference_audit.xlsx。检查正文引用、参考文献条目、缺失 DOI/volume/issue/pages、引用支持句子和人工核验项。不得编造缺失字段。
```

## Prompt 8：审稿自查

```text
请使用 peer-review 和 scientific-critical-thinking，生成 qc/reviewer_risk_audit.md。按 severity 排序列出 contribution、theory、method、statistics、causality、citation、figure/table、reproducibility、ethics 风险，并给出具体修改方案。
```

## Prompt 9：期刊排版流水线

```text
请生成 formatting/formatting_pipeline.md。目标格式：[APA 7 / 目标期刊]。依次处理 title page、abstract keywords、heading hierarchy、in-text citations、reference list、captions、statistics units、layout spacing、footnote/endnote、blind review。每步输出输入、规则、输出和人工核验项。
```

## Prompt 10：盲审匿名化

```text
请使用 journal-blind-review-anonymizer，生成 formatting/blind_review_checklist.md。检查作者姓名、单位、致谢、基金、自引、数据仓库、文件属性、补充材料和图表身份线索。输出 author version 与 blind version 差异清单。
```

## Prompt 11：Cover Letter

```text
请基于 manuscript_storyline.md 和 reviewer_risk_audit.md，生成 output/cover_letter.md。内容包括投稿期刊、论文标题、研究问题、核心贡献、方法和数据、为什么适合该期刊、原创性声明、利益冲突和通讯作者信息占位符。不要夸大贡献。
```
""",
        "manuscript_storyline_template.md": """# Manuscript Storyline Template

## Research Problem

What precise scholarly problem does the paper address?

## Practical Context

Why does this problem matter in a real setting?

## Theoretical Tension

Which theoretical explanation is incomplete, conflicting, or underspecified?

## Literature Gap

What is missing in prior work? Cite real sources only.

## Mechanism

Explain the causal or explanatory chain:

```text
context/stimulus -> mechanism -> outcome
```

## Hypotheses

List only hypotheses supported by theory and design.

## Method Fit

Why does the design answer the question?

## Findings

Report what the data actually show.

## Contribution

Separate theory, method, empirical and practical contributions.

## Limitations

State what cannot be claimed.
""",
        "reverse_outline_template.md": """# Reverse Outline Template

| Section | Paragraph | Purpose | Topic Sentence | Evidence | Citation | Link To Next | Risk |
|---|---:|---|---|---|---|---|---|
| Introduction | 1 | Open problem |  |  |  |  |  |
| Introduction | 2 | Establish tension |  |  |  |  |  |
| Literature Review | 1 | Theme A |  |  |  |  |  |
| Methods | 1 | Design overview |  |  |  |  |  |
| Results | 1 | Main finding |  |  |  |  |  |
| Discussion | 1 | Interpret finding |  |  |  |  |  |
""",
        "reviewer_risk_audit_template.md": """# Reviewer Risk Audit Template

| Severity | Risk | Why It Matters | Evidence | Required Fix | Owner |
|---|---|---|---|---|---|
| high | Contribution unclear | Reviewers may see no novelty |  |  |  |
| high | Causal overclaim | Design may not support causal language |  |  |  |
| medium | Citation support weak | Claim may not be supported |  |  |  |
| medium | Statistics incomplete | Results may be underreported |  |  |  |
| medium | Format noncompliance | Desk check may fail |  |  |  |

## Final Decision

- Ready for submission: yes/no
- Must fix:
- Optional fix:
""",
        "blind_review_checklist_template.md": """# Blind Review Checklist Template

| Location | Identity Risk | Found? | Action | Status |
|---|---|---|---|---|
| Title page | Author names and affiliations |  | remove from blind version |  |
| Acknowledgments | Funding or collaborators |  | mask or move to author version |  |
| Self-citation | Identifying phrasing |  | use neutral citation or [Author] if required |  |
| File metadata | Author property |  | inspect and remove |  |
| Data repository | GitHub/OSF owner identity |  | anonymized link if possible |  |
| Supplement | Author names in files |  | remove |  |
""",
        "formatting_pipeline_template.md": """# Formatting Pipeline Template

Target journal or style:

## Step 1 Title Page And Metadata

- Input:
- Rule:
- Output:
- Manual check:

## Step 2 Abstract And Keywords

- Word limit:
- Keyword rule:

## Step 3 Heading Hierarchy

- H1:
- H2:
- H3:

## Step 4 In-Text Citations

- Current style:
- Target style:

## Step 5 Reference List

- Sorting:
- DOI:
- Missing fields:

## Step 6 Table And Figure Captions

- Table title:
- Figure caption:
- Notes:

## Step 7 Statistics And Units

- p-value:
- CI:
- effect size:
- units:

## Step 8 Layout And Spacing

- margins:
- line spacing:
- indentation:

## Step 9 Footnote And Endnote

- conversion:
- numbering:

## Step 10 Blind Review

- author version:
- blind version:
- metadata:
""",
        "cover_letter_template.md": """# Cover Letter Template

Dear Editor,

We are pleased to submit our manuscript, "[Title]", for consideration in [Journal].

The manuscript examines [research problem]. Using [method/data], we find that [main finding]. The paper contributes to [field/theory] by [specific contribution].

We believe the manuscript fits [Journal] because [fit statement].

This manuscript is original, has not been published, and is not under consideration elsewhere. All authors have approved the submission. Any conflicts of interest, funding, data availability, and ethical approval information are disclosed in the manuscript.

Sincerely,

[Corresponding author]
""",
    }
    for name, text in files.items():
        (ROOT / name).write_text(text, encoding="utf-8")


def style_sheet(ws, widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    side = Side(style="thin", color="D2D8DE")
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, color="202A36")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="0D746E")
                cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7F8")
    ws.row_dimensions[1].height = 36
    for idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1] if widths and idx <= len(widths) else 20


def make_wb(filename, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows, widths in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_sheet(ws, widths)
    path = ROOT / filename
    wb.save(path)
    load_workbook(path).close()


def create_excels():
    outputs = {
        "manuscript_section_tracker_template.xlsx": [
            ("section_tracker", ["section", "target_words", "status", "input_files", "draft_file", "evidence_checked", "citation_checked", "stats_checked", "notes"], [["Abstract", "150-300", "to_check", "", "manuscript/title_abstract_keywords.md", "no", "no", "not_applicable", ""], ["Introduction", "", "to_check", "", "manuscript/introduction.md", "no", "no", "not_applicable", ""]], [24, 16, 18, 42, 42, 18, 18, 18, 44]),
            ("paragraph_outline", ["section", "paragraph_id", "purpose", "topic_sentence", "evidence", "citation", "risk", "status"], [["Introduction", "P1", "open problem", "", "", "", "", "to_check"]], [22, 16, 30, 48, 42, 34, 34, 18]),
        ],
        "reference_audit_template.xlsx": [
            ("reference_audit", ["citation_key", "in_text_sentence", "title", "authors", "year", "journal", "doi", "volume", "issue", "pages", "supports_claim", "missing_fields", "manual_check"], [["", "", "", "", "", "", "", "", "", "", "to_check", "", "required"]], [22, 60, 46, 32, 14, 34, 30, 14, 14, 18, 18, 30, 22]),
            ("bibtex_check", ["citation_key", "in_bib", "cited_in_text", "duplicate", "status", "notes"], [["", "to_check", "to_check", "to_check", "to_check", ""]], [24, 18, 18, 18, 18, 44]),
        ],
        "journal_formatting_audit_template.xlsx": [
            ("formatting_audit", ["module", "skill", "input", "target_rule", "status", "manual_check", "notes"], [["title page", "journal-title-page-metadata", "", "", "to_check", "required", ""], ["abstract keywords", "journal-abstract-keywords", "", "", "to_check", "required", ""], ["headings", "journal-heading-hierarchy", "", "", "to_check", "required", ""], ["in-text citation", "journal-intext-citation-style", "", "", "to_check", "required", ""], ["reference list", "journal-reference-list-format", "", "", "to_check", "required", ""], ["captions", "journal-table-figure-caption", "", "", "to_check", "required", ""], ["statistics units", "journal-statistics-units-style", "", "", "to_check", "required", ""], ["layout spacing", "journal-layout-spacing", "", "", "to_check", "required", ""], ["footnote endnote", "journal-footnote-endnote", "", "", "to_check", "required", ""], ["blind review", "journal-blind-review-anonymizer", "", "", "to_check", "required", ""]], [24, 34, 34, 42, 18, 18, 44]),
        ],
        "blind_review_anonymization_template.xlsx": [
            ("identity_inventory", ["location", "identity_signal", "found", "action", "status", "notes"], [["title page", "author names", "to_check", "remove in blind version", "to_check", ""]], [28, 36, 18, 42, 18, 44]),
            ("file_metadata", ["file", "metadata_field", "value_seen", "action", "status"], [["output/manuscript_blind_version.docx", "author", "", "remove", "to_check"]], [42, 24, 34, 28, 18]),
        ],
        "reviewer_risk_audit_template.xlsx": [
            ("reviewer_risks", ["severity", "risk_type", "issue", "why_it_matters", "evidence", "fix", "must_fix"], [["high", "causal overclaim", "", "", "", "", "yes"]], [16, 28, 42, 46, 34, 46, 16]),
            ("claims_audit", ["claim", "section", "evidence", "citation", "supported", "rewrite_needed"], [["", "", "", "", "to_check", "to_check"]], [58, 22, 42, 34, 18, 24]),
        ],
        "submission_package_tracker_template.xlsx": [
            ("submission_files", ["file", "purpose", "author_version", "blind_version", "status", "notes"], [["manuscript.docx", "main manuscript", "yes", "yes", "to_check", ""]], [42, 30, 18, 18, 18, 44]),
            ("required_statements", ["statement", "text", "included", "manual_check"], [["funding", "", "to_check", "required"], ["data availability", "", "to_check", "required"], ["ethics", "", "to_check", "required"], ["conflict of interest", "", "to_check", "required"]], [28, 60, 18, 22]),
        ],
    }
    for filename, sheets in outputs.items():
        make_wb(filename, sheets)
        wb = load_workbook(ROOT / filename)
        dv = DataValidation(type="list", formula1='"pass,revise,fail,to_check,yes,no,not_applicable,required"', allow_blank=True)
        for ws in wb.worksheets:
            ws.add_data_validation(dv)
            dv.add("A2:Z500")
        wb.save(ROOT / filename)
        load_workbook(ROOT / filename).close()


COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, level = item if isinstance(item, tuple) else (item, 0)
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=8.5):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 12 讲 · 论文写作、审稿自查与期刊排版", 0.45, 0.055, 4.3, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.0, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=9.5):
    gap = 0.12
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(bx - gap + 0.01), Inches(y + h / 2), Inches(bx - 0.02), Inches(y + h / 2))
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"])
    no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.16, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.18, accent=COLORS["amber"], body_size=11.2, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.72, size=8.0)
    textbox(slide, "人工核验", 4.65, 4.82, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.1)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 12 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "论文写作\n审稿自查\n期刊排版", 0.52, 1.70, 2.25, 1.65, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "从研究包到投稿包", 0.55, 6.10, 2.2, 0.35, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "把科研材料变成可核验、可匿名、可投稿的论文包", 3.55, 1.10, 8.75, 0.72, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "Storyline · Section Writing · Citation Audit · Peer Review · Formatting · Blind Review", 3.58, 1.95, 8.75, 0.35, size=14, color=COLORS["muted"])
    flow(slide, ["证据包", "故事线", "提纲", "章节", "引用", "审稿", "排版", "投稿"], 3.60, 3.10, 8.85, h=0.62, size=9.0)
    card(slide, "本讲交付", "evidence_pack、storyline、reverse_outline、section drafts、reference_audit、reviewer_risk_audit、formatting_pipeline、journal_formatting_audit、blind_review_checklist、cover_letter。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=11.2)
    card(slide, "底线", "引用不能编，数值不能改，排版不能改变学术内容。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def create_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, ["材料很多，但不知道怎样组织成论文主线", "AI 写作容易堆砌、空泛、贡献夸大", "引用看起来完整，但不一定真实支持句子", "结果表和正文统计值容易不一致", "人工适配期刊格式耗时且易出错", "盲审稿常残留作者身份和文件元数据"], 0.85, 1.55, 5.95, 3.20, size=15.0)
    card(slide, "本讲不做", "不让 AI 编论文、编引用、改数值、夸大贡献。", 7.05, 1.72, 5.10, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "让 AI 在证据边界内组织论证、审稿自查和执行确定性排版。", 7.05, 3.35, 5.10, 1.20, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["证据", "故事", "章节", "引用", "审稿", "排版", "匿名"], 0.95, 5.65, 11.35, h=0.62, size=10.5)

    slide = new("本讲最终产出文件")
    rows = [["阶段", "文件", "作用"], ["证据", "evidence_pack.md", "claim、证据、引用需求和风险绑定"], ["故事线", "manuscript_storyline.md", "problem、tension、gap、mechanism、contribution"], ["提纲", "reverse_outline.md", "每段功能、主题句、证据和连接"], ["写作", "section drafts / full_manuscript.md", "分章节写作和全文整合"], ["引用", "reference_audit.xlsx", "真实引用、缺失字段和句子支持矩阵"], ["审稿", "reviewer_risk_audit.md", "贡献、理论、方法、统计和因果风险"], ["排版", "formatting_pipeline / blind_review", "10 个格式模块和盲审提交包"]]
    table(slide, rows, 0.70, 1.46, 11.95, 4.98, col_widths=[1.55, 4.75, 5.65], font_size=8.6)
    card(slide, "提交底线", "没有引用核验、数值核验和盲审核验，不建议投稿。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [["功能", "Skills", "用途"], ["写作", "academic-paper-composer / scientific-writing", "按提纲分章节写作和质量门"], ["实证论文", "empirical-paper-writer", "标题、摘要、引言、方法、结果、结论"], ["审稿", "peer-review / scientific-critical-thinking", "方法、统计、因果、复现和伦理自查"], ["期刊", "venue-templates", "目标期刊与模板要求"], ["引用", "citation-management", "BibTeX、DOI、正文-文末一致性"], ["文档", "docx", "Word 稿件、作者版、盲审版"], ["排版", "10 个 journal-* Skills", "题名、摘要、引用、图表、统计、布局、盲审"]]
    table(slide, rows, 0.68, 1.42, 12.0, 4.75, col_widths=[1.55, 5.05, 5.40], font_size=8.8)
    card(slide, "配套模板", "本讲提供章节追踪、引用审计、格式审计、盲审匿名化、审稿风险和投稿包 Excel。", 0.85, 6.38, 11.55, 0.62, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=11.8)

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第12讲_论文写作审稿自查与期刊排版工作区"\ncd "D:\\AI科研训练营\\第12讲_论文写作审稿自查与期刊排版工作区"\n\nNew-Item -ItemType Directory -Force -Path input,manuscript,references,figures,tables,formatting,qc,output,skills\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md', 0.85, 1.55, 11.65, 2.30, size=9.0)
    card(slide, "检查标准", "input 放前面课程材料；manuscript 放草稿；references 放 bib 和审计；formatting 放期刊格式；output 放投稿包。", 0.95, 4.30, 11.45, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName academic-paper-composer,empirical-paper-writer,scientific-writing,peer-review,scientific-critical-thinking,venue-templates,journal-title-page-metadata,journal-abstract-keywords,journal-heading-hierarchy,journal-intext-citation-style,journal-reference-list-format,journal-table-figure-caption,journal-statistics-units-style,journal-layout-spacing,journal-footnote-endnote,journal-blind-review-anonymizer,citation-management,docx `\n  -Workspace "D:\\AI科研训练营\\第12讲_论文写作审稿自查与期刊排版工作区"', 0.80, 1.48, 11.75, 3.10, size=7.6)
    card(slide, "课堂提醒", "排版 Skills 只处理结构和格式，不负责改变结论、补造引用或修改统计值。", 0.95, 5.15, 11.1, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.3)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第12讲工作区 Agent 指令\n\n你是论文写作、审稿自查与期刊排版助手。\n\n原则：\n- 禁止编造引用、DOI、页码和统计结果。\n- 禁止修改统计数值、方向和样本量。\n- 禁止把非因果设计写成因果证明。\n- 排版只能改格式，不改学术内容。\n- 缺失字段标待人工核验。\n- 盲审前检查身份线索和文件元数据。", 0.85, 1.55, 6.45, 3.50, size=9.2)
    card(slide, "为什么要写", "最后一公里最容易出现引用、数值、格式和匿名化错误。", 7.62, 1.70, 4.45, 1.05, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson12_workspace_AGENTS_template.md。", 7.62, 3.20, 4.45, 0.80, accent=COLORS["blue"])
    card(slide, "不能省略", "没有边界时，AI 会把写作优化变成内容改写。", 7.62, 4.45, 4.45, 0.80, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    slide = new("从研究包到投稿包")
    flow(slide, ["研究问题", "文献矩阵", "方法计划", "结果表", "图表计划", "证据包", "章节", "审稿", "排版", "投稿"], 0.74, 1.48, 11.9, h=0.62, size=8.3)
    rows = [["输入缺失", "后果"], ["没有文献矩阵", "gap 和引用无法核验"], ["没有结果表", "正文可能改数值或编结果"], ["没有方法计划", "Methods 不可复现"], ["没有图表计划", "Results 结构散乱"], ["没有审计", "投稿前问题不可见"]]
    table(slide, rows, 1.05, 2.75, 11.10, 2.85, col_widths=[3.2, 7.9], font_size=10.2)
    card(slide, "原则", "先把证据装箱，再写论文；不要直接让 AI 写全文。", 1.05, 6.08, 11.10, 0.58, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12)

    step_slide(prs, 4, "实操 4：生成 Evidence Pack", "把所有 claim 与证据、引用和风险绑定。", "请读取 input/ 研究问题、文献、方法、结果表和图表计划，生成 manuscript/evidence_pack.md。每条包含 claim、evidence source、result/table/figure、citation needed、support strength、risk。", "manuscript/evidence_pack.md", ["claim 有证据", "结果有来源", "引用需求清楚", "风险标注", "不编造"], "没有 evidence pack，不要进入全文写作。")

    step_slide(prs, 5, "实操 5：生成论文故事线", "把研究材料组织成问题、张力、缺口、机制和贡献。", "请基于 evidence_pack.md，生成 manuscript/manuscript_storyline.md。包括 problem、context、theoretical tension、gap、mechanism、hypotheses、method fit、findings、contribution、limitations。", "manuscript_storyline.md", ["gap 有文献支撑", "机制连接变量", "贡献不空泛", "方法匹配问题", "限制清楚"], "不要把“研究少”当作充分 gap。")

    slide = new("好故事线的结构")
    flow(slide, ["问题", "张力", "缺口", "机制", "假设", "方法", "结果", "贡献"], 0.82, 1.45, 11.7, h=0.62, size=10)
    rows = [["坏写法", "修正"], ["数字化转型很重要", "具体到某个理论或行为问题"], ["相关研究较少", "指出已有解释无法覆盖的新机制"], ["本研究方法先进", "说明方法为何能回答理论问题"], ["结果显著", "说明结果如何支持机制边界"], ["贡献重大", "写出对哪条文献链的增量"]]
    table(slide, rows, 1.05, 2.65, 11.10, 3.25, col_widths=[4.1, 7.0], font_size=10.0)

    step_slide(prs, 6, "实操 6：Reverse Outline", "给每个段落分配功能，防止论文变成材料堆砌。", "请生成 manuscript/reverse_outline.md。按章节列出 paragraph purpose、topic sentence、evidence、citation、link to next paragraph、risk。每段只服务一个核心功能。", "manuscript/reverse_outline.md", ["每段目的明确", "主题句可见", "证据对应", "段落有推进", "风险可修"], "反向提纲比目录更能暴露逻辑断点。")

    slide = new("推荐写作顺序")
    rows = [["顺序", "章节", "原因"], ["1", "Methods", "最依赖事实，先写可防止虚构"], ["2", "Results", "先固定数值和图表叙事"], ["3", "Introduction", "围绕真实结果和 gap 组织"], ["4", "Literature/Theory", "解释机制和假设"], ["5", "Discussion", "解释发现、贡献、限制"], ["6", "Abstract", "最后总结全文"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.10, col_widths=[1.4, 3.0, 7.35], font_size=10.0)
    card(slide, "不要这样做", "不要第一步就让 Agent 写完整论文；这样最容易引用不实、结果漂移和贡献夸大。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(prs, 7, "实操 7：写 Methods", "先写最可核验、最不该发挥的章节。", "请基于 input/方法计划.md 和 evidence_pack.md，写 manuscript/methods.md。包含样本、程序、变量、测量、统计方法、伦理、排除规则。不确定处标待人工核验。", "manuscript/methods.md", ["可复现", "变量清楚", "排除规则清楚", "伦理声明", "不编造"], "Methods 缺细节，审稿人会质疑复现。")

    step_slide(prs, 8, "实操 8：写 Results", "从表格和统计输出生成结果正文，不改变数值。", "请基于 input/结果表.xlsx 和 input/统计输出.md，写 manuscript/results.md。所有数值必须来自输入。报告主结果、机制、稳健性、非显著结果、效应量和 CI。", "manuscript/results.md", ["数值一致", "非显著透明", "效应量/CI", "图表对应", "不写因果过度"], "Results 只报告证据，Discussion 再解释。")

    slide = new("统计报告底线")
    rows = [["项目", "要求"], ["p 值", "不改大小、不只写显著/不显著"], ["CI", "区间方向和数值一致"], ["effect size", "能报告就报告"], ["样本量", "与方法和表格一致"], ["非显著", "计划检验需透明报告"], ["因果语言", "由设计决定，不由显著性决定"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.10, col_widths=[2.4, 9.35], font_size=10.2)
    card(slide, "排版 Skill 边界", "journal-statistics-units-style 可以改统计符号格式，但绝对不能改数值。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(prs, 9, "实操 9：写 Introduction", "用问题、文献张力、gap、方案和贡献推进。", "请基于 manuscript_storyline.md、reverse_outline.md 和 reference_audit.xlsx，写 manuscript/introduction.md。引用必须真实支持句子，不确定引用标待人工核验。", "manuscript/introduction.md", ["问题具体", "gap 有证据", "贡献可支持", "引用真实", "段落推进"], "Introduction 不是背景介绍，而是问题设置。")

    step_slide(prs, 10, "实操 10：引用核验", "把每条引用从“像真的”变成“可核验”。", "请使用 citation-management。读取 manuscript/ 草稿和 references/references.bib，生成 references/reference_audit.xlsx。检查 in-text、bib、DOI、year、journal、volume、pages、support sentence。", "references/reference_audit.xlsx", ["正文文末一致", "DOI 不编造", "缺字段标注", "句子支持核验", "待人工项清楚"], "AI 最危险的错误之一是伪造引用。")

    slide = new("引用核验表")
    rows = [["字段", "作用"], ["citation_key", "正文和 BibTeX 对齐"], ["in_text_sentence", "核验引用支持哪句话"], ["DOI / journal / year", "确认文献真实"], ["missing_fields", "缺字段标待核验"], ["supports_claim", "防止引用不支持主张"], ["manual_check", "人工数据库复核"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.10, col_widths=[3.0, 8.75], font_size=10.2)
    card(slide, "底线", "缺失 DOI、期号或页码时，Agent 只能查证或标注，不能补造。", 0.95, 6.05, 11.35, 0.72, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(prs, 11, "实操 11：审稿人式自查", "在投稿前先模拟审稿人攻击论文。", "请使用 peer-review 和 scientific-critical-thinking，生成 qc/reviewer_risk_audit.md。按 severity 列 contribution、theory、method、statistics、causality、citation、figure/table、reproducibility、ethics 风险。", "qc/reviewer_risk_audit.md", ["严重问题排前", "原因具体", "修正可执行", "must-fix 清楚", "不粉饰"], "致命设计问题不能靠润色解决。")

    slide = new("审稿风险分类")
    rows = [["风险", "典型问题"], ["Contribution", "只是换场景，没有理论增量"], ["Theory", "理论标签化，不能解释机制"], ["Method", "设计不能回答问题"], ["Statistics", "模型、效应量、CI 或稳健性不足"], ["Causality", "非因果设计写成因果证明"], ["Citation", "引用不真实或不支持句子"], ["Reproducibility", "数据、代码、材料和流程不透明"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.70, col_widths=[2.8, 8.95], font_size=9.7)

    slide = new("期刊排版 10 个 Skills")
    rows = [["编号", "Skill", "处理内容"], ["1", "journal-title-page-metadata", "标题、作者、单位、通讯作者、基金"], ["2", "journal-abstract-keywords", "摘要字数、关键词格式"], ["3", "journal-heading-hierarchy", "一级、二级、三级标题"], ["4", "journal-intext-citation-style", "正文引用格式"], ["5", "journal-reference-list-format", "参考文献列表"], ["6", "journal-table-figure-caption", "图表标题和注释"], ["7", "journal-statistics-units-style", "统计符号、p 值、CI、单位"], ["8", "journal-layout-spacing", "页边距、行距、缩进"], ["9", "journal-footnote-endnote", "脚注/尾注转换"], ["10", "journal-blind-review-anonymizer", "盲审匿名化"]]
    table(slide, rows, 0.68, 1.34, 12.0, 5.50, col_widths=[0.8, 4.1, 7.1], font_size=8.4)

    step_slide(prs, 12, "实操 12：生成排版流水线", "把目标期刊要求拆成可执行格式模块。", "请生成 formatting/formatting_pipeline.md。目标格式：[APA 7/期刊名]。依次处理 title page、abstract、headings、citations、references、captions、statistics、layout、footnotes、blind review。", "formatting_pipeline.md", ["目标规则清楚", "每步输入输出", "不改学术内容", "人工核验项", "缺规则标注"], "投稿前必须看期刊官网最新指南。")

    slide = new("APA 7 和目标期刊适配")
    rows = [["模块", "APA 7 常见要求", "目标期刊核验"], ["Title page", "标题、作者、单位、作者注", "是否匿名、running head"], ["Abstract", "通常 150-250 words", "结构式/非结构式"], ["Headings", "五级标题规则", "期刊可能自定义"], ["References", "作者年份、句式和 DOI", "参考文献数量和格式"], ["Tables/Figures", "标题、注释、编号", "分辨率、上传格式"], ["Stats", "斜体符号、p 值格式", "专门统计指南"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.30, col_widths=[2.3, 4.6, 4.85], font_size=9.3)
    card(slide, "提醒", "APA 7 可做训练格式，但正式投稿必须按目标期刊 author guidelines 二次核验。", 0.95, 6.18, 11.35, 0.62, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12)

    step_slide(prs, 13, "实操 13：盲审匿名化", "生成作者版和盲审版差异清单，清除身份线索。", "请使用 journal-blind-review-anonymizer，生成 formatting/blind_review_checklist.md。检查 author names、affiliations、acknowledgments、funding、self-citations、file properties、supplement、repository links。", "blind_review_checklist.md", ["正文无身份", "自引处理", "文件元数据", "补充材料", "仓库匿名"], "盲审不只删 title page，还要查隐藏元数据。")

    slide = new("投稿包应该包含")
    rows = [["文件", "作用"], ["manuscript_author_version.docx", "含作者信息的正式版本"], ["manuscript_blind_version.docx", "双盲审稿版本"], ["cover_letter.md", "投稿信"], ["reference_audit.xlsx", "引用真实性和完整性"], ["journal_formatting_audit.xlsx", "格式模块核验"], ["blind_review_checklist.xlsx", "匿名化核验"], ["submission_checklist.md", "最终人工提交清单"]]
    table(slide, rows, 0.78, 1.48, 11.75, 4.40, col_widths=[4.0, 7.75], font_size=9.8)
    card(slide, "最终判断", "所有 must-fix 清零后，才建议进入投稿系统。", 0.95, 6.12, 11.35, 0.62, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.2)

    slide = new("课堂示范：从 BCI 结果到论文")
    rows = [["章节", "写法"], ["Introduction", "AI 信任与脑电智能建模的理论问题，不写成纯工程任务"], ["Methods", "EEG 采集、预处理、特征、split、模型和泄漏防护"], ["Results", "subject-wise 指标、baseline、confusion matrix、失败被试"], ["Discussion", "解释模型结果与理论机制的关系，但不夸大神经机制"], ["Limitations", "offline/online、样本量、跨被试泛化、模型解释边界"], ["Formatting", "统计符号、图表标题、参考文献和盲审匿名化"]]
    table(slide, rows, 0.78, 1.45, 11.75, 4.95, col_widths=[2.2, 9.55], font_size=9.4)
    card(slide, "最大风险", "把模型准确率写成对信任神经机制的证明。", 0.95, 6.55, 11.35, 0.50, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("常见错误与修正")
    rows = [["错误", "后果", "修正"], ["直接写全文", "逻辑和证据漂移", "先 evidence pack 和 outline"], ["引用未核验", "伪引用或错引", "reference audit"], ["结果正文改数值", "学术诚信风险", "结果表锁定"], ["贡献夸大", "审稿人攻击", "reviewer risk audit"], ["排版改内容", "数据完整性风险", "格式处理只改结构"], ["盲审只删作者", "元数据泄露", "匿名化清单"]]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.55, 4.10, 5.10], font_size=9.5)

    slide = new("课堂 180 分钟带做安排")
    rows = [["时间", "教师带做", "学生产出"], ["0-15 分钟", "研究包到投稿包逻辑", "理解底线"], ["15-30 分钟", "工作区、Skills、AGENTS", "可运行工作区"], ["30-50 分钟", "整理 input 和 evidence pack", "evidence_pack"], ["50-70 分钟", "故事线", "storyline"], ["70-90 分钟", "Reverse outline", "reverse_outline"], ["90-115 分钟", "Methods / Results 写作", "section drafts"], ["115-135 分钟", "Introduction 和引用核验", "intro / reference_audit"], ["135-155 分钟", "审稿自查", "reviewer_risk_audit"], ["155-175 分钟", "排版流水线和盲审", "formatting / blind_review"], ["175-180 分钟", "投稿包检查", "submission package"]]
    table(slide, rows, 0.70, 1.42, 11.95, 5.25, col_widths=[1.75, 5.0, 5.2], font_size=8.3)

    slide = new("课后提交要求")
    bullets(slide, ["提交文件夹：姓名_第12讲_论文写作审稿自查与期刊排版", "必须包含 evidence_pack、storyline、reverse_outline、至少 3 个章节草稿、reference_audit、reviewer_risk_audit、formatting_pipeline、blind_review_checklist、cover_letter", "目标期刊不明确时，用 APA 7 训练格式，并标注投稿前需二次核验", "附 400-600 字反思：最大拒稿风险、最需核验引用、最容易出错格式"], 0.90, 1.55, 11.3, 1.90, size=13.8)
    rows = [["评分项", "占比"], ["证据包与故事线", "20%"], ["章节写作质量", "25%"], ["引用真实性核验", "20%"], ["审稿自查深度", "20%"], ["排版与盲审检查", "15%"]]
    table(slide, rows, 2.15, 3.95, 8.85, 2.10, col_widths=[6.6, 2.25], font_size=11.3)
    card(slide, "一票否决", "编造引用、编造统计、排版改数据、非因果写成因果、盲审残留身份。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.1)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是让 AI 写完一篇论文，而是建立一条从证据到写作、从审稿到排版、从作者版到盲审版的投稿生产线。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["证据", "故事线", "章节", "引用", "数值", "审稿", "排版", "盲审", "投稿"], 1.00, 3.50, 11.20, h=0.70, size=9.0)
    card(slide, "结束标准", "每个主张、引用、数值、图表和格式都能被追踪和人工核验。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(ROOT / "第12讲_课件.pptx")


def main():
    write_texts()
    create_excels()
    create_ppt()
    print(ROOT)


if __name__ == "__main__":
    main()
