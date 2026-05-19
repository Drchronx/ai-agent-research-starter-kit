from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "第06讲_课件.pptx"

W, H = Inches(13.333), Inches(7.5)

COLORS = {
    "bg": RGBColor(248, 249, 247),
    "ink": RGBColor(32, 42, 54),
    "muted": RGBColor(88, 101, 114),
    "line": RGBColor(210, 216, 222),
    "white": RGBColor(255, 255, 255),
    "teal": RGBColor(13, 116, 110),
    "green": RGBColor(22, 120, 75),
    "blue": RGBColor(37, 99, 235),
    "amber": RGBColor(181, 93, 14),
    "red": RGBColor(185, 28, 28),
    "dark": RGBColor(31, 41, 55),
    "pale_teal": RGBColor(222, 246, 243),
    "pale_blue": RGBColor(226, 235, 255),
    "pale_amber": RGBColor(255, 244, 210),
    "pale_red": RGBColor(255, 229, 229),
    "panel": RGBColor(255, 255, 255),
}

FONT = "Microsoft YaHei"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=COLORS["line"], width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    fill(slide.background, COLORS["bg"])


def textbox(slide, text, x, y, w, h, size=16, color=COLORS["ink"], bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=14, color=COLORS["ink"], gap=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = text
        p.level = level
        p.font.name = FONT
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
    return box


def codebox(slide, text, x, y, w, h, size=10.2):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, COLORS["dark"])
    no_line(rect)
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Courier New"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(238, 242, 247)
    p.line_spacing = 1.0
    return rect


def header(slide, title, idx):
    bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.34))
    fill(bar, COLORS["teal"])
    no_line(bar)
    textbox(slide, "第 06 讲 · CNKI 中文选题", 0.45, 0.055, 2.6, 0.22, size=8.5, color=COLORS["white"], bold=True)
    textbox(slide, f"{idx:02d}", 12.20, 0.05, 0.55, 0.22, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, title, 0.58, 0.60, 11.8, 0.55, size=24, color=COLORS["ink"], bold=True)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.58), Inches(1.20), Inches(1.05), Inches(0.06))
    fill(accent, COLORS["amber"])
    no_line(accent)


def card(slide, title, body, x, y, w, h, accent=COLORS["teal"], body_size=12.3, fill_color=COLORS["panel"]):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(rect, fill_color)
    line(rect)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, title, x + 0.20, y + 0.14, w - 0.30, 0.32, size=14, color=accent, bold=True)
    textbox(slide, body, x + 0.20, y + 0.56, w - 0.33, h - 0.64, size=body_size, color=COLORS["ink"])
    return rect


def table(slide, rows, x, y, w, h, col_widths=None, font_size=9.4, header_fill=COLORS["teal"]):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(244, 247, 248))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size + 0.4 if r == 0 else font_size)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.alignment = PP_ALIGN.LEFT
    return shape


def flow(slide, labels, x, y, w, h=0.58, color=COLORS["pale_teal"], accent=COLORS["teal"], size=10):
    gap = 0.13
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, label in enumerate(labels):
        bx = x + i * (box_w + gap)
        rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y), Inches(box_w), Inches(h))
        fill(rect, color)
        line(rect, accent, 0.9)
        textbox(slide, label, bx + 0.04, y + 0.14, box_w - 0.08, h - 0.20, size=size, color=COLORS["ink"], bold=True, align=PP_ALIGN.CENTER)
        if i > 0:
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(bx - gap + 0.01),
                Inches(y + h / 2),
                Inches(bx - 0.02),
                Inches(y + h / 2),
            )
            conn.line.color.rgb = accent
            conn.line.width = Pt(1.0)


def step_slide(prs, idx, title, goal, prompt, output, checks, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, title, len(prs.slides))
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(1.36), Inches(1.30), Inches(0.30))
    fill(tag, COLORS["teal"])
    no_line(tag)
    textbox(slide, f"STEP {idx}", 0.68, 1.42, 1.10, 0.15, size=8.5, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    card(slide, "目标", goal, 0.72, 1.82, 3.55, 1.15, accent=COLORS["teal"], body_size=12)
    card(slide, "输出文件", output, 0.72, 3.15, 3.55, 1.05, accent=COLORS["blue"], body_size=12)
    if note:
        card(slide, "教师提醒", note, 0.72, 4.38, 3.55, 1.15, accent=COLORS["amber"], body_size=11.5, fill_color=COLORS["pale_amber"])
    textbox(slide, "复制给 Agent 的任务", 4.65, 1.48, 3.5, 0.30, size=14, color=COLORS["teal"], bold=True)
    codebox(slide, prompt, 4.65, 1.82, 7.85, 2.70, size=9.2)
    textbox(slide, "人工核验", 4.65, 4.80, 2.0, 0.30, size=14, color=COLORS["red"], bold=True)
    bullets(slide, checks, 4.68, 5.18, 7.70, 1.25, size=12.7)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(3.05), Inches(7.5))
    fill(left, COLORS["teal"])
    no_line(left)
    textbox(slide, "第 06 讲", 0.55, 0.85, 1.7, 0.45, size=24, color=COLORS["white"], bold=True)
    textbox(slide, "CNKI 中文选题\n与 C 刊迁移", 0.52, 1.70, 2.3, 1.35, size=22, color=COLORS["white"], bold=True)
    textbox(slide, "中文热点不是研究问题", 0.55, 6.10, 2.0, 0.35, size=11.5, color=COLORS["white"], bold=True)
    textbox(slide, "从中文热点到 C 刊选题，再到 SSCI/SCI 可迁移研究问题", 3.55, 1.12, 8.75, 0.66, size=27, color=COLORS["ink"], bold=True)
    textbox(slide, "CNKI Trend · C Journal Logic · English Frontier · CN-EN Bridge · Migration Audit", 3.58, 1.95, 8.75, 0.35, size=14.5, color=COLORS["muted"])
    flow(slide, ["中文关键词", "CNKI地图", "C刊语境", "英文前沿", "桥接表", "双版本选题", "审计"], 3.60, 3.10, 8.85, h=0.62, size=10.2)
    card(slide, "本讲交付", "cnki_search_strategy、cnki_topic_map、cnki_literature_screening、cn_en_bridge_table、c_journal_idea_bank、c_vs_ssci_topic_versions、topic_migration_audit。", 3.60, 4.45, 8.72, 1.15, accent=COLORS["amber"], body_size=12.2)
    card(slide, "底线", "不编造 CNKI 结果，不绕过验证码，不把中文热点或政策口号写成学术贡献。", 3.60, 5.92, 8.72, 0.86, accent=COLORS["red"], body_size=12.5, fill_color=COLORS["pale_red"])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)

    def new(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        header(slide, title, len(prs.slides))
        return slide

    slide = new("本讲解决什么问题")
    bullets(slide, [
        "中文选题只跟热点，没有学术问题",
        "C 刊写作只讲政策意义，缺少理论机制",
        "英文前沿只做翻译，不能进入中国语境",
        "中文关键词和英文构念无法对应",
        "CNKI 检索结果多而乱，不知道如何筛选",
        "不知道如何把中文议题改写成 SSCI/SCI 版本",
    ], 0.85, 1.60, 5.90, 3.10, size=16)
    card(slide, "本讲不做", "不追求抓最多 CNKI 文献；不把政策热词直接包装成贡献；不把英文题目机械翻译成中文题目。", 7.05, 1.72, 5.10, 1.35, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "本讲要做", "用 CNKI 判断中文研究生态，用英文前沿判断国际理论位置，再用桥接表生成 C 刊与 SSCI/SCI 两个版本。", 7.05, 3.55, 5.10, 1.35, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    flow(slide, ["热点", "文献", "问题", "理论", "机制", "方法", "版本"], 0.95, 5.75, 11.35, h=0.62, size=11.5)

    slide = new("本讲最终产出文件")
    rows = [
        ["阶段", "文件", "作用"],
        ["检索策略", "topic/cnki_search_strategy.md", "记录中文关键词、检索式、排序和年份"],
        ["中文地图", "topic/cnki_topic_map.md", "整理趋势、流派、方法和 C 刊问题意识"],
        ["文献筛选", "topic/cnki_literature_screening.xlsx", "筛选中文文献并标注理论、变量、方法"],
        ["桥接分析", "topic/cn_en_bridge_table.xlsx", "连接中文议题、国际构念、理论机制和方法"],
        ["选题库", "topic/c_journal_idea_bank.md", "生成 5-8 个 C 刊候选选题"],
        ["双版本", "topic/c_vs_ssci_topic_versions.md", "输出 C 刊版本与 SSCI/SCI 版本"],
        ["审计", "qc/topic_migration_audit.md", "检查热点误用、构念错配和贡献夸大"],
    ]
    table(slide, rows, 0.72, 1.48, 11.95, 4.85, col_widths=[1.55, 4.4, 6.0], font_size=9.4)
    card(slide, "提交底线", "所有 CNKI 和英文文献都标注已核验或待核验；C 刊与 SSCI/SCI 版本不能只是互译。", 0.85, 6.48, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.8)

    slide = new("需要哪些 Skills")
    rows = [
        ["功能", "Skills", "用途"],
        ["中文总控", "cnki-research-assistant", "热点、趋势、期刊匹配、研究空白"],
        ["趋势分析", "cnki-trend", "年度趋势、学科、期刊、机构和基金"],
        ["热榜追踪", "cnki-rank", "下载榜、热词榜、新上榜"],
        ["文献抓取", "cnki-crawler-literature / cnki-crawler-topic", "多排序抓取题录和主题文献池"],
        ["浏览器检索", "cnki-exp-search-automation", "有头浏览器下处理高级搜索"],
        ["英文前沿", "ai4scholar-research / academic-research-openalex", "国际构念、理论、顶刊文献"],
        ["引用核验", "citation-management", "核验题名、作者、年份、DOI"],
    ]
    table(slide, rows, 0.70, 1.48, 12.0, 4.85, col_widths=[1.55, 4.7, 5.75], font_size=9.2)
    card(slide, "合规提醒", "CNKI 自动化只能用于合法访问和元数据整理；遇到验证码或访问限制时，手动处理或停止，不绕过限制。", 0.85, 6.45, 11.55, 0.56, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=11.5)

    slide = new("实操 1：创建工作区")
    codebox(slide, 'New-Item -ItemType Directory -Force -Path "D:\\AI科研训练营\\第06讲_CNKI中文选题与C刊迁移工作区"\ncd "D:\\AI科研训练营\\第06讲_CNKI中文选题与C刊迁移工作区"\n\nNew-Item -ItemType Directory -Force -Path input,topic,reports,qc,output,skills\nNew-Item -ItemType Directory -Force -Path input\\lesson05_outputs,reports\\cnki_trend,reports\\cnki_auto,reports\\english_frontier\nNew-Item -ItemType File -Force -Path AGENTS.md,research_log.md', 0.85, 1.55, 11.65, 2.40, size=10.5)
    card(slide, "检查标准", "有 input、topic、reports、qc、skills 文件夹；有 AGENTS.md 和 research_log.md。", 0.95, 4.35, 5.45, 0.95, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "路径提醒", "如果没有 D 盘，换成自己的项目盘。路径可以有中文，但不要同时散落在多个目录。", 6.80, 4.35, 5.45, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"])

    slide = new("实操 2：复制本讲 Skills")
    codebox(slide, 'cd "D:\\desk\\AI agent科研资料\\本地Skills功能分类库"\n\npowershell -ExecutionPolicy Bypass -File .\\00_工具脚本\\copy_skill_to_workspace.ps1 `\n  -SkillName cnki-research-assistant,cnki-trend,cnki-rank,cnki-crawler-literature,cnki-crawler-topic,cnki-exp-search-automation,ai4scholar-research,academic-research-openalex,ai4scholar-auto-citation-bibtex,citation-management,literature-review,ssci-literature-review `\n  -Workspace "D:\\AI科研训练营\\第06讲_CNKI中文选题与C刊迁移工作区"', 0.80, 1.50, 11.75, 2.80, size=9.5)
    bullets(slide, [
        "复制后检查工作区 skills/ 文件夹。",
        "缺失 Skill 记录在 research_log.md。",
        "Skills 放当前工作区，不放 Codex 根目录。",
    ], 0.95, 4.75, 11.1, 1.0, size=15)
    card(slide, "课堂提醒", "CNKI 相关 Skill 可能依赖网络、Cookie、数据库或浏览器；失败时记录原因，使用手动导出替代。", 0.95, 6.05, 11.1, 0.65, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.5)

    slide = new("实操 3：写 AGENTS.md")
    codebox(slide, "# 第06讲工作区 Agent 指令\n\n你是 CNKI 中文选题、C 刊迁移与中英文文献桥接助手。\n\n原则：\n- 禁止编造 CNKI 结果、中文文献、英文文献、DOI。\n- 不绕过 CNKI 验证码，不非法批量下载全文。\n- 政策材料只能作为语境，不能替代学术证据。\n- C 刊版本必须有中国问题意识、理论机制和方法路径。\n- SSCI/SCI 版本必须有国际理论缺口，不能只是“中国样本”。\n- 中英文桥接是构念、理论、机制和方法桥接，不是词汇翻译。", 0.85, 1.55, 6.45, 3.55, size=9.9)
    card(slide, "为什么要写", "本讲涉及中文平台和政策语境，最容易出现“热点即贡献”和“引用未核验”。AGENTS.md 用来提前设定边界。", 7.62, 1.70, 4.45, 1.30, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "可直接复制", "使用 lesson06_workspace_AGENTS_template.md，学生不需要从零写。", 7.62, 3.55, 4.45, 0.85, accent=COLORS["blue"])
    card(slide, "不能省略", "没有约束时，Agent 可能把政策口号写成理论结论。", 7.62, 4.90, 4.45, 0.85, accent=COLORS["red"], fill_color=COLORS["pale_red"])

    slide = new("实操 4：准备中英文关键词")
    codebox(slide, "# input/cn_keywords.md\n## 主关键词\n## 同义词/近义词\n## 排除词\n## 学科范围\n## 目标发表\n\n# input/en_keywords.md\n## 直接翻译\n## 国际构念\n## 国际理论\n## 方法词\n## 目标英文期刊或领域", 0.85, 1.55, 5.40, 3.00, size=11)
    card(slide, "错误做法", "中文：新质生产力；英文：new quality productive forces。只做翻译，无法进入国际理论。", 6.70, 1.60, 5.50, 1.08, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "正确方向", "先找国际构念：innovation capability、digital transformation、technology-enabled productivity，再判断中国情境是否提供新边界。", 6.70, 3.05, 5.50, 1.25, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "人工核验", "英文关键词必须来自真实英文文献或领域惯用构念，不能凭中文直译。", 0.95, 5.25, 11.20, 0.75, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    slide = new("实操 5：构造 CNKI 检索策略")
    codebox(slide, "# topic/cnki_search_strategy.md\n\nV1 宽检索：\nSU=('数据要素'+'数据资产'+'数据资源')\n\nV2 学科场景：\nSU=('数据要素'+'数据资产') and KY=('企业'+'创新'+'治理')\n\nV3 排除无关：\nSU=('数据要素'+'数据资产') and KY=('企业'+'创新') not KY=('会计准则')", 0.85, 1.50, 6.35, 2.55, size=10.5)
    rows = [
        ["排序", "含义", "用途"],
        ["PT", "发表时间", "最新研究"],
        ["CF", "被引频次", "权威研究"],
        ["DFR", "下载频次", "热点研究"],
        ["ZH", "综合排序", "相关研究"],
    ]
    table(slide, rows, 7.55, 1.55, 4.45, 2.25, col_widths=[0.9, 1.65, 1.9], font_size=10)
    bullets(slide, [
        "字段代码以 CNKI 或本地 Skill 字段表为准。",
        "引号使用英文半角。",
        "年份用参数，不写入检索式。",
        "宽检索、学科检索、排除检索要分开保存。",
    ], 0.95, 4.65, 11.1, 1.25, size=14.5)

    slide = new("实操 6：运行 CNKI 趋势分析")
    codebox(slide, 'cd "D:\\AI科研训练营\\第06讲_CNKI中文选题与C刊迁移工作区\\skills\\cnki-trend"\npython -m pip install -r scripts\\requirements.txt\npython scripts\\cnki_keyword_trend_report.py "数据要素" --output-dir "..\\..\\reports\\cnki_trend\\数据要素" --print-report', 0.85, 1.50, 11.65, 1.75, size=10)
    card(slide, "看什么", "年度趋势、近三年变化、学科分布、期刊分布、机构与基金信号。当前年份是 year-to-date。", 0.95, 3.70, 5.60, 1.20, accent=COLORS["teal"], fill_color=COLORS["pale_teal"])
    card(slide, "不能说什么", "趋势上升不等于因果关系；热点不等于学术贡献；当前年份数据不能直接与完整年份比较。", 6.85, 3.70, 5.60, 1.20, accent=COLORS["red"], fill_color=COLORS["pale_red"])
    card(slide, "失败处理", "网络、Cookie 或验证码失败时，记录在 qc/cnki_source_audit.md，改用手动导出。", 0.95, 5.55, 11.35, 0.70, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    slide = new("实操 7：整理中文文献筛选表")
    rows = [
        ["字段", "用途"],
        ["title / authors / journal / year", "基础题录与来源核验"],
        ["source_type", "CSSCI/C刊/北核/普通期刊/学位论文等"],
        ["theory / variables / method / data", "判断是否有学术机制和可检验设计"],
        ["contribution_claim / limitation", "提取作者贡献与不足"],
        ["c_journal_relevance", "判断是否适合 C 刊选题参考"],
        ["verification_status", "已核验/待核验"],
    ]
    table(slide, rows, 0.80, 1.55, 11.75, 3.85, col_widths=[4.1, 7.65], font_size=10)
    card(slide, "模板文件", "cnki_literature_screening_template.xlsx 已放在本讲资料包中。学生复制后改名为 topic/cnki_literature_screening.xlsx。", 0.95, 5.80, 11.35, 0.70, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.2)

    step_slide(
        prs, 8, "实操 8：生成中文主题地图",
        "把 CNKI 趋势、自动分析和中文文献筛选表整合成中文研究生态地图。",
        "请基于 topic/cnki_literature_screening.xlsx、reports/cnki_trend 和 reports/cnki_auto，生成 topic/cnki_topic_map.md。\n\n按中文议题概览、年度趋势、学科分布、研究流派、理论方法、C 刊问题意识、中文研究不足、国际迁移可能性组织。\n不要逐篇摘要，不要把热点直接写成贡献。",
        "topic/cnki_topic_map.md",
        ["有趋势判断", "有研究流派", "有理论和方法", "有 C 刊问题意识", "有国际迁移可能性"],
        "主题地图不是关键词云。它要能解释中文文献围绕什么问题展开。"
    )

    slide = new("中文主题地图结构")
    rows = [
        ["模块", "要回答的问题"],
        ["议题概览", "这个中文议题是什么，边界在哪里"],
        ["趋势分布", "是否集中、扩散、上升或回落"],
        ["研究流派", "中文文献从哪些路线解释问题"],
        ["理论方法", "常用理论、变量、方法和数据是什么"],
        ["C 刊问题意识", "哪些问题符合 C 刊讨论方式"],
        ["研究不足", "理论、机制、测量、数据或方法哪里不足"],
        ["国际迁移", "哪些问题能进入国际构念和理论机制"],
    ]
    table(slide, rows, 0.80, 1.55, 11.75, 4.65, col_widths=[2.1, 9.65], font_size=10.2)
    card(slide, "判断标准", "读完主题地图，应能说清：这个议题不是因为热才值得做，而是因为现有解释仍有不足。", 0.95, 6.45, 11.35, 0.55, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12)

    step_slide(
        prs, 9, "实操 9：检索英文前沿",
        "从中文议题中抽出国际构念、理论机制和方法词，而不是直接翻译关键词。",
        "请基于 input/cn_keywords.md、topic/cnki_topic_map.md 和第05讲输出，生成英文前沿检索策略。\n\n请输出：直译词、国际构念词、理论机制词、方法词、推荐检索式、推荐数据库或平台、需要核验的代表文献。\n不要只使用中文直译词。",
        "reports/english_frontier/",
        ["有国际构念", "有理论机制词", "有代表英文文献", "引用标注已核验或待核验"],
        "英文词表质量决定能否做 SSCI/SCI 迁移。"
    )

    slide = new("中英文桥接不是翻译")
    card(slide, "错误桥接", "新质生产力 = new quality productive forces。只是词汇翻译，无法说明国际理论问题。", 0.90, 1.55, 5.40, 1.05, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)
    card(slide, "正确桥接", "追问它对应国际文献中的哪个构念、机制或方法问题：innovation capability、digital transformation、technology-enabled productivity。", 6.80, 1.55, 5.40, 1.25, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.2)
    rows = [
        ["中文议题", "可能国际桥接"],
        ["数据要素", "data governance, data sharing, data-driven innovation"],
        ["数字政府", "e-government, algorithmic governance, public sector digital transformation"],
        ["平台治理", "platform governance, algorithmic control, ecosystem governance"],
        ["智能制造", "smart manufacturing, human-AI collaboration, industrial AI"],
    ]
    table(slide, rows, 0.95, 3.35, 11.25, 2.65, col_widths=[2.5, 8.75], font_size=10.5)

    step_slide(
        prs, 10, "实操 10：生成中英文桥接表",
        "把中文议题、国际构念、共同机制、中国情境和双版本选题放在同一张表里。",
        "请基于 topic/cnki_topic_map.md 和 reports/english_frontier，生成 topic/cn_en_bridge_table.xlsx 的内容草稿。\n\n每一行说明：中文议题、中文语境、CNKI代表文献、国际构念、英文关键词、国际代表文献、国际理论、共同机制、中国情境新增了什么、C刊版本、SSCI/SCI版本、迁移风险。",
        "topic/cn_en_bridge_table.xlsx",
        ["不是词汇翻译", "有共同机制", "说明中国情境新增了什么", "有迁移风险", "有人工核验项"],
        "桥接表是本讲最核心产物。没有桥接表，双版本选题容易变成空话。"
    )

    slide = new("桥接表关键字段")
    rows = [
        ["字段", "作用"],
        ["中文议题", "中国语境中的问题对象"],
        ["国际构念", "英文文献可识别的学术概念"],
        ["共同机制", "中英文文献都能理解的解释链条"],
        ["中国情境新增了什么", "新边界、新机制、新数据或新制度条件"],
        ["C刊版本选题", "中国问题意识优先"],
        ["SSCI/SCI版本选题", "国际理论缺口优先"],
        ["迁移风险", "是否只是样本替换、构念错配或贡献夸大"],
    ]
    table(slide, rows, 0.80, 1.55, 11.75, 4.65, col_widths=[2.4, 9.35], font_size=10.2)
    card(slide, "模板文件", "cn_en_bridge_table_template.xlsx 已生成，学生可以直接复制到工作区 topic/ 下填写。", 0.95, 6.45, 11.35, 0.55, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12)

    step_slide(
        prs, 11, "实操 11：生成 C 刊选题库",
        "从中文主题地图和桥接表中生成 5-8 个可讨论的 C 刊选题候选。",
        "请基于 topic/cnki_topic_map.md 和 topic/cn_en_bridge_table.xlsx，生成 topic/c_journal_idea_bank.md。\n\n生成 5-8 个 C 刊选题候选。每个选题包含中文题目、问题意识、中国情境、学术缺口、理论机制、变量方法、数据来源、目标 C 刊方向、可迁移 SSCI/SCI 版本和审稿风险。\n禁止只写宏大题目。",
        "topic/c_journal_idea_bank.md",
        ["题目不宏大空泛", "有中国问题意识", "有理论机制", "有变量和方法", "有审稿风险"],
        "C 刊选题也必须有学术机制，不是政策作文。"
    )

    slide = new("弱选题与可修改版本")
    card(slide, "弱选题", "数字经济赋能企业高质量发展研究。", 0.95, 1.60, 5.15, 0.90, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=14)
    card(slide, "问题", "变量太泛；理论机制不清楚；方法无法确定；容易同质化。", 0.95, 2.90, 5.15, 1.05, accent=COLORS["red"], body_size=12.5)
    card(slide, "C 刊版本", "数据要素市场化配置如何影响企业数字创新绩效？基于跨组织协同不确定性的机制解释。", 6.65, 1.60, 5.60, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.5)
    card(slide, "SSCI/SCI 版本", "How data marketization reduces inter-organizational coordination uncertainty and improves digital innovation performance.", 6.65, 3.15, 5.60, 1.15, accent=COLORS["blue"], fill_color=COLORS["pale_blue"], body_size=12.2)
    card(slide, "仍需核验", "CNKI 是否已有类似题目；国际文献是否已有成熟解释；中国情境是否提供新制度边界；数据是否可得。", 0.95, 5.15, 11.30, 0.80, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=12.5)

    step_slide(
        prs, 12, "实操 12：生成双版本选题",
        "将最强选题分别写成 C 刊版本和 SSCI/SCI 版本，并说明差异。",
        "请基于 topic/c_journal_idea_bank.md，选择最强 2 个选题，生成 topic/c_vs_ssci_topic_versions.md。\n\n每个选题输出 C 刊版本和 SSCI/SCI 版本。C刊版本写中国问题意识、政策/实践语境、中文文献基础、理论机制、方法设计、贡献边界；SSCI/SCI版本写 international research question、theoretical gap、why Chinese context matters、mechanism、research design、contribution boundary。",
        "topic/c_vs_ssci_topic_versions.md",
        ["两个版本不是互译", "C 刊版本不是政策作文", "SSCI/SCI 版本不是中国样本替换", "都有贡献边界"],
        "同一中文议题可以有两种发表逻辑，但不能只有一套论证。"
    )

    slide = new("C 刊版本与 SSCI/SCI 版本差异")
    rows = [
        ["维度", "C 刊版本", "SSCI/SCI 版本"],
        ["起点", "中国实践、制度、政策、治理场景", "国际理论、普遍机制、跨情境解释"],
        ["贡献", "解释中国情境中的新问题", "推进理论机制或边界条件"],
        ["文献基础", "中文核心/CSSCI + 必要英文前沿", "SSCI/SCI 顶刊 + 可迁移中国情境"],
        ["方法重点", "严谨数据、识别和政策/实践解释", "理论贡献、识别策略和可推广性"],
        ["常见风险", "政策表述多，理论机制弱", "中国情境被认为只是样本替换"],
    ]
    table(slide, rows, 0.78, 1.55, 11.75, 4.30, col_widths=[1.55, 5.1, 5.1], font_size=9.8)
    card(slide, "课堂判断", "如果两个版本只是语言不同，说明你还没有完成学术迁移。", 0.95, 6.25, 11.35, 0.65, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=12.5)

    step_slide(
        prs, 13, "实操 13：迁移风险审计",
        "用导师和审稿人视角检查双版本选题是否站得住。",
        "请以严格导师和审稿人视角，审计 topic/c_vs_ssci_topic_versions.md，生成 qc/topic_migration_audit.md。\n\n检查：中文热点是否被误写成学术贡献、政策材料是否被误写成理论证据、C刊版本是否有问题意识和机制、SSCI/SCI版本是否只是中国样本替换、中英文构念是否真正对应、方法是否能支撑结论、引用是否真实可核验。",
        "qc/topic_migration_audit.md",
        ["列出致命问题", "列出重要问题", "推荐优先版本", "给出补充检索清单"],
        "迁移审计的目标不是否定选题，而是避免错误投稿路径。"
    )

    slide = new("CNKI 来源与合规审计")
    bullets(slide, [
        "不绕过验证码，不规避机构授权限制。",
        "不非法批量下载全文；课堂只整理题录、摘要和元数据。",
        "记录检索日期、检索式、排序、年份范围和数据来源。",
        "脚本失败时记录网络、Cookie、验证码或权限原因。",
        "CNKI 结果需要人工复核：来源类型、期刊、作者、年份、摘要。",
    ], 0.90, 1.55, 6.2, 2.70, size=15)
    codebox(slide, "# qc/cnki_source_audit.md\n\n- 检索日期：\n- 检索式：\n- 数据来源：CNKI / 手动导出 / Skill 报告\n- 是否遇到验证码：\n- 是否只整理元数据：\n- 来源核验问题：\n- 待补充核验：", 7.35, 1.65, 4.80, 2.75, size=10)
    card(slide, "底线", "技术能力不能替代合规边界。CNKI 的访问限制必须尊重。", 0.95, 5.25, 11.30, 0.70, accent=COLORS["red"], fill_color=COLORS["pale_red"], body_size=13)

    slide = new("常见错误与修正")
    rows = [
        ["错误", "后果", "修正"],
        ["把政策热词当研究问题", "只有现实意义，没有学术机制", "落到变量、机制、方法"],
        ["只看 CNKI 热度", "选题同质化", "结合英文前沿和目标期刊"],
        ["英文词只做直译", "找不到国际文献", "转成国际构念和理论机制"],
        ["C刊/SSCI版本互译", "发表逻辑错位", "分别写问题意识和理论缺口"],
        ["中国情境夸大", "被认为只是样本替换", "说明新边界或新机制"],
        ["引用未核验", "学术风险", "逐条核验题名、作者、年份、来源"],
    ]
    table(slide, rows, 0.78, 1.50, 11.75, 4.95, col_widths=[2.4, 4.05, 5.30], font_size=9.8)

    slide = new("课后提交要求")
    bullets(slide, [
        "提交文件夹：姓名_第06讲_CNKI中文选题与C刊迁移",
        "必须包含检索策略、中文主题地图、中文文献筛选表、中英文桥接表、C 刊选题库、双版本选题、合规审计和迁移审计",
        "附 300 字反思：中文议题如何被英文前沿、国际理论或 C 刊语境修正",
        "所有未核验引用必须保留“待核验”标记",
    ], 0.90, 1.55, 11.3, 1.75, size=15.5)
    rows = [
        ["评分项", "占比"],
        ["CNKI 检索策略与来源合规", "20%"],
        ["中文主题地图质量", "20%"],
        ["中英文桥接表质量", "25%"],
        ["双版本选题质量", "25%"],
        ["反思与审计", "10%"],
    ]
    table(slide, rows, 2.15, 3.80, 8.85, 2.20, col_widths=[6.6, 2.25], font_size=11.5)
    card(slide, "最低合格线", "能说清：中文议题为什么值得做，国际文献对应什么构念，中国情境新增了什么，两个版本分别怎么投稿。", 1.05, 6.35, 11.15, 0.62, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=12.3)

    slide = new("教师课堂收尾")
    card(slide, "今天真正完成的事", "不是让学生学会抓 CNKI，而是学会把中文议题变成可发表的问题结构，并判断它能否迁移到国际理论。", 1.00, 1.65, 11.20, 1.15, accent=COLORS["teal"], fill_color=COLORS["pale_teal"], body_size=15)
    flow(slide, ["中文热点", "CNKI地图", "C刊逻辑", "英文前沿", "桥接表", "双版本", "迁移审计"], 1.00, 3.50, 11.20, h=0.70, size=11)
    card(slide, "下一步", "补充真实文献核验，修改双版本选题，选择更适合当前资源和发表目标的版本优先推进。", 1.00, 5.18, 11.20, 0.95, accent=COLORS["amber"], fill_color=COLORS["pale_amber"], body_size=14)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
